from typing import List, Optional
from pathlib import Path
from datetime import datetime
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.models.store import db
from app.schemas.charge import ChargeRecord, SourceRef
from app.core.config import settings


PUBLIC_FIELDS = [
    ("record_no", "记录编号"),
    ("community_name", "社区名"),
    ("street_name", "街道名"),
    ("intersection", "路口"),
    ("address", "地址"),
    ("longitude", "经度"),
    ("latitude", "纬度"),
    ("time_period", "时段"),
    ("peak_type", "高峰类型"),
    ("scenario_label", "场景标注"),
    ("side_note", "侧边说明"),
    ("screenshot_note", "截图说明"),
    ("complaint_content", "投诉内容"),
    ("complaint_count", "投诉次数"),
    ("coord_status", "坐标状态"),
    ("status", "状态"),
]

EXCEPTION_FIELDS = [
    ("record_no", "记录编号"),
    ("community_name", "社区名"),
    ("intersection", "路口"),
    ("address", "地址"),
    ("status", "状态"),
    ("conflict_with", "冲突记录ID"),
    ("coord_deviation_meters", "坐标偏差(米)"),
    ("bad_data_flags", "坏数据标签"),
    ("source_refs_summary", "溯源摘要"),
]

MERGE_FIELDS = [
    ("group_id", "归并组号"),
    ("record_no", "记录编号"),
    ("community_name", "社区名"),
    ("intersection", "路口"),
    ("address", "地址"),
    ("complaint_count", "投诉次数"),
    ("merge_candidate_ids", "候选归并ID"),
    ("merge_status", "归并状态"),
]

SOURCE_REF_FIELDS = [
    ("record_no", "记录编号"),
    ("community_name", "社区名"),
    ("source_file", "源文件名"),
    ("sheet_name", "Sheet名"),
    ("row_number", "行号"),
    ("raw_content", "原始内容"),
]


def _autosize_column_widths(ws, header_count: int):
    for col_idx in range(1, header_count + 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for row_idx, row in enumerate(ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True), 1):
            cell_value = row[0]
            if cell_value is None:
                continue
            cell_str = str(cell_value)
            length = 0
            for ch in cell_str:
                if ord(ch) > 127:
                    length += 2
                else:
                    length += 1
            if length > max_length:
                max_length = length
        adjusted_width = min(max(max_length + 2, 10), 80)
        ws.column_dimensions[col_letter].width = adjusted_width


def _apply_header_style(ws, header_count: int):
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx in range(1, header_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = bold_font
        cell.alignment = center_align


def _get_record_value(rec: ChargeRecord, field: str) -> object:
    val = getattr(rec, field, None)
    if val is None:
        return ""
    return val


def _write_sheet1_public_list(ws, records: List[ChargeRecord]):
    headers = [h for _, h in PUBLIC_FIELDS]
    ws.append(headers)
    _apply_header_style(ws, len(headers))

    for rec in records:
        if rec.status and rec.status.startswith("merged_into_"):
            continue
        row = []
        for field, _ in PUBLIC_FIELDS:
            row.append(_get_record_value(rec, field))
        ws.append(row)

    _autosize_column_widths(ws, len(headers))


def _write_sheet2_exceptions(ws, records: List[ChargeRecord]):
    headers = [h for _, h in EXCEPTION_FIELDS]
    ws.append(headers)
    _apply_header_style(ws, len(headers))

    for rec in records:
        status = rec.status or ""
        coord_status = rec.coord_status or ""
        bad_flags = rec.bad_data_flags or []
        is_exception = (
            status in ("conflict", "suspended", "bad_data")
            or coord_status == "suspended"
            or len(bad_flags) > 0
        )
        if not is_exception:
            continue

        conflict_with = rec.conflict_with or []
        conflict_str = ",".join(conflict_with) if conflict_with else ""

        coord_dev = rec.coord_deviation_meters
        coord_dev_str = f"{coord_dev:.2f}" if coord_dev is not None else ""

        bad_flags_str = ";".join(bad_flags) if bad_flags else ""

        refs = rec.source_refs or []
        refs_summary_parts = []
        for ref in refs[:3]:
            part = f"{ref.source_file}:{ref.row_number}"
            refs_summary_parts.append(part)
        if len(refs) > 3:
            refs_summary_parts.append(f"...共{len(refs)}条")
        refs_summary = " | ".join(refs_summary_parts)

        row = [
            _get_record_value(rec, "record_no"),
            _get_record_value(rec, "community_name"),
            _get_record_value(rec, "intersection"),
            _get_record_value(rec, "address"),
            _get_record_value(rec, "status"),
            conflict_str,
            coord_dev_str,
            bad_flags_str,
            refs_summary,
        ]
        ws.append(row)

    _autosize_column_widths(ws, len(headers))


def _write_sheet3_merge_candidates(ws, records: List[ChargeRecord]):
    headers = [h for _, h in MERGE_FIELDS]
    ws.append(headers)
    _apply_header_style(ws, len(headers))

    candidate_records = [r for r in records if r.merge_status == "candidate"]
    processed_ids = set()
    group_id = 0

    for rec in candidate_records:
        if rec.id in processed_ids:
            continue
        candidates = rec.merge_candidate_ids or []
        if not candidates:
            continue

        group_id += 1

        group_ids = [rec.id] + [c for c in candidates if c not in processed_ids]
        group_records = [db.get_record(gid) for gid in group_ids]
        group_records = [g for g in group_records if g is not None]

        for g_rec in group_records:
            if not g_rec:
                continue
            processed_ids.add(g_rec.id)

            merge_candidate_ids = g_rec.merge_candidate_ids or []
            merge_candidates_str = ",".join(merge_candidate_ids) if merge_candidate_ids else ""

            row = [
                f"G{group_id:03d}",
                _get_record_value(g_rec, "record_no"),
                _get_record_value(g_rec, "community_name"),
                _get_record_value(g_rec, "intersection"),
                _get_record_value(g_rec, "address"),
                _get_record_value(g_rec, "complaint_count"),
                merge_candidates_str,
                _get_record_value(g_rec, "merge_status"),
            ]
            ws.append(row)

    _autosize_column_widths(ws, len(headers))


def _write_sheet4_source_refs(ws, records: List[ChargeRecord]):
    headers = [h for _, h in SOURCE_REF_FIELDS]
    ws.append(headers)
    _apply_header_style(ws, len(headers))

    for rec in records:
        refs: List[SourceRef] = rec.source_refs or []
        for ref in refs:
            raw_content = ref.raw_content or {}
            try:
                raw_str = json.dumps(raw_content, ensure_ascii=False)
            except Exception:
                raw_str = str(raw_content)

            row = [
                _get_record_value(rec, "record_no"),
                _get_record_value(rec, "community_name"),
                ref.source_file or "",
                ref.sheet_name or "",
                ref.row_number,
                raw_str,
            ]
            ws.append(row)

    _autosize_column_widths(ws, len(headers))


def export_public_list(
    record_ids: Optional[List[str]] = None,
    output_path: Optional[str] = None,
) -> str:
    settings.EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    if output_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = str(settings.EXPORT_DIR / f"公示清单_{timestamp}.xlsx")

    output = Path(output_path)
    if output.parent != settings.EXPORT_DIR:
        output = settings.EXPORT_DIR / output.name

    all_records = db.list_records()
    if record_ids:
        records_to_export = [r for r in all_records if r.id in record_ids]
    else:
        records_to_export = all_records

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "公示清单明细"
    _write_sheet1_public_list(ws1, records_to_export)

    ws2 = wb.create_sheet("异常待确认")
    _write_sheet2_exceptions(ws2, records_to_export)

    ws3 = wb.create_sheet("归并建议")
    _write_sheet3_merge_candidates(ws3, records_to_export)

    ws4 = wb.create_sheet("溯源引用")
    _write_sheet4_source_refs(ws4, records_to_export)

    wb.save(str(output))

    return str(output)
