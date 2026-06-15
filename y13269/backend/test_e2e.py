import sys, os, uuid, json
sys.path.insert(0, '.')

from app.core.config import settings
from app.models.store import db
from app.services.data_import import incremental_import, parse_excel_or_csv
from app.services.bad_data import detect_bad_data, scan_conflicts, resolve_conflict
from app.services.merge_service import detect_merge_candidates, propose_merge, execute_merge
from app.services.consistency_service import detect_peak_conflict
from app.services.coord_service import batch_verify, manual_confirm
from app.services.export_service import export_public_list
from app.schemas.charge import ImportBatch
from datetime import datetime

db.records.clear()
db.unified_notes.clear()
db.batches.clear()

print("=" * 60)
print("[步骤1] 导入第1周会议纪要")
print("=" * 60)
path1 = "/Users/mac/pro/solo/workspaces/y13269/sample_data/会议纪要_2026年第1周.xlsx"
rows, mapping = parse_excel_or_csv(path1)
print(f"解析 {len(rows)} 行")

batch1_id = uuid.uuid4().hex[:12]
batch1, recs1 = incremental_import(rows, os.path.basename(path1), batch1_id, os.path.getsize(path1))
db.add_batch(batch1)
print(f"新增 {batch1.new_records} 条，更新 {batch1.updated_records} 条")

print("  导入后执行坏数据检测...")
bd_res = detect_bad_data()
print(f"  坏数据: {len(bd_res)} 条")

for r in db.list_records():
    print(f"  [{r.status:8s}|coord={r.coord_status:8s}|merge={r.merge_status:9s}] {r.community_name[:8]:8s} | {r.peak_type or '':4s} | 投诉{r.complaint_count:2d}次 | 冲突:{r.conflict_with} | 坏:{r.bad_data_flags}")

print("\n" + "=" * 60)
print("[步骤2] 扫描早晚高峰口径冲突")
print("=" * 60)
detect_peak_conflict(db.list_records())
conflicts_list = scan_conflicts()
print(f"发现 {len(conflicts_list)} 组冲突记录：")
for a, b, fields in conflicts_list:
    print(f"  {a.community_name[:8]:8s}: A({a.peak_type}) vs B({b.peak_type}) 字段差异={fields}")

print("\n" + "=" * 60)
print("[步骤3] 批量坐标校验")
print("=" * 60)
verify_results = batch_verify()
print(f"坐标校验 {len(verify_results)} 条：")
for r in verify_results:
    dev = r.coord_deviation_meters or 0
    print(f"  {r.community_name[:8]:8s}: {r.coord_status:8s} | 偏差 {dev:6.1f}m | 反查={(r.coord_verified_address or '')[:20]}")
suspended = [r for r in db.list_records() if r.coord_status in ('suspended', 'suspicious')]
print(f"\n挂起需人工确认 {len(suspended)} 条：")
for r in suspended:
    print(f"  {r.community_name[:8]:8s} ({r.longitude},{r.latitude}) 疑似偏到隔壁街")

print("\n" + "=" * 60)
print("[步骤4] 扫描投诉归并候选")
print("=" * 60)
merge_groups = detect_merge_candidates()
print(f"发现 {len(merge_groups)} 组归并候选：")
for idx, group in enumerate(merge_groups):
    print(f"  组#{idx+1}: {[(r.community_name[:6], r.peak_type, r.id[:6]) for r in group]}")

print("\n" + "=" * 60)
print("[步骤5] 导入第2周会议纪要（后补材料不覆盖原判断）")
print("=" * 60)
path2 = "/Users/mac/pro/solo/workspaces/y13269/sample_data/会议纪要_2026年第2周.xlsx"
rows2, mapping2 = parse_excel_or_csv(path2)
batch2_id = uuid.uuid4().hex[:12]
batch2, recs2 = incremental_import(rows2, os.path.basename(path2), batch2_id, os.path.getsize(path2))
db.add_batch(batch2)
print(f"新增 {batch2.new_records} 条，更新 {batch2.updated_records} 条")

print("\n★ 关键验证-阳光花园早高峰场景是否保留最早判断？")
for r in db.list_records():
    if r.community_name == '阳光花园小区' and r.peak_type == '早高峰':
        print(f"  场景标注 = [{r.scenario_label}]")
        print(f"  溯源记录 {len(r.source_refs)} 条：")
        for s in r.source_refs:
            raw_scene = s.raw_content.get('场景标注', '')
            print(f"    来源: {s.source_file[:24]:24s} 行{s.row_number:2d} 场景=[{raw_scene}]")

print("\n" + "=" * 60)
print("[步骤6] 生成归并建议并执行（翠湖苑同街口多条投诉）")
print("=" * 60)
cuihu_groups = [g for g in detect_merge_candidates() if any(r.community_name == '翠湖苑' for r in g)]
if cuihu_groups:
    grp = cuihu_groups[0]
    a, b = grp[0], grp[1]
    print(f"候选 A(id={a.id[:8]}, peak={a.peak_type}, 投诉{a.complaint_count}次)")
    print(f"候选 B(id={b.id[:8]}, peak={b.peak_type}, 投诉{b.complaint_count}次)")
    prop = propose_merge(a.id, b.id)
    keep = prop['keep_record']
    remove = prop['remove_record']
    sugg = prop['proposed_data']
    print(f"建议保留 {keep.id[:8]} ({keep.peak_type})")
    print(f"建议字段: 场景=[{sugg.get('scenario_label','')[:30]}] 投诉数={sugg.get('complaint_count')} 溯源={len(keep.source_refs)+len(remove.source_refs)}")
    merged = execute_merge(keep.id, remove.id)
    print(f"执行归并 OK，保留记录 status={merged.status}, merge={merged.merge_status}, 投诉={merged.complaint_count}, 溯源={len(merged.source_refs)}")

print("\n" + "=" * 60)
print("[步骤7] 解决阳光花园早晚高峰口径冲突")
print("=" * 60)
yg_groups = [g for g in scan_conflicts() if any(r.community_name == '阳光花园小区' for r in g[:2])]
if yg_groups:
    a, b, fields = yg_groups[0]
    print(f"冲突 A ({a.peak_type}): 场景=[{a.scenario_label}] 侧边=[{a.side_note}]")
    print(f"冲突 B ({b.peak_type}): 场景=[{b.scenario_label}] 侧边=[{b.side_note}]")
    chosen = {
        "scenario_label": f"{a.peak_type}:{a.scenario_label} / {b.peak_type}:{b.scenario_label}",
        "side_note": "早晚高峰问题分别处理，统一口径后公示",
        "screenshot_note": "保留各时段截图说明"
    }
    res_a, res_b = resolve_conflict(a.id, b.id, chosen)
    print(f"冲突解决完成！双方 status={res_a.status}/{res_b.status}")
    print(f"  统一场景=[{res_a.scenario_label}]")
    print(f"  统一侧边=[{res_a.side_note}]")

print("\n" + "=" * 60)
print("[步骤8] 人工确认坐标挂起记录")
print("=" * 60)
for r in list(suspended):
    reloaded = db.get_record(r.id)
    if reloaded and reloaded.coord_status in ('suspended', 'suspicious'):
        confirmed = manual_confirm(reloaded.id, True, f"经算法值班人确认：{reloaded.community_name}坐标正确")
        if confirmed:
            print(f"  {confirmed.community_name[:8]:8s}: coord_status={confirmed.coord_status}, 已人工确认")

print("\n" + "=" * 60)
print("[步骤9] 多Sheet Excel 导出")
print("=" * 60)
out_path = export_public_list()
print(f"导出文件: {out_path}")
print(f"文件大小: {os.path.getsize(out_path)/1024:.1f} KB")

import openpyxl
wb = openpyxl.load_workbook(out_path)
print(f"Sheet列表: {wb.sheetnames}")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"  [{name}] {ws.max_row-1:3d}行 x {ws.max_column:2d}列")

print("\n" + "=" * 60)
print("[最终统计]")
print("=" * 60)
recs = db.list_records()
stats = {
    "total": len(recs),
    "normal": len([r for r in recs if r.status == 'normal']),
    "draft": len([r for r in recs if r.status == 'draft']),
    "conflict": len([r for r in recs if r.status == 'conflict']),
    "suspended": len([r for r in recs if r.status in ('suspended',)]),
    "bad_data": len([r for r in recs if r.bad_data_flags]),
    "merged": len([r for r in recs if str(r.status).startswith('merged')]),
    "merge_candidates": len([r for r in recs if r.merge_status == 'candidate']),
    "coord_verified": len([r for r in recs if r.coord_status == 'verified']),
    "coord_pending": len([r for r in recs if r.coord_status == 'pending']),
}
print(json.dumps(stats, ensure_ascii=False, indent=2))

db.save_to_disk()
print("\n✓ 持久化完成")
