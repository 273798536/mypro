import os
import pandas as pd
import numpy as np
from typing import Optional, List, Dict
from dataclasses import dataclass
import json
from datetime import datetime
from openpyxl import Workbook

from .data_loader import LoadedDataset
from .filter_sync import FilteredResult
from .quality_check import QualityReport, IssueType, Severity
from .unit_validation import UnitReport, UnitCategory
from .numerical_integration import IntegrationReport, METHOD_NAMES, METHOD_DESCRIPTIONS
from .charts import ChartOutput


@dataclass
class ReportOutput:
    file_path: str
    report_type: str
    title: str


def _translate_severity(sev: Severity) -> str:
    mapping = {
        Severity.ERROR: '需重点关注',
        Severity.WARNING: '建议关注',
        Severity.INFO: '一般提示',
    }
    return mapping.get(sev, sev.value)


def _describe_issue_type(itype: IssueType) -> Dict[str, str]:
    mapping = {
        IssueType.MISSING_SAMPLE: {
            'name': '采样中断',
            'explanation': '仪表在这段时间没有传回数据，可能是网络问题或仪表断电',
            'impact': '缺失时段的用电量是估算的，可能与实际有偏差',
        },
        IssueType.TIME_OUT_OF_ORDER: {
            'name': '时间顺序错误',
            'explanation': '数据传回的时间顺序不对，早的数据反而晚到',
            'impact': '如不修正会导致用电量计算错误',
        },
        IssueType.NEGATIVE_READING: {
            'name': '出现负数',
            'explanation': '用电量读数出现了负数，正常情况下用电不会产生负数',
            'impact': '可能是仪表故障或反向送电，需要现场核实',
        },
        IssueType.NULL_READING: {
            'name': '数据空缺',
            'explanation': '个别时间点没有读数',
            'impact': '影响不大，已按相邻数据自动补充',
        },
    }
    return mapping.get(itype, {'name': itype.value, 'explanation': '', 'impact': ''})


def generate_layperson_report(dataset: LoadedDataset,
                              filtered_result: FilteredResult,
                              quality_report: Optional[QualityReport] = None,
                              unit_report: Optional[UnitReport] = None,
                              charts: Optional[Dict[str, List[ChartOutput]]] = None,
                              output_dir: str = 'output/reports') -> ReportOutput:
    os.makedirs(output_dir, exist_ok=True)

    config = filtered_result.config.to_dict()
    integration_report = filtered_result.integration_report
    unit_report = unit_report or filtered_result.unit_report

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    file_path = os.path.join(output_dir, f'能耗估算报告_{timestamp}.xlsx')

    wb = Workbook()
    if wb.active is not None:
        wb.remove(wb.active)

    _generate_summary_sheet(wb, dataset, filtered_result, quality_report)
    _generate_result_sheet(wb, integration_report, unit_report, config)
    _generate_quality_sheet(wb, quality_report)
    _generate_calculation_sheet(wb, unit_report, integration_report, config)
    _generate_data_sheet(wb, filtered_result)

    wb.save(file_path)

    return ReportOutput(
        file_path=file_path,
        report_type='layperson',
        title='能耗估算报告（外行友好版）'
    )


def _generate_summary_sheet(wb, dataset, filtered_result, quality_report):
    ws = wb.create_sheet('报告概览', 0)

    ws['A1'] = '能耗估算报告'
    ws['A1'].font = ws['A1'].font.copy(size=18, bold=True)
    ws.merge_cells('A1:D1')

    ws['A3'] = '一、基本信息'
    ws['A3'].font = ws['A3'].font.copy(bold=True, size=12)

    basic_info = [
        ('数据来源文件', os.path.basename(dataset.source_file)),
        ('统计时间范围', f"{filtered_result.config.time_start.strftime('%Y-%m-%d %H:%M')} 至 {filtered_result.config.time_end.strftime('%Y-%m-%d %H:%M')}"),
        ('统计时长', f"{filtered_result.integration_report.total_duration_hours:.1f} 小时"),
        ('统计口径', filtered_result.config.to_dict()['group_by_name']),
        ('报告生成时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
    ]

    row = 4
    for label, value in basic_info:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    ws.cell(row=row + 1, column=1, value='二、核心结论')
    ws.cell(row=row + 1, column=1).font = ws.cell(row=row + 1, column=1).font.copy(bold=True, size=12)
    row += 2

    total_energy = filtered_result.integration_report.total_energy_all_columns
    energy_unit = filtered_result.unit_report.target_energy_unit

    ws.cell(row=row, column=1, value='总用电量')
    ws.cell(row=row, column=2, value=f'{total_energy:,.2f} {energy_unit}')
    ws.cell(row=row, column=2).font = ws.cell(row=row, column=2).font.copy(bold=True, size=14, color='008000')

    row += 2
    ws.cell(row=row, column=1, value='各回路用电量:')
    ws.cell(row=row, column=1).font = ws.cell(row=row, column=1).font.copy(bold=True)
    row += 1

    for col, result in filtered_result.integration_report.results.items():
        ws.cell(row=row, column=2, value=col)
        ws.cell(row=row, column=3, value=f'{result.total_energy:,.2f} {result.final_unit}')
        if len(filtered_result.integration_report.results) > 1 and total_energy > 0:
            pct = result.total_energy / total_energy * 100
            ws.cell(row=row, column=4, value=f'占比 {pct:.1f}%')
        row += 1

    row += 1
    ws.cell(row=row, column=1, value='三、数据质量说明')
    ws.cell(row=row, column=1).font = ws.cell(row=row, column=1).font.copy(bold=True, size=12)
    row += 1

    if quality_report and quality_report.issues:
        score = quality_report.data_quality_score
        ws.cell(row=row, column=1, value='数据质量评分')
        ws.cell(row=row, column=2, value=f'{score:.1f}/100')
        color = '008000' if score >= 90 else 'FFA500' if score >= 70 else 'FF0000'
        ws.cell(row=row, column=2).font = ws.cell(row=row, column=2).font.copy(bold=True, color=color)
        row += 2

        ws.cell(row=row, column=1, value='发现的问题:')
        ws.cell(row=row, column=1).font = ws.cell(row=row, column=1).font.copy(bold=True)
        row += 1

        for issue in quality_report.issues:
            desc = _describe_issue_type(issue.issue_type)
            sev = _translate_severity(issue.severity)
            ws.cell(row=row, column=2, value=f"【{sev}】{desc['name']}")
            ws.cell(row=row, column=2).font = ws.cell(row=row, column=2).font.copy(bold=True)
            row += 1
            ws.cell(row=row, column=3, value=f"数量: {issue.count} 处")
            row += 1
            ws.cell(row=row, column=3, value=f"说明: {desc['explanation']}")
            row += 1
            ws.cell(row=row, column=3, value=f"影响: {desc['impact']}")
            row += 2
    else:
        ws.cell(row=row, column=2, value='数据质量良好，未发现明显异常。')
        ws.cell(row=row, column=2).font = ws.cell(row=row, column=2).font.copy(color='008000')

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 15


def _generate_result_sheet(wb, integration_report, unit_report, config):
    ws = wb.create_sheet('用电量明细', 1)

    ws['A1'] = '各时段用电量明细'
    ws['A1'].font = ws['A1'].font.copy(size=14, bold=True)
    ws.merge_cells('A1:D1')

    ws['A3'] = f"计算方法: {config['integration_method_name']}"
    ws['A4'] = f"分组方式: {config['group_by_name']}"

    has_groups = any(r.group_results for r in integration_report.results.values())

    if has_groups:
        all_groups = set()
        for result in integration_report.results.values():
            all_groups.update(result.group_results.keys())
        all_groups = sorted(all_groups)

        headers = ['时间段'] + [f'{col} ({r.final_unit})' for col, r in integration_report.results.items()]
        if len(integration_report.results) > 1:
            headers.append('合计')

        ws.cell(row=6, column=1, value=headers[0])
        for i, h in enumerate(headers[1:], 1):
            ws.cell(row=6, column=i + 1, value=h)
        for cell in ws[6]:
            cell.font = cell.font.copy(bold=True)

        row = 7
        for group in all_groups:
            ws.cell(row=row, column=1, value=str(group))
            total_row = 0
            for col_idx, (col, result) in enumerate(integration_report.results.items(), 2):
                val = result.group_results.get(group, 0)
                ws.cell(row=row, column=col_idx, value=round(val, 2))
                total_row += val
            if len(integration_report.results) > 1:
                ws.cell(row=row, column=len(headers), value=round(total_row, 2))
            row += 1

        ws.cell(row=row, column=1, value='总计')
        ws.cell(row=row, column=1).font = ws.cell(row=row, column=1).font.copy(bold=True)
        for col_idx, (col, result) in enumerate(integration_report.results.items(), 2):
            ws.cell(row=row, column=col_idx, value=round(result.total_energy, 2))
            ws.cell(row=row, column=col_idx).font = ws.cell(row=row, column=col_idx).font.copy(bold=True)
        if len(integration_report.results) > 1:
            ws.cell(row=row, column=len(headers), value=round(integration_report.total_energy_all_columns, 2))
            ws.cell(row=row, column=len(headers)).font = ws.cell(row=row, column=len(headers)).font.copy(bold=True)

    else:
        headers = ['用电回路', '原单位', '标准单位', '总用电量', '说明']
        for i, h in enumerate(headers, 1):
            ws.cell(row=6, column=i, value=h)
        for cell in ws[6]:
            cell.font = cell.font.copy(bold=True)

        row = 7
        for col, result in integration_report.results.items():
            ws.cell(row=row, column=1, value=col)
            ws.cell(row=row, column=2, value=result.original_unit)
            ws.cell(row=row, column=3, value=result.final_unit)
            ws.cell(row=row, column=4, value=round(result.total_energy, 2))
            note = '; '.join(result.calculation_notes) if result.calculation_notes else '通过数值积分计算'
            ws.cell(row=row, column=5, value=note)
            row += 1

    ws.column_dimensions['A'].width = 20
    for col_idx in range(2, 6):
        ws.column_dimensions[chr(64 + col_idx)].width = 18


def _generate_quality_sheet(wb, quality_report):
    ws = wb.create_sheet('数据质量检查', 2)

    ws['A1'] = '数据质量检查详细报告'
    ws['A1'].font = ws['A1'].font.copy(size=14, bold=True)
    ws.merge_cells('A1:E1')

    if not quality_report or not quality_report.issues:
        ws['A3'] = '经检查，数据质量良好，未发现异常问题。'
        ws['A3'].font = ws['A3'].font.copy(color='008000')
        return

    headers = ['问题类型', '严重程度', '数量', '影响比例', '详细说明']
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    for cell in ws[3]:
        cell.font = cell.font.copy(bold=True)

    row = 4
    for issue in quality_report.issues:
        desc = _describe_issue_type(issue.issue_type)
        sev = _translate_severity(issue.severity)

        ws.cell(row=row, column=1, value=desc['name'])
        ws.cell(row=row, column=2, value=sev)
        ws.cell(row=row, column=3, value=issue.count)

        pct = issue.details.get('affected_rows_pct', 0)
        ws.cell(row=row, column=4, value=f'{pct:.2f}%')

        detail_parts = [desc['explanation'], desc['impact']]
        if issue.details.get('first_occurrence'):
            detail_parts.append(f"首次出现: {issue.details['first_occurrence']}")
        if issue.details.get('per_column'):
            col_info = '; '.join([f"{k} {v['count']}个" for k, v in issue.details['per_column'].items()])
            detail_parts.append(f"涉及回路: {col_info}")
        if issue.details.get('missing_gaps'):
            gaps = issue.details['missing_gaps'][:3]
            gap_desc = '; '.join([f"{g['gap_start']}~{g['gap_end']}({g['expected_missing_points']}个点)" for g in gaps])
            if len(issue.details['missing_gaps']) > 3:
                gap_desc += f" ... 共{len(issue.details['missing_gaps'])}处"
            detail_parts.append(f"缺失区间: {gap_desc}")

        ws.cell(row=row, column=5, value='；'.join(detail_parts))

        color = 'FF0000' if issue.severity == Severity.ERROR else 'FFA500' if issue.severity == Severity.WARNING else '0000FF'
        for col in range(1, 6):
            ws.cell(row=row, column=col).font = ws.cell(row=row, column=col).font.copy(color=color)

        row += 1

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 60


def _generate_calculation_sheet(wb, unit_report, integration_report, config):
    ws = wb.create_sheet('计算说明', 3)

    ws['A1'] = '计算方法说明'
    ws['A1'].font = ws['A1'].font.copy(size=14, bold=True)
    ws.merge_cells('A1:D1')

    row = 3
    ws.cell(row=row, column=1, value='一、单位转换说明')
    ws.cell(row=row, column=1).font = ws.cell(row=row, column=1).font.copy(bold=True, size=12)
    row += 1

    for col, result in unit_report.results.items():
        ws.cell(row=row, column=2, value=f"{col}:")
        ws.cell(row=row, column=2).font = ws.cell(row=row, column=2).font.copy(bold=True)
        row += 1

        ws.cell(row=row, column=3, value=f"原始单位: {result.original_unit}")
        row += 1
        ws.cell(row=row, column=3, value=f"标准单位: {result.standardized_unit}")
        row += 1

        if result.conversion_steps:
            ws.cell(row=row, column=3, value="转换过程:")
            row += 1
            for step in result.conversion_steps:
                ws.cell(row=row, column=4, value=f"公式: {step.formula}")
                row += 1
                ws.cell(row=row, column=4, value=f"示例: {step.example}")
                row += 1
        else:
            ws.cell(row=row, column=3, value="单位一致，无需转换")
            row += 1

        for note in result.validation_notes:
            ws.cell(row=row, column=3, value=f"备注: {note}")
            row += 1
        row += 1

    for note in unit_report.summary_notes:
        ws.cell(row=row, column=2, value=note)
        row += 1

    row += 2
    ws.cell(row=row, column=1, value='二、数值积分方法说明')
    ws.cell(row=row, column=1).font = ws.cell(row=row, column=1).font.copy(bold=True, size=12)
    row += 1

    method_name = config['integration_method_name']
    method_desc = METHOD_DESCRIPTIONS.get(integration_report.method_used, '')

    ws.cell(row=row, column=2, value=f"使用方法: {method_name}")
    ws.cell(row=row, column=2).font = ws.cell(row=row, column=2).font.copy(bold=True)
    row += 1
    ws.cell(row=row, column=3, value=method_desc)
    row += 2

    for col, result in integration_report.results.items():
        unit_result = unit_report.results.get(col)
        if unit_result and unit_result.category == UnitCategory.POWER and result.steps:
            ws.cell(row=row, column=2, value=f"{col} 前3步计算示例:")
            ws.cell(row=row, column=2).font = ws.cell(row=row, column=2).font.copy(bold=True)
            row += 1

            for step in result.steps[:3]:
                ws.cell(row=row, column=3, value=f"时段: {step.time_start} ~ {step.time_end}")
                row += 1
                ws.cell(row=row, column=4, value=f"公式: {step.formula}")
                row += 1
                ws.cell(row=row, column=4, value=f"计算: {step.calculation}")
                row += 1
            row += 1

            if result.calculation_notes:
                for note in result.calculation_notes:
                    ws.cell(row=row, column=3, value=f"注意: {note}")
                    row += 1
                row += 1

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 50


def _generate_data_sheet(wb, filtered_result):
    ws = wb.create_sheet('原始数据（备查）', 4)

    ws['A1'] = '原始用电数据（供核对使用）'
    ws['A1'].font = ws['A1'].font.copy(size=14, bold=True)
    ws.merge_cells('A1:E1')

    ws['A3'] = '说明: 此表保留原始读数，便于核对。处理后的数据用于实际计算。'
    ws['A3'].font = ws['A3'].font.copy(italic=True, color='666666')

    df = filtered_result.dataset.raw_df
    time_col = filtered_result.dataset.time_column
    reading_cols = filtered_result.dataset.reading_columns

    display_cols = [time_col] + reading_cols
    df_display = df[display_cols].copy()

    for i, col in enumerate(display_cols, 1):
        ws.cell(row=5, column=i, value=str(col))
    for cell in ws[5]:
        cell.font = cell.font.copy(bold=True)

    for row_idx, (_, row_data) in enumerate(df_display.iterrows(), 6):
        for col_idx, col in enumerate(display_cols, 1):
            val = row_data[col]
            if pd.isna(val):
                ws.cell(row=row_idx, column=col_idx, value='')
            elif isinstance(val, pd.Timestamp):
                ws.cell(row=row_idx, column=col_idx, value=val.strftime('%Y-%m-%d %H:%M:%S'))
            else:
                ws.cell(row=row_idx, column=col_idx, value=float(val) if isinstance(val, (int, float)) else str(val))

    for i, col in enumerate(display_cols, 1):
        ws.column_dimensions[chr(64 + i)].width = max(18, len(str(col)) * 2)


def generate_full_report(dataset: LoadedDataset,
                         filtered_result: FilteredResult,
                         quality_report: Optional[QualityReport] = None,
                         unit_report: Optional[UnitReport] = None,
                         charts: Optional[Dict[str, List[ChartOutput]]] = None,
                         output_dir: str = 'output/reports') -> Dict[str, ReportOutput]:
    outputs = {}

    outputs['layperson'] = generate_layperson_report(
        dataset, filtered_result, quality_report, unit_report, charts, output_dir
    )

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    detail_path = os.path.join(output_dir, f'能耗估算_计算细节_{timestamp}.json')
    with open(detail_path, 'w', encoding='utf-8') as f:
        detail_data = {
            'filter_config': filtered_result.config.to_dict(),
            'filter_description': filtered_result.filter_description,
            'data_points': {
                'before': filtered_result.data_points_before,
                'after': filtered_result.data_points_after,
            },
            'unit_validation': {
                col: {
                    'original_unit': r.original_unit,
                    'standardized_unit': r.standardized_unit,
                    'category': r.category.value,
                    'conversion_steps': [
                        {
                            'from': s.from_unit,
                            'to': s.to_unit,
                            'factor': s.conversion_factor,
                            'formula': s.formula,
                            'example': s.example,
                        } for s in r.conversion_steps
                    ],
                    'value_range': r.value_range,
                } for col, r in (unit_report or filtered_result.unit_report).results.items()
            },
            'integration_results': {
                col: {
                    'total_energy': r.total_energy,
                    'unit': r.final_unit,
                    'method': METHOD_NAMES.get(r.method, str(r.method)),
                    'group_results': r.group_results,
                    'calculation_notes': r.calculation_notes,
                    'sample_steps': [
                        {
                            'time_start': str(s.time_start),
                            'time_end': str(s.time_end),
                            'interval_hours': s.time_interval_hours,
                            'power_start': s.power_start,
                            'power_end': s.power_end,
                            'energy': s.energy_contribution,
                            'formula': s.formula,
                            'calculation': s.calculation,
                        } for s in r.steps[:5]
                    ],
                } for col, r in filtered_result.integration_report.results.items()
            },
            'quality_issues': [
                {
                    'type': iss.issue_type.value,
                    'severity': iss.severity.value,
                    'count': iss.count,
                    'description': iss.description,
                    'details': iss.details,
                } for iss in (quality_report.issues if quality_report else [])
            ],
        }
        json.dump(detail_data, f, ensure_ascii=False, indent=2, default=str)

    outputs['detail_json'] = ReportOutput(
        file_path=detail_path,
        report_type='detail_json',
        title='能耗估算计算细节（JSON）'
    )

    return outputs
