from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

from ..core import AuditResult
from ..versioning import VersionTracker


class ExcelExporter:
    def __init__(self):
        self.header_font = Font(bold=True, size=11, color="FFFFFF")
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.warning_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.info_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
        self.center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

    def export_audit_results(
        self,
        results: List[AuditResult],
        output_path: str,
        statistics: Optional[Dict[str, Any]] = None,
        include_details: bool = True,
    ):
        wb = Workbook()

        self._write_summary_sheet(wb, results, statistics)
        self._write_detail_sheet(wb, results)

        if include_details:
            self._write_explanation_sheet(wb)

        wb.save(output_path)

    def _write_summary_sheet(
        self,
        wb: Workbook,
        results: List[AuditResult],
        statistics: Optional[Dict[str, Any]],
    ):
        ws = wb.active
        ws.title = "审计总览"

        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = "长上下文截断审计报告"
        title_cell.font = Font(bold=True, size=16, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        title_cell.alignment = self.center_align
        ws.row_dimensions[1].height = 30

        ws["A3"] = "生成时间"
        ws["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws["A4"] = "记录总数"
        ws["B4"] = len(results)

        truncated = [r for r in results if r.is_truncated]
        ws["A5"] = "截断记录数"
        ws["B5"] = len(truncated)
        ws["A6"] = "截断率"
        ws["B6"] = f"{(len(truncated) / len(results) * 100):.1f}%" if results else "0%"

        for row in range(3, 7):
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"A{row}"].fill = self.info_fill

        start_row = 8
        ws[f"A{start_row}"] = "按截断原因统计"
        ws[f"A{start_row}"].font = Font(bold=True, size=12)
        ws.merge_cells(f"A{start_row}:B{start_row}")
        start_row += 1

        reason_counts = {}
        for r in truncated:
            reason_label = r.reason.value
            reason_counts[reason_label] = reason_counts.get(reason_label, 0) + 1

        headers = ["截断原因", "记录数量"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        start_row += 1
        for reason, count in sorted(reason_counts.items(), key=lambda x: -x[1]):
            ws.cell(row=start_row, column=1, value=reason).border = self.thin_border
            ws.cell(row=start_row, column=2, value=count).border = self.thin_border
            ws.cell(row=start_row, column=2).alignment = self.center_align
            start_row += 1

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 20
        ws.column_dimensions["F"].width = 30

    def _write_detail_sheet(self, wb: Workbook, results: List[AuditResult]):
        ws = wb.create_sheet("详细记录")

        headers = [
            "记录编号",
            "是否截断",
            "截断原因",
            "原因说明",
            "原始长度(字符)",
            "Token数量",
            "严重程度",
            "数据来源",
            "人工备注",
            "原始内容预览",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        for row_idx, result in enumerate(results, 2):
            data = result.to_dict()
            for col_idx, key in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=data.get(key, ""))
                cell.border = self.thin_border
                cell.alignment = self.left_align if col_idx in [4, 8, 9, 10] else self.center_align

                if result.is_truncated and key == "是否截断":
                    cell.fill = self.warning_fill
                    cell.font = Font(bold=True, color="9C0006")

                if result.is_truncated and key == "严重程度":
                    if result.severity == "high":
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        cell.font = Font(color="9C0006")
                    elif result.severity == "medium":
                        cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                        cell.font = Font(color="9C5700")

        column_widths = [15, 10, 18, 40, 15, 12, 12, 12, 25, 50]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"

    def _write_explanation_sheet(self, wb: Workbook):
        ws = wb.create_sheet("名词说明")

        ws["A1"] = "名词说明"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:B1")

        explanations = [
            ("截断原因", "说明"),
            ("token超限", "输入的文字转换成token后，数量超过了模型的限制，导致后面的内容被切掉"),
            ("上下文窗口不足", "输入内容太长，快接近模型能处理的最大长度了，可能会被截断"),
            ("工具调用参数错误", "在调用工具时参数传错了格式，导致内容处理失败"),
            ("人工截断标记", "人工在备注里标注了这是截断的或者不完整的内容"),
            ("输入本身不完整", "原始数据就是空的或者缺内容"),
            ("关键字段缺失", "数据里缺少了应该有的字段信息"),
            ("单位不明确", "数值后面没有写清楚单位，可能造成理解偏差"),
            ("未知原因", "检测到有截断迹象，但具体原因不确定"),
        ]

        for row_idx, (term, explanation) in enumerate(explanations, 3):
            ws.cell(row=row_idx, column=1, value=term).border = self.thin_border
            ws.cell(row=row_idx, column=2, value=explanation).border = self.thin_border
            ws.cell(row=row_idx, column=1).alignment = self.center_align
            ws.cell(row=row_idx, column=2).alignment = self.left_align

            if row_idx == 3:
                for col in [1, 2]:
                    ws.cell(row=row_idx, column=col).font = self.header_font
                    ws.cell(row=row_idx, column=col).fill = self.header_fill

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 60

    def export_version_comparison(
        self,
        tracker: VersionTracker,
        output_path: str,
    ):
        wb = Workbook()
        ws = wb.active
        ws.title = "版本变更记录"

        ws.merge_cells("A1:G1")
        title = ws["A1"]
        title.value = "人工修正前后对比"
        title.font = Font(bold=True, size=14, color="FFFFFF")
        title.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        title.alignment = self.center_align
        ws.row_dimensions[1].height = 25

        changes = tracker.get_changed_records()

        headers = [
            "记录编号",
            "变更字段",
            "修正前",
            "修正后",
            "影响程度",
            "审计人",
            "变更时间",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        row_idx = 4
        for change in changes:
            record_id = change["record_id"]
            latest = tracker.get_latest(record_id)
            for field_name, field_change in change["changes"].items():
                ws.cell(row=row_idx, column=1, value=record_id).border = self.thin_border
                ws.cell(row=row_idx, column=2, value=field_name).border = self.thin_border
                ws.cell(row=row_idx, column=3, value=str(field_change.get("旧值", ""))).border = self.thin_border
                ws.cell(row=row_idx, column=4, value=str(field_change.get("新值", ""))).border = self.thin_border

                impact = field_change.get("影响", "")
                impact_cell = ws.cell(row=row_idx, column=5, value=impact)
                impact_cell.border = self.thin_border
                if "核心" in impact or "高" in impact:
                    impact_cell.fill = self.warning_fill
                    impact_cell.font = Font(color="9C0006", bold=True)

                ws.cell(row=row_idx, column=6, value=latest.auditor if latest else "").border = self.thin_border
                ws.cell(row=row_idx, column=7, value=latest.timestamp if latest else "").border = self.thin_border

                for col in range(1, 8):
                    ws.cell(row=row_idx, column=col).alignment = self.left_align if col in [3, 4] else self.center_align

                row_idx += 1

        column_widths = [15, 12, 25, 25, 12, 12, 20]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width

        ws.freeze_panes = "A4"
        wb.save(output_path)
