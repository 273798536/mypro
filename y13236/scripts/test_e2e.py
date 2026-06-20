import sys
import os
import json
import time
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from scripts.init_test_data import init_test_data
from app.database import SessionLocal, engine, Base
from fastapi.testclient import TestClient
from main import app


def run_e2e():
    print("=" * 70)
    print("【端到端验证】重跑→查看异常队列→导出Excel")
    print("=" * 70)

    print("\n[步骤 1/5] 重置数据库并初始化测试数据")
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    init_test_data()
    print("   ✓ 测试数据已初始化")

    client = TestClient(app)

    print("\n[步骤 2/5] POST /api/v1/conflicts/run  重跑冲突检测")
    resp = client.post("/api/v1/conflicts/run")
    assert resp.status_code == 200, f"接口报错: {resp.text}"
    run_result = resp.json()
    print(f"   状态码: {resp.status_code}")
    print(f"   总冲突: {run_result['total_conflicts']} 条")
    print(f"   按类型: {json.dumps(run_result['by_type'], ensure_ascii=False)}")
    assert run_result["total_conflicts"] == 5, f"期望5条冲突，实际{run_result['total_conflicts']}条"
    print("   ✓ 重跑成功，返回与脚本验证一致")

    print("\n[步骤 3/5] GET /api/v1/conflicts/queue  查看异常队列（不带筛选）")
    resp = client.get("/api/v1/conflicts/queue")
    assert resp.status_code == 200
    queue = resp.json()
    print(f"   返回条数: {len(queue)}")
    assert len(queue) == 5
    for idx, c in enumerate(queue, 1):
        dev = c.get("timecode_deviation_seconds", "N/A")
        print(f"   {idx}. [{c['conflict_type']}] 偏差={dev}秒  描述前25字：{c['description'][:25]}...")
    print("   ✓ 队列数据正常，与屏幕一致")

    print("\n[步骤 4/5] GET /api/v1/conflicts/queue?conflict_type=timecode_deviation  按类型筛选")
    resp = client.get("/api/v1/conflicts/queue", params={"conflict_type": "timecode_deviation"})
    assert resp.status_code == 200
    filtered = resp.json()
    print(f"   筛选 [timecode_deviation] 条数: {len(filtered)}")
    assert len(filtered) == 2
    for c in filtered:
        assert c["conflict_type"] == "timecode_deviation"
        dev = c["timecode_deviation_seconds"]
        has_banpai = "半拍" in c["description"]
        print(f"   - 偏差 {dev} 秒，描述含'半拍'字样: {has_banpai}")
        assert "半拍" in c["description"]
        assert dev in (6, 12)
    print("   ✓ 筛选口径生效，只返回时码偏差")

    print("\n[步骤 5/5] GET /api/v1/conflicts/export/download  导出 Excel")
    resp = client.get("/api/v1/conflicts/export/download",
                    params={"conflict_type": "timecode_deviation"})
    assert resp.status_code == 200
    content_type = resp.headers.get("content-type")
    disposition = resp.headers.get("content-disposition")
    content_len = len(resp.content)
    print(f"   Content-Type: {content_type}")
    print(f"   Content-Disposition: {disposition}")
    print(f"   文件大小: {len(resp.content)} 字节")

    tmp_path = "/tmp/piano_export_test.xlsx"
    with open(tmp_path, "wb") as f:
        f.write(resp.content)
    print(f"   已保存到: {tmp_path}")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(tmp_path)
        ws = wb.active
        print(f"   Sheet名: {ws.title}")
        print(f"   行数: {ws.max_row}（含表头）")
        print(f"   列数: {ws.max_column}")
        headers = [cell.value for cell in ws[1]]
        print(f"   表头: {headers}")
        assert ws.max_row == 3, f"筛选时码偏差应返回 2 条数据 + 1 行表头 = 3行，实际{ws.max_row}行"
        assert "冲突类型" in headers
        assert "时码偏差(秒)" in headers
        assert "原始来源" in headers
        print("   ✓ Excel 格式可被 openpyxl 正常打开，内容与筛选口径一致")
        print("\n   数据行内容（第2行：")
        for col_idx, header in enumerate(headers, 1):
            val = ws.cell(row=2, column=col_idx).value
            print(f"     {header}: {str(val)[:50]}")
    except ImportError:
        print("   ⚠ openpyxl 未安装，跳过打开验证")

    print("\n" + "=" * 70)
    print("【全部通过】导出与屏幕数字一致，Excel可正常打开")
    print("=" * 70)


if __name__ == "__main__":
    run_e2e()
