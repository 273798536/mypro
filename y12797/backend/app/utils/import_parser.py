"""旧Excel表导入解析器：模拟实际工作中配方工程师拿到的历史遗留表。"""
import io
import re
from typing import Dict, Any, List, Tuple, Optional
from openpyxl import load_workbook
import pandas as pd


OLD_EXCEL_FIELD_ALIASES = {
    "反应条件": [
        ["条件", "项目", "反应条件", "condition"],
        ["数值", "value", "值", "设定值"],
        ["单位", "unit"],
        ["备注", "note", "说明"],
    ],
    "底物换算": [
        ["底物名称", "名称", "substrate", "化合物"],
        ["CAS号", "cas", "casno"],
        ["称样量", "质量", "mass", "称量"],
        ["称样量单位", "质量单位", "mass_unit"],
        ["体积", "volume", "定容体积"],
        ["体积单位", "volume_unit"],
        ["分子量", "MW", "molecular_weight", "摩尔质量"],
        ["纯度", "purity", "含量"],
        ["目标浓度单位", "target_unit"],
    ],
    "谱图数据": [
        ["峰名", "peak_name", "峰名称"],
        ["保留时间", "RT", "retention_time"],
        ["峰面积", "area", "peak_area"],
        ["峰高", "height", "peak_height"],
        ["谱图类型", "spectrum_type", "仪器"],
        ["波长", "detection_wavelength"],
        ["判读", "interpretation", "备注"],
    ],
}


def _match_col(header: str, aliases: List[List[str]]) -> bool:
    h = (header or "").strip().lower()
    for group in aliases:
        for a in group:
            if a.lower() in h:
                return True
    return False


def _find_section_columns(df_columns, alias_groups: List[List[str]]) -> List[int]:
    result = []
    for aliases in alias_groups:
        found = None
        for i, col in enumerate(df_columns):
            if _match_col(str(col), [aliases]):
                found = i
                break
        result.append(found)
    return result


def parse_old_excel(file_bytes: bytes, filename: str = "") -> Tuple[Dict[str, Any], List[str], List[str]]:
    """
    解析旧格式Excel文件。
    返回 (parsed_data, warnings, errors)
    parsed_data 结构同 build_sample_record 的返回结构。
    """
    warnings: List[str] = []
    errors: List[str] = []

    if not file_bytes:
        errors.append("文件为空")
        return {}, warnings, errors

    data: Dict[str, Any] = {
        "batch_no": "",
        "material_name": "",
        "source_file_name": filename,
        "source_format": "old_excel",
        "remark": "",
        "supplementary_note": "",
        "safety_note": "",
        "reaction_conditions": [],
        "substrate_conversions": [],
        "spectrum_data": [],
    }

    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        errors.append(f"Excel 文件无法解析：{str(e)}")
        return data, warnings, errors

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # 先尝试从 sheet 头部提取批次/物料信息
        header_scan = rows[:8]
        for row in header_scan:
            for cell in row:
                if not cell:
                    continue
                s = str(cell).strip()
                if not data["batch_no"] and ("批" in s or "Batch" in s or "Lot" in s):
                    m = re.search(r"[：:]\s*(\S+)", s)
                    if m:
                        data["batch_no"] = m.group(1)
                if not data["material_name"] and ("物料" in s or "名称" in s or "Material" in s):
                    m = re.search(r"[：:]\s*(.+)$", s)
                    if m:
                        data["material_name"] = m.group(1).strip()
                if not data["safety_note"] and ("安全" in s or "Safety" in s):
                    m = re.search(r"[：:]\s*(.+)$", s)
                    if m:
                        data["safety_note"] = m.group(1).strip()

        # 查找数据表格起始行
        data_start = None
        for i, row in enumerate(rows):
            non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if len(non_empty) >= 2 and any(
                kw in "".join(non_empty) for kw in ("条件", "底物", "峰", "数值", "CAS")
            ):
                data_start = i
                break
        if data_start is None:
            continue

        header_row = rows[data_start]
        body_rows = rows[data_start + 1:]

        # 反应条件表（按关键词组合判断）
        if any(_match_col(str(c), [OLD_EXCEL_FIELD_ALIASES["反应条件"][0]]) for c in header_row if c):
            cond_cols = _find_section_columns(header_row, OLD_EXCEL_FIELD_ALIASES["反应条件"])
            for idx, r in enumerate(body_rows):
                if not any(r):
                    continue
                name = str(r[cond_cols[0]]).strip() if cond_cols[0] is not None and r[cond_cols[0]] is not None else ""
                value = str(r[cond_cols[1]]).strip() if cond_cols[1] is not None and r[cond_cols[1]] is not None else ""
                unit = str(r[cond_cols[2]]).strip() if cond_cols[2] is not None and r[cond_cols[2]] is not None else ""
                note = str(r[cond_cols[3]]).strip() if cond_cols[3] is not None and r[cond_cols[3]] is not None else ""
                if not name and not value:
                    continue
                obj = {
                    "condition_name": name or f"条件{idx+1}",
                    "condition_value": value,
                    "unit": unit if unit and unit != "-" and unit != "无" else "",
                    "row_order": idx + 1,
                }
                if note:
                    obj["issue_description"] = note
                data["reaction_conditions"].append(obj)
                if not obj["unit"]:
                    warnings.append(f"反应条件「{obj['condition_name']}」未填写单位")

        # 底物换算表
        if any(_match_col(str(c), [OLD_EXCEL_FIELD_ALIASES["底物换算"][0]]) for c in header_row if c):
            sub_cols = _find_section_columns(header_row, OLD_EXCEL_FIELD_ALIASES["底物换算"])
            missing_unit_fields = []
            for idx, r in enumerate(body_rows):
                if not any(r):
                    continue

                def _cell(pos, default=None, conv=None):
                    if pos is None or r[pos] is None:
                        return default
                    v = r[pos]
                    if conv:
                        try:
                            return conv(v)
                        except Exception:
                            return default
                    return str(v).strip() if isinstance(v, str) else v

                name = _cell(sub_cols[0], f"底物{idx+1}")
                mass = _cell(sub_cols[2], None, float)
                mass_unit = _cell(sub_cols[3], "")
                volume = _cell(sub_cols[4], None, float)
                vol_unit = _cell(sub_cols[5], "")
                mw = _cell(sub_cols[6], None, float)
                purity = _cell(sub_cols[7], 100.0, float)
                target_unit = _cell(sub_cols[8], "mmol/L")

                obj = {
                    "substrate_name": name,
                    "cas_no": _cell(sub_cols[1], ""),
                    "initial_mass": mass,
                    "initial_mass_unit": mass_unit if mass_unit and mass_unit != "-" else "",
                    "volume": volume,
                    "volume_unit": vol_unit if vol_unit and vol_unit != "-" else "",
                    "molecular_weight": mw,
                    "purity": purity,
                    "final_concentration_unit": target_unit if target_unit else "mmol/L",
                    "row_order": idx + 1,
                }

                if mass is not None and not obj["initial_mass_unit"]:
                    missing_unit_fields.append(f"底物「{name}」称样量单位")
                if volume is not None and not obj["volume_unit"]:
                    missing_unit_fields.append(f"底物「{name}」体积单位")
                if mw is None:
                    warnings.append(f"底物「{name}」缺失分子量，无法完成自动换算")

                data["substrate_conversions"].append(obj)

            for f in missing_unit_fields:
                warnings.append(f"{f}未填写")

        # 谱图数据表
        if any(_match_col(str(c), [OLD_EXCEL_FIELD_ALIASES["谱图数据"][0]]) for c in header_row if c):
            spec_cols = _find_section_columns(header_row, OLD_EXCEL_FIELD_ALIASES["谱图数据"])
            for idx, r in enumerate(body_rows):
                if not any(r):
                    continue

                def _c(pos, default=None, conv=None):
                    if pos is None or r[pos] is None:
                        return default
                    v = r[pos]
                    if conv:
                        try:
                            return conv(v)
                        except Exception:
                            return default
                    return str(v).strip() if isinstance(v, str) else v

                obj = {
                    "peak_name": _c(spec_cols[0], f"峰{idx+1}"),
                    "retention_time": _c(spec_cols[1], None, float),
                    "peak_area": _c(spec_cols[2], None, float),
                    "peak_height": _c(spec_cols[3], None, float),
                    "spectrum_type": _c(spec_cols[4], "HPLC"),
                    "detection_wavelength": _c(spec_cols[5], ""),
                    "interpretation": _c(spec_cols[6], ""),
                    "row_order": idx + 1,
                    "interpretation_linked": True,
                }
                data["spectrum_data"].append(obj)

    if not data["reaction_conditions"] and not data["substrate_conversions"] and not data["spectrum_data"]:
        warnings.append("未从文件中识别到任何结构化表格，请检查表头是否包含「条件/底物/峰」等关键词")

    return data, warnings, errors
