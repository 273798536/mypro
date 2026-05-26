import json
from datetime import datetime
from decimal import Decimal
from typing import Optional
from pathlib import Path

from sqlalchemy.orm import Session

from app.models.models import (
    Reconciliation, Enterprise, CorrectionLog, ReconciliationItem
)
from app.config import DATA_DIR


class DecimalEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, Decimal):
            return float(obj)
        if isinstance(obj, datetime):
            return obj.isoformat()
        return super().default(obj)


def generate_audit_report(db: Session, reconciliation_id: int) -> dict:
    rec = db.query(Reconciliation).filter(Reconciliation.id == reconciliation_id).first()
    if not rec:
        return {"error": "reconciliation not found"}

    items = []
    for item in rec.items:
        ent = db.query(Enterprise).filter(Enterprise.id == item.enterprise_id).first()
        items.append({
            "enterprise_code": ent.enterprise_code if ent else None,
            "enterprise_name": ent.name if ent else None,
            "opening_balance": item.opening_balance,
            "period_in": item.period_in,
            "period_out": item.period_out,
            "red_flush_adjustment": item.red_flush_adjustment,
            "calculated_closing": item.calculated_closing,
            "reported_closing": item.reported_closing,
            "balance_diff": item.balance_diff,
            "cross_month_readings": item.cross_month_readings,
            "duplicate_credits": item.duplicate_credits,
            "red_flush_unapplied": item.red_flush_unapplied,
            "receipts_unverified": item.receipts_unverified,
            "has_anomaly": item.has_anomaly,
        })

    anomalies = []
    for a in rec.anomalies:
        anomalies.append({
            "id": a.id,
            "type": a.anomaly_type,
            "severity": a.severity,
            "description": a.description,
            "status": a.status,
        })

    critical_count = sum(1 for a in rec.anomalies if a.severity == "critical" and a.status == "open")
    warning_count = sum(1 for a in rec.anomalies if a.severity == "warning" and a.status == "open")
    resolved_count = sum(1 for a in rec.anomalies if a.status == "resolved")

    report = {
        "reconciliation_id": rec.id,
        "period": rec.period,
        "generated_at": datetime.now().isoformat(),
        "summary": {
            "total_enterprises": rec.total_enterprises,
            "anomalies_found": rec.anomalies_found,
            "critical_open": critical_count,
            "warning_open": warning_count,
            "resolved": resolved_count,
            "status": rec.status,
        },
        "items": items,
        "anomalies": anomalies,
    }
    return report


def export_report_json(db: Session, reconciliation_id: int) -> str:
    report = generate_audit_report(db, reconciliation_id)
    if "error" in report:
        return json.dumps(report, ensure_ascii=False, indent=2)

    filepath = DATA_DIR / f"audit_report_{reconciliation_id}_{report['period']}.json"
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2, cls=DecimalEncoder)

    return json.dumps(report, ensure_ascii=False, indent=2, cls=DecimalEncoder)


def export_report_excel(db: Session, reconciliation_id: int) -> Optional[str]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return None

    rec = db.query(Reconciliation).filter(Reconciliation.id == reconciliation_id).first()
    if not rec:
        return None

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "汇总"
    headers = ["项目", "值"]
    ws_summary.append(headers)
    ws_summary.append(["对账周期", rec.period])
    ws_summary.append(["企业总数", rec.total_enterprises])
    ws_summary.append(["异常总数", rec.anomalies_found])
    ws_summary.append(["状态", rec.status])
    ws_summary.append(["开始时间", str(rec.started_at)])
    ws_summary.append(["完成时间", str(rec.completed_at)])

    ws_items = wb.create_sheet("对账明细")
    item_headers = [
        "企业编号", "企业名称", "期初余额", "本期收入", "本期支出",
        "红冲调整", "计算期末", "申报期末", "差额",
        "跨月读数", "重复流水", "红冲未处理", "回执未核验", "是否异常"
    ]
    ws_items.append(item_headers)

    for item in rec.items:
        ent = db.query(Enterprise).filter(Enterprise.id == item.enterprise_id).first()
        ws_items.append([
            ent.enterprise_code if ent else "",
            ent.name if ent else "",
            float(item.opening_balance),
            float(item.period_in),
            float(item.period_out),
            float(item.red_flush_adjustment),
            float(item.calculated_closing),
            float(item.reported_closing),
            float(item.balance_diff),
            item.cross_month_readings,
            item.duplicate_credits,
            item.red_flush_unapplied,
            item.receipts_unverified,
            "是" if item.has_anomaly else "否",
        ])

    ws_anomalies = wb.create_sheet("异常清单")
    anomaly_headers = ["ID", "类型", "严重度", "描述", "状态", "解决备注"]
    ws_anomalies.append(anomaly_headers)
    for a in rec.anomalies:
        ws_anomalies.append([
            a.id, a.anomaly_type, a.severity, a.description,
            a.status, a.resolution_note or ""
        ])

    filepath = DATA_DIR / f"audit_report_{reconciliation_id}_{rec.period}.xlsx"
    wb.save(str(filepath))
    return str(filepath)


def get_enterprise_history(db: Session, enterprise_id: int) -> dict:
    ent = db.query(Enterprise).filter(Enterprise.id == enterprise_id).first()
    if not ent:
        return {"error": "enterprise not found"}

    corrections = (
        db.query(CorrectionLog)
        .filter(CorrectionLog.enterprise_id == enterprise_id)
        .order_by(CorrectionLog.created_at.desc())
        .all()
    )

    rec_items = (
        db.query(ReconciliationItem)
        .filter(ReconciliationItem.enterprise_id == enterprise_id)
        .order_by(ReconciliationItem.created_at.desc())
        .all()
    )

    return {
        "enterprise": {
            "id": ent.id,
            "enterprise_code": ent.enterprise_code,
            "name": ent.name,
            "current_balance": ent.current_balance,
            "status": ent.status,
        },
        "corrections": [
            {
                "id": c.id,
                "target_table": c.target_table,
                "target_id": c.target_id,
                "field_name": c.field_name,
                "old_value": c.old_value,
                "new_value": c.new_value,
                "reason": c.reason,
                "operator": c.operator,
                "created_at": str(c.created_at),
            }
            for c in corrections
        ],
        "reconciliations": [
            {
                "id": ri.id,
                "reconciliation_id": ri.reconciliation_id,
                "calculated_closing": ri.calculated_closing,
                "reported_closing": ri.reported_closing,
                "balance_diff": ri.balance_diff,
                "has_anomaly": ri.has_anomaly,
                "created_at": str(ri.created_at),
            }
            for ri in rec_items
        ],
    }