import os
import zipfile
import tempfile
import re
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from sqlalchemy.orm import Session

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

from app.config import settings
from app.models.base import Batch, Contract, Attachment, BatchStatus, DuplicateStrategy, PaymentNode, AcceptanceEmail, PriceChange
from app.schemas.contract import ContractCreate, PaymentNodeCreate, AcceptanceEmailCreate
from app.services.state_machine import ContractStateMachine
from app.services.task_service import TaskService
from app.models.base import TaskStatus


class FileService:
    def __init__(self, db: Session):
        self.db = db
        self.state_machine = ContractStateMachine(db)
        self.task_service = TaskService(db)

    def _save_uploaded_file(self, file_content: bytes, filename: str, batch_id: int, 
                           attachment_type: str, uploaded_by: str) -> Attachment:
        batch_dir = settings.UPLOAD_DIR / f"batch_{batch_id}"
        batch_dir.mkdir(parents=True, exist_ok=True)
        
        safe_filename = re.sub(r'[^\w\s.-]', '_', filename)
        file_path = batch_dir / safe_filename
        file_path.write_bytes(file_content)
        
        attachment = Attachment(
            batch_id=batch_id,
            file_name=safe_filename,
            file_path=str(file_path),
            file_type=filename.split('.')[-1].lower() if '.' in filename else 'unknown',
            file_size=len(file_content),
            uploaded_by=uploaded_by,
            attachment_type=attachment_type
        )
        self.db.add(attachment)
        self.db.commit()
        self.db.refresh(attachment)
        return attachment

    def _parse_pdf_contract(self, file_path: Path) -> Dict[str, Any]:
        if PdfReader is None:
            return {"error": "pypdf not installed"}
        
        try:
            reader = PdfReader(str(file_path))
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            
            result = {
                "total_pages": len(reader.pages),
                "raw_text": text[:2000]
            }
            
            contract_no = None
            keywords = ["合同编号", "合同号", "Contract No", "合同编号：", "合同号："]
            for keyword in keywords:
                if keyword in text:
                    idx = text.find(keyword)
                    line = text[idx:idx+100].split('\n')[0]
                    result["contract_no_hint"] = line
                    match = re.search(r'[A-Za-z0-9-]{5,}', line)
                    if match:
                        contract_no = match.group(0)
                        result["contract_no"] = contract_no
                        break
            
            amount_match = re.search(r'[￥¥$]?\s*([\d,]+\.?\d*)\s*(?:元|人民币|USD)?', text)
            if amount_match:
                try:
                    amount_str = amount_match.group(1).replace(',', '')
                    result["total_amount"] = float(amount_str)
                except:
                    pass
            
            return result
        except Exception as e:
            return {"error": str(e)}

    def _parse_excel_payment_nodes(self, file_path: Path) -> List[Dict[str, Any]]:
        if openpyxl is None:
            return []
        
        try:
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            ws = wb.active
            nodes = []
            
            headers = []
            for cell in ws[1]:
                headers.append(str(cell.value).strip() if cell.value else "")
            
            header_mapping = {}
            for i, h in enumerate(headers):
                if h:
                    header_mapping[h.lower()] = i
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_values = list(row)
                if not any(v is not None and v != "" for v in row_values):
                    continue
                
                node_data = {}
                
                for h, idx in header_mapping.items():
                    if idx < len(row_values):
                        val = row_values[idx]
                        if 'node' in h or '名称' in h:
                            node_data['node_name'] = str(val) if val else ""
                        elif 'type' in h or '类型' in h:
                            node_data['node_type'] = str(val) if val else "milestone"
                        elif 'ratio' in h or '比例' in h:
                            try:
                                node_data['payment_ratio'] = float(val) if val else 0
                            except:
                                node_data['payment_ratio'] = 0
                        elif 'amount' in h or '金额' in h:
                            try:
                                node_data['payment_amount'] = float(val) if val else 0
                            except:
                                node_data['payment_amount'] = 0
                        elif 'due' in h or '到期' in h or '日期' in h:
                            node_data['due_date'] = val
                        elif 'status' in h or '状态' in h:
                            node_data['status'] = str(val) if val else "pending"
                
                if node_data.get('node_name'):
                    nodes.append(node_data)
            
            return nodes
        except Exception as e:
            return [{"error": str(e)}]

    def _parse_price_change_excel(self, file_path: Path) -> List[Dict[str, Any]]:
        if openpyxl is None:
            return []
        
        try:
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            ws = wb.active
            price_changes = []
            
            headers = []
            for cell in ws[1]:
                headers.append(str(cell.value).strip() if cell.value else "")
            
            header_mapping = {}
            for i, h in enumerate(headers):
                if h:
                    header_mapping[h.lower()] = i
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_values = list(row)
                if not any(v is not None and v != "" for v in row_values):
                    continue
                
                pc_data = {}
                
                for h, idx in header_mapping.items():
                    if idx < len(row_values):
                        val = row_values[idx]
                        if 'original' in h or '原价' in h:
                            try:
                                pc_data['original_price'] = float(val) if val else 0
                            except:
                                pc_data['original_price'] = 0
                        elif 'new' in h or '新价' in h:
                            try:
                                pc_data['new_price'] = float(val) if val else 0
                            except:
                                pc_data['new_price'] = 0
                        elif 'reason' in h or '原因' in h:
                            pc_data['change_reason'] = str(val) if val else ""
                        elif 'approved' in h or '审批' in h:
                            pc_data['approved_by'] = str(val) if val else ""
                        elif 'effective' in h or '生效' in h:
                            pc_data['effective_date'] = val
                
                if pc_data.get('original_price') is not None or pc_data.get('new_price') is not None:
                    price_changes.append(pc_data)
            
            return price_changes
        except Exception as e:
            return []

    def _parse_eml_file(self, file_path: Path) -> Dict[str, Any]:
        try:
            import email
            from email import policy
            
            with open(file_path, 'rb') as f:
                msg = email.message_from_bytes(f.read(), policy=policy.default)
            
            subject = msg['subject'] or ''
            from_addr = msg['from'] or ''
            to_addr = msg['to'] or ''
            date_str = msg['date']
            
            email_date = datetime.utcnow()
            if date_str:
                try:
                    from email.utils import parsedate_to_datetime
                    email_date = parsedate_to_datetime(date_str)
                except:
                    pass
            
            content = ""
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == 'text/plain':
                        payload = part.get_payload(decode=True)
                        if payload:
                            content += payload.decode('utf-8', errors='ignore')
            else:
                payload = msg.get_payload(decode=True)
                if payload:
                    content = payload.decode('utf-8', errors='ignore')
            
            acceptance_result = "pending"
            acceptance_amount = 0.0
            
            if '验收通过' in content or '验收合格' in content or 'approved' in subject.lower():
                acceptance_result = "approved"
            elif '验收不通过' in content or 'rejected' in subject.lower():
                acceptance_result = "rejected"
            
            amount_match = re.search(r'[￥¥$]?\s*([\d,]+\.?\d*)\s*(?:元|人民币)?', content)
            if amount_match:
                try:
                    amount_str = amount_match.group(1).replace(',', '')
                    acceptance_amount = float(amount_str)
                except:
                    pass
            
            return {
                "subject": subject,
                "from": from_addr,
                "to": to_addr,
                "date": date_str,
                "email_date": email_date,
                "content": content[:2000],
                "acceptance_result": acceptance_result,
                "acceptance_amount": acceptance_amount
            }
        except Exception as e:
            return {"error": str(e)}

    def _extract_contract_key(self, filename: str, parsed: Dict[str, Any]) -> str:
        contract_no = parsed.get('contract_no') or parsed.get('contract_no_hint')
        if contract_no:
            return str(contract_no)
        base_name = Path(filename).stem
        return base_name

    def _apply_duplicate_strategy(
        self, 
        batch_id: int, 
        contract_key: str, 
        duplicate_strategy: DuplicateStrategy,
        uploaded_by: str
    ) -> Tuple[Optional[Contract], str]:
        existing = self.db.query(Contract).filter(
            Contract.batch_id == batch_id,
            (Contract.contract_no == contract_key) | 
            (Contract.contract_name.like(f"%{contract_key}%"))
        ).first()
        
        if existing:
            if duplicate_strategy == DuplicateStrategy.IGNORE:
                return existing, "ignored"
            elif duplicate_strategy == DuplicateStrategy.OVERWRITE:
                self.db.delete(existing)
                self.db.flush()
                return None, "overwritten"
            else:
                return None, "appended"
        return None, "new"

    def process_contract_pdf(self, batch_id: int, file_content: bytes, filename: str, 
                            uploaded_by: str, duplicate_strategy: DuplicateStrategy = DuplicateStrategy.APPEND,
                            task_id: Optional[str] = None) -> Dict[str, Any]:
        try:
            attachment = self._save_uploaded_file(file_content, filename, batch_id, "contract_pdf", uploaded_by)
            parsed = self._parse_pdf_contract(Path(attachment.file_path))
            
            contract_key = self._extract_contract_key(filename, parsed)
            existing_contract, action = self._apply_duplicate_strategy(
                batch_id, contract_key, duplicate_strategy, uploaded_by
            )
            
            if action == "ignored":
                return {
                    "contract_id": existing_contract.id,
                    "attachment_id": attachment.id,
                    "action": "ignored",
                    "parsed": parsed,
                    "message": "Contract already exists, skipped"
                }
            
            contract_no = parsed.get('contract_no') or f"PDF_{batch_id}_{attachment.id}"
            contract_data = ContractCreate(
                contract_no=contract_no,
                contract_name=filename.replace('.pdf', ''),
                total_amount=parsed.get('total_amount', 0),
                metadata_={"parsed_pdf": parsed, "source": "pdf", "action": action}
            )
            
            contract = self.state_machine.create_contract(batch_id, contract_data, uploaded_by)
            attachment.contract_id = contract.id
            self.db.commit()
            
            if task_id:
                self.task_service.update_task(task_id, progress=100)
            
            return {
                "contract_id": contract.id,
                "attachment_id": attachment.id,
                "action": action,
                "parsed": parsed
            }
        except Exception as e:
            if task_id:
                self.task_service.fail_task(task_id, str(e), "pdf_parse_error", retry=True)
            raise

    def process_payment_excel(self, batch_id: int, contract_id: int, file_content: bytes, 
                             filename: str, uploaded_by: str,
                             task_id: Optional[str] = None) -> Dict[str, Any]:
        try:
            attachment = self._save_uploaded_file(file_content, filename, batch_id, "payment_excel", uploaded_by)
            attachment.contract_id = contract_id
            
            nodes = self._parse_excel_payment_nodes(Path(attachment.file_path))
            
            contract = self.db.query(Contract).filter(Contract.id == contract_id).first()
            if contract:
                for node_data in nodes:
                    due_date = node_data.get('due_date')
                    if isinstance(due_date, str):
                        try:
                            due_date = datetime.fromisoformat(due_date)
                        except:
                            due_date = None
                    elif not isinstance(due_date, datetime):
                        due_date = None
                    
                    node = PaymentNode(
                        contract_id=contract_id,
                        node_name=node_data.get('node_name', ''),
                        node_type=node_data.get('node_type', 'milestone'),
                        payment_ratio=node_data.get('payment_ratio', 0),
                        payment_amount=node_data.get('payment_amount', 0),
                        due_date=due_date,
                        status=node_data.get('status', 'pending')
                    )
                    self.db.add(node)
            
            self.db.commit()
            
            if task_id:
                self.task_service.update_task(task_id, progress=100)
            
            return {"nodes_count": len(nodes), "attachment_id": attachment.id}
        except Exception as e:
            if task_id:
                self.task_service.fail_task(task_id, str(e), "excel_parse_error", retry=True)
            raise

    def process_acceptance_email(self, batch_id: int, contract_id: int, file_content: bytes,
                                filename: str, uploaded_by: str,
                                task_id: Optional[str] = None) -> Dict[str, Any]:
        try:
            attachment = self._save_uploaded_file(file_content, filename, batch_id, "acceptance_email", uploaded_by)
            attachment.contract_id = contract_id
            
            parsed = self._parse_eml_file(Path(attachment.file_path))
            
            email = AcceptanceEmail(
                contract_id=contract_id,
                email_subject=parsed.get('subject', ''),
                email_from=parsed.get('from', ''),
                email_to=parsed.get('to', ''),
                email_date=parsed.get('email_date', datetime.utcnow()),
                acceptance_result=parsed.get('acceptance_result', 'pending'),
                acceptance_amount=parsed.get('acceptance_amount', 0),
                content=parsed.get('content', '')
            )
            self.db.add(email)
            self.db.commit()
            
            if task_id:
                self.task_service.update_task(task_id, progress=100)
            
            return {"email_id": email.id, "attachment_id": attachment.id, "parsed": parsed}
        except Exception as e:
            if task_id:
                self.task_service.fail_task(task_id, str(e), "eml_parse_error", retry=False)
            raise

    def process_price_change_excel(self, batch_id: int, contract_id: int, file_content: bytes,
                                  filename: str, uploaded_by: str,
                                  task_id: Optional[str] = None) -> Dict[str, Any]:
        try:
            attachment = self._save_uploaded_file(file_content, filename, batch_id, "price_change", uploaded_by)
            attachment.contract_id = contract_id
            
            price_changes = self._parse_price_change_excel(Path(attachment.file_path))
            
            count = 0
            for pc_data in price_changes:
                effective_date = pc_data.get('effective_date')
                if not isinstance(effective_date, datetime):
                    effective_date = datetime.utcnow()
                
                pc = PriceChange(
                    contract_id=contract_id,
                    original_price=pc_data.get('original_price', 0),
                    new_price=pc_data.get('new_price', 0),
                    change_reason=pc_data.get('change_reason', ''),
                    approved_by=pc_data.get('approved_by', uploaded_by),
                    approved_date=datetime.utcnow(),
                    effective_date=effective_date,
                    is_manual=True
                )
                self.db.add(pc)
                count += 1
            
            self.db.commit()
            
            if task_id:
                self.task_service.update_task(task_id, progress=100)
            
            return {"price_changes_count": count, "attachment_id": attachment.id}
        except Exception as e:
            if task_id:
                self.task_service.fail_task(task_id, str(e), "price_excel_error", retry=True)
            raise

    def process_zip_archive(self, batch_id: int, file_content: bytes, filename: str,
                           uploaded_by: str, duplicate_strategy: DuplicateStrategy,
                           task_id: Optional[str] = None) -> Dict[str, Any]:
        if task_id:
            self.task_service.update_task(task_id, status=TaskStatus.PROCESSING, progress=0)
        
        attachment = self._save_uploaded_file(file_content, filename, batch_id, "archive", uploaded_by)
        
        temp_dir = tempfile.mkdtemp()
        zip_path = Path(temp_dir) / filename
        
        with open(zip_path, 'wb') as f:
            f.write(file_content)
        
        results = {
            "pdfs": 0,
            "payment_excels": 0,
            "price_changes": 0,
            "emls": 0,
            "contracts_created": 0,
            "contracts_ignored": 0,
            "contracts_overwritten": 0,
            "payment_nodes_added": 0,
            "price_changes_added": 0,
            "emails_added": 0,
            "errors": [],
            "warnings": []
        }
        
        contract_map = {}
        
        try:
            with zipfile.ZipFile(zip_path, 'r') as zf:
                file_list = zf.infolist()
                total_files = len([f for f in file_list if not f.is_dir()])
                processed = 0
                
                for file_info in file_list:
                    if file_info.is_dir():
                        continue
                    
                    inner_filename = file_info.filename
                    inner_content = zf.read(file_info)
                    
                    processed += 1
                    if task_id:
                        progress = int((processed / total_files) * 100)
                        self.task_service.update_task(task_id, progress=progress)
                    
                    try:
                        lower_name = inner_filename.lower()
                        base_name = Path(inner_filename).stem.lower()
                        
                        if lower_name.endswith('.pdf'):
                            parsed = self._parse_pdf_contract_from_bytes(inner_content, inner_filename)
                            contract_key = self._extract_contract_key(inner_filename, parsed)
                            
                            existing_contract, action = self._apply_duplicate_strategy(
                                batch_id, contract_key, duplicate_strategy, uploaded_by
                            )
                            
                            if action == "ignored":
                                results["contracts_ignored"] += 1
                                results["pdfs"] += 1
                                continue
                            
                            pdf_attachment = self._save_temp_as_attachment(
                                inner_content, inner_filename, batch_id, "contract_pdf", uploaded_by
                            )
                            
                            contract_no = parsed.get('contract_no') or f"ZIP_{batch_id}_{pdf_attachment.id}"
                            contract_data = ContractCreate(
                                contract_no=contract_no,
                                contract_name=Path(inner_filename).stem,
                                total_amount=parsed.get('total_amount', 0),
                                metadata_={"parsed_pdf": parsed, "source": "zip_archive", "action": action}
                            )
                            
                            contract = self.state_machine.create_contract(batch_id, contract_data, uploaded_by)
                            pdf_attachment.contract_id = contract.id
                            self.db.commit()
                            
                            contract_map[base_name] = contract.id
                            results["pdfs"] += 1
                            results["contracts_created"] += 1
                            if action == "overwritten":
                                results["contracts_overwritten"] += 1
                        
                        elif lower_name.endswith(('.xlsx', '.xls')):
                            if 'payment' in base_name or '付款' in base_name or '节点' in base_name:
                                contract_id = self._find_contract_for_file(contract_map, base_name, batch_id)
                                if contract_id:
                                    nodes = self._parse_excel_payment_nodes_from_bytes(inner_content)
                                    for node_data in nodes:
                                        self._add_payment_node(contract_id, node_data)
                                    results["payment_nodes_added"] += len(nodes)
                                    results["payment_excels"] += 1
                                else:
                                    results["warnings"].append(f"{inner_filename}: No matching contract found")
                            
                            elif 'price' in base_name or '改价' in base_name or '调价' in base_name:
                                contract_id = self._find_contract_for_file(contract_map, base_name, batch_id)
                                if contract_id:
                                    price_changes = self._parse_price_change_from_bytes(inner_content)
                                    for pc_data in price_changes:
                                        self._add_price_change(contract_id, pc_data, uploaded_by)
                                    results["price_changes_added"] += len(price_changes)
                                    results["price_changes"] += 1
                                else:
                                    results["warnings"].append(f"{inner_filename}: No matching contract found")
                            else:
                                results["warnings"].append(f"{inner_filename}: Unrecognized Excel file type")
                        
                        elif lower_name.endswith('.eml'):
                            contract_id = self._find_contract_for_file(contract_map, base_name, batch_id)
                            if contract_id:
                                parsed = self._parse_eml_from_bytes(inner_content)
                                self._add_acceptance_email(contract_id, parsed)
                                results["emails_added"] += 1
                                results["emls"] += 1
                            else:
                                results["warnings"].append(f"{inner_filename}: No matching contract found")
                    
                    except Exception as e:
                        results["errors"].append(f"{inner_filename}: {str(e)}")
            
            if task_id:
                self.task_service.update_task(
                    task_id, 
                    status=TaskStatus.SUCCESS,
                    progress=100,
                    result=results
                )
        
        except zipfile.BadZipFile as e:
            error_msg = f"Invalid ZIP file: {str(e)}"
            results["errors"].append(error_msg)
            if task_id:
                self.task_service.fail_task(task_id, error_msg, "bad_zip_file", retry=False)
        
        except Exception as e:
            error_msg = f"Archive processing failed: {str(e)}"
            results["errors"].append(error_msg)
            if task_id:
                self.task_service.fail_task(task_id, error_msg, "archive_processing_error", retry=True)
        
        finally:
            import shutil
            shutil.rmtree(temp_dir, ignore_errors=True)
        
        return results

    def _parse_pdf_contract_from_bytes(self, content: bytes, filename: str) -> Dict[str, Any]:
        if PdfReader is None:
            return {}
        
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        try:
            temp_file.write(content)
            temp_file.close()
            return self._parse_pdf_contract(Path(temp_file.name))
        finally:
            os.unlink(temp_file.name)

    def _parse_excel_payment_nodes_from_bytes(self, content: bytes) -> List[Dict[str, Any]]:
        if openpyxl is None:
            return []
        
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        try:
            temp_file.write(content)
            temp_file.close()
            return self._parse_excel_payment_nodes(Path(temp_file.name))
        finally:
            os.unlink(temp_file.name)

    def _parse_price_change_from_bytes(self, content: bytes) -> List[Dict[str, Any]]:
        if openpyxl is None:
            return []
        
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        try:
            temp_file.write(content)
            temp_file.close()
            return self._parse_price_change_excel(Path(temp_file.name))
        finally:
            os.unlink(temp_file.name)

    def _parse_eml_from_bytes(self, content: bytes) -> Dict[str, Any]:
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.eml')
        try:
            temp_file.write(content)
            temp_file.close()
            return self._parse_eml_file(Path(temp_file.name))
        finally:
            os.unlink(temp_file.name)

    def _save_temp_as_attachment(self, content: bytes, filename: str, batch_id: int,
                                attachment_type: str, uploaded_by: str) -> Attachment:
        batch_dir = settings.UPLOAD_DIR / f"batch_{batch_id}"
        batch_dir.mkdir(parents=True, exist_ok=True)
        
        safe_filename = re.sub(r'[^\w\s.-]', '_', filename)
        file_path = batch_dir / safe_filename
        file_path.write_bytes(content)
        
        attachment = Attachment(
            batch_id=batch_id,
            file_name=safe_filename,
            file_path=str(file_path),
            file_type=filename.split('.')[-1].lower() if '.' in filename else 'unknown',
            file_size=len(content),
            uploaded_by=uploaded_by,
            attachment_type=attachment_type
        )
        self.db.add(attachment)
        self.db.commit()
        self.db.refresh(attachment)
        return attachment

    def _find_contract_for_file(self, contract_map: Dict[str, int], base_name: str, batch_id: int) -> Optional[int]:
        for key, cid in contract_map.items():
            if key in base_name or base_name in key:
                return cid
        
        contracts = self.db.query(Contract).filter(Contract.batch_id == batch_id).order_by(Contract.id.desc()).first()
        if contracts:
            return contracts.id
        return None

    def _add_payment_node(self, contract_id: int, node_data: Dict[str, Any]):
        due_date = node_data.get('due_date')
        if isinstance(due_date, str):
            try:
                due_date = datetime.fromisoformat(due_date)
            except:
                due_date = None
        elif not isinstance(due_date, datetime):
            due_date = None
        
        node = PaymentNode(
            contract_id=contract_id,
            node_name=node_data.get('node_name', ''),
            node_type=node_data.get('node_type', 'milestone'),
            payment_ratio=node_data.get('payment_ratio', 0),
            payment_amount=node_data.get('payment_amount', 0),
            due_date=due_date,
            status=node_data.get('status', 'pending')
        )
        self.db.add(node)

    def _add_price_change(self, contract_id: int, pc_data: Dict[str, Any], uploaded_by: str):
        effective_date = pc_data.get('effective_date')
        if not isinstance(effective_date, datetime):
            effective_date = datetime.utcnow()
        
        pc = PriceChange(
            contract_id=contract_id,
            original_price=pc_data.get('original_price', 0),
            new_price=pc_data.get('new_price', 0),
            change_reason=pc_data.get('change_reason', ''),
            approved_by=pc_data.get('approved_by', uploaded_by),
            approved_date=datetime.utcnow(),
            effective_date=effective_date,
            is_manual=True
        )
        self.db.add(pc)

    def _add_acceptance_email(self, contract_id: int, parsed: Dict[str, Any]):
        email = AcceptanceEmail(
            contract_id=contract_id,
            email_subject=parsed.get('subject', ''),
            email_from=parsed.get('from', ''),
            email_to=parsed.get('to', ''),
            email_date=parsed.get('email_date', datetime.utcnow()),
            acceptance_result=parsed.get('acceptance_result', 'pending'),
            acceptance_amount=parsed.get('acceptance_amount', 0),
            content=parsed.get('content', '')
        )
        self.db.add(email)

    def async_import_archive(self, batch_id: int, file_content: bytes, filename: str,
                            uploaded_by: str, duplicate_strategy: DuplicateStrategy) -> str:
        task = self.task_service.create_task(
            task_type="archive_import",
            created_by=uploaded_by,
            batch_id=batch_id
        )
        
        import threading
        
        def import_worker():
            try:
                self.process_zip_archive(
                    batch_id, file_content, filename, uploaded_by, duplicate_strategy, task.task_id
                )
            except Exception as e:
                self.task_service.fail_task(task.task_id, str(e), "worker_exception", retry=True)
        
        thread = threading.Thread(target=import_worker, daemon=True)
        thread.start()
        
        return task.task_id

    def get_attachments(self, batch_id: int, contract_id: Optional[int] = None) -> List[Attachment]:
        query = self.db.query(Attachment).filter(Attachment.batch_id == batch_id)
        if contract_id:
            query = query.filter(Attachment.contract_id == contract_id)
        return query.all()
