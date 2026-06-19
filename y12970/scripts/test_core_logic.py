#!/usr/bin/env python3
import os
import sys
import json
import time
import subprocess
import urllib.request
import urllib.error

BASE_DIR = os.path.abspath(os.path.dirname(os.path.dirname(__file__)))
SAMPLE_DIR = os.path.join(BASE_DIR, 'sample_data')
sys.path.insert(0, BASE_DIR)

BASE_URL = 'http://127.0.0.1:5001'


def print_sep(title):
    print()
    print('=' * 70)
    print(f'  {title}')
    print('=' * 70)


def print_ok(msg):
    print(f'  ✅ {msg}')


def print_fail(msg):
    print(f'  ❌ {msg}')


def load_json(filename):
    path = os.path.join(SAMPLE_DIR, filename)
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def test_api(cmd, **kwargs):
    try:
        url = f'{BASE_URL}{cmd}'
        data = kwargs.get('data')
        method = kwargs.get('method', 'GET')
        
        if data is not None:
            req = urllib.request.Request(
                url,
                data=json.dumps(data).encode('utf-8'),
                headers={'Content-Type': 'application/json'},
                method=method
            )
        else:
            req = urllib.request.Request(url, method=method)
        
        with urllib.request.urlopen(req, timeout=10) as resp:
            body = resp.read().decode('utf-8')
            return resp.status, json.loads(body)
    except urllib.error.HTTPError as e:
        body = e.read().decode('utf-8')
        try:
            return e.code, json.loads(body)
        except Exception:
            return e.code, {'error': body}
    except Exception as e:
        return None, {'error': str(e)}


def run_all_tests():
    print_sep('初始化测试环境 - 重置数据库')
    db_path = os.path.join(BASE_DIR, 'audit_platform.db')
    if os.path.exists(db_path):
        os.remove(db_path)
    report_dir = os.path.join(BASE_DIR, 'reports')
    if os.path.exists(report_dir):
        for f in os.listdir(report_dir):
            os.remove(os.path.join(report_dir, f))
    print_ok('数据库重置完成')

    print_sep('启动服务...')
    venv_python = os.path.join(BASE_DIR, 'venv', 'bin', 'python')
    proc = subprocess.Popen(
        [venv_python, os.path.join(BASE_DIR, 'app.py')],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        cwd=BASE_DIR
    )
    time.sleep(3)

    try:
        print_ok('服务启动等待中...')
        for _ in range(10):
            try:
                with urllib.request.urlopen(f'{BASE_URL}/api/health', timeout=2) as r:
                    print_ok('服务健康检查通过')
                    break
            except Exception:
                time.sleep(1)
        else:
            print_fail('服务启动失败')
            proc.terminate()
            return False

        print_sep('1. 权限导入功能测试')
        perm_data = load_json('permissions_batch1.json')
        status, resp = test_api('/api/import/permissions', data=perm_data, method='POST')
        assert status == 200 and resp.get('success'), f'导入失败: {resp}'
        print(f'  - 导入结果: {resp["summary"]["success"]} 条成功, {resp["summary"]["duplicate"]} 条重复')
        batch_no_1 = resp['batch_no']
        assert resp['summary']['success'] == 8, f'期望成功8条，实际{resp["summary"]["success"]}条'
        print_ok('首次权限导入成功')

        print_sep('2. 重复导入检测测试 - 完整重复导入')
        status, resp = test_api('/api/import/permissions', data=perm_data, method='POST')
        assert status == 200 and resp.get('success')
        dup = resp['summary']['duplicate']
        print(f'  - 重复导入结果: {resp["summary"]["success"]} 条成功, {dup} 条重复')
        assert dup == 8, f'期望8条重复，实际{dup}条'
        assert 'duplicate_warning' in resp, '应有重复警告'
        print_ok(f'正确检测到批量重复导入: {resp["duplicate_warning"]["message"]}')
        print(f'  - 可操作建议: {resp["duplicate_warning"]["suggestion"]}')

        print_sep('2.1 部分重复导入测试 - 混合新旧记录')
        partial_data = {
            'source_system': 'core_db_prod',
            'operator': 'tester',
            'records': perm_data['records'][:3] + [
                {
                    'permission_id': 'PERM-NEW001',
                    'user_account': 'newuser',
                    'user_name': '新用户',
                    'department': '审计部',
                    'resource': 'new_table',
                    'resource_type': 'table',
                    'permission_level': 'read',
                    'action': 'select',
                    'granted_by': 'admin'
                }
            ]
        }
        status, resp = test_api('/api/import/permissions', data=partial_data, method='POST')
        assert status == 200 and resp.get('success')
        print(f'  - 部分重复结果: {resp["summary"]["success"]} 条新导入, {resp["summary"]["duplicate"]} 条重复')
        assert resp['summary']['success'] == 1
        assert resp['summary']['duplicate'] == 3
        print_ok('部分重复场景检测通过')

        print_sep('3. 越权检测功能测试')
        detect_data = {'source_system': 'core_db_prod', 'operator': 'tester'}
        status, resp = test_api('/api/detect/violations', data=detect_data, method='POST')
        assert status == 200 and resp.get('success')
        detected = resp['detected_count']
        print(f'  - 检测到 {detected} 条违规记录')
        assert detected > 0, '应检测到至少1条违规'
        print(f'  - 风险分布: 严重={resp["by_risk"]["critical"]}, 高={resp["by_risk"]["high"]}, 中={resp["by_risk"]["medium"]}')
        print(f'  - 索引失效类: {resp["index_failure_count"]} 条')
        print(f'  - 需工程师复核: {resp["needs_engineer_review_count"]} 条')
        print_ok('越权检测完成')

        print_sep('4. 查看违规列表 - 可用性标记验证')
        status, resp = test_api('/api/violations')
        assert status == 200 and resp.get('success')
        vios = resp['violations']
        assert len(vios) > 0
        has_need_review = any(v['needs_engineer_review'] for v in vios)
        has_index_fail = any(v['is_index_failure'] for v in vios)
        print(f'  - 共 {len(vios)} 条违规，其中')
        need_review_cnt = sum(1 for v in vios if v['needs_engineer_review'])
        print(f'  - 需工程师复核: {need_review_cnt} 条')
        print(f'  - 可用性标记样例:')
        for v in vios[:3]:
            print(f'    - {v["usability"]} | {v["violation_type"]} | {v["user_account"]} -> {v["resource"]}')
        assert has_need_review or has_index_fail, '应有需工程师复核或索引失效的记录'
        print_ok('违规列表可用性标记正确')

        print_sep('5. 备份校验 - 完整备份')
        backup_valid = load_json('backup_valid.json')
        status, resp = test_api('/api/backup/register', data=backup_valid, method='POST')
        assert status == 200 and resp.get('success')
        print_ok(f'备份注册成功: {resp["backup_id"]}')
        verify_data = {
            'expected_checksum': backup_valid['checksum'],
            'expected_record_count': backup_valid['record_count'],
            'metadata_only': True
        }
        status, resp = test_api(
            f'/api/backup/{backup_valid["backup_id"]}/verify',
            data=verify_data,
            method='POST'
        )
        assert status == 200 and resp.get('success')
        assert resp['is_valid'] is True
        print_ok('完整备份校验通过')

        print_sep('5.1 备份校验 - 不完整备份')
        backup_inc = load_json('backup_incomplete.json')
        status, resp = test_api('/api/backup/register', data=backup_inc, method='POST')
        assert status == 200 and resp.get('success')
        print_ok('不完整备份注册成功')
        status, resp = test_api(f'/api/backup/{backup_inc["backup_id"]}/verify', data={}, method='POST')
        assert status == 200 and resp.get('success')
        assert resp['is_valid'] is False
        print(f'  - 检测到问题: {len(resp["issues"])} 项')
        for issue in resp['issues']:
            print(f'    * {issue}')
        print(f'  - 可操作建议: {resp["suggestion"]}')
        assert 'suggestion' in resp, '应有可操作的修复建议'
        print_ok('不完整备份校验提示正确，给出可操作建议')

        print_sep('5.2 错误处理 - 不存在的备份ID')
        status, resp = test_api('/api/backup/NONEXISTENT/verify', data={}, method='POST')
        assert status == 404
        assert 'suggestion' in resp
        print(f'  - 错误消息: {resp["error"]}')
        print(f'  - 可操作建议: {resp["suggestion"]}')
        print_ok('不存在备份给出可操作提示')

        print_sep('6. Schema对比功能测试')
        schema_data = load_json('schema_compare.json')
        status, resp = test_api('/api/schema/compare', data=schema_data, method='POST')
        assert status == 200 and resp.get('success')
        print(f'  - 差异总数: {resp["total_differences"]}')
        print(f'  - 严重程度分布: {resp["by_severity"]}')
        print(f'  - 审计结论: {resp["conclusion"]}')
        assert resp['total_differences'] > 0
        assert 'conclusion' in resp
        print_ok('Schema对比完成，结论已关联来源材料')

        print_sep('7. 审计报告导出测试')
        req = urllib.request.Request(f'{BASE_URL}/api/report/export')
        with urllib.request.urlopen(req, timeout=15) as r:
            data = r.read()
            report_path = os.path.join(BASE_DIR, 'reports', 'test_report.xlsx')
            with open(report_path, 'wb') as f:
                f.write(data)
        size = os.path.getsize(report_path)
        assert size > 0
        print_ok(f'报告导出成功，大小: {size} 字节')
        print(f'  - 保存至: {report_path}')

        import openpyxl
        wb = openpyxl.load_workbook(report_path)
        sheets = wb.sheetnames
        print(f'  - Sheet列表: {sheets}')
        expected_sheets = ['审计概览', '越权明细', '备份校验', 'Schema对比']
        for s in expected_sheets:
            assert s in sheets, f'缺少Sheet: {s}'
        print_ok('报告包含所有必需Sheet页')

        ws_vio = wb['越权明细']
        has_red = 0
        has_orange = 0
        for row in ws_vio.iter_rows(min_row=2):
            fill = None
            if row[0].fill and row[0].fill.start_color:
                fill = row[0].fill.start_color.rgb
            if fill == '00FFC7CE':
                has_red += 1
            elif fill == '00FCD5B4':
                has_orange += 1
        print(f'  - 红色标记(需工程师复核): {has_red} 行')
        print(f'  - 橙色标记(待审计确认): {has_orange} 行')
        assert has_red > 0 and has_orange > 0, '应有颜色区分可用性'
        print_ok('越权明细可用性颜色标记正确')

        print_sep('8. 批次回滚测试')
        status, resp = test_api(f'/api/import/{batch_no_1}/rollback', data={}, method='POST')
        assert status == 200 and resp.get('success')
        print_ok(f'批次回滚成功，删除 {resp["rolled_back_permissions"]} 条记录')

        print()
        print('=' * 70)
        print('  🎉 🌟🌈  所有测试通过！  🌈🌟 �')
        print('=' * 70)
        print()

        return True

    finally:
        proc.terminate()
        proc.wait()


if __name__ == '__main__':
    try:
        success = run_all_tests()
        sys.exit(0 if success else 1)
    except AssertionError as e:
        print()
        print_fail(f'测试失败: {e}')
        import traceback
        traceback.print_exc()
        sys.exit(1)
    except Exception as e:
        print_fail(f'异常: {e}')
        import traceback
        traceback.print_exc()
        sys.exit(1)
