import json
from typing import List, Optional
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app import models
from app.services import RoleBasedViewService


class ExportService:
    def __init__(self, db: Session, user_role: str = "viewer"):
        self.db = db
        self.role_service = RoleBasedViewService(user_role)

    def export_sample_labels_to_excel(self, status_filter: Optional[str] = None) -> BytesIO:
        query = self.db.query(models.SampleLabel).filter(models.SampleLabel.is_active == True)
        if status_filter:
            query = query.filter(models.SampleLabel.status == status_filter)
        
        records = query.all()
        output = BytesIO()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "留样标签台账"

        headers = ["批次号", "产品名称", "生产时间", "留样时间", "留样人员", 
                   "存储位置", "状态", "版本", "创建时间", "更新时间", "创建人", "更新人"]
        ws.append(headers)
        
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")

        for record in records:
            record_dict = {
                "batch_no": record.batch_no,
                "product_name": record.product_name,
                "production_time": record.production_time.strftime("%Y-%m-%d %H:%M:%S") if record.production_time else "",
                "sample_time": record.sample_time.strftime("%Y-%m-%d %H:%M:%S") if record.sample_time else "",
                "sampler": record.sampler,
                "storage_location": record.storage_location,
                "status": self._translate_status(record.status),
                "version": record.version,
                "created_at": record.created_at.strftime("%Y-%m-%d %H:%M:%S") if record.created_at else "",
                "updated_at": record.updated_at.strftime("%Y-%m-%d %H:%M:%S") if record.updated_at else "",
                "created_by": record.created_by,
                "updated_by": record.updated_by
            }
            
            masked = self.role_service.mask_sensitive_fields(record_dict)
            
            ws.append([
                masked["batch_no"],
                masked["product_name"],
                masked["production_time"],
                masked["sample_time"],
                masked["sampler"],
                masked["storage_location"],
                masked["status"],
                masked["version"],
                masked["created_at"],
                masked["updated_at"],
                masked["created_by"],
                masked["updated_by"]
            ])

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(output)
        output.seek(0)
        return output

    def export_temperature_records_to_excel(self, sample_label_id: Optional[int] = None) -> BytesIO:
        query = self.db.query(models.TemperatureRecord).filter(models.TemperatureRecord.is_active == True)
        if sample_label_id:
            query = query.filter(models.TemperatureRecord.sample_label_id == sample_label_id)
        
        records = query.all()
        output = BytesIO()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "温度记录"

        headers = ["记录ID", "关联批次ID", "记录时间", "温度(°C)", "记录人员", "状态", "版本", "创建时间"]
        ws.append(headers)
        
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")

        for record in records:
            record_dict = {
                "id": record.id,
                "sample_label_id": record.sample_label_id,
                "record_time": record.record_time.strftime("%Y-%m-%d %H:%M:%S") if record.record_time else "",
                "temperature": record.temperature,
                "recorder": record.recorder,
                "status": self._translate_status(record.status),
                "version": record.version,
                "created_at": record.created_at.strftime("%Y-%m-%d %H:%M:%S") if record.created_at else ""
            }
            
            masked = self.role_service.mask_sensitive_fields(record_dict)
            
            ws.append([
                masked["id"],
                masked["sample_label_id"],
                masked["record_time"],
                masked["temperature"],
                masked["recorder"],
                masked["status"],
                masked["version"],
                masked["created_at"]
            ])

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(output)
        output.seek(0)
        return output

    def export_batch_trace_to_excel(self, batch_no: str) -> BytesIO:
        sample_label = self.db.query(models.SampleLabel).filter(
            models.SampleLabel.batch_no == batch_no
        ).first()

        if not sample_label:
            raise ValueError(f"批次 {batch_no} 不存在")

        output = BytesIO()
        wb = Workbook()

        ws1 = wb.active
        ws1.title = "批次基本信息"
        ws1.append(["批次号", sample_label.batch_no])
        ws1.append(["产品名称", sample_label.product_name])
        ws1.append(["生产时间", sample_label.production_time.strftime("%Y-%m-%d %H:%M:%S") if sample_label.production_time else ""])
        ws1.append(["留样时间", sample_label.sample_time.strftime("%Y-%m-%d %H:%M:%S") if sample_label.sample_time else ""])
        ws1.append(["状态", self._translate_status(sample_label.status)])
        ws1.append(["版本", sample_label.version])

        ws2 = wb.create_sheet("关联门店")
        ws2.append(["门店ID", "门店名称", "扫码次数", "总数量"])
        seen_stores = set()
        for scan in sample_label.scan_records:
            if scan.store_id not in seen_stores:
                seen_stores.add(scan.store_id)
                count = len([s for s in sample_label.scan_records if s.store_id == scan.store_id])
                total_qty = sum(s.quantity for s in sample_label.scan_records if s.store_id == scan.store_id)
                ws2.append([scan.store_id, scan.store_name, count, total_qty])

        ws3 = wb.create_sheet("温度记录")
        ws3.append(["记录时间", "温度(°C)", "记录人员", "状态"])
        for temp in sample_label.temperature_records:
            record_dict = {
                "record_time": temp.record_time.strftime("%Y-%m-%d %H:%M:%S") if temp.record_time else "",
                "temperature": temp.temperature,
                "recorder": temp.recorder,
                "status": self._translate_status(temp.status)
            }
            masked = self.role_service.mask_sensitive_fields(record_dict)
            ws3.append([masked["record_time"], masked["temperature"], masked["recorder"], masked["status"]])

        ws4 = wb.create_sheet("门店投诉")
        ws4.append(["门店名称", "投诉类型", "投诉描述", "投诉时间", "处理人", "状态"])
        for complaint in sample_label.complaints:
            record_dict = {
                "store_name": complaint.store_name,
                "complaint_type": complaint.complaint_type,
                "complaint_desc": complaint.complaint_desc,
                "complaint_time": complaint.complaint_time.strftime("%Y-%m-%d %H:%M:%S") if complaint.complaint_time else "",
                "handler": complaint.handler,
                "status": self._translate_status(complaint.status)
            }
            masked = self.role_service.mask_sensitive_fields(record_dict)
            ws4.append([
                masked["store_name"], masked["complaint_type"], masked["complaint_desc"],
                masked["complaint_time"], masked["handler"], masked["status"]
            ])

        ws5 = wb.create_sheet("扫码明细")
        ws5.append(["门店名称", "扫码时间", "扫码人员", "数量", "状态"])
        for scan in sample_label.scan_records:
            record_dict = {
                "store_name": scan.store_name,
                "scan_time": scan.scan_time.strftime("%Y-%m-%d %H:%M:%S") if scan.scan_time else "",
                "scanner": scan.scanner,
                "quantity": scan.quantity,
                "status": self._translate_status(scan.status)
            }
            masked = self.role_service.mask_sensitive_fields(record_dict)
            ws5.append([masked["store_name"], masked["scan_time"], masked["scanner"], masked["quantity"], masked["status"]])

        ws6 = wb.create_sheet("操作审计")
        ws6.append(["操作类型", "原状态", "新状态", "变更原因", "操作人", "操作人角色", "操作时间"])
        for audit in sample_label.audit_logs:
            record_dict = {
                "action_type": audit.action_type,
                "old_status": self._translate_status(audit.old_status) if audit.old_status else "",
                "new_status": self._translate_status(audit.new_status) if audit.new_status else "",
                "change_reason": audit.change_reason,
                "operator": audit.operator,
                "operator_role": audit.operator_role,
                "created_at": audit.created_at.strftime("%Y-%m-%d %H:%M:%S") if audit.created_at else ""
            }
            masked = self.role_service.mask_sensitive_fields(record_dict)
            ws6.append([
                masked["action_type"], masked["old_status"], masked["new_status"],
                masked["change_reason"], masked["operator"], masked["operator_role"], masked["created_at"]
            ])

        for ws in [ws1, ws2, ws3, ws4, ws5, ws6]:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 40)
                ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(output)
        output.seek(0)
        return output

    def _translate_status(self, status: Optional[str]) -> str:
        status_map = {
            "draft": "草稿",
            "submitted": "已提交",
            "rejected": "已驳回",
            "confirmed": "已确认",
            None: ""
        }
        return status_map.get(status, status or "")
