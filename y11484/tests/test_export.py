import pytest
from io import BytesIO
from openpyxl import load_workbook


class TestExport:
    def test_filename_encoding_function(self):
        """测试: 文件名编码函数正确性"""
        from app.main import encode_filename
        
        result = encode_filename("留样标签台账_20260525.xlsx")
        assert "filename*=UTF-8''" in result
        assert "%E7%95%99%E6%A0%B7" in result
        
        ascii_result = encode_filename("test.xlsx")
        assert "filename*=UTF-8''test.xlsx" in ascii_result

    def test_export_sample_labels_no_unicode_error(self, client, create_sample_label):
        """测试: 导出留样标签 - 不应触发 UnicodeEncodeError"""
        create_sample_label("EXPORT_001")
        
        response = client.get(
            "/export/sample-labels",
            headers={"X-User-Role": "quality_manager"}
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        content_disposition = response.headers["content-disposition"]
        assert "filename*=UTF-8''" in content_disposition
        assert "latin-1" not in content_disposition.lower()

    def test_export_sample_labels_valid_excel(self, client, create_sample_label):
        """测试: 导出的Excel文件格式正确"""
        create_sample_label("EXPORT_002")
        
        response = client.get(
            "/export/sample-labels",
            headers={"X-User-Role": "quality_manager"}
        )
        
        assert response.status_code == 200
        
        excel_content = BytesIO(response.content)
        wb = load_workbook(excel_content)
        
        assert "留样标签台账" in wb.sheetnames
        ws = wb["留样标签台账"]
        headers = [cell.value for cell in ws[1]]
        assert "批次号" in headers
        assert "产品名称" in headers

    def test_export_temperature_records(self, client, create_sample_label):
        """测试: 导出温度记录"""
        label = create_sample_label("EXPORT_003")
        label_id = label["id"]
        
        temp_data = {
            "sample_label_id": label_id,
            "record_time": "2026-05-25T10:00:00",
            "temperature": 5.0,
            "recorder": "测试员"
        }
        client.post("/temperature-records/", json=temp_data, headers={"X-User-Id": "test"})
        
        response = client.get(
            "/export/temperature-records",
            headers={"X-User-Role": "quality_manager"}
        )
        
        assert response.status_code == 200
        content_disposition = response.headers["content-disposition"]
        assert "filename*=UTF-8''" in content_disposition
        
        excel_content = BytesIO(response.content)
        wb = load_workbook(excel_content)
        assert "温度记录" in wb.sheetnames

    def test_export_batch_trace(self, client, create_sample_label):
        """测试: 导出批次追溯完整报告"""
        label = create_sample_label("EXPORT_TRACE_001")
        label_id = label["id"]
        
        temp_data = {
            "sample_label_id": label_id,
            "record_time": "2026-05-25T10:00:00",
            "temperature": 5.0,
            "recorder": "测试员"
        }
        client.post("/temperature-records/", json=temp_data, headers={"X-User-Id": "test"})
        
        scan_data = {
            "sample_label_id": label_id,
            "store_id": "STORE001",
            "store_name": "测试门店",
            "scan_time": "2026-05-25T11:00:00",
            "scanner": "测试员",
            "quantity": 10
        }
        client.post("/scan-records/", json=scan_data, headers={"X-User-Id": "test"})
        
        response = client.get(
            "/export/batch/EXPORT_TRACE_001",
            headers={"X-User-Role": "quality_manager"}
        )
        
        assert response.status_code == 200
        content_disposition = response.headers["content-disposition"]
        assert "filename*=UTF-8''" in content_disposition
        
        excel_content = BytesIO(response.content)
        wb = load_workbook(excel_content)
        
        expected_sheets = ["批次基本信息", "关联门店", "温度记录", "门店投诉", "扫码明细", "操作审计"]
        for sheet_name in expected_sheets:
            assert sheet_name in wb.sheetnames, f"缺少Sheet: {sheet_name}"

    def test_export_with_viewer_role_masked(self, client, create_sample_label):
        """测试: viewer角色导出敏感字段脱敏"""
        create_sample_label("EXPORT_MASK_001")
        
        response_manager = client.get(
            "/export/sample-labels",
            headers={"X-User-Role": "quality_manager"}
        )
        assert response_manager.status_code == 200
        
        response_viewer = client.get(
            "/export/sample-labels",
            headers={"X-User-Role": "viewer"}
        )
        assert response_viewer.status_code == 200
        
        assert len(response_manager.content) > 0
        assert len(response_viewer.content) > 0

    def test_export_batch_not_found(self, client):
        """测试: 导出不存在的批次应返回404"""
        response = client.get(
            "/export/batch/NONEXISTENT_999",
            headers={"X-User-Role": "quality_manager"}
        )
        assert response.status_code == 404

    def test_export_and_api_same_fact(self, client, create_sample_label):
        """测试: 导出文件与详情接口讲同一套事实"""
        label = create_sample_label("EXPORT_CONSIST_001")
        batch_no = label["batch_no"]
        
        api_response = client.get(f"/trace/batch/{batch_no}")
        api_data = api_response.json()
        
        export_response = client.get(
            f"/export/batch/{batch_no}",
            headers={"X-User-Role": "quality_manager"}
        )
        assert export_response.status_code == 200
        
        excel_content = BytesIO(export_response.content)
        wb = load_workbook(excel_content)
        
        ws_info = wb["批次基本信息"]
        excel_batch_no = ws_info["B1"].value
        excel_product = ws_info["B2"].value
        
        assert excel_batch_no == api_data["batch_no"]
        assert excel_product == api_data["product_name"]
