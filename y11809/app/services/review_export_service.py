from sqlalchemy.orm import Session
from sqlalchemy import and_
from typing import List, Dict, Any, Optional
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from app.models.models import (
    StudentCard, RechargeRecord, ConsumeRevoke, RefundRecord, RefundDetail,
    RechargeAmendment
)
from app.services.refund_service import RefundRuleEngine, BalanceService
from app.models.schemas import (
    RefundResult, RefundDetailItem, RefundRecordResponse, ReportRow
)


class RefundSettlementService:

    @staticmethod
    def batch_calculate_refunds(db: Session, card_nos: List[str], operator: str) -> Dict[str, Any]:
        """批量毕业季退卡结算"""
        batch_no = f"BT{datetime.now().strftime('%Y%m%d%H%M%S')}"
        results: List[RefundResult] = []
        failed_cards: List[Dict[str, Any]] = []
        success_count = 0

        for card_no in card_nos:
            try:
                result = RefundRuleEngine.calculate_refund(db, card_no, operator)

                refund_record = RefundRecord(
                    card_no=card_no,
                    refund_batch_no=batch_no,
                    refund_no=result.refund_no,
                    total_refund=result.total_refund,
                    subsidy_refund=result.subsidy_refund,
                    recharge_refund=result.recharge_refund,
                    revoke_refund=result.revoke_refund,
                    blocked_subsidy=result.blocked_subsidy,
                    has_blocked_subsidy=result.has_blocked_subsidy,
                    has_cross_day_revoke=result.has_cross_day_revoke,
                    has_merged_card=result.has_merged_card,
                    status=result.status,
                    review_status="pending",
                    operator=operator
                )
                db.add(refund_record)
                db.flush()

                detail_items = []
                for d in result.details:
                    detail = RefundDetail(
                        refund_id=refund_record.id,
                        item_type=d.item_type,
                        item_no=d.item_no,
                        amount=d.amount,
                        is_blocked=d.is_blocked,
                        block_reason=d.block_reason
                    )
                    detail_items.append(detail)

                db.add_all(detail_items)
                db.flush()

                results.append(result)
                success_count += 1

            except Exception as e:
                import traceback
                failed_cards.append({
                    "card_no": card_no,
                    "reason": str(e),
                    "traceback": traceback.format_exc()[:500]
                })
                db.rollback()
                continue

        db.commit()

        return {
            "batch_no": batch_no,
            "total_count": len(card_nos),
            "success_count": success_count,
            "failed_count": len(failed_cards),
            "results": results,
            "failed_cards": failed_cards
        }

    @staticmethod
    def get_refund_records(
        db: Session,
        batch_no: Optional[str] = None,
        status: Optional[str] = None,
        has_blocked_subsidy: Optional[bool] = None,
        has_cross_day_revoke: Optional[bool] = None,
        has_merged_card: Optional[bool] = None,
        review_status: Optional[str] = None,
    ) -> List[RefundRecordResponse]:
        """查询退费记录，支持多维度筛选"""
        query = db.query(RefundRecord)

        if batch_no:
            query = query.filter(RefundRecord.refund_batch_no == batch_no)
        if status:
            query = query.filter(RefundRecord.status == status)
        if has_blocked_subsidy is not None:
            query = query.filter(RefundRecord.has_blocked_subsidy == has_blocked_subsidy)
        if has_cross_day_revoke is not None:
            query = query.filter(RefundRecord.has_cross_day_revoke == has_cross_day_revoke)
        if has_merged_card is not None:
            query = query.filter(RefundRecord.has_merged_card == has_merged_card)
        if review_status:
            query = query.filter(RefundRecord.review_status == review_status)

        records = query.order_by(RefundRecord.created_at.desc()).all()

        return [
            RefundRecordResponse(
                id=r.id,
                card_no=r.card_no,
                refund_batch_no=r.refund_batch_no,
                refund_no=r.refund_no,
                total_refund=r.total_refund,
                subsidy_refund=r.subsidy_refund,
                recharge_refund=r.recharge_refund,
                revoke_refund=r.revoke_refund,
                blocked_subsidy=r.blocked_subsidy,
                has_blocked_subsidy=r.has_blocked_subsidy,
                has_cross_day_revoke=r.has_cross_day_revoke,
                has_merged_card=r.has_merged_card,
                status=r.status,
                review_status=r.review_status,
                refund_time=r.refund_time,
                operator=r.operator,
                reviewer=r.reviewer,
            )
            for r in records
        ]


class ReviewAndExportService:

    @staticmethod
    def batch_review(
        db: Session,
        batch_no: str,
        refund_nos: List[str],
        action: str,
        reviewer: str,
        remark: Optional[str] = None,
    ) -> Dict[str, Any]:
        """批量复核：通过/拒绝"""
        approved = 0
        rejected = 0
        reviewed = 0

        for refund_no in refund_nos:
            record = db.query(RefundRecord).filter(
                and_(
                    RefundRecord.refund_no == refund_no,
                    RefundRecord.refund_batch_no == batch_no
                )
            ).first()

            if not record:
                continue

            if action == "approve":
                if record.has_blocked_subsidy:
                    record.status = "blocked"
                    continue
                record.status = "approved"
                record.review_status = "reviewed"
                approved += 1
            elif action == "reject":
                record.status = "rejected"
                record.review_status = "reviewed"
                rejected += 1

            record.reviewer = reviewer
            record.review_time = datetime.now()
            if remark:
                record.remark = remark

            reviewed += 1

        db.commit()

        return {
            "batch_no": batch_no,
            "reviewed_count": reviewed,
            "approved_count": approved,
            "rejected_count": rejected,
        }

    @staticmethod
    def generate_report_data(
        db: Session,
        report_type: str,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None,
        batch_no: Optional[str] = None,
        include_blocked: bool = True,
        include_cross_day: bool = True,
        include_merged: bool = True,
    ) -> List[ReportRow]:
        """生成报表数据（日常/月底口径一致"""
        query = db.query(RefundRecord)

        if batch_no:
            query = query.filter(RefundRecord.refund_batch_no == batch_no)

        if start_date:
            query = query.filter(RefundRecord.refund_time >= start_date)
        if end_date:
            query = query.filter(RefundRecord.refund_time <= end_date)

        if not include_blocked:
            query = query.filter(RefundRecord.has_blocked_subsidy == False)
        if not include_cross_day:
            query = query.filter(RefundRecord.has_cross_day_revoke == False)
        if not include_merged:
            query = query.filter(RefundRecord.has_merged_card == False)

        records = query.order_by(RefundRecord.created_at.desc()).all()

        report_rows = []
        for r in records:
            card = db.query(StudentCard).filter(StudentCard.card_no == r.card_no).first()
            report_rows.append(
                ReportRow(
                    card_no=r.card_no,
                    student_name=card.student_name if card else None,
                    department=card.department if card else None,
                    total_balance=r.total_refund + r.blocked_subsidy,
                    refundable_amount=r.total_refund,
                    blocked_amount=r.blocked_subsidy,
                    has_blocked_subsidy=r.has_blocked_subsidy,
                    has_cross_day_revoke=r.has_cross_day_revoke,
                    has_merged_card=r.has_merged_card,
                    status=r.status,
                    review_status=r.review_status,
                )
            )

        return report_rows

    @staticmethod
    def export_to_excel(
        db: Session,
        report_type: str,
        report_rows: List[ReportRow],
    ) -> BytesIO:
        """导出Excel报表"""
        wb = Workbook()
        ws = wb.active
        ws.title = "沉淀金报表"

        headers = [
            "卡号", "姓名", "院系", "总余额", "可退金额", "拦截金额",
            "补贴拦截", "跨日撤销", "卡号合并", "状态", "复核状态"
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, row_data in enumerate(report_rows, 2):
            ws.cell(row=row_idx, column=1, value=row_data.card_no)
            ws.cell(row=row_idx, column=2, value=row_data.student_name or "")
            ws.cell(row=row_idx, column=3, value=row_data.department or "")
            ws.cell(row=row_idx, column=4, value=row_data.total_balance)
            ws.cell(row=row_idx, column=5, value=row_data.refundable_amount)
            ws.cell(row=row_idx, column=6, value=row_data.blocked_amount)
            ws.cell(row=row_idx, column=7, value="是" if row_data.has_blocked_subsidy else "否")
            ws.cell(row=row_idx, column=8, value="是" if row_data.has_cross_day_revoke else "否")
            ws.cell(row=row_idx, column=9, value="是" if row_data.has_merged_card else "否")
            ws.cell(row=row_idx, column=10, value=row_data.status)
            ws.cell(row=row_idx, column=10).number_format = '0.00'
            ws.cell(row=row_idx, column=10).number_format = '0.00'
            ws.cell(row=row_idx, column=11, value=row_data.review_status)

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
