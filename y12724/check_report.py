from openpyxl import load_workbook
import sys

wb = load_workbook(sys.argv[1])
print('Sheets:', wb.sheetnames)

ws = wb['候选清单']
print('\n候选清单 表头:', [c.value for c in ws[1]])

print('\n--- 不可用记录（红色高亮）---')
for row in ws.iter_rows(min_row=2, values_only=False):
    avail = row[5].value
    if avail == '不可用':
        fill = row[5].fill.start_color.rgb if row[5].fill and row[5].fill.start_color else None
        reason = row[6].value
        print(f'  {row[1].value}: 可用={avail}, 高亮={fill}')
        print(f'    拦截原因: {reason[:80] if reason else None}')

print('\n--- 受边界影响/延迟到达（黄色高亮）---')
for row in ws.iter_rows(min_row=2, values_only=False):
    late = row[8].value
    affected = row[9].value
    if late == '是' or (affected and str(affected).strip()):
        fill = row[0].fill.start_color.rgb if row[0].fill and row[0].fill.start_color else None
        print(f'  {row[1].value}: delayed={late}, 高亮={fill}')
        print(f'    受影响: {str(affected)[:80] if affected else None}')

ws2 = wb['不可用记录详情']
print(f'\n--- 不可用记录详情 sheet ({ws2.max_row-1} 条) ---')
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row[0]:
        print(f'  行{row[0]} {row[1]}: 类型={row[3]}')
        print(f'    {str(row[4])[:80]}')

ws3 = wb['计算草稿历史']
print(f'\n--- 计算草稿历史 sheet ({ws3.max_row-1} 条) ---')
for row in ws3.iter_rows(min_row=2, values_only=True):
    if row[0]:
        print(f'  {row[0]} | {row[1]} {row[2]} | 计算={row[5]} | 差异={row[7]}')
