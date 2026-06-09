import io
import json
from datetime import datetime
from typing import List, Optional
from collections import defaultdict

from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import StreamingResponse, JSONResponse
from sqlalchemy.orm import Session
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from database import init_db, get_db, Batch, Record, StatusHistory, CalculationDraft
from schemas import (
    BatchOut, BatchListOut, BatchStatusUpdate, RecordOut, RecordDetail,
    RecordStatusUpdate, CalculationDraftCreate, CalculationDraftOut,
    CalculationDraftUpdate, ImportResult, ValidationIssue, StatusHistoryOut
)
from validators import (
    validate_record, can_transition, describe_status, evaluate_formula,
    format_block_reason_for_report
)

init_db()
app = FastAPI(title="数独候选削减器", version="1.0.0")


def _add_status_history(db: Session, record_id: int, from_status: str, to_status: str,
                        changed_by: str = "system", note: str = ""):
    if from_status == to_status:
        return
    history = StatusHistory(
        record_id=record_id,
        from_status=from_status,
        to_status=to_status,
        changed_by=changed_by,
        changed_at=datetime.utcnow(),
        note=note
    )
    db.add(history)


@app.get("/api/batches", response_model=List[BatchListOut], summary="查询所有批次（重启后仍可见）")
def list_batches(db: Session = Depends(get_db)):
    batches = db.query(Batch).order_by(Batch.uploaded_at.desc()).all()
    return batches


@app.get("/api/batches/{batch_id}", response_model=BatchOut, summary="查询批次详情")
def get_batch(batch_id: int, db: Session = Depends(get_db)):
    batch = db.query(Batch).filter(Batch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")
    return batch


@app.patch("/api/batches/{batch_id}", response_model=BatchOut, summary="推进批次状态")
def update_batch_status(batch_id: int, body: BatchStatusUpdate, db: Session = Depends(get_db)):
    batch = db.query(Batch).filter(Batch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")
    batch.status = body.status
    if body.note:
        batch.note = (batch.note + "\n" if batch.note else "") + f"[{datetime.now().strftime('%Y-%m-%d %H:%M')}] {body.note}"
    db.commit()
    db.refresh(batch)
    return batch


@app.post("/api/import", response_model=ImportResult, summary="导入候选数据（Excel/CSV/JSON）")
async def import_records(
    file: UploadFile = File(...),
    batch_name: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    content = await file.read()
    filename = file.filename.lower()
    rows = []

    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        df = pd.read_excel(io.BytesIO(content))
        rows = df.where(pd.notnull(df), None).to_dict(orient="records")
    elif filename.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(content))
        rows = df.where(pd.notnull(df), None).to_dict(orient="records")
    elif filename.endswith(".json"):
        data = json.loads(content.decode("utf-8"))
        rows = data if isinstance(data, list) else [data]
    else:
        raise HTTPException(status_code=400, detail="仅支持 xlsx / csv / json 格式")

    name = batch_name or f"{file.filename} - {datetime.now().strftime('%Y%m%d %H%M')}"
    batch = Batch(name=name, uploaded_at=datetime.utcnow(), status="imported")
    db.add(batch)
    db.flush()

    issues: List[ValidationIssue] = []
    seen_keys = set()
    valid_count = 0
    invalid_count = 0

    for idx, row in enumerate(rows, start=1):
        is_valid, block_reason, cleaned = validate_record(row, row_no=idx)

        key = (cleaned["candidate_name"], cleaned["value"])
        if cleaned["value"] is not None and key in seen_keys:
            is_valid = False
            block_reason = (block_reason + "；" if block_reason else "") + "重复记录：同一批次中存在同名候选且数值完全一致。"
            issues.append(ValidationIssue(
                row_no=idx, candidate_name=cleaned["candidate_name"],
                issue_type="duplicate", issue_detail="同一批次存在同名且数值一致的重复记录"
            ))
        else:
            if cleaned["value"] is not None:
                seen_keys.add(key)

        initial_status = "blocked" if not is_valid else "pending_review"

        record = Record(
            batch_id=batch.id,
            row_no=cleaned["row_no"],
            candidate_name=cleaned["candidate_name"],
            value=cleaned["value"],
            unit=cleaned["unit"],
            raw_data=cleaned["raw_data"],
            is_valid=is_valid,
            block_reason=block_reason,
            status=initial_status,
            is_boundary=cleaned["is_boundary"],
            boundary_arrived_late=False,
            previous_conclusion="",
            affected_conclusions=[],
        )
        db.add(record)
        db.flush()

        _add_status_history(db, record.id, "", initial_status,
                            changed_by="system",
                            note=block_reason if block_reason else "导入成功")

        if is_valid:
            valid_count += 1
        else:
            invalid_count += 1
            if "单位缺失" in block_reason:
                issues.append(ValidationIssue(
                    row_no=idx, candidate_name=cleaned["candidate_name"],
                    issue_type="unit_missing",
                    issue_detail="单位缺失：无法进行跨记录汇总对比，已被拦截。请补录单位后再提交复核。"
                ))
            if "数值缺失" in block_reason:
                issues.append(ValidationIssue(
                    row_no=idx, candidate_name=cleaned["candidate_name"],
                    issue_type="value_missing",
                    issue_detail="数值缺失：缺少核心数值字段。"
                ))
            if "数值格式异常" in block_reason:
                issues.append(ValidationIssue(
                    row_no=idx, candidate_name=cleaned["candidate_name"],
                    issue_type="value_invalid",
                    issue_detail="数值格式异常：无法解析为有效数字。"
                ))

    batch.total_records = len(rows)
    batch.valid_records = valid_count
    batch.invalid_records = invalid_count
    db.commit()
    db.refresh(batch)

    return ImportResult(
        batch_id=batch.id,
        batch_name=batch.name,
        total_count=len(rows),
        valid_count=valid_count,
        invalid_count=invalid_count,
        issues=issues,
    )


@app.get("/api/batches/{batch_id}/records", response_model=List[RecordOut], summary="查询批次下的所有记录")
def list_records(batch_id: int, status: Optional[str] = None, only_invalid: bool = False,
                 db: Session = Depends(get_db)):
    batch = db.query(Batch).filter(Batch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")
    q = db.query(Record).filter(Record.batch_id == batch_id)
    if status:
        q = q.filter(Record.status == status)
    if only_invalid:
        q = q.filter(Record.is_valid == False)
    return q.order_by(Record.row_no.asc()).all()


@app.get("/api/records/{record_id}", response_model=RecordDetail, summary="查询单条记录详情（含状态历史和计算草稿）")
def get_record(record_id: int, db: Session = Depends(get_db)):
    record = db.query(Record).filter(Record.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")
    return record


@app.patch("/api/records/{record_id}", response_model=RecordOut, summary="更新记录状态（复核推进）")
def update_record_status(record_id: int, body: RecordStatusUpdate, db: Session = Depends(get_db)):
    record = db.query(Record).filter(Record.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")
    if not can_transition(record.status, body.status):
        raise HTTPException(status_code=400,
                            detail=f"无法从「{describe_status(record.status)}」推进到「{describe_status(body.status)}」")
    old_status = record.status
    record.status = body.status
    record.reviewed_by = body.reviewed_by or "user"
    record.reviewed_at = datetime.utcnow()
    if body.status in ("confirmed", "reopened"):
        record.is_valid = True
        record.block_reason = ""
    _add_status_history(db, record.id, old_status, body.status,
                        changed_by=body.reviewed_by or "user",
                        note=body.note or "")
    db.commit()
    db.refresh(record)
    return record


@app.patch("/api/records/{record_id}/fix", summary="补录缺失字段（如单位），自动更新状态")
def fix_record_fields(
    record_id: int,
    unit: Optional[str] = Form(None),
    value: Optional[float] = Form(None),
    reviewed_by: str = Form("user"),
    db: Session = Depends(get_db)
):
    record = db.query(Record).filter(Record.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")

    changed_fields = []
    old_status = record.status

    if unit is not None and unit.strip() != "":
        record.unit = unit.strip()
        changed_fields.append(f"单位补录为「{unit.strip()}」")

    if value is not None:
        record.value = value
        changed_fields.append(f"数值补录为 {value}")

    if record.unit and record.value is not None:
        record.is_valid = True
        new_reasons = []
        for r in (record.block_reason or "").split("；"):
            if "单位缺失" not in r and "数值缺失" not in r and "数值格式异常" not in r:
                if r.strip():
                    new_reasons.append(r.strip())
        record.block_reason = "；".join(new_reasons)
        if not record.block_reason and old_status == "blocked":
            record.status = "pending_review"
            record.reviewed_by = reviewed_by
            record.reviewed_at = datetime.utcnow()
            _add_status_history(db, record.id, old_status, "pending_review",
                                changed_by=reviewed_by,
                                note="；".join(changed_fields) + "，解除拦截")
        else:
            _add_status_history(db, record.id, old_status, old_status,
                                changed_by=reviewed_by,
                                note="；".join(changed_fields))

    db.commit()
    db.refresh(record)
    return RecordOut.model_validate(record)


@app.post("/api/calculations", response_model=CalculationDraftOut, summary="新增或补录计算草稿（自动历史对比）")
def create_calculation(body: CalculationDraftCreate, db: Session = Depends(get_db)):
    record = db.query(Record).filter(Record.id == body.record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")

    prev = (db.query(CalculationDraft)
            .filter(CalculationDraft.record_id == body.record_id,
                    CalculationDraft.formula_name == body.formula_name)
            .order_by(CalculationDraft.version.desc())
            .first())

    new_version = (prev.version + 1) if prev else 1
    prev_version_id = prev.id if prev else 0

    context = {"value": record.value} if record.value is not None else {}
    if record.raw_data and isinstance(record.raw_data, dict):
        for k, v in record.raw_data.items():
            try:
                context[k] = float(v)
            except (TypeError, ValueError):
                pass

    computed = None
    if body.draft_value is not None and str(body.draft_value).strip() != "":
        try:
            computed = float(body.draft_value)
        except (TypeError, ValueError):
            computed = evaluate_formula(body.formula_expr, context)
    else:
        computed = evaluate_formula(body.formula_expr, context)

    diff_note = body.diff_note
    if prev and prev.computed_value is not None and computed is not None:
        delta = computed - prev.computed_value
        direction = "上升" if delta > 0 else ("下降" if delta < 0 else "持平")
        diff_note = (diff_note + "；" if diff_note else "") + f"相对 v{prev.version} {direction} {abs(delta):.4f}"
    elif body.is_supplement:
        diff_note = (diff_note + "；" if diff_note else "") + "草稿补录"

    draft = CalculationDraft(
        record_id=body.record_id,
        formula_name=body.formula_name,
        formula_expr=body.formula_expr,
        draft_value=body.draft_value,
        computed_value=computed,
        is_supplement=body.is_supplement,
        created_at=datetime.utcnow(),
        updated_at=datetime.utcnow(),
        version=new_version,
        prev_version_id=prev_version_id,
        diff_note=diff_note,
    )
    db.add(draft)

    if record.status == "pending_review":
        old = record.status
        record.status = "reviewing"
        _add_status_history(db, record.id, old, "reviewing",
                            changed_by="system",
                            note=f"计算草稿 {body.formula_name} v{new_version} 已补录")

    db.commit()
    db.refresh(draft)
    return draft


@app.get("/api/records/{record_id}/calculations", response_model=List[CalculationDraftOut], summary="查看记录的计算草稿历史")
def list_calculations(record_id: int, db: Session = Depends(get_db)):
    record = db.query(Record).filter(Record.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")
    return (db.query(CalculationDraft)
            .filter(CalculationDraft.record_id == record_id)
            .order_by(CalculationDraft.formula_name.asc(), CalculationDraft.version.desc())
            .all())


@app.patch("/api/calculations/{draft_id}", response_model=CalculationDraftOut, summary="更新计算草稿（生成新版本，保留历史对比）")
def update_calculation(draft_id: int, body: CalculationDraftUpdate, db: Session = Depends(get_db)):
    old_draft = db.query(CalculationDraft).filter(CalculationDraft.id == draft_id).first()
    if not old_draft:
        raise HTTPException(status_code=404, detail="计算草稿不存在")

    record = db.query(Record).filter(Record.id == old_draft.record_id).first()

    new_version = old_draft.version + 1
    context = {"value": record.value} if record and record.value is not None else {}
    if record and record.raw_data and isinstance(record.raw_data, dict):
        for k, v in record.raw_data.items():
            try:
                context[k] = float(v)
            except (TypeError, ValueError):
                pass

    draft_value = body.draft_value if body.draft_value is not None else old_draft.draft_value
    computed = None
    if draft_value is not None and str(draft_value).strip() != "":
        try:
            computed = float(draft_value)
        except (TypeError, ValueError):
            computed = evaluate_formula(old_draft.formula_expr, context)
    else:
        computed = evaluate_formula(old_draft.formula_expr, context)

    if body.computed_value is not None:
        computed = body.computed_value

    delta_note = ""
    if old_draft.computed_value is not None and computed is not None:
        delta = computed - old_draft.computed_value
        direction = "上升" if delta > 0 else ("下降" if delta < 0 else "持平")
        delta_note = f"相对 v{old_draft.version} {direction} {abs(delta):.4f}"

    diff_note = body.diff_note
    if delta_note:
        diff_note = (diff_note + "；" if diff_note else "") + delta_note

    new_draft = CalculationDraft(
        record_id=old_draft.record_id,
        formula_name=old_draft.formula_name,
        formula_expr=old_draft.formula_expr,
        draft_value=draft_value,
        computed_value=computed,
        is_supplement=old_draft.is_supplement,
        created_at=datetime.utcnow(),
        updated_at=datetime.utcnow(),
        version=new_version,
        prev_version_id=old_draft.id,
        diff_note=diff_note,
    )
    db.add(new_draft)
    db.commit()
    db.refresh(new_draft)
    return new_draft


@app.post("/api/batches/{batch_id}/boundary-append", summary="追加延迟到达的边界样例，返回受影响的结论")
async def append_boundary_samples(
    batch_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    batch = db.query(Batch).filter(Batch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")

    content = await file.read()
    filename = file.filename.lower()
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        df = pd.read_excel(io.BytesIO(content))
        rows = df.where(pd.notnull(df), None).to_dict(orient="records")
    elif filename.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(content))
        rows = df.where(pd.notnull(df), None).to_dict(orient="records")
    elif filename.endswith(".json"):
        rows = json.loads(content.decode("utf-8"))
        rows = rows if isinstance(rows, list) else [rows]
    else:
        raise HTTPException(status_code=400, detail="仅支持 xlsx / csv / json 格式")

    existing = db.query(Record).filter(Record.batch_id == batch_id).all()
    existing_by_name = defaultdict(list)
    for r in existing:
        existing_by_name[r.candidate_name].append(r)

    affected = []
    start_row = (batch.total_records or 0) + 1

    for idx, row in enumerate(rows):
        is_valid, block_reason, cleaned = validate_record(row, row_no=start_row + idx)
        cleaned["is_boundary"] = True

        initial_status = "blocked" if not is_valid else "pending_review"
        record = Record(
            batch_id=batch.id,
            row_no=cleaned["row_no"],
            candidate_name=cleaned["candidate_name"],
            value=cleaned["value"],
            unit=cleaned["unit"],
            raw_data=cleaned["raw_data"],
            is_valid=is_valid,
            block_reason=block_reason,
            status=initial_status,
            is_boundary=True,
            boundary_arrived_late=True,
            previous_conclusion="",
            affected_conclusions=[],
        )
        db.add(record)
        db.flush()
        _add_status_history(db, record.id, "", initial_status,
                            changed_by="system",
                            note=(block_reason + "；" if block_reason else "") + "延迟到达的边界样例")

        same_name_records = existing_by_name.get(cleaned["candidate_name"], [])
        for same in same_name_records:
            if same.status in ("confirmed", "reviewing"):
                old_conclusion = describe_status(same.status)
                affected_msgs = []
                if same.value is not None and cleaned["value"] is not None:
                    if cleaned["value"] < same.value:
                        affected_msgs.append(f"边界样例数值 {cleaned['value']} 低于已确认值 {same.value}，原排名/阈值结论可能失效")
                    else:
                        affected_msgs.append(f"边界样例数值 {cleaned['value']} 已补充，需重新复核 {same.candidate_name} 的候选结论")
                if not affected_msgs:
                    affected_msgs.append(f"新增同名边界样例，需重新评估 {same.candidate_name} 的结论")

                same.previous_conclusion = (same.previous_conclusion + "；" if same.previous_conclusion else "") + old_conclusion
                existing_affected = same.affected_conclusions or []
                existing_affected.extend(affected_msgs)
                same.affected_conclusions = list(dict.fromkeys(existing_affected))

                if same.status == "confirmed":
                    old_st = same.status
                    same.status = "reviewing"
                    _add_status_history(db, same.id, old_st, "reviewing",
                                        changed_by="system",
                                        note="；".join(affected_msgs))

                affected.append({
                    "record_id": same.id,
                    "candidate_name": same.candidate_name,
                    "from_status": old_conclusion,
                    "to_status": describe_status(same.status),
                    "affected_messages": affected_msgs,
                })

    batch.total_records = (batch.total_records or 0) + len(rows)
    db.commit()
    db.refresh(batch)

    return {
        "appended_count": len(rows),
        "affected_records": affected,
        "hint": "以上记录的结论可能因延迟到达的边界样例而受影响，系统未自动覆盖旧结果，已退回复核中。",
    }


@app.get("/api/batches/{batch_id}/issues", summary="查询批次异常清单（重点：单位缺失）")
def list_batch_issues(batch_id: int, db: Session = Depends(get_db)):
    batch = db.query(Batch).filter(Batch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")
    invalid_records = db.query(Record).filter(Record.batch_id == batch_id, Record.is_valid == False).all()
    grouped = defaultdict(list)
    for r in invalid_records:
        if "单位缺失" in r.block_reason:
            grouped["unit_missing"].append(r)
        if "数值缺失" in r.block_reason:
            grouped["value_missing"].append(r)
        if "数值格式异常" in r.block_reason:
            grouped["value_invalid"].append(r)
        if "重复记录" in r.block_reason:
            grouped["duplicate"].append(r)
    return {
        "batch_id": batch_id,
        "batch_name": batch.name,
        "summary": {
            "unit_missing_count": len(grouped.get("unit_missing", [])),
            "value_missing_count": len(grouped.get("value_missing", [])),
            "value_invalid_count": len(grouped.get("value_invalid", [])),
            "duplicate_count": len(grouped.get("duplicate", [])),
        },
        "unit_missing_records": [
            {"id": r.id, "row_no": r.row_no, "candidate_name": r.candidate_name,
             "value": r.value, "block_reason": r.block_reason}
            for r in grouped.get("unit_missing", [])
        ],
    }


@app.get("/api/batches/{batch_id}/export", summary="导出分析报告 Excel（学生视角：突出不可用记录及拦截原因）")
def export_report(batch_id: int, db: Session = Depends(get_db)):
    batch = db.query(Batch).filter(Batch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")

    records = db.query(Record).filter(Record.batch_id == batch_id).order_by(Record.row_no.asc()).all()

    wb = Workbook()

    ws_main = wb.active
    ws_main.title = "候选清单"

    headers = ["序号", "候选名称", "数值", "单位", "状态", "是否可用", "不可用原因（学生必读）",
               "是否边界样例", "延迟到达", "受影响结论提示"]
    ws_main.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws_main.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006", bold=True)
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    yellow_font = Font(color="9C5700")

    for r in records:
        row = [
            r.row_no,
            r.candidate_name,
            r.value if r.value is not None else "",
            r.unit if r.unit else "",
            describe_status(r.status),
            "可用" if r.is_valid else "不可用",
            format_block_reason_for_report(r.block_reason) if not r.is_valid else "",
            "是" if r.is_boundary else "否",
            "是" if r.boundary_arrived_late else "否",
            "\n".join(r.affected_conclusions or []),
        ]
        ws_main.append(row)
        current_row = ws_main.max_row
        if not r.is_valid:
            for col_idx in range(1, len(headers) + 1):
                ws_main.cell(row=current_row, column=col_idx).fill = red_fill
            ws_main.cell(row=current_row, column=6).font = red_font
            ws_main.cell(row=current_row, column=7).font = red_font
        elif r.boundary_arrived_late or r.affected_conclusions:
            for col_idx in range(1, len(headers) + 1):
                ws_main.cell(row=current_row, column=col_idx).fill = yellow_fill
            ws_main.cell(row=current_row, column=10).font = yellow_font

    col_widths = [8, 22, 12, 10, 12, 10, 60, 12, 10, 50]
    for i, w in enumerate(col_widths, start=1):
        ws_main.column_dimensions[get_column_letter(i)].width = w
    for row in ws_main.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws_invalid = wb.create_sheet("不可用记录详情")
    ws_invalid.append(["序号", "候选名称", "数值", "问题类型", "详细说明（为什么不能用）"])
    for col_idx in range(1, 6):
        cell = ws_invalid.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    invalid_count = 0
    for r in records:
        if not r.is_valid:
            invalid_count += 1
            reasons = []
            if "单位缺失" in r.block_reason:
                reasons.append("单位缺失")
            if "数值缺失" in r.block_reason:
                reasons.append("数值缺失")
            if "数值格式异常" in r.block_reason:
                reasons.append("数值格式异常")
            if "重复记录" in r.block_reason:
                reasons.append("重复记录")
            ws_invalid.append([
                r.row_no, r.candidate_name, r.value if r.value is not None else "",
                "、".join(reasons) if reasons else "其他",
                format_block_reason_for_report(r.block_reason),
            ])
    for i, w in enumerate([8, 22, 12, 18, 65], start=1):
        ws_invalid.column_dimensions[get_column_letter(i)].width = w
    for row in ws_invalid.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    ws_calc = wb.create_sheet("计算草稿历史")
    ws_calc.append(["候选名称", "公式名", "版本", "公式", "草稿值", "计算结果",
                    "是否补录", "相对上版差异", "说明"])
    for col_idx in range(1, 10):
        cell = ws_calc.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in records:
        drafts = (db.query(CalculationDraft)
                  .filter(CalculationDraft.record_id == r.id)
                  .order_by(CalculationDraft.formula_name.asc(), CalculationDraft.version.asc())
                  .all())
        for d in drafts:
            ws_calc.append([
                r.candidate_name, d.formula_name, f"v{d.version}", d.formula_expr,
                d.draft_value if d.draft_value is not None else "",
                d.computed_value if d.computed_value is not None else "",
                "是" if d.is_supplement else "否",
                d.diff_note,
                "",
            ])
    for i, w in enumerate([22, 18, 8, 30, 14, 14, 10, 28, 20], start=1):
        ws_calc.column_dimensions[get_column_letter(i)].width = w
    for row in ws_calc.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws_summary = wb.create_sheet("报告摘要", 0)
    ws_summary.append(["数独候选削减器 - 分析报告"])
    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"].font = Font(bold=True, size=14)
    ws_summary.append([])
    ws_summary.append(["批次名称", batch.name])
    ws_summary.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws_summary.append(["总记录数", batch.total_records])
    ws_summary.append(["可用记录数", batch.valid_records])
    ws_summary.append(["不可用记录数", invalid_count])
    ws_summary.append([])
    ws_summary.append(["重点关注（给学生）："])
    ws_summary.append(["", f"1. 共有 {invalid_count} 条记录被标记为不可用，请查看「不可用记录详情」页。"])
    ws_summary.append(["", "2. 单位缺失的记录因口径不一无法汇总，已自动拦截，补录单位后可重开。"])
    ws_summary.append(["", "3. 黄色高亮行表示受边界样例影响，结论需重新复核。"])
    ws_summary.column_dimensions["A"].width = 20
    ws_summary.column_dimensions["B"].width = 70

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"sudoku_report_batch{batch_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.get("/", summary="健康检查")
def root():
    return {"service": "数独候选削减器", "version": "1.0.0",
            "docs": "/docs",
            "hint": "启动后访问 /docs 查看所有接口；使用说明见同目录的 使用说明.md"}
