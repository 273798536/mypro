import os
import pandas as pd
import plotly.graph_objects as go
from typing import Dict, List, Optional
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    PageBreak, Image as RLImage
)
from .models import (
    ReagentBatch, Sample, SequencingResult, PedigreeMember,
    ValidationResult, ConflictRecord, ImportRecord, ValidationStatus
)
from .dashboard import DashboardData
from .database import db

class ReportExporter:
    def __init__(self, db_session=None, export_dir: str = None):
        self.db = db_session or db.session
        self.export_dir = export_dir
        self.dashboard = DashboardData(db_session)

        self.color_map = {
            'pass': '22c55e',
            'fail': 'ef4444',
            'review': 'f59e0b',
            'warning': 'f59e0b',
            'error': 'ef4444',
            'critical': 'dc2626',
            'info': '3b82f6',
            'header': '1e3a5f',
            'ready': '22c55e',
            'not_ready': 'ef4444'
        }

    def export_to_excel(self, batch_id: int, output_path: str = None) -> Dict:
        batch = ReagentBatch.query.get(batch_id)
        if not batch:
            return {'success': False, 'error': '批次不存在'}

        if not output_path and self.export_dir:
            output_path = os.path.join(
                self.export_dir,
                f'pedigree_report_{batch.batch_number}_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
            )

        details = self.dashboard.get_batch_details(batch_id)
        usability = self.dashboard.get_usability_assessment(batch_id)
        validation = self.dashboard.get_validation_charts(batch_id)

        wb = Workbook()

        self._create_summary_sheet(wb, batch, usability, validation, details)
        self._create_samples_sheet(wb, details['samples'])
        self._create_sequencing_sheet(wb, details['sequencing'])
        self._create_pedigree_sheet(wb, details['pedigree'])
        self._create_conflicts_sheet(wb, details['conflicts'])
        self._create_validations_sheet(wb, details['validations'])
        self._create_imports_sheet(wb, details['imports'])
        self._create_usability_sheet(wb, usability)

        wb.save(output_path)

        return {
            'success': True,
            'output_path': output_path,
            'file_name': os.path.basename(output_path),
            'file_size': os.path.getsize(output_path)
        }

    def _create_summary_sheet(self, wb, batch, usability, validation, details):
        ws = wb.active
        ws.title = '汇总'

        header_font = Font(bold=True, color='FFFFFF', size=14)
        header_fill = PatternFill(start_color=self.color_map['header'],
                                  end_color=self.color_map['header'],
                                  fill_type='solid')
        label_font = Font(bold=True)
        warning_fill = PatternFill(start_color='fff3cd', end_color='fff3cd', fill_type='solid')

        row = 1
        ws.cell(row=row, column=1, value=f'遗传家系谱系校验报告 - {batch.batch_number}')
        ws.cell(row=row, column=1).font = header_font
        ws.cell(row=row, column=1).fill = header_fill
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)

        row += 2
        assessment_label = usability['assessment_label']
        ws.cell(row=row, column=1, value='可用性评估')
        ws.cell(row=row, column=1).font = label_font
        ws.cell(row=row, column=2, value=assessment_label)
        if usability['assessment'] == 'ready':
            ws.cell(row=row, column=2).font = Font(bold=True, color=self.color_map['pass'])
        elif usability['assessment'] == 'review':
            ws.cell(row=row, column=2).font = Font(bold=True, color=self.color_map['warning'])
        else:
            ws.cell(row=row, column=2).font = Font(bold=True, color=self.color_map['fail'])

        row += 1
        ws.cell(row=row, column=1, value='主要结论')
        ws.cell(row=row, column=1).font = label_font
        ws.cell(row=row, column=2, value=usability['primary_issue'])

        row += 2
        ws.cell(row=row, column=1, value='批次信息')
        ws.cell(row=row, column=1).font = label_font

        row += 1
        ws.cell(row=row, column=1, value='批次号')
        ws.cell(row=row, column=2, value=batch.batch_number)
        row += 1
        ws.cell(row=row, column=1, value='名称')
        ws.cell(row=row, column=2, value=batch.name or '')
        row += 1
        ws.cell(row=row, column=1, value='创建时间')
        ws.cell(row=row, column=2, value=batch.created_at.strftime('%Y-%m-%d %H:%M:%S') if batch.created_at else '')
        row += 1
        ws.cell(row=row, column=1, value='源文件')
        ws.cell(row=row, column=2, value=batch.source_file or '')

        row += 2
        ws.cell(row=row, column=1, value='数据统计')
        ws.cell(row=row, column=1).font = label_font

        summary_data = details['summary']
        row += 1
        ws.cell(row=row, column=1, value='样本数')
        ws.cell(row=row, column=2, value=summary_data['sample_count'])
        row += 1
        ws.cell(row=row, column=1, value='测序结果数')
        ws.cell(row=row, column=2, value=summary_data['sequencing_count'])
        row += 1
        ws.cell(row=row, column=1, value='家系成员数')
        ws.cell(row=row, column=2, value=summary_data['pedigree_count'])

        row += 2
        ws.cell(row=row, column=1, value='校验结果')
        ws.cell(row=row, column=1).font = label_font

        status_data = validation.get('status_data', {})
        row += 1
        ws.cell(row=row, column=1, value='通过')
        ws.cell(row=row, column=2, value=status_data.get('pass', 0))
        row += 1
        ws.cell(row=row, column=1, value='失败')
        ws.cell(row=row, column=2, value=status_data.get('fail', 0))
        if status_data.get('fail', 0) > 0:
            ws.cell(row=row, column=2).fill = warning_fill
        row += 1
        ws.cell(row=row, column=1, value='需复核')
        ws.cell(row=row, column=2, value=status_data.get('review', 0))
        if status_data.get('review', 0) > 0:
            ws.cell(row=row, column=2).fill = warning_fill

        row += 2
        ws.cell(row=row, column=1, value='冲突记录')
        ws.cell(row=row, column=1).font = label_font

        row += 1
        ws.cell(row=row, column=1, value='总数')
        ws.cell(row=row, column=2, value=summary_data['conflict_count'])
        row += 1
        ws.cell(row=row, column=1, value='严重冲突')
        ws.cell(row=row, column=2, value=summary_data['critical_conflicts'])
        if summary_data['critical_conflicts'] > 0:
            ws.cell(row=row, column=2).fill = warning_fill

        for col in range(1, 7):
            ws.column_dimensions[get_column_letter(col)].width = 25

    def _create_samples_sheet(self, wb, samples):
        ws = wb.create_sheet('样本清单')
        headers = ['样本ID', '姓名', '性别', '样本类型', '原始行号', '来源文件', '来源Sheet']
        self._write_table_header(ws, headers)

        for i, sample in enumerate(samples, start=2):
            ws.cell(row=i, column=1, value=sample['sample_id'])
            ws.cell(row=i, column=2, value=sample['name'] or '')
            ws.cell(row=i, column=3, value=sample['gender'] or '')
            ws.cell(row=i, column=4, value=sample['sample_type'] or '')
            ws.cell(row=i, column=5, value=sample['original_row_number'] or '')
            ws.cell(row=i, column=6, value=sample['source_file'] or '')
            ws.cell(row=i, column=7, value=sample['source_sheet'] or '')

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

    def _create_sequencing_sheet(self, wb, sequencing):
        ws = wb.create_sheet('测序结果')
        headers = ['样本ID', '测序ID', '基因', '变异', '基因型', '质量值', '深度', '原始行号', '来源文件']
        self._write_table_header(ws, headers)

        for i, seq in enumerate(sequencing, start=2):
            ws.cell(row=i, column=1, value=seq['sample_id'])
            ws.cell(row=i, column=2, value=seq['sequencing_id'] or '')
            ws.cell(row=i, column=3, value=seq['gene'] or '')
            ws.cell(row=i, column=4, value=seq['variant'] or '')
            ws.cell(row=i, column=5, value=seq['genotype'] or '')
            ws.cell(row=i, column=6, value=seq['quality_score'] or '')
            ws.cell(row=i, column=7, value=seq['depth'] or '')
            ws.cell(row=i, column=8, value=seq['original_row_number'] or '')
            ws.cell(row=i, column=9, value=seq['source_file'] or '')

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    def _create_pedigree_sheet(self, wb, pedigree):
        ws = wb.create_sheet('家系信息')
        headers = ['家系ID', '个体ID', '父亲ID', '母亲ID', '性别', '患病状态', '关系', '代次', '原始行号', '来源文件']
        self._write_table_header(ws, headers)

        for i, p in enumerate(pedigree, start=2):
            ws.cell(row=i, column=1, value=p['family_id'] or '')
            ws.cell(row=i, column=2, value=p['individual_id'])
            ws.cell(row=i, column=3, value=p['father_id'] or '')
            ws.cell(row=i, column=4, value=p['mother_id'] or '')
            ws.cell(row=i, column=5, value=p['gender'] or '')
            ws.cell(row=i, column=6, value=p['affection_status'] or '')
            ws.cell(row=i, column=7, value=p['relationship'] or '')
            ws.cell(row=i, column=8, value=p['generation'] or '')
            ws.cell(row=i, column=9, value=p['original_row_number'] or '')
            ws.cell(row=i, column=10, value=p['source_file'] or '')

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    def _create_conflicts_sheet(self, wb, conflicts):
        ws = wb.create_sheet('冲突记录')
        headers = ['ID', '类型', '严重程度', '优先级', '消息', '期望值', '实际值', '来源记录', '是否严重']
        self._write_table_header(ws, headers)

        for i, c in enumerate(conflicts, start=2):
            ws.cell(row=i, column=1, value=c['id'])
            ws.cell(row=i, column=2, value=c['conflict_type'])
            ws.cell(row=i, column=3, value=c['severity'])
            ws.cell(row=i, column=4, value=c['priority'])
            ws.cell(row=i, column=5, value=c['message'])
            ws.cell(row=i, column=6, value=c['expected_value'] or '')
            ws.cell(row=i, column=7, value=c['actual_value'] or '')
            ws.cell(row=i, column=8, value=str(c['source_records']) if c['source_records'] else '')
            ws.cell(row=i, column=9, value='是' if c['is_critical'] else '否')

            if c['is_critical']:
                for col in range(1, len(headers) + 1):
                    ws.cell(row=i, column=col).fill = PatternFill(
                        start_color='ffebee', end_color='ffebee', fill_type='solid'
                    )

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    def _create_validations_sheet(self, wb, validations):
        ws = wb.create_sheet('校验结果')
        headers = ['ID', '校验类型', '状态', '严重程度', '消息', '期望值', '实际值', '来源记录']
        self._write_table_header(ws, headers)

        for i, v in enumerate(validations, start=2):
            ws.cell(row=i, column=1, value=v['id'])
            ws.cell(row=i, column=2, value=v['validation_type'])
            ws.cell(row=i, column=3, value=v['status'])
            ws.cell(row=i, column=4, value=v['severity'] or '')
            ws.cell(row=i, column=5, value=v['message'])
            ws.cell(row=i, column=6, value=v['expected_value'] or '')
            ws.cell(row=i, column=7, value=v['actual_value'] or '')
            ws.cell(row=i, column=8, value=str(v['source_records']) if v['source_records'] else '')

            if v['status'] == ValidationStatus.FAIL:
                fill_color = 'ffebee'
            elif v['status'] == ValidationStatus.REVIEW:
                fill_color = 'fff3cd'
            else:
                fill_color = None

            if fill_color:
                for col in range(1, len(headers) + 1):
                    ws.cell(row=i, column=col).fill = PatternFill(
                        start_color=fill_color, end_color=fill_color, fill_type='solid'
                    )

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

    def _create_imports_sheet(self, wb, imports):
        ws = wb.create_sheet('导入记录')
        headers = ['导入ID', '文件名', '类型', '导入时间', '总行数', '已导入', '跳过', '重复', '状态']
        self._write_table_header(ws, headers)

        for i, imp in enumerate(imports, start=2):
            ws.cell(row=i, column=1, value=imp['id'])
            ws.cell(row=i, column=2, value=imp['file_name'])
            ws.cell(row=i, column=3, value=imp['import_type'])
            ws.cell(row=i, column=4, value=imp['import_time'])
            ws.cell(row=i, column=5, value=imp['total_rows'])
            ws.cell(row=i, column=6, value=imp['imported_rows'])
            ws.cell(row=i, column=7, value=imp['skipped_rows'])
            ws.cell(row=i, column=8, value=imp['duplicate_rows'])
            ws.cell(row=i, column=9, value=imp['status'])

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

    def _create_usability_sheet(self, wb, usability):
        ws = wb.create_sheet('可用性分类')

        header_font = Font(bold=True, color='FFFFFF', size=12)
        header_fill = PatternFill(start_color=self.color_map['header'],
                                  end_color=self.color_map['header'],
                                  fill_type='solid')

        row = 1
        ws.cell(row=row, column=1, value='样本可用性分类')
        ws.cell(row=row, column=1).font = header_font
        ws.cell(row=row, column=1).fill = header_fill
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)

        section_font = Font(bold=True, size=11)

        details = usability['details']

        row += 2
        ws.cell(row=row, column=1, value=f'✅ 可以直接使用 ({len(details["ready_samples"])} 个样本)')
        ws.cell(row=row, column=1).font = section_font
        ws.cell(row=row, column=1).font.color = self.color_map['pass']

        row += 1
        if details['ready_samples']:
            headers = ['样本ID', '姓名', '性别', '来源']
            self._write_table_header(ws, headers, start_row=row)
            for i, s in enumerate(details['ready_samples'], start=row + 1):
                ws.cell(row=i, column=1, value=s['sample_id'])
                ws.cell(row=i, column=2, value=s['name'] or '')
                ws.cell(row=i, column=3, value=s['gender'] or '')
                ws.cell(row=i, column=4, value=s['source'])
            row += len(details['ready_samples']) + 2
        else:
            ws.cell(row=row, column=1, value='无')
            row += 2

        ws.cell(row=row, column=1, value=f'⚠️ 需要复核 ({len(details["needs_review_samples"])} 个样本)')
        ws.cell(row=row, column=1).font = section_font
        ws.cell(row=row, column=1).font.color = self.color_map['warning']

        row += 1
        if details['needs_review_samples']:
            headers = ['样本ID', '姓名', '性别', '来源', '问题']
            self._write_table_header(ws, headers, start_row=row)
            for i, s in enumerate(details['needs_review_samples'], start=row + 1):
                ws.cell(row=i, column=1, value=s['sample_id'])
                ws.cell(row=i, column=2, value=s['name'] or '')
                ws.cell(row=i, column=3, value=s['gender'] or '')
                ws.cell(row=i, column=4, value=s['source'])
                ws.cell(row=i, column=5, value='; '.join(s['issues']))
            row += len(details['needs_review_samples']) + 2
        else:
            ws.cell(row=row, column=1, value='无')
            row += 2

        ws.cell(row=row, column=1, value=f'❌ 不能使用 ({len(details["not_usable_samples"])} 个样本)')
        ws.cell(row=row, column=1).font = section_font
        ws.cell(row=row, column=1).font.color = self.color_map['fail']

        row += 1
        if details['not_usable_samples']:
            headers = ['样本ID', '姓名', '性别', '来源', '问题']
            self._write_table_header(ws, headers, start_row=row)
            for i, s in enumerate(details['not_usable_samples'], start=row + 1):
                ws.cell(row=i, column=1, value=s['sample_id'])
                ws.cell(row=i, column=2, value=s['name'] or '')
                ws.cell(row=i, column=3, value=s['gender'] or '')
                ws.cell(row=i, column=4, value=s['source'])
                ws.cell(row=i, column=5, value='; '.join(s['issues']))
                for col in range(1, len(headers) + 1):
                    ws.cell(row=i, column=col).fill = PatternFill(
                        start_color='ffebee', end_color='ffebee', fill_type='solid'
                    )
        else:
            ws.cell(row=row, column=1, value='无')

        for col in range(1, 7):
            ws.column_dimensions[get_column_letter(col)].width = 25

    def _write_table_header(self, ws, headers, start_row=1):
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color=self.color_map['header'],
                                  end_color=self.color_map['header'],
                                  fill_type='solid')
        for i, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=i, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

    def export_to_pdf(self, batch_id: int, output_path: str = None,
                      annotated_images: List[str] = None) -> Dict:
        batch = ReagentBatch.query.get(batch_id)
        if not batch:
            return {'success': False, 'error': '批次不存在'}

        if not output_path and self.export_dir:
            output_path = os.path.join(
                self.export_dir,
                f'pedigree_report_{batch.batch_number}_{datetime.now().strftime("%Y%m%d%H%M%S")}.pdf'
            )

        details = self.dashboard.get_batch_details(batch_id)
        usability = self.dashboard.get_usability_assessment(batch_id)

        doc = SimpleDocTemplate(output_path, pagesize=A4,
                               leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                               topMargin=0.75 * inch, bottomMargin=0.75 * inch)

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'],
                                    fontSize=18, spaceAfter=20, textColor=self._get_rgb_color('header'))
        section_style = ParagraphStyle('SectionTitle', parent=styles['Heading2'],
                                      fontSize=14, spaceAfter=10, textColor=self._get_rgb_color('header'))
        normal_style = styles['Normal']

        story = []

        story.append(Paragraph(f'遗传家系谱系校验报告', title_style))
        story.append(Paragraph(f'批次: {batch.batch_number}', styles['Heading2']))
        story.append(Paragraph(f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', normal_style))
        story.append(Spacer(1, 0.3 * inch))

        assessment_color = {
            'ready': colors.green,
            'review': colors.orange,
            'not_ready': colors.red
        }[usability['assessment']]

        assessment_text = f'<font size="14" color="{assessment_color.hexval()}"><b>{usability["assessment_label"]}</b></font>'
        story.append(Paragraph(assessment_text, normal_style))
        story.append(Paragraph(f'<i>{usability["primary_issue"]}</i>', normal_style))
        story.append(Spacer(1, 0.3 * inch))

        story.append(Paragraph('数据统计', section_style))
        summary_data = [
            ['项目', '数量'],
            ['样本数', str(details['summary']['sample_count'])],
            ['测序结果数', str(details['summary']['sequencing_count'])],
            ['家系成员数', str(details['summary']['pedigree_count'])],
            ['冲突数', str(details['summary']['conflict_count'])],
            ['严重冲突数', str(details['summary']['critical_conflicts'])],
        ]
        t = Table(summary_data, colWidths=[2 * inch, 1.5 * inch])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), self._get_rgb_color('header')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey])
        ]))
        story.append(t)
        story.append(Spacer(1, 0.3 * inch))

        details_us = usability['details']

        if details_us['ready_samples']:
            story.append(Paragraph(f'✅ 可以直接使用 ({len(details_us["ready_samples"])} 个)', section_style))
            ready_data = [['样本ID', '姓名', '性别', '来源']]
            for s in details_us['ready_samples'][:20]:
                ready_data.append([s['sample_id'], s['name'] or '', s['gender'] or '', s['source']])
            t = Table(ready_data, colWidths=[1 * inch, 1 * inch, 0.6 * inch, 2 * inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), self._get_rgb_color('pass')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.2 * inch))

        if details_us['needs_review_samples']:
            story.append(Paragraph(f'⚠️ 需要复核 ({len(details_us["needs_review_samples"])} 个)', section_style))
            review_data = [['样本ID', '姓名', '性别', '问题']]
            for s in details_us['needs_review_samples'][:10]:
                review_data.append([s['sample_id'], s['name'] or '', s['gender'] or '',
                                   '; '.join(s['issues'])[:50]])
            t = Table(review_data, colWidths=[0.8 * inch, 0.8 * inch, 0.5 * inch, 2.5 * inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), self._get_rgb_color('warning')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.2 * inch))

        if details_us['not_usable_samples']:
            story.append(Paragraph(f'❌ 不能使用 ({len(details_us["not_usable_samples"])} 个)', section_style))
            unusable_data = [['样本ID', '姓名', '性别', '问题']]
            for s in details_us['not_usable_samples'][:10]:
                unusable_data.append([s['sample_id'], s['name'] or '', s['gender'] or '',
                                     '; '.join(s['issues'])[:50]])
            t = Table(unusable_data, colWidths=[0.8 * inch, 0.8 * inch, 0.5 * inch, 2.5 * inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), self._get_rgb_color('fail')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#ffebee')),
            ]))
            story.append(t)

        if annotated_images:
            story.append(PageBreak())
            story.append(Paragraph('标注图片', section_style))
            for img_path in annotated_images[:3]:
                if os.path.exists(img_path):
                    try:
                        story.append(RLImage(img_path, width=6 * inch, height=4 * inch))
                        story.append(Spacer(1, 0.2 * inch))
                    except:
                        pass

        story.append(PageBreak())
        story.append(Paragraph('冲突记录详情', section_style))

        if details['conflicts']:
            conflict_data = [['类型', '严重程度', '消息', '来源']]
            for c in details['conflicts'][:20]:
                conflict_data.append([
                    c['conflict_type'],
                    c['severity'],
                    Paragraph(c['message'][:100], normal_style),
                    str(c['source_records'])[:80] if c['source_records'] else ''
                ])
            t = Table(conflict_data, colWidths=[1 * inch, 0.6 * inch, 2.5 * inch, 1.5 * inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), self._get_rgb_color('header')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 7),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ]))
            story.append(t)

        doc.build(story)

        return {
            'success': True,
            'output_path': output_path,
            'file_name': os.path.basename(output_path),
            'file_size': os.path.getsize(output_path)
        }

    def _get_rgb_color(self, color_name: str):
        hex_color = self.color_map.get(color_name, '3b82f6')
        return colors.HexColor(hex_color)

    def export_chart_image(self, figure_dict: Dict, output_path: str) -> Optional[str]:
        try:
            import plotly.io as pio
            fig = go.Figure(figure_dict)
            pio.write_image(fig, output_path, width=800, height=500, scale=2)
            return output_path
        except Exception as e:
            print(f"导出图表失败: {e}")
            return None

    def export_all(self, batch_id: int, include_images: bool = False) -> Dict:
        results = {}

        excel_result = self.export_to_excel(batch_id)
        results['excel'] = excel_result

        annotated_images = []
        if include_images:
            from .image_annotator import ImageAnnotator
            annotator = ImageAnnotator(self.db, self.export_dir)
            images = annotator.get_annotated_images(batch_id)
            for img_info in images:
                img_path = os.path.join(self.export_dir, img_info['image_file'])
                annotated = annotator.render_annotated_image(img_path)
                if annotated:
                    annotated_images.append(annotated)

        pdf_result = self.export_to_pdf(batch_id, annotated_images=annotated_images)
        results['pdf'] = pdf_result

        return {
            'success': True,
            'batch_id': batch_id,
            'exports': results,
            'timestamp': datetime.now().isoformat()
        }
