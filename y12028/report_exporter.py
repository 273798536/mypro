from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from models import (
    SessionSplit, SplitDetail, GameSession, Order, DataConflict,
    AnomalyRecord, DM, Script, ScriptAuthorization, CouponVerification
)


class ReportExporter:
    def __init__(self):
        self.header_font = Font(bold=True, color='FFFFFF', size=11)
        self.header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        self.warning_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
        self.error_fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
        self.success_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        self.center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def export_monthly_report(self, start_date=None, end_date=None):
        wb = Workbook()
        wb.remove(wb.active)

        summary_ws = wb.create_sheet('月度汇总', 0)
        self._write_summary_sheet(summary_ws, start_date, end_date)

        split_ws = wb.create_sheet('场次分账明细', 1)
        self._write_split_sheet(split_ws, start_date, end_date)

        conflict_ws = wb.create_sheet('数据冲突记录', 2)
        self._write_conflict_sheet(conflict_ws, start_date, end_date)

        anomaly_ws = wb.create_sheet('异常记录', 3)
        self._write_anomaly_sheet(anomaly_ws, start_date, end_date)

        dm_ws = wb.create_sheet('DM分账汇总', 4)
        self._write_dm_summary_sheet(dm_ws, start_date, end_date)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def _write_summary_sheet(self, ws, start_date, end_date):
        ws['A1'] = f'剧本杀门店分账月度报告'
        ws['A1'].font = Font(bold=True, size=16, color='4472C4')
        ws.merge_cells('A1:F1')
        ws['A1'].alignment = self.center_align

        ws['A2'] = f'统计期间: {start_date or "不限"} 至 {end_date or "不限"}'
        ws['A2'].font = Font(italic=True)
        ws.merge_cells('A2:F2')

        ws['A3'] = f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
        ws['A3'].font = Font(italic=True)
        ws.merge_cells('A3:F3')

        splits = self._get_splits(start_date, end_date)
        conflicts = DataConflict.query.all()
        anomalies = AnomalyRecord.query.filter_by(status='open').all()

        total_revenue = sum(s.total_revenue for s in splits)
        total_dm_fee = sum(s.dm_fee for s in splits)
        total_auth_fee = sum(s.authorization_fee for s in splits)
        total_store_share = sum(s.store_share for s in splits)
        total_coupon_discount = sum(s.total_coupon_discount for s in splits)

        summary_data = [
            ['指标', '数值', '说明'],
            ['场次总数', len(splits), '已完成分账的场次数量'],
            ['总营收（实收）', total_revenue, '所有订单实际收款金额'],
            ['优惠券优惠总额', total_coupon_discount, '使用优惠券减免的总金额'],
            ['DM分账总额', total_dm_fee, '所有DM应得报酬'],
            ['剧本授权费总额', total_auth_fee, '支付给剧本授权方的费用'],
            ['门店净收入', total_store_share, '营收扣除DM费和授权费'],
            ['待处理数据冲突', len([c for c in conflicts if c.status == 'pending']), '需要人工确认的数据不一致'],
            ['待处理异常', len(anomalies), '需要处理的异常记录'],
        ]

        row = 5
        for i, row_data in enumerate(summary_data):
            for j, value in enumerate(row_data):
                cell = ws.cell(row=row + i, column=j + 1, value=value)
                cell.border = self.thin_border
                if i == 0:
                    cell.font = self.header_font
                    cell.fill = self.header_fill
                    cell.alignment = self.center_align
                else:
                    cell.alignment = self.left_align if j == 2 else self.center_align
                    if j == 1 and isinstance(value, (int, float)) and '总额' in summary_data[i][0]:
                        cell.number_format = '¥#,##0.00'

        for col in ['A', 'B', 'C']:
            ws.column_dimensions[col].width = 25
        ws.row_dimensions[1].height = 30

    def _write_split_sheet(self, ws, start_date, end_date):
        ws['A1'] = '场次分账明细'
        ws['A1'].font = Font(bold=True, size=14, color='4472C4')
        ws.merge_cells('A1:L1')
        ws['A1'].alignment = self.center_align

        headers = [
            '场次编号', '日期', '剧本名称', '实际DM',
            '订单数', '总营收', '优惠券优惠', '门店分成',
            'DM费用', '授权费', '分账备注', '分账时间'
        ]

        row = 3
        for j, header in enumerate(headers):
            cell = ws.cell(row=row, column=j + 1, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        splits = self._get_splits(start_date, end_date)
        for i, split in enumerate(splits):
            session = split.session
            row_data = [
                session.session_no,
                session.session_date.strftime('%Y-%m-%d'),
                session.script.name if session.script else '',
                session.actual_dm.name if session.actual_dm else (session.scheduled_dm.name if session.scheduled_dm else ''),
                len(session.orders),
                split.total_revenue,
                split.total_coupon_discount,
                split.store_share,
                split.dm_fee,
                split.authorization_fee,
                split.remark,
                split.calculated_at.strftime('%Y-%m-%d %H:%M')
            ]
            for j, value in enumerate(row_data):
                cell = ws.cell(row=row + 1 + i, column=j + 1, value=value)
                cell.border = self.thin_border
                if j in [5, 6, 7, 8, 9]:
                    cell.number_format = '¥#,##0.00'
                    cell.alignment = self.center_align
                else:
                    cell.alignment = self.center_align

                if split.remark and 'DM代班' in split.remark:
                    cell.fill = self.warning_fill

        widths = [15, 12, 18, 10, 8, 12, 12, 12, 12, 12, 25, 18]
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w

    def _write_conflict_sheet(self, ws, start_date, end_date):
        ws['A1'] = '数据冲突记录'
        ws['A1'].font = Font(bold=True, size=14, color='4472C4')
        ws.merge_cells('A1:I1')
        ws['A1'].alignment = self.center_align

        type_names = {
            'dm_mismatch': 'DM不匹配',
            'player_count_mismatch': '人数不匹配',
            'authorization_mismatch': '授权不匹配'
        }

        headers = [
            '冲突类型', '关联场次', '冲突字段',
            'A方记录', 'A方维护人', 'A方数值',
            'B方记录', 'B方维护人', 'B方数值', '状态'
        ]

        row = 3
        for j, header in enumerate(headers):
            cell = ws.cell(row=row, column=j + 1, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        conflicts = DataConflict.query.all()
        for i, conflict in enumerate(conflicts):
            session = GameSession.query.get(conflict.session_id)
            row_data = [
                type_names.get(conflict.conflict_type, conflict.conflict_type),
                session.session_no if session else '',
                conflict.field_name,
                conflict.record_a_type,
                conflict.record_a_maintainer,
                conflict.record_a_value,
                conflict.record_b_type,
                conflict.record_b_maintainer,
                conflict.record_b_value,
                '待处理' if conflict.status == 'pending' else '已解决'
            ]
            for j, value in enumerate(row_data):
                cell = ws.cell(row=row + 1 + i, column=j + 1, value=value)
                cell.border = self.thin_border
                cell.alignment = self.center_align
                if conflict.status == 'pending':
                    cell.fill = self.warning_fill

        widths = [12, 15, 12, 15, 15, 18, 15, 15, 18, 10]
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w

    def _write_anomaly_sheet(self, ws, start_date, end_date):
        ws['A1'] = '异常记录（人话版）'
        ws['A1'].font = Font(bold=True, size=14, color='4472C4')
        ws.merge_cells('A1:G1')
        ws['A1'].alignment = self.center_align

        type_names = {
            'duplicate_coupon': '券重复核销',
            'dm_substitution': 'DM代班',
            'missing_authorization': '缺少授权',
            'unauthorized_fee': '多扣授权费',
            'amount_mismatch': '金额不匹配'
        }

        severity_names = {
            'error': '严重',
            'warning': '警告',
            'info': '提示'
        }

        headers = [
            '异常类型', '严重程度', '关联场次',
            '系统描述', '人话解释', '状态', '检测时间'
        ]

        row = 3
        for j, header in enumerate(headers):
            cell = ws.cell(row=row, column=j + 1, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        anomalies = AnomalyRecord.query.all()
        for i, anomaly in enumerate(anomalies):
            session = GameSession.query.get(anomaly.session_id)
            row_data = [
                type_names.get(anomaly.anomaly_type, anomaly.anomaly_type),
                severity_names.get(anomaly.severity, anomaly.severity),
                session.session_no if session else '',
                anomaly.description,
                anomaly.plain_explanation,
                '待处理' if anomaly.status == 'open' else '已处理',
                anomaly.detected_at.strftime('%Y-%m-%d %H:%M')
            ]
            for j, value in enumerate(row_data):
                cell = ws.cell(row=row + 1 + i, column=j + 1, value=value)
                cell.border = self.thin_border
                cell.alignment = self.left_align if j in [3, 4] else self.center_align

                if anomaly.severity == 'error' and anomaly.status == 'open':
                    cell.fill = self.error_fill
                elif anomaly.severity == 'warning' and anomaly.status == 'open':
                    cell.fill = self.warning_fill
                elif anomaly.status == 'resolved':
                    cell.fill = self.success_fill

        widths = [14, 10, 15, 40, 50, 10, 18]
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w

    def _write_dm_summary_sheet(self, ws, start_date, end_date):
        ws['A1'] = 'DM分账汇总'
        ws['A1'].font = Font(bold=True, size=14, color='4472C4')
        ws.merge_cells('A1:F1')
        ws['A1'].alignment = self.center_align

        headers = [
            'DM姓名', '带场次数', '代班次数', '基础费合计', '分成合计', '总报酬'
        ]

        row = 3
        for j, header in enumerate(headers):
            cell = ws.cell(row=row, column=j + 1, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        splits = self._get_splits(start_date, end_date)
        dm_stats = {}

        for split in splits:
            session = split.session
            dm_id = session.actual_dm_id or session.scheduled_dm_id
            dm = DM.query.get(dm_id)
            if not dm:
                continue

            if dm_id not in dm_stats:
                dm_stats[dm_id] = {
                    'name': dm.name,
                    'session_count': 0,
                    'substitute_count': 0,
                    'total_fee': 0.0
                }

            dm_stats[dm_id]['session_count'] += 1
            if session.actual_dm_id and session.actual_dm_id != session.scheduled_dm_id:
                dm_stats[dm_id]['substitute_count'] += 1

            for detail in split.details:
                if detail.split_type == 'dm_fee' and detail.recipient_id == dm_id:
                    dm_stats[dm_id]['total_fee'] += detail.amount

        for i, (dm_id, stats) in enumerate(dm_stats.items()):
            dm = DM.query.get(dm_id)
            base_fee_total = stats['session_count'] * (dm.base_fee if dm else 0)
            split_total = stats['total_fee'] - base_fee_total

            row_data = [
                stats['name'],
                stats['session_count'],
                stats['substitute_count'],
                base_fee_total,
                max(0, split_total),
                stats['total_fee']
            ]
            for j, value in enumerate(row_data):
                cell = ws.cell(row=row + 1 + i, column=j + 1, value=value)
                cell.border = self.thin_border
                cell.alignment = self.center_align
                if j in [3, 4, 5]:
                    cell.number_format = '¥#,##0.00'

        widths = [12, 10, 10, 14, 14, 14]
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w

    def _get_splits(self, start_date, end_date):
        query = SessionSplit.query
        if start_date:
            query = query.filter(SessionSplit.split_date >= start_date)
        if end_date:
            query = query.filter(SessionSplit.split_date <= end_date)
        return query.order_by(SessionSplit.split_date).all()
