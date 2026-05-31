from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from cash_shortage.config import FIELD_SPECS
from cash_shortage.checks import AuditResult, AuditFinding


@dataclass
class ReportContext:
    report_date: date
    filter_description: str = ""
    prepared_by: str = "连锁门店稽核"
    output_dir: str = "output"


class ReportExporter:
    def __init__(self, ctx: ReportContext):
        self.ctx = ctx
        self._header_font = Font(bold=True, size=11, color="FFFFFF")
        self._header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self._title_font = Font(bold=True, size=14)
        self._sub_title_font = Font(bold=True, size=11)
        self._wrap = Alignment(wrap_text=True, vertical="center")
        self._thin = Side(border_style="thin", color="BFBFBF")
        self._border = Border(left=self._thin, right=self._thin, top=self._thin, bottom=self._thin)

    def export(
        self,
        audit_result: AuditResult,
        filtered_df: pd.DataFrame,
        raw_dataset=None,
        filename: Optional[str] = None,
    ) -> str:
        Path(self.ctx.output_dir).mkdir(parents=True, exist_ok=True)

        if not filename:
            filename = f"门店现金短款调查报告_{self.ctx.report_date.isoformat()}.xlsx"
        out_path = str(Path(self.ctx.output_dir) / filename)

        wb = Workbook()

        self._write_summary(wb, audit_result)
        self._write_findings_by_type(wb, audit_result)
        self._write_all_findings_detail(wb, audit_result)
        self._write_filtered_data(wb, filtered_df)
        if raw_dataset is not None and len(raw_dataset.raw_df.columns) > 0:
            self._write_raw_data(wb, raw_dataset.raw_df)
        self._write_caliber_explanation(wb)

        wb.save(out_path)
        return out_path

    def _style_header_row(self, ws, row: int, cols: int):
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = self._header_font
            cell.fill = self._header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._border

    def _style_data_area(self, ws, start_row: int, end_row: int, cols: int):
        for r in range(start_row, end_row + 1):
            for c in range(1, cols + 1):
                cell = ws.cell(row=r, column=c)
                cell.alignment = self._wrap
                cell.border = self._border

    def _autosize_cols(self, ws, max_col: int, max_width: int = 40):
        for c in range(1, max_col + 1):
            letter = get_column_letter(c)
            max_len = 0
            for cell in ws[letter]:
                if cell.value:
                    text = str(cell.value)
                    for line in text.split("\n"):
                        if len(line) > max_len:
                            max_len = len(line)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), max_width)

    def _display_name(self, canonical_key: str) -> str:
        if canonical_key in FIELD_SPECS:
            return FIELD_SPECS[canonical_key].display_name
        return canonical_key

    def _write_summary(self, wb: Workbook, audit_result: AuditResult):
        ws = wb.active
        ws.title = "概览"

        ws.cell(row=1, column=1, value="门店现金短款调查稽核报告").font = self._title_font
        ws.merge_cells("A1:E1")

        ws.cell(row=2, column=1, value=f"报告日期: {self.ctx.report_date.isoformat()}")
        ws.cell(row=3, column=1, value=f"编制: {self.ctx.prepared_by}")
        if self.ctx.filter_description:
            ws.cell(row=4, column=1, value=f"筛选条件: {self.ctx.filter_description}")

        row = 6
        ws.cell(row=row, column=1, value="本次稽核总体结论").font = self._sub_title_font
        row += 1

        total_issues = audit_result.total_issues()
        total_amount = audit_result.total_amount()

        if total_issues == 0:
            conclusion = f"在本次筛选范围内，未发现退款漏签、备用金错班、流水重复三类风险。"
        else:
            conclusion = (
                f"在本次筛选范围内，共发现 {total_issues} 笔风险记录，"
                f"涉及金额 {round(total_amount, 2)} 元。"
                f"其中："
                + "，".join(
                    f"{f.check_label} {f.finding_count} 笔(涉及 {round(f.total_affected_amount, 2)} 元)"
                    for f in audit_result.findings if f.finding_count > 0
                )
                + "。详细明细见后续工作表。"
            )

        cell = ws.cell(row=row, column=1, value=conclusion)
        cell.alignment = self._wrap
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 2

        ws.cell(row=row, column=1, value="分项检查汇总").font = self._sub_title_font
        row += 1

        headers = ["检查项", "风险级别", "发现笔数", "涉及金额(元)", "结论说明"]
        for c, h in enumerate(headers, 1):
            ws.cell(row=row, column=c, value=h)
        self._style_header_row(ws, row, len(headers))
        row += 1

        for f in audit_result.findings:
            ws.cell(row=row, column=1, value=f.check_label)
            ws.cell(row=row, column=2, value=f.severity)
            ws.cell(row=row, column=3, value=f.finding_count)
            ws.cell(row=row, column=4, value=round(f.total_affected_amount, 2))
            ws.cell(row=row, column=5, value=f.summary)
            row += 1

        self._style_data_area(ws, 7, row - 1, len(headers))
        self._autosize_cols(ws, len(headers))

    def _write_findings_by_type(self, wb: Workbook, audit_result: AuditResult):
        for f in audit_result.findings:
            self._write_single_finding_sheet(wb, f)

    def _write_single_finding_sheet(self, wb: Workbook, finding: AuditFinding):
        ws = wb.create_sheet(title=finding.check_label)

        ws.cell(row=1, column=1, value=f"【{finding.check_label}】检查明细").font = self._title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)

        ws.cell(row=2, column=1, value=f"风险级别: {finding.severity}")
        ws.cell(row=3, column=1, value=f"发现笔数: {finding.finding_count}")
        ws.cell(row=4, column=1, value=f"涉及金额: {round(finding.total_affected_amount, 2)} 元")
        ws.cell(row=5, column=1, value=f"检查说明: {finding.summary}")

        row = 7
        ws.cell(row=row, column=1, value="认定理由说明（逐条）").font = self._sub_title_font
        row += 1

        display_cols = [self._display_name(k) for k in finding.evidence_columns]
        headers = ["序号"] + display_cols + ["认定理由"]
        for c, h in enumerate(headers, 1):
            ws.cell(row=row, column=c, value=h)
        self._style_header_row(ws, row, len(headers))
        row += 1

        start_data_row = row

        for idx, (_, record) in enumerate(finding.finding_df.iterrows(), 1):
            ws.cell(row=row, column=1, value=idx)
            for c2, key in enumerate(finding.evidence_columns, 2):
                val = record[key]
                if key == "amount" and pd.notna(val):
                    val = round(float(val), 2)
                ws.cell(row=row, column=c2, value=val)
            reason = finding.reason_template.format(**record.to_dict())
            ws.cell(row=row, column=len(headers), value=reason)
            row += 1

        self._style_data_area(ws, start_data_row, row - 1, len(headers))
        self._autosize_cols(ws, len(headers))

    def _write_all_findings_detail(self, wb: Workbook, audit_result: AuditResult):
        ws = wb.create_sheet(title="全部风险明细汇总")

        if len(audit_result.raw_audit_detail) == 0:
            ws.cell(row=1, column=1, value="本次筛选范围内未发现风险记录。")
            return

        ws.cell(row=1, column=1, value="全部风险明细汇总").font = self._title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)

        row = 3
        key_cols = ["store_code", "store_name", "trans_date", "shift", "cashier", "trans_id"]
        display_cols = [self._display_name(k) for k in key_cols]
        headers = ["序号", "检查类型", "风险级别"] + display_cols + ["金额(元)", "认定理由"]
        for c, h in enumerate(headers, 1):
            ws.cell(row=row, column=c, value=h)
        self._style_header_row(ws, row, len(headers))
        row += 1

        start_data_row = row

        for idx, (_, r) in enumerate(audit_result.raw_audit_detail.iterrows(), 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=r["_check_label"])
            ws.cell(row=row, column=3, value=r["_severity"])
            for c2, key in enumerate(key_cols, 4):
                ws.cell(row=row, column=c2, value=r[key])
            ws.cell(row=row, column=len(headers) - 1, value=round(float(r["amount"]), 2))
            ws.cell(row=row, column=len(headers), value=r["_reason"])
            row += 1

        self._style_data_area(ws, start_data_row, row - 1, len(headers))
        self._autosize_cols(ws, len(headers))

    def _write_filtered_data(self, wb: Workbook, filtered_df: pd.DataFrame):
        ws = wb.create_sheet(title="筛选后标准流水(备查)")

        ws.cell(row=1, column=1, value="筛选后的标准流水（口径统一后）").font = self._sub_title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(filtered_df.columns))

        headers = [self._display_name(c) for c in filtered_df.columns]
        for c, h in enumerate(headers, 1):
            ws.cell(row=2, column=c, value=h)
        self._style_header_row(ws, 2, len(headers))

        for r_idx, (_, record) in enumerate(filtered_df.iterrows(), 3):
            for c_idx, col in enumerate(filtered_df.columns, 1):
                val = record[col]
                if col == "amount" and pd.notna(val):
                    val = round(float(val), 2)
                ws.cell(row=r_idx, column=c_idx, value=val)

        end_row = 2 + len(filtered_df)
        self._style_data_area(ws, 3, end_row, len(headers))
        self._autosize_cols(ws, len(headers))

    def _write_raw_data(self, wb: Workbook, raw_df: pd.DataFrame):
        ws = wb.create_sheet(title="原始流水(转同事核对用)")

        ws.cell(row=1, column=1, value="原始收银流水（未经过口径转换，转同事核对用）").font = self._sub_title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(raw_df.columns))

        headers = list(raw_df.columns)
        for c, h in enumerate(headers, 1):
            ws.cell(row=2, column=c, value=h)
        self._style_header_row(ws, 2, len(headers))

        for r_idx, (_, record) in enumerate(raw_df.iterrows(), 3):
            for c_idx, col in enumerate(headers, 1):
                ws.cell(row=r_idx, column=c_idx, value=record[col])

        end_row = 2 + len(raw_df)
        self._style_data_area(ws, 3, end_row, len(headers))
        self._autosize_cols(ws, len(headers), max_width=50)

    def _write_caliber_explanation(self, wb: Workbook):
        ws = wb.create_sheet(title="口径说明")

        ws.cell(row=1, column=1, value="字段口径与检查规则说明").font = self._title_font
        ws.merge_cells("A1:C1")

        row = 3
        ws.cell(row=row, column=1, value="一、字段口径说明").font = self._sub_title_font
        row += 1

        headers = ["字段显示名", "原始列名", "说明"]
        for c, h in enumerate(headers, 1):
            ws.cell(row=row, column=c, value=h)
        self._style_header_row(ws, row, 3)
        row += 1

        for spec in FIELD_SPECS.values():
            ws.cell(row=row, column=1, value=spec.display_name)
            ws.cell(row=row, column=2, value=spec.raw_column)
            ws.cell(row=row, column=3, value=spec.description)
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="二、检查规则说明").font = self._sub_title_font
        row += 1

        headers = ["检查项", "风险级别", "判定规则"]
        for c, h in enumerate(headers, 1):
            ws.cell(row=row, column=c, value=h)
        self._style_header_row(ws, row, 3)
        row += 1

        rules = [
            ("退款漏签", "高", "交易类型为「退款」且「是否签字确认」不等于「是/已签字/Y」"),
            ("备用金错班", "中", "交易类型含「备用金」，且备用金类型中标注的班次（早班/中班/晚班）与实际当班班次不一致"),
            ("流水重复", "高", "同一流水号出现多次，或同一门店、同一日期、同一班次、同一金额、同一交易类型的记录重复出现"),
        ]
        for label, severity, rule in rules:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=severity)
            ws.cell(row=row, column=3, value=rule)
            row += 1

        self._style_data_area(ws, 4, row - 1, 3)
        self._autosize_cols(ws, 3, max_width=60)
