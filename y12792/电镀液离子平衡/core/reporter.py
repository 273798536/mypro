import os
from datetime import datetime
from typing import List, Optional
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import BatchRecord
from .analyzer import IonBalanceAnalyzer


class ReportGenerator:
    def __init__(self, output_dir: str = "output"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
        self.analyzer = IonBalanceAnalyzer()
        
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", size=12)
        self.abnormal_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.warning_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        self.normal_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        self.title_font = Font(bold=True, size=14)
        self.section_font = Font(bold=True, size=12)
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def _generate_filename(self, prefix: str, batch_no: Optional[str] = None,
                          include_run_id: bool = False, run_id: str = "") -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        parts = [prefix]
        if batch_no:
            parts.append(f"批号{batch_no}")
        if include_run_id and run_id:
            parts.append(f"运行{run_id}")
        parts.append(timestamp)
        return "_".join(parts) + ".xlsx"
    
    def generate_batch_report(self, record: BatchRecord, include_spectrum: bool = True) -> str:
        filename = self._generate_filename(
            "电镀液离子平衡报告",
            batch_no=record.batch_no,
            include_run_id=not record.is_latest,
            run_id=record.run_id
        )
        filepath = os.path.join(self.output_dir, filename)
        
        wb = Workbook()
        
        self._write_summary_sheet(wb.active, record)
        
        self._write_ion_results_sheet(wb.create_sheet("离子浓度检测结果"), record)
        
        self._write_reagent_sheet(wb.create_sheet("试剂使用记录"), record)
        
        self._write_weighing_sheet(wb.create_sheet("样品称量记录"), record)
        
        if include_spectrum and record.spectrum_data:
            self._write_spectrum_sheet(wb.create_sheet("原始谱图数据"), record)
        
        self._write_trace_sheet(wb.create_sheet("异常追溯明细"), record)
        
        wb.save(filepath)
        return filepath
    
    def generate_summary_report(self, records: List[BatchRecord]) -> str:
        filename = self._generate_filename("电镀液离子平衡汇总报告")
        filepath = os.path.join(self.output_dir, filename)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "批次汇总"
        
        ws.merge_cells("A1:I1")
        ws["A1"] = "电镀液离子平衡检测汇总报告"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws["A3"] = f"报告生成时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M:%S')}"
        ws["A4"] = f"本次共分析 {len(records)} 个批次"
        
        abnormal_count = sum(1 for r in records if r.overall_status == "abnormal")
        warning_count = sum(1 for r in records if r.overall_status == "warning")
        ws["A5"] = f"其中异常 {abnormal_count} 批，需关注 {warning_count} 批，正常 {len(records)-abnormal_count-warning_count} 批"
        
        headers = [
            "批号", "样品名称", "检测时间", "操作员",
            "温度(℃)", "pH值", "整体状态", "异常离子数", "备注"
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=7, column=col, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border
        
        for row, record in enumerate(records, 8):
            abnormal_ions = sum(1 for r in record.analysis_results if r.status == "abnormal")
            
            status_text = {
                "normal": "正常",
                "warning": "需关注",
                "abnormal": "异常",
                "pending": "待检测"
            }.get(record.overall_status, record.overall_status)
            
            values = [
                record.batch_no,
                record.sample_name,
                record.run_time,
                record.operator,
                record.temperature,
                record.ph_value,
                status_text,
                abnormal_ions,
                record.remark or ""
            ]
            
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = self.thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
                if col == 7:
                    if record.overall_status == "abnormal":
                        cell.fill = self.abnormal_fill
                    elif record.overall_status == "warning":
                        cell.fill = self.warning_fill
                    elif record.overall_status == "normal":
                        cell.fill = self.normal_fill
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        
        ws2 = wb.create_sheet("异常批次清单")
        self._write_abnormal_list_sheet(ws2, records)
        
        wb.save(filepath)
        return filepath
    
    def _write_summary_sheet(self, ws, record: BatchRecord):
        ws.title = "报告概览"
        
        ws.merge_cells("A1:D1")
        ws["A1"] = "电镀液离子平衡检测报告"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws["A3"] = "基本信息"
        ws["A3"].font = self.section_font
        
        balance_info = self.analyzer.calculate_ion_balance(record)
        
        info_items = [
            ("批号", record.batch_no),
            ("样品名称", record.sample_name or "未填写"),
            ("检测时间", record.run_time),
            ("运行编号", record.run_id),
            ("操作员", record.operator or "未填写"),
            ("工作温度", f"{record.temperature} {record.temperature_unit}"),
            ("pH值", record.ph_value or "未检测"),
            ("数据来源", record.source_file or "未记录"),
        ]
        
        for i, (label, value) in enumerate(info_items, 4):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=value)
        
        ws["A14"] = "离子平衡状态"
        ws["A14"].font = self.section_font
        
        balance_items = [
            ("阳离子总浓度", f"{balance_info['total_cations']} g/L"),
            ("阴离子总浓度", f"{balance_info['total_anions']} g/L"),
            ("阴阳离子比", balance_info['balance_ratio']),
            ("平衡状态", self._translate_balance_status(balance_info['balance_status'])),
            ("平衡说明", balance_info['balance_remark']),
        ]
        
        for i, (label, value) in enumerate(balance_items, 15):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=value)
        
        ws["A22"] = "整体结论"
        ws["A22"].font = self.section_font
        
        status_text = {
            "normal": "✅ 正常",
            "warning": "⚠️ 需关注",
            "abnormal": "❌ 异常",
            "pending": "⏳ 待检测"
        }.get(record.overall_status, record.overall_status)
        
        ws["A23"] = f"检测状态：{status_text}"
        ws["A23"].font = Font(bold=True, size=12)
        if record.overall_status == "abnormal":
            ws["A23"].fill = self.abnormal_fill
        elif record.overall_status == "warning":
            ws["A23"].fill = self.warning_fill
        elif record.overall_status == "normal":
            ws["A23"].fill = self.normal_fill
        
        ws["A25"] = "分析意见"
        ws["A25"].font = self.section_font
        ws["A26"] = record.analysis_opinion or "暂无"
        ws.merge_cells("A26:D35")
        ws["A26"].alignment = Alignment(wrap_text=True, vertical="top")
        
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 20
    
    def _translate_balance_status(self, status: str) -> str:
        return {
            "balanced": "基本平衡",
            "anion_excess": "阴离子偏多",
            "cation_excess": "阳离子偏多"
        }.get(status, status)
    
    def _write_ion_results_sheet(self, ws, record: BatchRecord):
        headers = [
            "离子名称", "检测浓度(g/L)", "正常范围下限(g/L)", "正常范围上限(g/L)",
            "状态", "异常说明", "特征波长(nm)", "计算方法说明"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.thin_border
        
        normal_ranges = {
            "铜离子(Cu²⁺)": (150, 220),
            "镍离子(Ni²⁺)": (60, 100),
            "锌离子(Zn²⁺)": (20, 50),
            "铬离子(Cr⁶⁺)": (5, 20),
            "硫酸根(SO₄²⁻)": (100, 180),
        }
        
        wavelength_map = {
            "铜离子(Cu²⁺)": 810,
            "镍离子(Ni²⁺)": 395,
            "锌离子(Zn²⁺)": 620,
            "铬离子(Cr⁶⁺)": 540,
            "硫酸根(SO₄²⁻)": 420,
        }
        
        for row, result in enumerate(record.analysis_results, 2):
            normal_low, normal_high = normal_ranges.get(result.ion_name, (0, 9999))
            status_text = "正常" if result.status == "normal" else "异常"
            
            values = [
                result.ion_name,
                result.concentration,
                normal_low,
                normal_high,
                status_text,
                result.remark or "无",
                wavelength_map.get(result.ion_name, "-"),
                result.calculation_formula
            ]
            
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = self.thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                
                if col == 5:
                    if result.status == "abnormal":
                        cell.fill = self.abnormal_fill
                    else:
                        cell.fill = self.normal_fill
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        
        ws.row_dimensions[1].height = 30
    
    def _write_reagent_sheet(self, ws, record: BatchRecord):
        headers = [
            "试剂名称", "试剂批号", "浓度", "浓度单位", "使用量(mL)",
            "供应商", "有效期", "备注"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border
        
        for row, reagent in enumerate(record.reagent_records, 2):
            values = [
                reagent.reagent_name,
                reagent.reagent_batch,
                reagent.concentration,
                reagent.unit,
                reagent.volume_used,
                reagent.supplier or "未填写",
                reagent.expiry_date or "未填写",
                ""
            ]
            
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = self.thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
    
    def _write_weighing_sheet(self, ws, record: BatchRecord):
        headers = [
            "样品名称", "称量重量", "重量单位", "称量人", "称量时间", "备注"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border
        
        for row, weighing in enumerate(record.weighing_records, 2):
            values = [
                weighing.sample_name,
                weighing.weight,
                weighing.weight_unit,
                weighing.operator or "未填写",
                weighing.weigh_time or "未填写",
                ""
            ]
            
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = self.thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
    
    def _write_spectrum_sheet(self, ws, record: BatchRecord):
        ws.merge_cells("A1:C1")
        ws["A1"] = f"原始谱图数据 - 批号 {record.batch_no}"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")
        
        headers = ["波长(nm)", "吸光度(Abs)", "浓度计算值(g/L)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border
        
        for row, point in enumerate(record.spectrum_data, 4):
            values = [
                point.wavelength,
                point.absorbance,
                point.concentration if point.concentration is not None else "未计算"
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = self.thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 20
    
    def _write_trace_sheet(self, ws, record: BatchRecord):
        ws.merge_cells("A1:D1")
        ws["A1"] = "异常追溯明细"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")
        
        ws["A3"] = "使用说明：从异常离子浓度出发，向下可追溯到原始谱图数据、试剂使用记录和处理意见，方便复核。"
        ws.merge_cells("A3:D3")
        ws["A3"].alignment = Alignment(wrap_text=True)
        
        current_row = 5
        abnormal_results = [r for r in record.analysis_results if r.status == "abnormal"]
        
        if not abnormal_results:
            ws["A5"] = "✅ 本次检测无异常项目，所有离子浓度均在正常范围内。"
            ws["A5"].fill = self.normal_fill
            return
        
        for result in abnormal_results:
            trace = self.analyzer.get_abnormal_trace(record, result.ion_name)
            if not trace:
                continue
            
            ws.cell(row=current_row, column=1, value=f"异常项目：{result.ion_name}")
            ws.cell(row=current_row, column=1).font = Font(bold=True, size=12)
            ws.cell(row=current_row, column=1).fill = self.abnormal_fill
            ws.merge_cells(f"A{current_row}:D{current_row}")
            current_row += 1
            
            ws.cell(row=current_row, column=1, value="异常详情：").font = Font(bold=True)
            ws.cell(row=current_row, column=2, value=f"浓度 {result.concentration} {result.unit}，{result.remark}")
            ws.merge_cells(f"B{current_row}:D{current_row}")
            current_row += 1
            
            ws.cell(row=current_row, column=1, value="相关谱图数据：").font = Font(bold=True)
            current_row += 1
            
            ws.cell(row=current_row, column=1, value="波长(nm)")
            ws.cell(row=current_row, column=2, value="吸光度")
            ws.cell(row=current_row, column=3, value="说明")
            for c in range(1, 4):
                ws.cell(row=current_row, column=c).font = Font(bold=True)
                ws.cell(row=current_row, column=c).border = self.thin_border
            current_row += 1
            
            for point in result.spectrum_points:
                ws.cell(row=current_row, column=1, value=point.wavelength)
                ws.cell(row=current_row, column=2, value=point.absorbance)
                ws.cell(row=current_row, column=3,
                       value="特征峰附近" if abs(point.wavelength - trace.get('spectrum_points', [point])[0].wavelength) < 5 else "邻域点")
                for c in range(1, 4):
                    ws.cell(row=current_row, column=c).border = self.thin_border
                current_row += 1
            
            ws.cell(row=current_row, column=1, value="计算方法：").font = Font(bold=True)
            ws.cell(row=current_row, column=2, value=result.calculation_formula)
            ws.merge_cells(f"B{current_row}:D{current_row}")
            current_row += 1
            
            ws.cell(row=current_row, column=1, value="温度条件：").font = Font(bold=True)
            ws.cell(row=current_row, column=2, value=f"{trace['temperature']} {trace['temperature_unit']}")
            current_row += 1
            
            ws.cell(row=current_row, column=1, value="可能相关试剂：").font = Font(bold=True)
            current_row += 1
            
            for reagent in record.reagent_records:
                ws.cell(row=current_row, column=1, value=reagent.reagent_name)
                ws.cell(row=current_row, column=2, value=f"批号：{reagent.reagent_batch}")
                ws.cell(row=current_row, column=3, value=f"用量：{reagent.volume_used}mL")
                for c in range(1, 4):
                    ws.cell(row=current_row, column=c).border = self.thin_border
                current_row += 1
            
            ws.cell(row=current_row, column=1, value="分析处理意见：").font = Font(bold=True)
            ws.cell(row=current_row, column=2, value=trace['analysis_opinion'])
            ws.merge_cells(f"B{current_row}:D{current_row + 4}")
            ws.cell(row=current_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            current_row += 6
            
            ws.cell(row=current_row, column=1, value=f"数据来源文件：{trace['source_file']}")
            ws.merge_cells(f"A{current_row}:D{current_row}")
            current_row += 1
            ws.cell(row=current_row, column=1, value=f"运行编号：{trace['run_id']}，运行时间：{trace['run_time']}")
            ws.merge_cells(f"A{current_row}:D{current_row}")
            current_row += 2
        
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 20
    
    def _write_abnormal_list_sheet(self, ws, records: List[BatchRecord]):
        ws.merge_cells("A1:G1")
        ws["A1"] = "异常批次清单"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")
        
        headers = [
            "批号", "样品名称", "异常离子", "异常浓度",
            "正常范围", "异常说明", "检测时间"
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border
        
        row = 4
        for record in records:
            for result in record.analysis_results:
                if result.status == "abnormal":
                    normal_ranges = {
                        "铜离子(Cu²⁺)": "150-220 g/L",
                        "镍离子(Ni²⁺)": "60-100 g/L",
                        "锌离子(Zn²⁺)": "20-50 g/L",
                        "铬离子(Cr⁶⁺)": "5-20 g/L",
                        "硫酸根(SO₄²⁻)": "100-180 g/L",
                    }
                    
                    values = [
                        record.batch_no,
                        record.sample_name or "未填写",
                        result.ion_name,
                        f"{result.concentration} {result.unit}",
                        normal_ranges.get(result.ion_name, "参考资料"),
                        result.remark,
                        record.run_time
                    ]
                    
                    for col, value in enumerate(values, 1):
                        cell = ws.cell(row=row, column=col, value=value)
                        cell.border = self.thin_border
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.fill = self.abnormal_fill
                    
                    row += 1
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
