#!/usr/bin/env python3
"""导出一致性测试 - 验证批次统计与导出数据一致性"""
import sys
import json
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from customs_statemachine.config import settings
from customs_statemachine.database import SessionLocal, engine
from customs_statemachine.models import Base, Batch, Package, TrackingNode, TaxNotice, TempRecord
from customs_statemachine.schemas import BatchCreate, PackageCreate, TrackingNodeCreate, TaxNoticeCreate, TempRecordCreate, IdempotencyMode
from customs_statemachine.state_machine import StateMachine, IdempotencyHandler
from customs_statemachine.audit_service import AuditService
from customs_statemachine.export_service import ExportService


def test_export_consistency():
    print("=" * 70)
    print("  导出一致性测试")
    print("=" * 70)

    # 清理并重建数据库
    if Path(settings.DATABASE_URL.replace("sqlite:///", "")).exists():
        Path(settings.DATABASE_URL.replace("sqlite:///", "")).unlink()
    Base.metadata.create_all(bind=engine)

    db = SessionLocal()
    try:
        state_machine = StateMachine(db)
        audit = AuditService(db)
        idempotency = IdempotencyHandler(db)

        # 1. 创建带多数据源的批次
        print("\n1. 创建批次（多数据源：5个包裹）")
        batch_data = BatchCreate(
            batch_no="BATCH-CONSISTENCY-TEST",
            created_by="test_operator",
            source_type="mixed",
            customs_code="CN021",
            packages=[
                PackageCreate(package_no=f"PKG-TEST-{i:03d}", waybill_no=f"WB-TEST-{i:03d}",
                            declared_value=100.0 * (i + 1), tax_amount=10.0 * (i + 1),
                            source="declaration_form")
                for i in range(5)
            ],
            tracking_nodes=[
                TrackingNodeCreate(package_no=f"PKG-TEST-{i:03d}", node_type="clearance",
                                status="正常", location=f"海关-{i}", node_time=f"2024-05-{10+i:02d}T10:00:00")
                for i in range(5)
            ],
            tax_notices=[
                TaxNoticeCreate(notice_no=f"TAX-TEST-{i:03d}", package_no=f"PKG-TEST-{i:03d}",
                              tax_amount=12.0 * (i + 1))
                for i in range(5)
            ],
            temp_records=[
                TempRecordCreate(record_type="manual_entry", package_no=f"PKG-TEST-{i:03d}",
                               content=f"补录记录-{i}", recorded_by=f"officer-{i}")
                for i in range(5)
            ],
            idempotency_mode=IdempotencyMode.IGNORE
        )

        batch, result = state_machine.create_batch(batch_data)
        print(f"  ✓ 批次ID: {batch.id}")
        print(f"  ✓ 创建结果: {result}")

        # 2. 立即检查批次统计
        print("\n2. 验证批次统计字段")
        db.refresh(batch)
        print(f"  批次.total_packages = {batch.total_packages}")
        print(f"  批次.total_tax_amount = {batch.total_tax_amount}")

        # 3. 直接查询数据库中的实际数据数量
        actual_packages = db.query(Package).filter(Package.batch_id == batch.id).count()
        actual_nodes = db.query(TrackingNode).filter(TrackingNode.batch_id == batch.id).count()
        actual_notices = db.query(TaxNotice).filter(TaxNotice.batch_id == batch.id).count()
        actual_temp = db.query(TempRecord).filter(TempRecord.batch_id == batch.id).count()

        print(f"\n3. 数据库实际记录数")
        print(f"  packages: {actual_packages}")
        print(f"  tracking_nodes: {actual_nodes}")
        print(f"  tax_notices: {actual_notices}")
        print(f"  temp_records: {actual_temp}")

        # 4. 验证统计一致性
        print("\n4. 验证一致性")
        checks = []
        checks.append(("批次.total_packages == 实际packages数", batch.total_packages == actual_packages))
        checks.append(("批次.total_tax_amount > 0", batch.total_tax_amount > 0))

        for check_name, passed in checks:
            status = "✓ PASS" if passed else "✗ FAIL"
            print(f"  {status}: {check_name}")
            if not passed:
                print(f"    期望: {actual_packages}, 实际: {batch.total_packages}")

        # 5. 导出数据并验证
        print("\n5. 导出数据验证")
        export_svc = ExportService(db)
        export_result = export_svc.export_to_excel(batch=batch, include_history=True, include_packages=True)
        print(f"  ✓ 导出文件: {export_result['file_name']}")

        # 读取导出的Excel文件验证概览数据
        import openpyxl
        wb = openpyxl.load_workbook(export_result['file_path'])

        # 检查批次概览表
        ws_overview = wb["批次概览"]
        overview_data = {}
        for row in ws_overview.iter_rows(min_row=1, max_row=ws_overview.max_row, values_only=True):
            if row[0] and row[1]:
                overview_data[row[0]] = row[1]

        print(f"\n6. Excel导出内容验证")
        print(f"  批次概览 - 批次号: {overview_data.get('批次号', 'N/A')}")
        print(f"  批次概览 - 包裹总数: {overview_data.get('包裹总数', 'N/A')}")
        print(f"  批次概览 - 税费总额: {overview_data.get('税费总额', 'N/A')}")

        # 检查包裹明细
        ws_packages = wb["包裹明细"]
        package_rows = list(ws_packages.iter_rows(min_row=2, max_row=ws_packages.max_row, values_only=True))
        print(f"\n  包裹明细 - 行数: {len(package_rows)}")

        # 验证导出一致性
        export_total = overview_data.get('包裹总数', 0)
        checks.append(("Excel概览包裹总数 == 实际packages数", export_total == actual_packages))
        checks.append(("Excel包裹明细行数 == 实际packages数", len(package_rows) == actual_packages))

        for check_name, passed in checks[2:]:
            status = "✓ PASS" if passed else "✗ FAIL"
            print(f"  {status}: {check_name}")
            if not passed:
                if "包裹总数" in check_name:
                    print(f"    期望: {actual_packages}, Excel显示: {export_total}")

        wb.close()

        # 7. 测试 APPEND 模式追加数据后再次验证
        print("\n7. APPEND模式追加数据后验证")
        from customs_statemachine.schemas import BatchDataAppend

        append_data = BatchDataAppend(
            changed_by="test_operator",
            packages=[
                PackageCreate(package_no="PKG-APPEND-001", waybill_no="WB-APPEND-001",
                            declared_value=500.0, tax_amount=50.0, source="temp_record")
            ],
            tracking_nodes=[
                TrackingNodeCreate(package_no="PKG-APPEND-001", node_type="clearance",
                                status="正常", location="追加海关", node_time="2024-05-20T10:00:00")
            ],
            tax_notices=[
                TaxNoticeCreate(notice_no="TAX-APPEND-001", package_no="PKG-APPEND-001", tax_amount=55.0)
            ],
            temp_records=[
                TempRecordCreate(record_type="manual_entry", package_no="PKG-APPEND-001",
                               content="追加记录", recorded_by="append-officer")
            ],
            idempotency_mode=IdempotencyMode.APPEND
        )

        state_machine.append_data(batch, append_data)
        db.refresh(batch)

        actual_after_append = db.query(Package).filter(Package.batch_id == batch.id).count()
        print(f"  追加后实际packages数: {actual_after_append}")
        print(f"  批次.total_packages: {batch.total_packages}")

        checks.append(("APPEND后批次.total_packages == 实际packages数", batch.total_packages == actual_after_append))

        # 再次导出验证
        export_result2 = export_svc.export_to_excel(batch=batch, include_history=True, include_packages=True)
        wb2 = openpyxl.load_workbook(export_result2['file_path'])
        ws_overview2 = wb2["批次概览"]
        overview2 = {}
        for row in ws_overview2.iter_rows(min_row=1, max_row=ws_overview2.max_row, values_only=True):
            if row[0] and row[1]:
                overview2[row[0]] = row[1]
        ws_packages2 = wb2["包裹明细"]
        package_rows2 = list(ws_packages2.iter_rows(min_row=2, max_row=ws_packages2.max_row, values_only=True))

        export_total2 = overview2.get('包裹总数', 0)
        checks.append(("APPEND后Excel概览包裹总数 == 实际packages数", export_total2 == actual_after_append))
        checks.append(("APPEND后Excel包裹明细行数 == 实际packages数", len(package_rows2) == actual_after_append))

        wb2.close()

        # 8. 总结
        print("\n" + "=" * 70)
        print("  测试总结")
        print("=" * 70)

        passed_count = sum(1 for _, passed in checks if passed)
        for check_name, passed in checks:
            status = "✓ PASS" if passed else "✗ FAIL"
            print(f"  {status}: {check_name}")

        print(f"\n  总计: {passed_count}/{len(checks)} 项测试通过")

        if all(passed for _, passed in checks):
            print("\n  🎉 所有导出一致性测试通过！")
            return 0
        else:
            print(f"\n  ❌ {len(checks) - passed_count} 项测试失败")
            return 1

    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(test_export_consistency())
