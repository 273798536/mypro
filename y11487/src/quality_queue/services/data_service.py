from datetime import datetime, date
from typing import Dict, Any, List, Optional, Tuple
from sqlalchemy.orm import Session
import json
import csv
from io import StringIO
import uuid

from ..models import (
    Inspection, ReworkOrder, MachineShift, ExceptionRecord,
    RecordSource, CompensationQueue, SessionLocal, DirtyRecord
)
from ..utils.json_utils import to_json_serializable
from .data_validator import DataValidator


class DataService:
    def __init__(self, db: Session):
        self.db = db
        self.validator = DataValidator(db)

    def _generate_draft_id(self, prefix: str) -> str:
        return f"{prefix}_DRAFT_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:4]}"

    def _link_dirty_records_to_draft(self, dirty_records: List[DirtyRecord],
                                     draft_id: int):
        for dr in dirty_records:
            dr.draft_record_id = draft_id
        self.db.flush()

    def _is_draft_complete(self, record, required_fields: List[str]) -> bool:
        for field in required_fields:
            value = getattr(record, field, None)
            if value in (None, "", 0):
                return False
        return True

    def create_machine_shift(self, data: Dict[str, Any]) -> Tuple[Optional[MachineShift], List[Any]]:
        is_valid, dirty_records = self.validator.validate_machine_shift(data)

        if isinstance(data.get("shift_date"), str):
            from dateutil import parser
            try:
                data["shift_date"] = parser.parse(data["shift_date"]).date()
            except:
                pass
        if "start_time" in data and isinstance(data["start_time"], str):
            from dateutil import parser
            try:
                data["start_time"] = parser.parse(data["start_time"])
            except:
                pass
        if "end_time" in data and isinstance(data["end_time"], str):
            from dateutil import parser
            try:
                data["end_time"] = parser.parse(data["end_time"])
            except:
                pass

        shift_code = data.get("shift_code") or self._generate_draft_id("SH")
        is_draft = not data.get("shift_code")

        shift = MachineShift(
            shift_code=shift_code,
            machine_no=data.get("machine_no", ""),
            shift_date=data.get("shift_date"),
            shift_type=data.get("shift_type", ""),
            operator=data.get("operator", ""),
            team_leader=data.get("team_leader", ""),
            start_time=data.get("start_time"),
            end_time=data.get("end_time"),
            target_quantity=data.get("target_quantity", 0),
            actual_quantity=data.get("actual_quantity", 0),
            defect_quantity=data.get("defect_quantity", 0),
            base_yield_rate=data.get("base_yield_rate", 0.0),
            raw_data=to_json_serializable(data),
            is_draft=is_draft
        )

        self.db.add(shift)
        self.db.flush()
        self._link_dirty_records_to_draft(dirty_records, shift.id)
        return shift, dirty_records

    def create_inspection(self, data: Dict[str, Any]) -> Tuple[Optional[Inspection], List[Any]]:
        is_valid, dirty_records = self.validator.validate_inspection(data)

        shift = None
        if "shift_code" in data and data["shift_code"]:
            shift = self.db.query(MachineShift).filter(
                MachineShift.shift_code == data["shift_code"]
            ).first()

        if isinstance(data.get("inspection_date"), str):
            from dateutil import parser
            try:
                data["inspection_date"] = parser.parse(data["inspection_date"]).date()
            except:
                pass

        inspection_no = data.get("inspection_no") or self._generate_draft_id("IN")
        is_draft = not data.get("inspection_no")

        inspection = Inspection(
            inspection_no=inspection_no,
            shift_id=shift.id if shift else None,
            machine_no=data.get("machine_no", ""),
            inspection_date=data.get("inspection_date"),
            inspector=data.get("inspector", ""),
            batch_no=data.get("batch_no", ""),
            style_no=data.get("style_no", ""),
            fabric_type=data.get("fabric_type", ""),
            sample_size=data.get("sample_size", 0),
            defect_count=data.get("defect_count", 0),
            defect_rate=data.get("defect_rate", 0.0),
            is_qualified=data.get("is_qualified", True),
            raw_data=to_json_serializable(data),
            is_draft=is_draft
        )

        self.db.add(inspection)
        self.db.flush()
        self._link_dirty_records_to_draft(dirty_records, inspection.id)
        return inspection, dirty_records

    def create_rework_order(self, data: Dict[str, Any]) -> Tuple[Optional[ReworkOrder], List[Any]]:
        is_valid, dirty_records = self.validator.validate_rework_order(data)

        inspection = None
        if "inspection_no" in data and data["inspection_no"]:
            inspection = self.db.query(Inspection).filter(
                Inspection.inspection_no == data["inspection_no"]
            ).first()

        shift = None
        if "shift_code" in data and data["shift_code"]:
            shift = self.db.query(MachineShift).filter(
                MachineShift.shift_code == data["shift_code"]
            ).first()

        parent_rework = None
        if "parent_rework_no" in data and data["parent_rework_no"]:
            parent_rework = self.db.query(ReworkOrder).filter(
                ReworkOrder.rework_no == data["parent_rework_no"]
            ).first()

        if isinstance(data.get("rework_date"), str):
            from dateutil import parser
            try:
                data["rework_date"] = parser.parse(data["rework_date"]).date()
            except:
                pass

        rework_no = data.get("rework_no") or self._generate_draft_id("RW")
        is_draft = not data.get("rework_no")

        rework_order = ReworkOrder(
            rework_no=rework_no,
            inspection_id=inspection.id if inspection else None,
            shift_id=shift.id if shift else None,
            parent_rework_id=parent_rework.id if parent_rework else None,
            rework_date=data.get("rework_date"),
            machine_no=data.get("machine_no", ""),
            defect_type=data.get("defect_type", ""),
            defect_description=data.get("defect_description", ""),
            rework_quantity=data.get("rework_quantity", 0),
            rework_operator=data.get("rework_operator", ""),
            rework_team=data.get("rework_team", ""),
            compensation_amount=data.get("compensation_amount", 0.0),
            is_reworked=data.get("is_reworked", False),
            rework_pass_count=data.get("rework_pass_count", 0),
            rework_fail_count=data.get("rework_fail_count", 0),
            rework_count=data.get("rework_count", 1),
            final_yield_rate=data.get("final_yield_rate", 0.0),
            responsible_shift_code=data.get("responsible_shift_code", ""),
            raw_data=to_json_serializable(data),
            is_draft=is_draft
        )

        self.db.add(rework_order)
        self.db.flush()
        self._link_dirty_records_to_draft(dirty_records, rework_order.id)
        return rework_order, dirty_records

    def create_exception_record(self, data: Dict[str, Any]) -> Tuple[ExceptionRecord, List[Any]]:
        if "source_type" in data:
            data = dict(data)
            source_type = data.pop("source_type")
            data["source"] = source_type

        source_value = data.get("source", "sms")
        if isinstance(source_value, str):
            source = RecordSource(source_value)
        else:
            source = source_value

        inspection = None
        if "inspection_no" in data and data["inspection_no"]:
            inspection = self.db.query(Inspection).filter(
                Inspection.inspection_no == data["inspection_no"]
            ).first()

        rework_order = None
        if "rework_no" in data and data["rework_no"]:
            rework_order = self.db.query(ReworkOrder).filter(
                ReworkOrder.rework_no == data["rework_no"]
            ).first()

        shift = None
        if "shift_code" in data and data["shift_code"]:
            shift = self.db.query(MachineShift).filter(
                MachineShift.shift_code == data["shift_code"]
            ).first()

        if isinstance(data.get("record_date"), str):
            from dateutil import parser
            try:
                data["record_date"] = parser.parse(data["record_date"]).date()
            except:
                pass

        if isinstance(data.get("sent_at"), str):
            from dateutil import parser
            try:
                data["sent_at"] = parser.parse(data["sent_at"])
            except:
                pass

        record_no = data.get("record_no") or self._generate_draft_id("EX")
        is_draft = not data.get("record_no")

        exception = ExceptionRecord(
            record_no=record_no,
            source=source,
            inspection_id=inspection.id if inspection else None,
            rework_order_id=rework_order.id if rework_order else None,
            shift_id=shift.id if shift else None,
            record_date=data.get("record_date"),
            machine_no=data.get("machine_no", ""),
            defect_type=data.get("defect_type", ""),
            description=data.get("description", data.get("sms_content", "")),
            photo_path=data.get("photo_path", ""),
            sms_content=data.get("sms_content", ""),
            sender=data.get("sender", ""),
            sent_at=data.get("sent_at"),
            is_verified=data.get("is_verified", False),
            verified_by=data.get("verified_by", ""),
            verified_at=data.get("verified_at"),
            raw_data=to_json_serializable(data),
            is_draft=is_draft
        )

        self.db.add(exception)
        self.db.flush()
        return exception, []


class ExportService:
    def __init__(self, db: Session):
        self.db = db

    def export_queue_to_csv(self, status: Optional[str] = None) -> str:
        query = self.db.query(CompensationQueue)
        if status:
            from ..models import QueueStatus
            query = query.filter(CompensationQueue.status == QueueStatus(status))

        items = query.all()

        output = StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "ID", "状态", "重试分类", "缺陷类型", "机台号",
            "责任班次", "原始班次", "补偿金额", "补偿数量",
            "重试次数", "创建时间", "关闭时间"
        ])

        for item in items:
            writer.writerow([
                item.id,
                item.status.value if item.status else "",
                item.retry_category.value if item.retry_category else "",
                item.defect_type or "",
                item.machine_no or "",
                item.responsible_shift_code or "",
                item.original_shift_code or "",
                item.compensation_amount or 0,
                item.compensation_quantity or 0,
                item.retry_count or 0,
                item.created_at.strftime("%Y-%m-%d %H:%M:%S") if item.created_at else "",
                item.closed_at.strftime("%Y-%m-%d %H:%M:%S") if item.closed_at else ""
            ])

        return output.getvalue()

    def export_queue_to_excel(self, filepath: str, status: Optional[str] = None) -> bool:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return False

        query = self.db.query(CompensationQueue)
        if status:
            from ..models import QueueStatus
            query = query.filter(CompensationQueue.status == QueueStatus(status))

        items = query.all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "补偿队列"

        headers = [
            "ID", "状态", "重试分类", "缺陷类型", "机台号",
            "责任班次", "原始班次", "补偿金额", "补偿数量",
            "重试次数", "创建时间", "关闭时间"
        ]

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row, item in enumerate(items, 2):
            ws.cell(row=row, column=1, value=item.id)
            ws.cell(row=row, column=2, value=item.status.value if item.status else "")
            ws.cell(row=row, column=3, value=item.retry_category.value if item.retry_category else "")
            ws.cell(row=row, column=4, value=item.defect_type or "")
            ws.cell(row=row, column=5, value=item.machine_no or "")
            ws.cell(row=row, column=6, value=item.responsible_shift_code or "")
            ws.cell(row=row, column=7, value=item.original_shift_code or "")
            ws.cell(row=row, column=8, value=item.compensation_amount or 0)
            ws.cell(row=row, column=9, value=item.compensation_quantity or 0)
            ws.cell(row=row, column=10, value=item.retry_count or 0)
            ws.cell(row=row, column=11, value=item.created_at.strftime("%Y-%m-%d %H:%M:%S") if item.created_at else "")
            ws.cell(row=row, column=12, value=item.closed_at.strftime("%Y-%m-%d %H:%M:%S") if item.closed_at else "")

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15

        wb.save(filepath)
        return True
