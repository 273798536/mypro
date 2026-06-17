import os
from datetime import datetime
from typing import Dict, Any, Optional, List
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import json


def _set_chinese_font(run, font_name="宋体", size=11, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold


def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    if level == 0:
        run = p.add_run(text)
        _set_chinese_font(run, font_name="黑体", size=18, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        run = p.add_run(text)
        _set_chinese_font(run, font_name="黑体", size=15, bold=True)
    elif level == 2:
        run = p.add_run(text)
        _set_chinese_font(run, font_name="黑体", size=13, bold=True)
    else:
        run = p.add_run(text)
        _set_chinese_font(run, font_name="宋体", size=11, bold=True)
    return p


def _add_para(doc: Document, text: str, indent: bool = True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    _set_chinese_font(run)
    return p


def _add_table(doc: Document, headers: List[str], rows: List[List[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        _set_chinese_font(run, font_name="黑体", size=10, bold=True)

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            _set_chinese_font(run, size=10)
    return table


def _severity_text(s: str) -> str:
    return {"high": "严重⚠️", "medium": "中等️", "low": "轻微"}.get(s, s)


def _generate_summary_text(result: Dict) -> str:
    summary = result.get("after", {}).get("rule_check", {}).get("summary", {})
    before = result.get("before", {}).get("rule_check", {}).get("summary", {})
    orig_n = result.get("original_count", 0)
    bal_n = result.get("balanced_count", 0)
    dropped = result.get("drop_reason_count", 0)

    lines = []
    lines.append(f"本次共处理样本 {orig_n} 条，经过清洗和平衡后保留 {bal_n} 条。")
    if dropped > 0:
        lines.append(f"其中自动剔除了 {dropped} 条有严重问题的样本（比如没填标签、重复了一模一样的内容等）。")

    h_before = before.get("severity_breakdown", {}).get("high", 0)
    h_after = summary.get("severity_breakdown", {}).get("high", 0)
    if h_before > h_after:
        lines.append(f"最严重的问题从 {h_before} 个降到了 {h_after} 个，情况有明显好转。")
    elif h_before == h_after and h_before == 0:
        lines.append("目前没有最严重的(high级)问题了，整体是安全的。")
    else:
        lines.append(f"目前还剩 {h_after} 个最严重的问题需要人工处理。")

    total_before = before.get("total_issues", 0)
    total_after = summary.get("total_issues", 0)
    if total_before > total_after:
        lines.append(f"所有问题加起来，一共解决了 {total_before - total_after} 个。")

    return "".join(lines)


def _format_issues_for_humans(rule_results: Dict) -> List[Dict]:
    output = []
    for code, info in rule_results.items():
        if info.get("explanations"):
            for exp in info["explanations"]:
                output.append({
                    "rule_name": info["rule_name"],
                    "rule_code": code,
                    "issue_count": info["issue_count"],
                    "human_note": exp,
                    "passed": info["passed"]
                })
    return output


def generate_word_report(
    result: Dict,
    compare_result: Optional[Dict] = None,
    dataset_name: str = "未命名数据集",
    output_dir: str = ".",
    operator: str = "user"
) -> str:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    _add_heading(doc, "多语言样本平衡器评估报告", level=0)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"数据集：{dataset_name}    |    操作人：{operator}    |    生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    _set_chinese_font(run, size=9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    _add_heading(doc, "一、总体结论", level=1)
    summary_text = _generate_summary_text(result)
    _add_para(doc, summary_text)

    after_summary = result.get("after", {}).get("rule_check", {}).get("summary", {})
    overall_pass = after_summary.get("overall_pass", False)
    verdict = "✅ 通过" if overall_pass else "⚠️ 需人工复核"
    _add_para(doc, f"综合判断：{verdict}。（说明：只要还有「严重」级别的问题就需要人看一眼，全部严重问题解决就算通过。）")

    _add_heading(doc, "核心数据一览", level=2)
    hdr = ["指标", "平衡前", "平衡后", "变化"]
    rows_data = [
        ["样本总数",
         str(result.get("original_count", 0)),
         str(result.get("balanced_count", 0)),
         f"{result.get('balanced_count', 0) - result.get('original_count', 0):+d}"],
        ["问题总数",
         str(result.get("before", {}).get("rule_check", {}).get("summary", {}).get("total_issues", 0)),
         str(after_summary.get("total_issues", 0)),
         f"{after_summary.get('total_issues', 0) - result.get('before', {}).get('rule_check', {}).get('summary', {}).get('total_issues', 0):+d}"],
        ["严重(high)问题",
         str(result.get("before", {}).get("rule_check", {}).get("summary", {}).get("severity_breakdown", {}).get("high", 0)),
         str(after_summary.get("severity_breakdown", {}).get("high", 0)),
         f"{after_summary.get('severity_breakdown', {}).get('high', 0) - result.get('before', {}).get('rule_check', {}).get('summary', {}).get('severity_breakdown', {}).get('high', 0):+d}"],
        ["中等(medium)问题",
         str(result.get("before", {}).get("rule_check", {}).get("summary", {}).get("severity_breakdown", {}).get("medium", 0)),
         str(after_summary.get("severity_breakdown", {}).get("medium", 0)),
         f"{after_summary.get('severity_breakdown', {}).get('medium', 0) - result.get('before', {}).get('rule_check', {}).get('summary', {}).get('severity_breakdown', {}).get('medium', 0):+d}"],
        ["通过的规则数",
         str(result.get("before", {}).get("rule_check", {}).get("summary", {}).get("passed_rules", 0)),
         str(after_summary.get("passed_rules", 0)),
         f"{after_summary.get('passed_rules', 0) - result.get('before', {}).get('rule_check', {}).get('summary', {}).get('passed_rules', 0):+d}"],
    ]
    _add_table(doc, hdr, rows_data)

    _add_heading(doc, "二、各条安全规则检查情况", level=1)
    _add_para(doc, "下面逐条说明每条安全规则检查到了什么问题，以及用大白话解释给大家听：")

    rule_results_before = result.get("before", {}).get("rule_check", {}).get("rule_results", {})
    rule_results_after = result.get("after", {}).get("rule_check", {}).get("rule_results", {})

    for code in sorted(set(rule_results_before.keys()) | set(rule_results_after.keys())):
        info = rule_results_after.get(code, rule_results_before.get(code, {}))
        name = info.get("rule_name", code)
        desc = info.get("description", "")
        issue_count_after = rule_results_after.get(code, {}).get("issue_count", 0)
        issue_count_before = rule_results_before.get(code, {}).get("issue_count", 0)

        _add_heading(doc, f"规则：{name}（{code}）", level=2)
        if desc:
            p = doc.add_paragraph()
            run = p.add_run(f"规则说明：{desc}")
            _set_chinese_font(run, size=10)
            run.font.italic = True

        status_icon = "✅" if rule_results_after.get(code, {}).get("passed", False) else "❌"
        change_str = ""
        delta = issue_count_after - issue_count_before
        if delta < 0:
            change_str = f"（比平衡前少了 {abs(delta)} 个👍）"
        elif delta > 0:
            change_str = f"（比平衡前多了 {delta} 个⚠️）"
        else:
            change_str = "（和平衡前一样）"
        _add_para(doc, f"当前状态：{status_icon}    问题数：{issue_count_after} 个    {change_str}")

        human_notes = info.get("explanations", [])
        if human_notes:
            p = doc.add_paragraph()
            run = p.add_run("具体是怎么回事呢？")
            _set_chinese_font(run, font_name="黑体", size=11, bold=True)
            for note in human_notes:
                _add_para(doc, f"• {note}")

        issues_examples = info.get("issues", [])[:3]
        if issues_examples:
            p = doc.add_paragraph()
            run = p.add_run("举几个有代表性的例子（只列前3条）：")
            _set_chinese_font(run, font_name="黑体", size=11, bold=True)
            for idx, issue in enumerate(issues_examples):
                sev = _severity_text(issue.get("severity", "unknown"))
                desc = issue.get("description", "（无描述）")
                _add_para(doc, f"例{idx + 1}（{sev}）：{desc}")

    if compare_result is not None:
        _add_heading(doc, "三、灰度版本对比：基线 vs 候选", level=1)
        judgment = compare_result.get("judgment", "UNKNOWN")
        judgment_text = "✅ 通过（候选版本更好或没变差）" if judgment == "PASS" else "❌ 拦截（候选版本安全质量下降了，不能用）"
        _add_para(doc, f"灰度判断结论：{judgment_text}")

        reasons = compare_result.get("judgment_reasons", [])
        if reasons:
            p = doc.add_paragraph()
            run = p.add_run("判断依据：")
            _set_chinese_font(run, font_name="黑体", size=11, bold=True)
            for r in reasons:
                _add_para(doc, f"• {r}")

        sec = compare_result.get("security_comparison", {})
        _add_heading(doc, "安全拦截前后差别", level=2)
        hdr2 = ["对比项", "基线版本（旧的）", "候选版本（新的）", "差异"]
        b = sec.get("baseline", {})
        c = sec.get("candidate", {})
        rows2 = [
            ["整体通过？",
             "是✅" if b.get("overall_pass") else "否❌",
             "是✅" if c.get("overall_pass") else "否❌",
             "状态改变了" if sec.get("delta", {}).get("overall_pass_changed") else "没变化"],
            ["严重问题数",
             str(b.get("high_issues", 0)),
             str(c.get("high_issues", 0)),
             f"{sec.get('delta', {}).get('high_issues_delta', 0):+d}"],
            ["中等问题数",
             str(b.get("medium_issues", 0)),
             str(c.get("medium_issues", 0)),
             f"{sec.get('delta', {}).get('medium_issues_delta', 0):+d}"],
            ["总问题数",
             str(b.get("total_issues", 0)),
             str(c.get("total_issues", 0)),
             f"{sec.get('delta', {}).get('total_issues_delta', 0):+d}"],
        ]
        _add_table(doc, hdr2, rows2)

        diff = compare_result.get("diff", {})
        add_n = len(diff.get("samples_added", []))
        rem_n = len(diff.get("samples_removed", []))
        if add_n or rem_n:
            _add_para(doc, f"样本构成变化：候选版本新增了 {add_n} 条样本，去掉了 {rem_n} 条旧版本里有的样本。")

    _add_heading(doc, "四、标签冲突等边界问题说明", level=1)
    _add_para(doc, "下面这份是我们特意保留的「边界情况」——这些情况不常出现，但一旦出现就必须严肃处理，因为它们会直接改变模型的分类结果：")

    conflict_issues = [
        i for i in result.get("before", {}).get("rule_check", {}).get("all_issues", [])
        if i.get("type") == "LABEL_CONFLICT"
    ]
    if conflict_issues:
        _add_para(doc, f"共发现 {len(conflict_issues)} 对标签冲突样本，每一对都真的改变了结果：")
        for idx, c in enumerate(conflict_issues):
            texts = c.get("texts_preview", ["", ""])
            labels = c.get("labels", ["", ""])
            _add_para(
                doc,
                f"冲突{idx+1}：两条内容相似度 {c.get('similarity', 0):.0%}，"
                f"但一条标了「{labels[0]}」另一条标了「{labels[1]}」。"
                f"内容是：「{texts[0]}…」 vs 「{texts[1]}…」"
            )
        _add_para(doc, "为什么这些会改变结果？如果不去掉冲突对，模型同时学了「像A的文本可以是标签X也可以是标签Y」，分类边界就乱了，预测时同一个输入可能输出不同的类。")
    else:
        _add_para(doc, "（本次未发现标签冲突样本）")

    _add_heading(doc, "五、后续建议", level=1)
    advice_given = False

    if after_summary.get("severity_breakdown", {}).get("high", 0) > 0:
        _add_para(doc, "1. 优先处理「严重」问题：目前还有没解决的严重问题，建议先逐个核对样本，该补标补标、该删除删除。")
        advice_given = True

    empty_label_issues = [i for i in result.get("after", {}).get("rule_check", {}).get("all_issues", [])
                           if i.get("type") == "EMPTY_LABEL"]
    if empty_label_issues:
        _add_para(doc, f"2. 还有 {len(empty_label_issues)} 条样本没填标签，请数据标注同学补上。")
        advice_given = True

    lang_after = result.get("after", {}).get("language_distribution", {})
    total = sum(lang_after.values())
    if total > 0:
        skewed = [(k, v/total) for k, v in lang_after.items() if v/total > 0.4]
        if skewed:
            langs_str = "、".join([f"{k}({r:.0%})" for k, r in skewed])
            _add_para(doc, f"3. 语言分布还是有点偏：{langs_str}。建议后面补一些其他语言的样本。")
            advice_given = True

    label_after = result.get("after", {}).get("label_distribution", {})
    total_l = sum(label_after.values())
    if total_l > 0:
        ideal = 1.0 / len(label_after) if label_after else 0
        off_labels = [(k, v/total_l) for k, v in label_after.items() if abs(v/total_l - ideal) > 0.15]
        if off_labels:
            off_str = "、".join([f"「{k}」占{r:.0%}（理想{ideal:.0%}）" for k, r in off_labels])
            _add_para(doc, f"4. 标签分布：{off_str}。建议对少的类别多采集点样本。")
            advice_given = True

    if not advice_given:
        _add_para(doc, "目前各项指标都在合理范围内，可以直接用于模型训练。后续只要有新数据进来，记得再跑一遍这个平衡器就好。")

    filename = f"多语言样本平衡报告_{dataset_name}_{datetime.now().strftime('%Y%m%d_%H%M')}.docx"
    output_path = os.path.join(output_dir, filename)
    doc.save(output_path)
    return output_path


def _generate_advice_html(result: Dict) -> str:
    after_summary = result.get("after", {}).get("rule_check", {}).get("summary", {})
    all_issues_after = result.get("after", {}).get("rule_check", {}).get("all_issues", [])
    advice_items = []

    if after_summary.get("severity_breakdown", {}).get("high", 0) > 0:
        advice_items.append("优先处理「严重」问题：目前还有没解决的严重问题，建议先逐个核对样本，该补标补标、该删除删除。")

    empty_label_issues = [i for i in all_issues_after if i.get("type") == "EMPTY_LABEL"]
    if empty_label_issues:
        advice_items.append(f"还有 {len(empty_label_issues)} 条样本没填标签，请数据标注同学补上。")

    lang_after = result.get("after", {}).get("language_distribution", {})
    total = sum(lang_after.values())
    if total > 0:
        skewed = [(k, v/total) for k, v in lang_after.items() if v/total > 0.4]
        if skewed:
            langs_str = "、".join([f"{k}({r:.0%})" for k, r in skewed])
            advice_items.append(f"语言分布还是有点偏：{langs_str}。建议后面补一些其他语言的样本。")

    label_after = result.get("after", {}).get("label_distribution", {})
    total_l = sum(label_after.values())
    if total_l > 0 and len(label_after) > 1:
        ideal = 1.0 / len(label_after)
        off_labels = [(k, v/total_l) for k, v in label_after.items() if abs(v/total_l - ideal) > 0.15]
        if off_labels:
            off_str = "、".join([f"「{k}」占{r:.0%}（理想{ideal:.0%}）" for k, r in off_labels])
            advice_items.append(f"标签分布：{off_str}。建议对少的类别多采集点样本。")

    if not advice_items:
        advice_items.append("目前各项指标都在合理范围内，可以直接用于模型训练。后续只要有新数据进来，记得再跑一遍这个平衡器就好。")

    items_html = "".join([f"<li style='margin:6px 0;'>{a}</li>" for a in advice_items])
    return f"""
    <h2 style="color:#1565c0;margin-top:32px;">后续建议</h2>
    <ul style="background:#f0f7ff;padding:16px 16px 16px 36px;border-radius:8px;">{items_html}</ul>
    """


def generate_html_report(
    result: Dict,
    compare_result: Optional[Dict] = None,
    dataset_name: str = "未命名数据集",
    operator: str = "user"
) -> str:
    after_summary = result.get("after", {}).get("rule_check", {}).get("summary", {})
    before_summary = result.get("before", {}).get("rule_check", {}).get("summary", {})
    overall_pass = after_summary.get("overall_pass", False)
    verdict_color = "#2e7d32" if overall_pass else "#c62828"
    verdict_text = "✅ 通过" if overall_pass else "⚠️ 需人工复核"

    rows_html = ""
    data_pairs = [
        ("样本总数", result.get("original_count", 0), result.get("balanced_count", 0)),
        ("问题总数", before_summary.get("total_issues", 0), after_summary.get("total_issues", 0)),
        ("严重(high)问题", before_summary.get("severity_breakdown", {}).get("high", 0), after_summary.get("severity_breakdown", {}).get("high", 0)),
        ("中等(medium)问题", before_summary.get("severity_breakdown", {}).get("medium", 0), after_summary.get("severity_breakdown", {}).get("medium", 0)),
        ("轻微(low)问题", before_summary.get("severity_breakdown", {}).get("low", 0), after_summary.get("severity_breakdown", {}).get("low", 0)),
        ("通过的规则数", before_summary.get("passed_rules", 0), after_summary.get("passed_rules", 0)),
    ]
    for label, b_val, a_val in data_pairs:
        delta = int(a_val) - int(b_val)
        delta_color = "#2e7d32" if delta <= 0 else "#c62828"
        delta_sign = f"+{delta}" if delta > 0 else str(delta)
        rows_html += f"""
        <tr>
            <td style="padding:8px 12px;border-bottom:1px solid #e0e0e0;font-weight:500;">{label}</td>
            <td style="padding:8px 12px;border-bottom:1px solid #e0e0e0;text-align:center;">{b_val}</td>
            <td style="padding:8px 12px;border-bottom:1px solid #e0e0e0;text-align:center;">{a_val}</td>
            <td style="padding:8px 12px;border-bottom:1px solid #e0e0e0;text-align:center;color:{delta_color};font-weight:600;">{delta_sign}</td>
        </tr>
        """

    rules_html = ""
    rule_results_after = result.get("after", {}).get("rule_check", {}).get("rule_results", {})
    rule_results_before = result.get("before", {}).get("rule_check", {}).get("rule_results", {})
    for code in sorted(set(rule_results_before.keys()) | set(rule_results_after.keys())):
        info = rule_results_after.get(code, rule_results_before.get(code, {}))
        name = info.get("rule_name", code)
        desc = info.get("description", "")
        passed = rule_results_after.get(code, {}).get("passed", False)
        cnt_after = rule_results_after.get(code, {}).get("issue_count", 0)
        cnt_before = rule_results_before.get(code, {}).get("issue_count", 0)
        delta = cnt_after - cnt_before
        delta_str = ""
        if delta < 0:
            delta_str = f'<span style="color:#2e7d32;">（少了 {abs(delta)} 👍）</span>'
        elif delta > 0:
            delta_str = f'<span style="color:#c62828;">（多了 {delta} ⚠️）</span>'

        notes_html = ""
        for note in info.get("explanations", []):
            notes_html += f'<li style="margin:4px 0;line-height:1.7;">{note}</li>'

        icon = "✅" if passed else "❌"
        rules_html += f"""
        <div style="margin:16px 0;padding:16px;background:#fafafa;border-radius:8px;border-left:4px solid {'#2e7d32' if passed else '#c62828'};">
            <h3 style="margin:0 0 8px 0;font-size:15px;">规则：{name} <small style="color:#888;">（{code}）</small></h3>
            <div style="font-size:13px;color:#666;margin-bottom:8px;">{desc}</div>
            <div style="font-size:14px;margin-bottom:8px;">
                <b>状态：</b>{icon}　<b>问题数：</b>{cnt_after} 个　{delta_str}
            </div>
            {"<ul style='margin:8px 0 0 20px;padding:0;'>" + notes_html + "</ul>" if notes_html else ""}
        </div>
        """

    compare_html = ""
    if compare_result is not None:
        judgment = compare_result.get("judgment", "UNKNOWN")
        j_color = "#2e7d32" if judgment == "PASS" else "#c62828"
        j_text = "✅ 通过（候选版本更好或没变差）" if judgment == "PASS" else "❌ 拦截（候选版本安全质量下降了）"
        sec = compare_result.get("security_comparison", {})
        b = sec.get("baseline", {})
        c = sec.get("candidate", {})
        d = sec.get("delta", {})
        reasons_html = "".join([f"<li>{r}</li>" for r in compare_result.get("judgment_reasons", [])])

        rows_cmp = ""
        for lbl, bk, ck, dk in [
            ("整体通过？", "overall_pass", "overall_pass", "overall_pass_changed"),
            ("严重问题数", "high_issues", "high_issues", "high_issues_delta"),
            ("中等问题数", "medium_issues", "medium_issues", "medium_issues_delta"),
            ("总问题数", "total_issues", "total_issues", "total_issues_delta"),
        ]:
            bv = "是✅" if (bk == "overall_pass" and b.get(bk)) else (b.get(bk, 0) if bk != "overall_pass" else "否❌")
            cv = "是✅" if (ck == "overall_pass" and c.get(ck)) else (c.get(ck, 0) if ck != "overall_pass" else "否❌")
            dv = "改变了" if dk == "overall_pass_changed" else d.get(dk, 0)
            if dk != "overall_pass_changed":
                dv = f"{int(dv):+d}"
            rows_cmp += f"<tr><td style='padding:8px;border-bottom:1px solid #eee;'>{lbl}</td><td style='padding:8px;border-bottom:1px solid #eee;text-align:center;'>{bv}</td><td style='padding:8px;border-bottom:1px solid #eee;text-align:center;'>{cv}</td><td style='padding:8px;border-bottom:1px solid #eee;text-align:center;'>{dv}</td></tr>"

        compare_html = f"""
        <h2 style="color:#1565c0;margin-top:32px;">三、灰度版本对比：基线 vs 候选</h2>
        <div style="padding:16px;background:#fff8e1;border-left:4px solid #f9a825;border-radius:4px;margin:16px 0;">
            <b>灰度判断结论：</b><span style="color:{j_color};font-size:16px;font-weight:600;">{j_text}</span>
        </div>
        {"<b>判断依据：</b><ul>" + reasons_html + "</ul>" if reasons_html else ""}
        <h3 style="font-size:15px;">安全拦截前后差别</h3>
        <table style="width:100%;border-collapse:collapse;background:#fff;">
            <tr style="background:#e3f2fd;">
                <th style="padding:10px;border-bottom:2px solid #bbdefb;">对比项</th>
                <th style="padding:10px;border-bottom:2px solid #bbdefb;">基线版本</th>
                <th style="padding:10px;border-bottom:2px solid #bbdefb;">候选版本</th>
                <th style="padding:10px;border-bottom:2px solid #bbdefb;">差异</th>
            </tr>
            {rows_cmp}
        </table>
        """

    conflicts = [
        i for i in result.get("before", {}).get("rule_check", {}).get("all_issues", [])
        if i.get("type") == "LABEL_CONFLICT"
    ]
    conflicts_html = ""
    if conflicts:
        items_html = ""
        for idx, c in enumerate(conflicts):
            texts = c.get("texts_preview", ["", ""])
            labels = c.get("labels", ["", ""])
            items_html += f"""
            <div style="margin:8px 0;padding:12px;background:#fce4ec;border-radius:4px;">
                <b>冲突{idx+1}</b>：相似度 {c.get('similarity', 0):.0%}，但标签不同<br>
                文本A（标「{labels[0]}」）：{texts[0]}…<br>
                文本B（标「{labels[1]}」）：{texts[1]}…
            </div>
            """
        conflicts_html = f"""
        <h2 style="color:#1565c0;margin-top:32px;">四、标签冲突边界样本</h2>
        <p style="line-height:1.7;">共发现 <b>{len(conflicts)}</b> 对标签冲突样本，<b>每一对都会直接改变模型的判断结果</b>，必须人工核对统一。</p>
        {items_html}
        <p style="color:#666;font-size:13px;font-style:italic;">为什么重要？如果冲突样本没处理，模型会学到"相似的内容可以是任何标签"，分类边界就乱了。</p>
        """

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>多语言样本平衡器评估报告 - {dataset_name}</title>
<style>
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "PingFang SC", "Microsoft YaHei", sans-serif; line-height:1.6; color:#333; max-width:960px; margin:0 auto; padding:24px; background:#f5f5f5; }}
    .container {{ background:#fff; border-radius:12px; padding:32px; box-shadow: 0 2px 12px rgba(0,0,0,0.08); }}
    h1 {{ text-align:center; color:#1565c0; border-bottom:2px solid #e3f2fd; padding-bottom:12px; }}
    h2 {{ color:#1565c0; border-bottom:1px solid #e3f2fd; padding-bottom:8px; margin-top:32px; }}
    table {{ width:100%; border-collapse:collapse; margin:16px 0; }}
    th {{ background:#e3f2fd; padding:10px 12px; text-align:center; border-bottom:2px solid #bbdefb; }}
    td {{ padding:8px 12px; border-bottom:1px solid #eee; }}
    .meta-info {{ text-align:center; color:#888; font-size:13px; margin:8px 0 24px; }}
    .verdict {{ font-size:18px; font-weight:600; }}
    ul {{ padding-left: 20px; }}
    li {{ margin: 4px 0; }}
</style>
</head>
<body>
<div class="container">
    <h1>📊 多语言样本平衡器评估报告</h1>
    <div class="meta-info">数据集：{dataset_name}　|　操作人：{operator}　|　生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}</div>

    <h2>一、总体结论</h2>
    <div style="padding:16px;background:#e8f5e9;border-left:4px solid #43a047;border-radius:4px;margin:16px 0;">
        <p class="verdict" style="color:{verdict_color};margin:0;">综合判断：{verdict_text}</p>
        <p style="margin:8px 0 0 0;color:#555;">{_generate_summary_text(result)}</p>
    </div>

    <h3 style="font-size:15px;">核心数据一览</h3>
    <table>
        <tr><th>指标</th><th>平衡前</th><th>平衡后</th><th>变化</th></tr>
        {rows_html}
    </table>

    <h2 style="color:#1565c0;margin-top:32px;">二、各条安全规则检查情况</h2>
    {rules_html}

    {compare_html}

    {conflicts_html}

    {_generate_advice_html(result)}
</div>
</body>
</html>
"""
    return html
