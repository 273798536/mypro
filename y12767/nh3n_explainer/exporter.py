import os
from typing import List, Dict, Any

import pandas as pd

from .models import (
    MonitorRecord,
    ReagentRecord,
    AnomalyRecord,
    BatchRecord,
    ANOMALY_TYPES,
    SEVERITY_LEVELS,
)


def _severity_cn(level: str) -> str:
    return SEVERITY_LEVELS.get(level, level)


def _anomaly_cn(at: str) -> str:
    return ANOMALY_TYPES.get(at, at)


def _fmt_date(dt):
    if dt is None:
        return ""
    try:
        return dt.strftime("%Y-%m-%d")
    except Exception:
        return str(dt)


def build_anomalies_dataframe(
    anomalies: List[AnomalyRecord],
    monitors: List[MonitorRecord],
    reagents: List[ReagentRecord],
) -> pd.DataFrame:
    monitor_map = {m.record_id: m for m in monitors}
    reagent_map = {r.reagent_id: r for r in reagents}
    rows = []
    for a in anomalies:
        m = monitor_map.get(a.monitor_record_id)
        reagent_info_list = []
        for rid in a.reagent_evidence:
            r = reagent_map.get(rid)
            if r:
                reagent_info_list.append(
                    f"{r.name}(批号{r.batch_no}, {r.manufacturer})"
                )
            else:
                reagent_info_list.append(f"{rid}(台账中未找到)")

        balance_concentration = ""
        if a.balance_calc:
            bc = a.balance_calc
            balance_concentration = f"{bc.get('concentration_mg_l', '')} {bc.get('concentration_unit', '')}"

        rows.append({
            "异常编号": a.anomaly_id,
            "严重程度": _severity_cn(a.severity),
            "异常类型": _anomaly_cn(a.anomaly_type),
            "关联监测记录编号": a.monitor_record_id,
            "样品名称": m.sample_name if m else "",
            "样品编号": m.sample_id if m else "",
            "监测日期": _fmt_date(m.monitor_date) if m else "",
            "检测人员": m.operator if m else "",
            "问题描述": a.description,
            "普通话解释": a.plain_explanation,
            "关联试剂台账": "；".join(reagent_info_list),
            "配平计算浓度": balance_concentration,
            "复测建议": a.retest_suggestion or "",
            "处理意见": a.action_suggestion or "",
        })
    return pd.DataFrame(rows)


def build_monitors_dataframe(
    monitors: List[MonitorRecord],
    balance_results: Dict[str, Dict[str, Any]],
    anomalies: List[AnomalyRecord],
) -> pd.DataFrame:
    anomaly_by_monitor: Dict[str, List[AnomalyRecord]] = {}
    for a in anomalies:
        if a.monitor_record_id not in anomaly_by_monitor:
            anomaly_by_monitor[a.monitor_record_id] = []
        anomaly_by_monitor[a.monitor_record_id].append(a)

    rows = []
    for m in monitors:
        b = balance_results.get(m.record_id, {})
        a_list = anomaly_by_monitor.get(m.record_id, [])
        anomaly_types = "；".join([_anomaly_cn(a.anomaly_type) for a in a_list])
        max_severity = ""
        if a_list:
            order = {"critical": 4, "major": 3, "minor": 2, "info": 1}
            a_list_sorted = sorted(a_list, key=lambda x: order.get(x.severity, 0), reverse=True)
            max_severity = _severity_cn(a_list_sorted[0].severity)

        rows.append({
            "监测记录编号": m.record_id,
            "样品名称": m.sample_name,
            "样品编号": m.sample_id,
            "监测日期": _fmt_date(m.monitor_date),
            "空白对照值": m.blank_control_value if m.blank_control_value is not None else "",
            "空白对照单位": m.blank_control_unit or "",
            "样品检测值": m.sample_value if m.sample_value is not None else "",
            "样品单位": m.sample_unit or "",
            "校正吸光度": b.get("corrected_absorbance", ""),
            "计算浓度(mg/L)": b.get("concentration_mg_l", ""),
            "稀释倍数": b.get("dilution_factor", ""),
            "标准曲线编号": m.standard_curve_id or "",
            "检测人员": m.operator or "",
            "审核人员": m.reviewer or "",
            "使用试剂编号": "、".join(m.reagent_ids),
            "备注": m.remarks or "",
            "是否存在异常": "是" if a_list else "否",
            "最严重级别": max_severity,
            "异常类型": anomaly_types,
        })
    return pd.DataFrame(rows)


def build_reagents_dataframe(
    reagents: List[ReagentRecord],
    monitors: List[MonitorRecord],
) -> pd.DataFrame:
    monitor_by_reagent: Dict[str, List[str]] = {}
    for m in monitors:
        for rid in m.reagent_ids:
            if rid not in monitor_by_reagent:
                monitor_by_reagent[rid] = []
            monitor_by_reagent[rid].append(f"{m.sample_name}({m.sample_id})")

    from datetime import datetime
    now = datetime.now()
    rows = []
    for r in reagents:
        expiry_status = "正常"
        if r.expiry_date:
            days_left = (r.expiry_date - now).days
            if days_left < 0:
                expiry_status = f"已过期（超期{-days_left}天）"
            elif days_left <= 30:
                expiry_status = f"即将过期（剩{days_left}天）"
        used_by = "、".join(monitor_by_reagent.get(r.reagent_id, []))
        rows.append({
            "试剂编号": r.reagent_id,
            "试剂名称": r.name,
            "批号": r.batch_no,
            "生产厂家": r.manufacturer,
            "开瓶日期": _fmt_date(r.open_date),
            "有效期至": _fmt_date(r.expiry_date),
            "状态": expiry_status,
            "使用量": r.volume_used if r.volume_used is not None else "",
            "单位": r.volume_unit or "",
            "使用人": r.operator or "",
            "用于检测样品": used_by,
            "备注": r.remarks or "",
        })
    return pd.DataFrame(rows)


def build_trace_dataframe(
    anomalies: List[AnomalyRecord],
    monitors: List[MonitorRecord],
    reagents: List[ReagentRecord],
) -> pd.DataFrame:
    monitor_map = {m.record_id: m for m in monitors}
    reagent_map = {r.reagent_id: r for r in reagents}
    rows = []
    for a in anomalies:
        m = monitor_map.get(a.monitor_record_id)
        for rid in (a.reagent_evidence or []):
            r = reagent_map.get(rid)
            rows.append({
                "异常编号": a.anomaly_id,
                "严重程度": _severity_cn(a.severity),
                "异常类型": _anomaly_cn(a.anomaly_type),
                "样品名称": m.sample_name if m else "",
                "样品编号": m.sample_id if m else "",
                "监测日期": _fmt_date(m.monitor_date) if m else "",
                "检测人员": m.operator if m else "",
                "问题描述": a.description,
                "试剂编号": rid,
                "试剂名称": r.name if r else "(未找到)",
                "试剂批号": r.batch_no if r else "",
                "生产厂家": r.manufacturer if r else "",
                "试剂开瓶日期": _fmt_date(r.open_date) if r else "",
                "试剂有效期": _fmt_date(r.expiry_date) if r else "",
                "试剂使用人": r.operator if r else "",
                "试剂备注": r.remarks if r else "",
                "处理意见": a.action_suggestion or "",
            })
    return pd.DataFrame(rows)


def export_to_excel(
    batch: BatchRecord,
    monitors: List[MonitorRecord],
    reagents: List[ReagentRecord],
    anomalies: List[AnomalyRecord],
    balance_results: Dict[str, Dict[str, Any]],
    output_dir: str,
) -> str:
    os.makedirs(output_dir, exist_ok=True)
    filename = f"氨氮监测异常解释_{batch.batch_id}.xlsx"
    filepath = os.path.join(output_dir, filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        summary_rows = [
            {"项目": "批次编号", "内容": batch.batch_id},
            {"项目": "批次名称", "内容": batch.batch_name},
            {"项目": "生成时间", "内容": batch.created_at.strftime("%Y-%m-%d %H:%M:%S")},
            {"项目": "输入目录", "内容": batch.input_dir},
            {"项目": "监测记录数", "内容": len(monitors)},
            {"项目": "试剂台账数", "内容": len(reagents)},
            {"项目": "检出异常数", "内容": len(anomalies)},
        ]
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="批次概览", index=False)

        df_anom = build_anomalies_dataframe(anomalies, monitors, reagents)
        if df_anom.empty:
            df_anom = pd.DataFrame(columns=["说明"])
            df_anom.loc[0] = ["本批次未检出异常"]
        df_anom.to_excel(writer, sheet_name="异常详情", index=False)

        df_mon = build_monitors_dataframe(monitors, balance_results, anomalies)
        df_mon.to_excel(writer, sheet_name="监测记录+配平", index=False)

        df_reag = build_reagents_dataframe(reagents, monitors)
        df_reag.to_excel(writer, sheet_name="试剂台账", index=False)

        df_trace = build_trace_dataframe(anomalies, monitors, reagents)
        if df_trace.empty:
            df_trace = pd.DataFrame(columns=["说明"])
            df_trace.loc[0] = ["本批次未检出异常，无需追溯"]
        df_trace.to_excel(writer, sheet_name="异常追溯链路", index=False)

        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for col_idx, column_cells in enumerate(ws.columns, start=1):
                max_length = 20
                for cell in column_cells:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    try:
                        if cell.value:
                            length = min(max(len(str(line)) for line in str(cell.value).split("\n")) + 2, 60)
                            max_length = max(max_length, length)
                    except Exception:
                        pass
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)

            if sheet_name == "异常详情":
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                    severity_cell = row[1]
                    if severity_cell.value == "严重":
                        fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    elif severity_cell.value == "主要":
                        fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                    elif severity_cell.value in ("次要", "提示"):
                        fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    else:
                        fill = None
                    if fill:
                        for cell in row:
                            cell.fill = fill

            ws.freeze_panes = "A2"

    return filepath
