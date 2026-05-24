import io
import csv
import json
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill

from app.database import get_db
from app.models import User, ReturnApplication, LedgerRecord, DirtyRecord
from app.security import get_current_user, require_action, UserRole, RolePermission, mask_sensitive_value
from app.config import settings

router = APIRouter()


def mask_export_data(data: dict, role: UserRole) -> dict:
    if role == UserRole.SUPERVISOR:
        return data
    
    cn_to_en_fields = {
        "联系电话": "phone",
        "联系人": "contact",
        "申请人": "applicant"
    }
    
    masked = data.copy()
    for key, value in masked.items():
        if value:
            field_key = cn_to_en_fields.get(key, key)
            masked[key] = mask_sensitive_value(str(value), field_key, role)
    return masked


def application_to_dict(app: ReturnApplication) -> dict:
    ledger = app.ledger
    return {
        "申请单号": app.application_no,
        "批次号": app.batch_no,
        "SKU编码": app.sku_code,
        "SKU名称": app.sku_name,
        "供应商ID": app.supplier_id,
        "供应商名称": app.supplier_name,
        "联系人": app.supplier_contact,
        "联系电话": app.supplier_phone,
        "退供数量": app.return_quantity,
        "退供原因": app.return_reason,
        "申请日期": app.application_date.strftime("%Y-%m-%d") if app.application_date else "",
        "申请人": app.applicant,
        "仓库": app.warehouse_name,
        "状态": app.status,
        "质检照片数": len(app.inspection_photos),
        "物流回单数": len(app.logistics_receipts),
        "退款单数": len(app.refund_records),
        "盘点差异数量": ledger.inventory_diff_quantity if ledger else 0,
        "剩余货品数量": ledger.remaining_goods_quantity if ledger else 0,
        "剩余货品状态": ledger.remaining_goods_status if ledger else "",
        "创建时间": app.created_at.strftime("%Y-%m-%d %H:%M:%S")
    }


@router.get("/applications/csv")
async def export_applications_csv(
    status: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_action("export"))
):
    query = db.query(ReturnApplication)
    
    if status:
        query = query.filter(ReturnApplication.status == status)
    if start_date:
        query = query.filter(ReturnApplication.created_at >= start_date)
    if end_date:
        query = query.filter(ReturnApplication.created_at <= end_date)
    
    applications = query.all()
    
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=application_to_dict(applications[0]).keys() if applications else [])
    writer.writeheader()
    
    for app in applications:
        row = application_to_dict(app)
        row = mask_export_data(row, current_user.role)
        writer.writerow(row)
    
    output.seek(0)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"return_applications_{timestamp}.csv"
    
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/applications/excel")
async def export_applications_excel(
    status: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_action("export"))
):
    query = db.query(ReturnApplication)
    
    if status:
        query = query.filter(ReturnApplication.status == status)
    if start_date:
        query = query.filter(ReturnApplication.created_at >= start_date)
    if end_date:
        query = query.filter(ReturnApplication.created_at <= end_date)
    
    applications = query.all()
    
    wb = openpyxl.Workbook()
    
    ws1 = wb.active
    ws1.title = "退供申请"
    
    if applications:
        headers = list(application_to_dict(applications[0]).keys())
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        
        for row_idx, app in enumerate(applications, 2):
            row_data = application_to_dict(app)
            row_data = mask_export_data(row_data, current_user.role)
            for col_idx, (key, value) in enumerate(row_data.items(), 1):
                ws1.cell(row=row_idx, column=col_idx, value=value)
    
    ws2 = wb.create_sheet("脏记录")
    dirty_headers = ["申请单号", "脏记录类型", "字段名", "原始值", "当前值", "处理意见", "是否已解决", "创建时间"]
    for col, header in enumerate(dirty_headers, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    dirty_records = db.query(DirtyRecord).all()
    for row_idx, dr in enumerate(dirty_records, 2):
        ws2.cell(row=row_idx, column=1, value=dr.application.application_no if dr.application else "")
        ws2.cell(row=row_idx, column=2, value=dr.dirty_type)
        ws2.cell(row=row_idx, column=3, value=dr.field_name)
        ws2.cell(row=row_idx, column=4, value=dr.original_value)
        ws2.cell(row=row_idx, column=5, value=dr.current_value)
        ws2.cell(row=row_idx, column=6, value=dr.handling_opinion)
        ws2.cell(row=row_idx, column=7, value="是" if dr.is_resolved else "否")
        ws2.cell(row=row_idx, column=8, value=dr.created_at.strftime("%Y-%m-%d %H:%M:%S"))
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"return_applications_{timestamp}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/ledger/excel")
async def export_ledger_excel(
    is_closed: Optional[bool] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_action("export"))
):
    query = db.query(LedgerRecord)
    if is_closed is not None:
        query = query.filter(LedgerRecord.is_closed == is_closed)
    
    ledgers = query.all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "台账汇总"
    
    headers = [
        "申请单号", "SKU编码", "SKU名称", "供应商",
        "申请数量", "物流实收", "供应商确认", "盘点差异",
        "剩余货品数量", "剩余货品状态", "处理人", "是否关闭", "更新时间"
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    for row_idx, ledger in enumerate(ledgers, 2):
        app = ledger.application
        summary = ledger.summary or {}
        
        ws.cell(row=row_idx, column=1, value=app.application_no if app else "")
        ws.cell(row=row_idx, column=2, value=app.sku_code if app else "")
        ws.cell(row=row_idx, column=3, value=app.sku_name if app else "")
        ws.cell(row=row_idx, column=4, value=app.supplier_name if app else "")
        ws.cell(row=row_idx, column=5, value=summary.get("application_quantity", 0))
        ws.cell(row=row_idx, column=6, value=summary.get("logistics_received_quantity", 0))
        ws.cell(row=row_idx, column=7, value=summary.get("supplier_approved_quantity", 0))
        ws.cell(row=row_idx, column=8, value=ledger.inventory_diff_quantity)
        ws.cell(row=row_idx, column=9, value=ledger.remaining_goods_quantity)
        ws.cell(row=row_idx, column=10, value=ledger.remaining_goods_status)
        ws.cell(row=row_idx, column=11, value=ledger.remaining_handler)
        ws.cell(row=row_idx, column=12, value="是" if ledger.is_closed else "否")
        ws.cell(row=row_idx, column=13, value=ledger.updated_at.strftime("%Y-%m-%d %H:%M:%S") if ledger.updated_at else "")
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"ledger_summary_{timestamp}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
