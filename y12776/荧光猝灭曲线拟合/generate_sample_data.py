import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import os

np.random.seed(42)

out_dir = os.path.join(os.path.dirname(__file__), 'data')
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, '样例数据_2024质检_荧光猝灭.xlsx')

n_normal = 12
concentrations = np.array([0.0, 0.5, 1.0, 2.0, 4.0, 6.0, 8.0, 10.0, 12.0, 15.0, 20.0, 25.0])

k_sv_true = 0.085
i0_true = 980.0
noise = np.random.normal(0, 12.0, n_normal)
intensities = i0_true / (1 + k_sv_true * concentrations) + noise

records = []

records.append({
    '序号': 1,
    '样品批号': 'FQ-2024-0315-A',
    '材料名称': '异硫氰酸荧光素',
    '猝灭剂浓度': 0.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[0],
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-15',
    '操作员': '张三',
    '备注': ''
})

records.append({
    '序号': 2,
    '样品批号': 'FQ-2024-0315-A',
    '材料名称': '异硫氰酸荧光素',
    '猝灭剂浓度': 0.5,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[1],
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-15',
    '操作员': '张三',
    '备注': '重复检测'
})

records.append({
    '序号': 3,
    '样品批号': 'FQ-2024-0315-B',
    '材料名称': '异硫氰酸荧光素',
    '猝灭剂浓度': 1.0,
    '浓度单位': '',
    '荧光强度': intensities[2],
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-15',
    '操作员': '张三',
    '备注': ''
})

records.append({
    '序号': 4,
    '样品批号': 'FQ-2024-0316-A',
    '材料名称': '异硫氰酸荧光素',
    '猝灭剂浓度': 2.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[3],
    '反应时间': None,
    '时间单位': 'min',
    '测试日期': '2024-03-16',
    '操作员': '李四',
    '备注': '补录，昨天仪器记录丢失'
})

records.append({
    '序号': 5,
    '样品批号': 'FQ-2024-0316-B',
    '材料名称': '异硫氰酸荧光素',
    '猝灭剂浓度': 4.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[4],
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-16',
    '操作员': '李四',
    '备注': ''
})

records.append({
    '序号': 6,
    '样品批号': 'FQ-2024-0316-B',
    '材料名称': '异硫氰酸荧光素',
    '猝灭剂浓度': 4.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[4] * 0.92,
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-16',
    '操作员': '李四',
    '备注': '平行样'
})

records.append({
    '序号': 7,
    '样品批号': 'FQ-2024-0317-A',
    '材料名称': '罗丹明B',
    '猝灭剂浓度': 6.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[5],
    '反应时间': '',
    '时间单位': '',
    '测试日期': '2024-03-17',
    '操作员': '王五',
    '备注': ''
})

records.append({
    '序号': 8,
    '样品批号': 'FQ-2024-0317-B',
    '材料名称': '罗丹明B',
    '猝灭剂浓度': 8.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[6],
    '反应时间': 30,
    '时间单位': 'min',
    '测试日期': '2024-03-17',
    '操作员': '王五',
    '备注': ''
})

records.append({
    '序号': 9,
    '样品批号': 'FQ-2024-0318-A',
    '材料名称': '罗丹明B',
    '猝灭剂浓度': 10.0,
    '浓度单位': 'mmol/L',
    '荧光强度': 1560.0,
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-18',
    '操作员': '赵六',
    '备注': ''
})

records.append({
    '序号': 10,
    '样品批号': 'FQ-2024-0318-B',
    '材料名称': '罗丹明B',
    '猝灭剂浓度': 12.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[8],
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-18',
    '操作员': '赵六',
    '备注': ''
})

records.append({
    '序号': 11,
    '样品批号': 'FQ-2024-0319-A',
    '材料名称': '罗丹明6G',
    '猝灭剂浓度': 15.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[9],
    '反应时间': None,
    '时间单位': 'min',
    '测试日期': '2024-03-19',
    '操作员': '张三',
    '备注': ''
})

records.append({
    '序号': 12,
    '样品批号': 'FQ-2024-0319-B',
    '材料名称': '罗丹明6G',
    '猝灭剂浓度': 20.0,
    '浓度单位': '',
    '荧光强度': intensities[10],
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-19',
    '操作员': '张三',
    '备注': '旧表，单位后续补上'
})

records.append({
    '序号': 13,
    '样品批号': 'FQ-2024-0319-C',
    '材料名称': '罗丹明6G',
    '猝灭剂浓度': 25.0,
    '浓度单位': 'mmol/L',
    '荧光强度': intensities[11],
    '反应时间': 15,
    '时间单位': 'min',
    '测试日期': '2024-03-19',
    '操作员': '张三',
    '备注': ''
})

df = pd.DataFrame(records)

wb = Workbook()
ws = wb.active
ws.title = '荧光猝灭原始数据'

header_font = Font(bold=True, size=12, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center')
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        cell.border = thin_border
        if r_idx == 1:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        else:
            cell.alignment = Alignment(horizontal='center', vertical='center')

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 18
ws.column_dimensions['D'].width = 14
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 14
ws.column_dimensions['G'].width = 12
ws.column_dimensions['H'].width = 10
ws.column_dimensions['I'].width = 14
ws.column_dimensions['J'].width = 10
ws.column_dimensions['K'].width = 28

warning_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
warning_font = Font(color='9C0006')

bad_row = 9
for col in range(1, 12):
    ws.cell(row=bad_row + 1, column=col).fill = warning_fill
    ws.cell(row=bad_row + 1, column=col).font = warning_font

ws2 = wb.create_sheet('说明')
ws2['A1'] = '荧光猝灭测试原始数据表'
ws2['A1'].font = Font(bold=True, size=14)
ws2['A2'] = '本数据包含日常检测中常见问题样例，用于算法验证：'
ws2['A3'] = '1. 反应时间漏记（序号4、7、11）'
ws2['A4'] = '2. 批号重复（FQ-2024-0315-A、FQ-2024-0316-B）'
ws2['A5'] = '3. 浓度单位漏填（序号3、12）'
ws2['A6'] = '4. 补录数据带备注（序号4、12）'
ws2['A7'] = '5. 异常荧光强度值（序号9，约1560，明显偏离正常范围）'
ws2['A8'] = '6. 不同材料混合（异硫氰酸荧光素、罗丹明B、罗丹明6G）'
ws2['A9'] = '7. 反应时间不一致（序号8为30min，其它为15min）'
ws2['A10'] = '8. 旧表格式残留（序号12备注"旧表"）'

for r in range(2, 11):
    ws2.cell(row=r, column=1).alignment = Alignment(wrap_text=True)

ws2.column_dimensions['A'].width = 60

wb.save(out_path)
print(f'样例数据已生成: {out_path}')
print(f'共 {len(df)} 条记录')
print('问题分布：')
print('  - 反应时间漏记: 序号 4, 7, 11')
print('  - 批号重复: FQ-2024-0315-A (序号1,2), FQ-2024-0316-B (序号5,6)')
print('  - 单位漏填: 序号 3, 12')
print('  - 异常强度值: 序号 9 (1560.0)')
print('  - 补录备注: 序号 4, 12')
