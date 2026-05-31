import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional, Union
import io

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

from ..core.models import DiagnosisResult, TripRecord, AnomalyType
from ..core.config import EXPORTS_DIR
from .chart_generator import generate_charts, ANOMALY_LABELS


ANOMALY_TYPE_NAMES = {
    AnomalyType.LOW_TEMPERATURE: '低温影响',
    AnomalyType.FAST_CHARGING_EXCESS: '快充过多',
    AnomalyType.ABNORMAL_TRIP: '行程异常',
    AnomalyType.BATTERY_DEGRADATION: '电池衰减',
    AnomalyType.DRIVING_HABIT: '驾驶习惯',
}

SEVERITY_NAMES = {
    'low': '低',
    'medium': '中',
    'high': '高',
    'critical': '严重',
}


class Exporter:
    def __init__(self):
        self.export_dir = EXPORTS_DIR
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def _get_filename(self, diagnosis_id: str, extension: str) -> Path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return self.export_dir / f"diagnosis_{diagnosis_id}_{timestamp}.{extension}"

    def export_to_json(
        self,
        diagnosis_result: DiagnosisResult,
        trips: List[TripRecord]
    ) -> Path:
        filepath = self._get_filename(diagnosis_result.diagnosis_id, 'json')

        export_data = {
            'export_time': datetime.now().isoformat(),
            'diagnosis': diagnosis_result.model_dump(mode='json'),
            'trips': [t.model_dump(mode='json') for t in trips],
        }

        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, ensure_ascii=False, indent=2)

        return filepath

    def export_to_csv(
        self,
        diagnosis_result: DiagnosisResult,
        trips: List[TripRecord]
    ) -> List[Path]:
        files = []

        main_file = self._get_filename(diagnosis_result.diagnosis_id, 'csv')
        files.append(main_file)

        with open(main_file, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['电动车续航衰减诊断报告'])
            writer.writerow([])
            writer.writerow(['诊断ID', diagnosis_result.diagnosis_id])
            writer.writerow(['车辆VIN', diagnosis_result.vin])
            writer.writerow(['诊断时间', diagnosis_result.created_at.isoformat()])
            writer.writerow(['数据包ID', diagnosis_result.packet_id])
            writer.writerow(['数据质量评分', f"{diagnosis_result.data_quality_score:.2f}"])
            writer.writerow([])
            writer.writerow(['=== 续航估算结果 ==='])
            re = diagnosis_result.range_estimate
            writer.writerow(['标称续航(km)', re.nominal_range_km])
            writer.writerow(['实际估算续航(km)', re.actual_estimated_range_km])
            writer.writerow(['电池健康度(%)', re.battery_health_percent])
            writer.writerow(['基础能耗(kWh/100km)', re.base_consumption_kwh_100km])
            writer.writerow(['调整后能耗(kWh/100km)', re.adjusted_consumption_kwh_100km])
            writer.writerow(['置信度', re.confidence_score])
            writer.writerow(['估算方法', re.method])
            writer.writerow([])
            writer.writerow(['=== 异常检测结果 ==='])
            writer.writerow(['异常类型', '严重程度', '描述', '影响续航(km)', '关联行程数'])
            for anomaly in diagnosis_result.anomalies:
                writer.writerow([
                    ANOMALY_TYPE_NAMES.get(anomaly.anomaly_type, anomaly.anomaly_type.value),
                    SEVERITY_NAMES.get(anomaly.severity, anomaly.severity),
                    anomaly.description,
                    anomaly.affected_range_km,
                    len(anomaly.trip_ids)
                ])
            writer.writerow([])
            writer.writerow(['=== 改进建议 ==='])
            for i, suggestion in enumerate(diagnosis_result.recommendations, 1):
                writer.writerow([f"{i}.", suggestion])

        trips_file = self._get_filename(diagnosis_result.diagnosis_id + '_trips', 'csv')
        files.append(trips_file)

        with open(trips_file, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([
                '行程ID', '开始时间', '结束时间', '起始SOC(%)', '结束SOC(%)',
                '行驶距离(km)', '平均速度(km/h)', '平均温度(°C)',
                '7天快充次数', '来源文件', '来源行号'
            ])
            for trip in trips:
                writer.writerow([
                    trip.trip_id,
                    trip.start_time.isoformat(),
                    trip.end_time.isoformat(),
                    trip.start_soc,
                    trip.end_soc,
                    trip.distance_km,
                    trip.avg_speed_kmh,
                    trip.avg_temp_c,
                    trip.fast_charging_count_7d,
                    trip.source_file,
                    trip.source_line or ''
                ])

        return files

    def export_to_excel(
        self,
        diagnosis_result: DiagnosisResult,
        trips: List[TripRecord],
        include_charts: bool = True
    ) -> Optional[Path]:
        if not EXCEL_AVAILABLE:
            return None

        filepath = self._get_filename(diagnosis_result.diagnosis_id, 'xlsx')

        wb = Workbook()

        title_font = Font(name='微软雅黑', size=14, bold=True, color='FFFFFF')
        header_font = Font(name='微软雅黑', size=11, bold=True)
        normal_font = Font(name='微软雅黑', size=10)
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        title_fill = PatternFill(start_color='203864', end_color='203864', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

        ws_summary = wb.active
        ws_summary.title = '诊断摘要'

        ws_summary.merge_cells('A1:F1')
        cell = ws_summary['A1']
        cell.value = '电动车续航衰减诊断报告'
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = center_align

        re = diagnosis_result.range_estimate

        summary_data = [
            ['诊断ID', diagnosis_result.diagnosis_id, '', '', '', ''],
            ['车辆VIN', diagnosis_result.vin, '', '', '', ''],
            ['诊断时间', diagnosis_result.created_at.strftime('%Y-%m-%d %H:%M:%S'), '', '', '', ''],
            ['数据包ID', diagnosis_result.packet_id, '', '', '', ''],
            ['数据质量评分', f"{diagnosis_result.data_quality_score:.2f}", '', '', '', ''],
            ['', '', '', '', '', ''],
            ['标称续航', f"{re.nominal_range_km:.0f} km", '实际估算续航', f"{re.actual_estimated_range_km:.0f} km", '', ''],
            ['电池健康度', f"{re.battery_health_percent:.1f} %", '续航衰减', f"{re.nominal_range_km - re.actual_estimated_range_km:.0f} km", '', ''],
            ['基础能耗', f"{re.base_consumption_kwh_100km:.2f} kWh/100km", '调整后能耗', f"{re.adjusted_consumption_kwh_100km:.2f} kWh/100km", '', ''],
            ['置信度', f"{re.confidence_score:.2f}", '估算方法', re.method, '', ''],
        ]

        row_offset = 3
        for row_data in summary_data:
            for col, value in enumerate(row_data, 1):
                cell = ws_summary.cell(row=row_offset, column=col, value=value)
                cell.font = normal_font
                cell.alignment = left_align
                cell.border = thin_border
                if col in [1, 3]:
                    cell.font = header_font
            row_offset += 1

        if include_charts:
            charts = generate_charts(diagnosis_result, trips, to_base64=False)

            chart_row = row_offset + 2
            for chart_name, chart_path in charts.items():
                if chart_path and isinstance(chart_path, Path) and chart_path.exists():
                    try:
                        img = XLImage(chart_path)
                        img.width = 500
                        img.height = 300
                        ws_summary.add_image(img, f'A{chart_row}')
                        chart_row += 20
                    except Exception:
                        pass

        ws_anomalies = wb.create_sheet('异常检测')
        ws_anomalies.merge_cells('A1:G1')
        cell = ws_anomalies['A1']
        cell.value = '异常检测结果'
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = center_align

        headers = ['序号', '异常类型', '严重程度', '描述', '影响续航(km)', '关联行程数', '来源引用']
        for col, header in enumerate(headers, 1):
            cell = ws_anomalies.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, anomaly in enumerate(diagnosis_result.anomalies, 4):
            ws_anomalies.cell(row=row, column=1, value=row - 3).font = normal_font
            ws_anomalies.cell(row=row, column=2, value=ANOMALY_TYPE_NAMES.get(anomaly.anomaly_type, anomaly.anomaly_type.value)).font = normal_font
            ws_anomalies.cell(row=row, column=3, value=SEVERITY_NAMES.get(anomaly.severity, anomaly.severity)).font = normal_font
            ws_anomalies.cell(row=row, column=4, value=anomaly.description).font = normal_font
            ws_anomalies.cell(row=row, column=5, value=anomaly.affected_range_km).font = normal_font
            ws_anomalies.cell(row=row, column=6, value=len(anomaly.trip_ids)).font = normal_font
            ws_anomalies.cell(row=row, column=7, value='; '.join(anomaly.source_refs)).font = normal_font

            for col in range(1, 8):
                cell = ws_anomalies.cell(row=row, column=col)
                cell.alignment = center_align if col != 4 else left_align
                cell.border = thin_border

        ws_factors = wb.create_sheet('因素分解')
        ws_factors.merge_cells('A1:E1')
        cell = ws_factors['A1']
        cell.value = '续航衰减因素分解'
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = center_align

        headers = ['序号', '影响因素', '贡献度(%)', '影响续航(km)', '证据摘要']
        for col, header in enumerate(headers, 1):
            cell = ws_factors.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, factor in enumerate(diagnosis_result.factor_breakdown, 4):
            ws_factors.cell(row=row, column=1, value=row - 3).font = normal_font
            ws_factors.cell(row=row, column=2, value=ANOMALY_TYPE_NAMES.get(factor.factor, factor.factor.value)).font = normal_font
            ws_factors.cell(row=row, column=3, value=f"{factor.contribution_percent:.1f}%").font = normal_font
            ws_factors.cell(row=row, column=4, value=f"{factor.impact_range_km:.1f} km").font = normal_font
            ws_factors.cell(row=row, column=5, value=factor.evidence[0] if factor.evidence else '').font = normal_font

            for col in range(1, 6):
                cell = ws_factors.cell(row=row, column=col)
                cell.alignment = center_align if col != 5 else left_align
                cell.border = thin_border

        ws_suggestions = wb.create_sheet('改进建议')
        ws_suggestions.merge_cells('A1:B1')
        cell = ws_suggestions['A1']
        cell.value = '改进建议'
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = center_align

        headers = ['序号', '建议内容']
        for col, header in enumerate(headers, 1):
            cell = ws_suggestions.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, suggestion in enumerate(diagnosis_result.recommendations, 4):
            ws_suggestions.cell(row=row, column=1, value=row - 3).font = normal_font
            ws_suggestions.cell(row=row, column=2, value=suggestion).font = normal_font

            for col in range(1, 3):
                cell = ws_suggestions.cell(row=row, column=col)
                cell.alignment = center_align if col == 1 else left_align
                cell.border = thin_border

        ws_trips = wb.create_sheet('行程明细')
        ws_trips.merge_cells('A1:J1')
        cell = ws_trips['A1']
        cell.value = '行程数据明细'
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = center_align

        headers = [
            '行程ID', '开始时间', '结束时间', '起始SOC(%)', '结束SOC(%)',
            '行驶距离(km)', '平均速度(km/h)', '平均温度(°C)',
            '7天快充次数', '数据来源'
        ]
        for col, header in enumerate(headers, 1):
            cell = ws_trips.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, trip in enumerate(trips, 4):
            ws_trips.cell(row=row, column=1, value=trip.trip_id).font = normal_font
            ws_trips.cell(row=row, column=2, value=trip.start_time.strftime('%Y-%m-%d %H:%M')).font = normal_font
            ws_trips.cell(row=row, column=3, value=trip.end_time.strftime('%Y-%m-%d %H:%M')).font = normal_font
            ws_trips.cell(row=row, column=4, value=trip.start_soc).font = normal_font
            ws_trips.cell(row=row, column=5, value=trip.end_soc).font = normal_font
            ws_trips.cell(row=row, column=6, value=trip.distance_km).font = normal_font
            ws_trips.cell(row=row, column=7, value=trip.avg_speed_kmh).font = normal_font
            ws_trips.cell(row=row, column=8, value=trip.avg_temp_c).font = normal_font
            ws_trips.cell(row=row, column=9, value=trip.fast_charging_count_7d).font = normal_font
            ws_trips.cell(row=row, column=10, value=f"{trip.source_file}#L{trip.source_line}").font = normal_font

            for col in range(1, 11):
                cell = ws_trips.cell(row=row, column=col)
                cell.alignment = center_align
                cell.border = thin_border

        for ws in [ws_summary, ws_anomalies, ws_factors, ws_suggestions, ws_trips]:
            for col in range(1, 15):
                ws.column_dimensions[get_column_letter(col)].width = 18

        wb.save(filepath)
        return filepath

    def export(
        self,
        diagnosis_result: DiagnosisResult,
        trips: List[TripRecord],
        format: str = 'excel'
    ) -> Union[Path, List[Path], None]:
        format = format.lower()
        if format == 'json':
            return self.export_to_json(diagnosis_result, trips)
        elif format == 'csv':
            return self.export_to_csv(diagnosis_result, trips)
        elif format == 'excel':
            return self.export_to_excel(diagnosis_result, trips)
        else:
            return None


def export_result(
    diagnosis_result: DiagnosisResult,
    trips: List[TripRecord],
    format: str = 'excel'
) -> Union[Path, List[Path], None]:
    exporter = Exporter()
    return exporter.export(diagnosis_result, trips, format)
