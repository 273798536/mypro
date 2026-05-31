import pandas as pd
from datetime import datetime, date
from typing import List, Dict, Optional, Union
import os
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import EXPORT_DIR, ROOM_STATUS, VERIFICATION_CATEGORY
from models import (
    Contract, StayRecord, Reschedule, Cancellation,
    VerificationSheet, DisputeNote, AuditLog, VerificationResult
)
from core import DataCleaner, DataValidator


class DataImporter:
    COLUMN_MAPPINGS = {
        'contract': {
            'contract_no': ['合同编号', '合同号', '单号', 'contract_no', 'contractNo'],
            'hotel_name': ['酒店名称', '酒店', 'hotel', 'hotel_name'],
            'room_type': ['房型', '房间类型', 'room_type', 'roomType'],
            'checkin_date': ['入住日期', '入住日', '入住时间', 'checkin', 'checkin_date'],
            'checkout_date': ['离店日期', '离店日', '离店时间', 'checkout', 'checkout_date'],
            'contracted_nights': ['房晚数', '房晚', '晚数', ' nights', 'nights'],
            'contracted_rate': ['协议价', '房价', '单价', 'rate', 'price'],
            'guest_name': ['客人姓名', '姓名', '客人', 'guest', 'guest_name'],
            'id_card': ['身份证号', '身份证', '证件号', 'id_card', 'idCard'],
            'contact_phone': ['联系电话', '电话', '手机号', 'phone', 'mobile'],
            'status': ['状态', 'status'],
            'old_remark': ['旧备注', '历史备注', '原备注', 'old_remark'],
            'remark': ['备注', '说明', 'remark']
        },
        'stay': {
            'record_no': ['入住单号', '入住号', '记录号', 'record_no'],
            'hotel_name': ['酒店名称', '酒店', 'hotel'],
            'room_type': ['房型', 'room_type'],
            'room_number': ['房号', '房间号', 'room_number'],
            'guest_name': ['客人姓名', '姓名', 'guest'],
            'id_card': ['身份证号', '身份证', 'id_card'],
            'checkin_date': ['入住日期', '实际入住日期', 'checkin'],
            'checkout_date': ['离店日期', '实际离店日期', 'checkout'],
            'actual_nights': ['实际房晚', '房晚数', '晚数', 'nights'],
            'actual_rate': ['实际房价', '房价', '单价', 'rate'],
            'is_self_booked': ['是否自营补房', '自营补房', '自营', 'self_booked'],
            'is_no_show': ['是否晚到', '晚到', 'noshow', 'no_show'],
            'old_remark': ['旧备注', '历史备注', 'old_remark'],
            'remark': ['备注', 'remark']
        },
        'reschedule': {
            'reschedule_no': ['改期单号', '改期号', '变更单号', 'reschedule_no'],
            'hotel_name': ['酒店名称', '酒店', 'hotel'],
            'guest_name': ['客人姓名', '姓名', 'guest'],
            'original_checkin': ['原入住日期', '原入住', 'original_checkin'],
            'original_checkout': ['原离店日期', '原离店', 'original_checkout'],
            'new_checkin': ['新入住日期', '新入住', '改后入住', 'new_checkin'],
            'new_checkout': ['新离店日期', '新离店', '改后离店', 'new_checkout'],
            'reschedule_nights': ['改期房晚', '房晚数', '晚数', 'nights'],
            'is_cross_week': ['是否跨周', '跨周', 'cross_week'],
            'is_revised': ['是否修正', '已修正', 'revised'],
            'operator': ['操作人', '操作员', 'operator'],
            'operate_time': ['操作时间', '改期时间', 'operate_time'],
            'old_remark': ['旧备注', '历史备注', 'old_remark'],
            'remark': ['备注', 'remark']
        },
        'cancel': {
            'cancel_no': ['取消单号', '取消号', 'cancel_no'],
            'hotel_name': ['酒店名称', '酒店', 'hotel'],
            'guest_name': ['客人姓名', '姓名', 'guest'],
            'checkin_date': ['原入住日期', '入住日期', 'checkin'],
            'checkout_date': ['原离店日期', '离店日期', 'checkout'],
            'cancel_nights': ['取消房晚', '房晚数', '晚数', 'nights'],
            'cancel_type': ['取消类型', '类型', 'cancel_type'],
            'cancel_time': ['取消时间', '操作时间', 'cancel_time'],
            'operator': ['操作人', 'operator'],
            'penalty_amount': ['罚金', '罚款金额', 'penalty'],
            'is_penalty_applied': ['是否扣罚', '扣罚', 'penalty_applied'],
            'old_remark': ['旧备注', '历史备注', 'old_remark'],
            'remark': ['备注', 'remark']
        }
    }

    def __init__(self, db: Session):
        self.db = db
        self.errors = []
        self.warnings = []

    def _map_columns(self, df_columns: List[str], mapping_type: str) -> Dict[str, str]:
        mappings = self.COLUMN_MAPPINGS.get(mapping_type, {})
        result = {}
        df_cols_lower = {str(col).strip().lower(): col for col in df_columns}
        for field, possible_names in mappings.items():
            for name in possible_names:
                name_lower = str(name).strip().lower()
                if name_lower in df_cols_lower:
                    result[field] = df_cols_lower[name_lower]
                    break
                if name in df_columns:
                    result[field] = name
                    break
        return result

    def _safe_get(self, row: pd.Series, col_name: Optional[str], default=None):
        if col_name and col_name in row.index and not pd.isna(row[col_name]):
            return row[col_name]
        return default

    def import_contracts(self, file_path: str) -> Dict:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f'文件不存在: {file_path}')
        df = pd.read_excel(file_path) if file_path.endswith('.xlsx') else pd.read_csv(file_path)
        col_map = self._map_columns(df.columns, 'contract')
        filename = os.path.basename(file_path)
        imported = 0
        skipped = 0
        for idx, row in df.iterrows():
            try:
                raw_data = DataCleaner.serialize_raw(row.to_dict())
                contract_kwargs = {
                    'source_file': filename,
                    'raw_data': raw_data
                }
                for field, col in col_map.items():
                    value = self._safe_get(row, col)
                    if field in ['checkin_date', 'checkout_date']:
                        value = DataCleaner.clean_date(value, field, idx)
                    elif field == 'contracted_nights':
                        value = DataCleaner.clean_integer(value, field, idx)
                    elif field == 'contracted_rate':
                        value = DataCleaner.clean_float(value, field, idx)
                    else:
                        value = DataCleaner.clean_string(value, field)
                    contract_kwargs[field] = value
                if not contract_kwargs.get('guest_name'):
                    self.warnings.append(f'第{idx+1}行缺少客人姓名，跳过')
                    skipped += 1
                    continue
                if not contract_kwargs.get('checkin_date'):
                    self.warnings.append(f'第{idx+1}行缺少入住日期，跳过')
                    skipped += 1
                    continue
                if contract_kwargs.get('remark'):
                    parsed = DataCleaner.parse_remark(contract_kwargs['remark'])
                    if parsed.get('extracted_info'):
                        if 'nights' in parsed['extracted_info'] and not contract_kwargs.get('contracted_nights'):
                            contract_kwargs['contracted_nights'] = parsed['extracted_info']['nights']
                if contract_kwargs.get('old_remark') and contract_kwargs.get('remark'):
                    contract_kwargs['remark'] = DataCleaner.merge_remarks(
                        contract_kwargs['old_remark'], contract_kwargs['remark']
                    )
                valid, msg = DataValidator.validate_nights(
                    contract_kwargs.get('checkin_date'),
                    contract_kwargs.get('checkout_date'),
                    contract_kwargs.get('contracted_nights')
                )
                if not valid:
                    self.warnings.append(f'第{idx+1}行: {msg}')
                if not contract_kwargs.get('contract_no'):
                    contract_kwargs['contract_no'] = DataValidator.generate_sheet_no('CT')
                existing = self.db.query(Contract).filter(
                    Contract.contract_no == contract_kwargs['contract_no']
                ).first()
                if existing:
                    for k, v in contract_kwargs.items():
                        if v is not None and k not in ['id', 'created_at']:
                            setattr(existing, k, v)
                else:
                    contract = Contract(**contract_kwargs)
                    self.db.add(contract)
                imported += 1
            except Exception as e:
                self.errors.append(f'第{idx+1}行导入失败: {str(e)}')
                skipped += 1
        self.db.commit()
        return {
            'total': len(df),
            'imported': imported,
            'skipped': skipped,
            'errors': self.errors,
            'warnings': self.warnings
        }

    def import_stays(self, file_path: str) -> Dict:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f'文件不存在: {file_path}')
        df = pd.read_excel(file_path) if file_path.endswith('.xlsx') else pd.read_csv(file_path)
        col_map = self._map_columns(df.columns, 'stay')
        filename = os.path.basename(file_path)
        imported = 0
        skipped = 0
        for idx, row in df.iterrows():
            try:
                raw_data = DataCleaner.serialize_raw(row.to_dict())
                stay_kwargs = {
                    'source_file': filename,
                    'raw_data': raw_data
                }
                for field, col in col_map.items():
                    value = self._safe_get(row, col)
                    if field in ['checkin_date', 'checkout_date']:
                        value = DataCleaner.clean_date(value, field, idx)
                    elif field in ['actual_nights']:
                        value = DataCleaner.clean_integer(value, field, idx)
                    elif field == 'actual_rate':
                        value = DataCleaner.clean_float(value, field, idx)
                    elif field in ['is_self_booked', 'is_no_show']:
                        value = DataCleaner.clean_boolean(value, field, idx)
                    else:
                        value = DataCleaner.clean_string(value, field)
                    stay_kwargs[field] = value
                if not stay_kwargs.get('guest_name'):
                    self.warnings.append(f'第{idx+1}行缺少客人姓名，跳过')
                    skipped += 1
                    continue
                if not stay_kwargs.get('checkin_date'):
                    self.warnings.append(f'第{idx+1}行缺少入住日期，跳过')
                    skipped += 1
                    continue
                if stay_kwargs.get('remark'):
                    parsed = DataCleaner.parse_remark(stay_kwargs['remark'])
                    if parsed['is_no_show']:
                        stay_kwargs['is_no_show'] = True
                    if parsed['is_self_booked']:
                        stay_kwargs['is_self_booked'] = True
                    if parsed.get('extracted_info'):
                        if 'nights' in parsed['extracted_info'] and not stay_kwargs.get('actual_nights'):
                            stay_kwargs['actual_nights'] = parsed['extracted_info']['nights']
                contract = self.db.query(Contract).filter(
                    Contract.guest_name == stay_kwargs['guest_name'],
                    Contract.checkin_date == stay_kwargs['checkin_date']
                ).first()
                if contract:
                    stay_kwargs['contract_id'] = contract.id
                if not stay_kwargs.get('record_no'):
                    stay_kwargs['record_no'] = DataValidator.generate_sheet_no('ST')
                stay = StayRecord(**stay_kwargs)
                self.db.add(stay)
                imported += 1
            except Exception as e:
                self.errors.append(f'第{idx+1}行导入失败: {str(e)}')
                skipped += 1
        self.db.commit()
        return {
            'total': len(df),
            'imported': imported,
            'skipped': skipped,
            'errors': self.errors,
            'warnings': self.warnings
        }

    def import_reschedules(self, file_path: str) -> Dict:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f'文件不存在: {file_path}')
        df = pd.read_excel(file_path) if file_path.endswith('.xlsx') else pd.read_csv(file_path)
        col_map = self._map_columns(df.columns, 'reschedule')
        filename = os.path.basename(file_path)
        imported = 0
        skipped = 0
        for idx, row in df.iterrows():
            try:
                raw_data = DataCleaner.serialize_raw(row.to_dict())
                reschedule_kwargs = {
                    'source_file': filename,
                    'raw_data': raw_data
                }
                for field, col in col_map.items():
                    value = self._safe_get(row, col)
                    if field in ['original_checkin', 'original_checkout', 'new_checkin', 'new_checkout']:
                        value = DataCleaner.clean_date(value, field, idx)
                    elif field == 'operate_time':
                        value = DataCleaner.clean_datetime(value, field, idx)
                    elif field == 'reschedule_nights':
                        value = DataCleaner.clean_integer(value, field, idx)
                    elif field in ['is_cross_week', 'is_revised']:
                        value = DataCleaner.clean_boolean(value, field, idx)
                    else:
                        value = DataCleaner.clean_string(value, field)
                    reschedule_kwargs[field] = value
                if not reschedule_kwargs.get('guest_name'):
                    self.warnings.append(f'第{idx+1}行缺少客人姓名，跳过')
                    skipped += 1
                    continue
                if not reschedule_kwargs.get('original_checkin') or not reschedule_kwargs.get('new_checkin'):
                    self.warnings.append(f'第{idx+1}行缺少日期信息，跳过')
                    skipped += 1
                    continue
                is_cross, weeks_diff = DataValidator.detect_cross_week(
                    reschedule_kwargs['original_checkin'],
                    reschedule_kwargs['new_checkin']
                )
                reschedule_kwargs['is_cross_week'] = is_cross or reschedule_kwargs.get('is_cross_week', False)
                if reschedule_kwargs.get('remark'):
                    parsed = DataCleaner.parse_remark(reschedule_kwargs['remark'])
                    if parsed.get('extracted_info'):
                        if 'nights' in parsed['extracted_info'] and not reschedule_kwargs.get('reschedule_nights'):
                            reschedule_kwargs['reschedule_nights'] = parsed['extracted_info']['nights']
                contract = self.db.query(Contract).filter(
                    Contract.guest_name == reschedule_kwargs['guest_name'],
                    Contract.checkin_date == reschedule_kwargs['original_checkin']
                ).first()
                if contract:
                    reschedule_kwargs['contract_id'] = contract.id
                if not reschedule_kwargs.get('reschedule_no'):
                    reschedule_kwargs['reschedule_no'] = DataValidator.generate_sheet_no('RS')
                reschedule = Reschedule(**reschedule_kwargs)
                self.db.add(reschedule)
                imported += 1
            except Exception as e:
                self.errors.append(f'第{idx+1}行导入失败: {str(e)}')
                skipped += 1
        self.db.commit()
        return {
            'total': len(df),
            'imported': imported,
            'skipped': skipped,
            'errors': self.errors,
            'warnings': self.warnings
        }

    def import_cancellations(self, file_path: str) -> Dict:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f'文件不存在: {file_path}')
        df = pd.read_excel(file_path) if file_path.endswith('.xlsx') else pd.read_csv(file_path)
        col_map = self._map_columns(df.columns, 'cancel')
        filename = os.path.basename(file_path)
        imported = 0
        skipped = 0
        for idx, row in df.iterrows():
            try:
                raw_data = DataCleaner.serialize_raw(row.to_dict())
                cancel_kwargs = {
                    'source_file': filename,
                    'raw_data': raw_data
                }
                for field, col in col_map.items():
                    value = self._safe_get(row, col)
                    if field in ['checkin_date', 'checkout_date']:
                        value = DataCleaner.clean_date(value, field, idx)
                    elif field == 'cancel_time':
                        value = DataCleaner.clean_datetime(value, field, idx)
                    elif field == 'cancel_nights':
                        value = DataCleaner.clean_integer(value, field, idx)
                    elif field in ['penalty_amount']:
                        value = DataCleaner.clean_float(value, field, idx)
                    elif field == 'is_penalty_applied':
                        value = DataCleaner.clean_boolean(value, field, idx)
                    else:
                        value = DataCleaner.clean_string(value, field)
                    cancel_kwargs[field] = value
                if not cancel_kwargs.get('guest_name'):
                    self.warnings.append(f'第{idx+1}行缺少客人姓名，跳过')
                    skipped += 1
                    continue
                if not cancel_kwargs.get('checkin_date'):
                    self.warnings.append(f'第{idx+1}行缺少入住日期，跳过')
                    skipped += 1
                    continue
                if cancel_kwargs.get('remark'):
                    parsed = DataCleaner.parse_remark(cancel_kwargs['remark'])
                    if parsed['is_no_show'] or '晚到' in str(cancel_kwargs.get('cancel_type', '')):
                        cancel_kwargs['cancel_type'] = 'NO_SHOW'
                if not cancel_kwargs.get('cancel_type'):
                    cancel_kwargs['cancel_type'] = 'NORMAL'
                contract = self.db.query(Contract).filter(
                    Contract.guest_name == cancel_kwargs['guest_name'],
                    Contract.checkin_date == cancel_kwargs['checkin_date']
                ).first()
                if contract:
                    cancel_kwargs['contract_id'] = contract.id
                if not cancel_kwargs.get('cancel_no'):
                    cancel_kwargs['cancel_no'] = DataValidator.generate_sheet_no('CN')
                cancel = Cancellation(**cancel_kwargs)
                self.db.add(cancel)
                imported += 1
            except Exception as e:
                self.errors.append(f'第{idx+1}行导入失败: {str(e)}')
                skipped += 1
        self.db.commit()
        return {
            'total': len(df),
            'imported': imported,
            'skipped': skipped,
            'errors': self.errors,
            'warnings': self.warnings
        }


class ExcelExporter:
    HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
    WARN_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    DISPUTE_FILL = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    REVISED_FILL = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    CROSS_WEEK_FILL = PatternFill(start_color='FFD699', end_color='FFD699', fill_type='solid')
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    def __init__(self):
        self.export_dir = EXPORT_DIR

    def _apply_style(self, cell, is_header=False, highlight_type=None):
        cell.border = self.THIN_BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if is_header:
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
        elif highlight_type == 'warning':
            cell.fill = self.WARN_FILL
        elif highlight_type == 'dispute':
            cell.fill = self.DISPUTE_FILL
        elif highlight_type == 'revised':
            cell.fill = self.REVISED_FILL
        elif highlight_type == 'cross_week':
            cell.fill = self.CROSS_WEEK_FILL

    def export_verification_report(self, verify_result: Dict,
                                    filename: str = None) -> str:
        if not filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f'核销报告_{timestamp}.xlsx'
        filepath = os.path.join(self.export_dir, filename)
        wb = Workbook()
        ws1 = wb.active
        ws1.title = '核销明细'
        results = verify_result['results']
        if results:
            sample = results[0].to_dict()
            headers = list(sample.keys())
            for col, header in enumerate(headers, 1):
                cell = ws1.cell(row=1, column=col, value=header)
                self._apply_style(cell, is_header=True)
            for row_idx, result in enumerate(results, 2):
                row_data = result.to_dict()
                highlight = None
                if result.is_revised:
                    highlight = 'revised'
                elif result.is_cross_week:
                    highlight = 'cross_week'
                elif result.has_dispute:
                    highlight = 'dispute'
                for col, key in enumerate(headers, 1):
                    cell = ws1.cell(row=row_idx, column=col, value=row_data.get(key, ''))
                    self._apply_style(cell, highlight_type=highlight if col <= 10 else None)
                    if key == '房晚差异' and row_data.get(key, 0) != 0:
                        self._apply_style(cell, highlight_type='warning')
            for col in range(1, len(headers) + 1):
                ws1.column_dimensions[get_column_letter(col)].width = 15
            ws1.column_dimensions['B'].width = 12
            ws1.column_dimensions['C'].width = 15
            ws1.column_dimensions['D'].width = 12
            ws1.column_dimensions['E'].width = 20
            ws1.column_dimensions['F'].width = 12
            ws1.column_dimensions['G'].width = 12
            ws1.freeze_panes = 'A2'
        ws2 = wb.create_sheet('汇总统计')
        summary = verify_result['summary']
        summary_data = [
            ['项目', '数值'],
            ['合同总数', verify_result['total_contracts']],
            ['核销记录数', verify_result['total_results']],
            ['核销房晚总数', summary['total_nights']],
            ['核销金额总计', f"¥{summary['total_amount']:,.2f}"],
            ['跨周改期数', summary['cross_week_count']],
            ['修正后核销数', summary['revised_count']],
            ['晚到未住数', summary['no_show_count']],
            ['自营补房数', summary['self_booked_count']],
            ['有争议记录数', summary['dispute_count']],
            ['', ''],
            ['按核销分类统计', ''],
        ]
        for cat, stats in summary['by_category'].items():
            cat_name = VERIFICATION_CATEGORY.get(cat, cat)
            summary_data.append([
                f'  {cat_name}',
                f"{stats['count']}条 / {stats['nights']}晚 / ¥{stats['amount']:,.2f}"
            ])
        summary_data.append(['', ''])
        summary_data.append(['按状态统计', ''])
        for status, stats in summary['by_status'].items():
            status_name = ROOM_STATUS.get(status, status)
            summary_data.append([
                f'  {status_name}',
                f"{stats['count']}条 / {stats['nights']}晚"
            ])
        for row_idx, row_data in enumerate(summary_data, 1):
            for col, value in enumerate(row_data, 1):
                cell = ws2.cell(row=row_idx, column=col, value=value)
                is_header = (row_idx == 1) or (col == 1 and not str(value).startswith('  '))
                self._apply_style(cell, is_header=is_header)
        ws2.column_dimensions['A'].width = 25
        ws2.column_dimensions['B'].width = 40
        ws3 = wb.create_sheet('警告信息')
        warnings = verify_result['warnings']
        ws3.append(['序号', '时间', '警告信息', '上下文'])
        for col in range(1, 5):
            self._apply_style(ws3.cell(row=1, column=col), is_header=True)
        for idx, warn in enumerate(warnings, 1):
            ctx_str = str(warn.get('context', {}))
            ws3.append([
                idx,
                warn.get('time', '').strftime('%Y-%m-%d %H:%M:%S') if hasattr(warn.get('time'), 'strftime') else '',
                warn.get('message', ''),
                ctx_str
            ])
        for col in range(1, 5):
            ws3.column_dimensions[get_column_letter(col)].width = 20
        ws3.column_dimensions['C'].width = 50
        ws4 = wb.create_sheet('错误信息')
        errors = verify_result['errors']
        ws4.append(['序号', '时间', '错误信息', '上下文'])
        for col in range(1, 5):
            self._apply_style(ws4.cell(row=1, column=col), is_header=True)
        for idx, err in enumerate(errors, 1):
            ctx_str = str(err.get('context', {}))
            ws4.append([
                idx,
                err.get('time', '').strftime('%Y-%m-%d %H:%M:%S') if hasattr(err.get('time'), 'strftime') else '',
                err.get('message', ''),
                ctx_str
            ])
        for col in range(1, 5):
            ws4.column_dimensions[get_column_letter(col)].width = 20
        ws4.column_dimensions['C'].width = 50
        ws5 = wb.create_sheet('重复房晚')
        duplicates = verify_result['duplicates']
        ws5.append(['序号', '客人姓名', '入住日期', '涉及合同', '记录数'])
        for col in range(1, 6):
            self._apply_style(ws5.cell(row=1, column=col), is_header=True)
        for idx, dup in enumerate(duplicates, 1):
            ws5.append([
                idx,
                dup.get('guest_name', ''),
                dup.get('checkin_date', '').strftime('%Y-%m-%d') if hasattr(dup.get('checkin_date'), 'strftime') else '',
                ', '.join(dup.get('contracts', [])),
                dup.get('record_count', 0)
            ])
        for col in range(1, 6):
            ws5.column_dimensions[get_column_letter(col)].width = 18
        ws5.column_dimensions['D'].width = 40
        wb.save(filepath)
        return filepath

    def export_audit_log(self, dispute_notes: List[DisputeNote] = None,
                          db: Session = None, filename: str = None) -> str:
        if not filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f'审计日志_{timestamp}.xlsx'
        filepath = os.path.join(self.export_dir, filename)
        if db and not dispute_notes:
            dispute_notes = db.query(DisputeNote).all()
        wb = Workbook()
        ws = wb.active
        ws.title = '争议与审计'
        headers = ['争议ID', '客人姓名', '争议类型', '争议内容', '处理人',
                   '处理结果', '处理时间', '是否解决', '备注', '操作记录']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, is_header=True)
        for row_idx, note in enumerate(dispute_notes, 2):
            audit_text = ''
            for log in note.audit_logs:
                audit_text += f"[{log.created_at.strftime('%Y-%m-%d %H:%M')}] {log.operator}: {log.action}"
                if log.field_name:
                    audit_text += f" {log.field_name}: {log.old_value} -> {log.new_value}"
                if log.remark:
                    audit_text += f" ({log.remark})"
                audit_text += '\n'
            row_data = [
                note.id,
                note.stay_record.guest_name if note.stay_record else '',
                note.dispute_type,
                note.dispute_content,
                note.handler,
                note.handle_result,
                note.handle_time.strftime('%Y-%m-%d %H:%M') if note.handle_time else '',
                '是' if note.is_resolved else '否',
                note.remark or '',
                audit_text.strip()
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                if not note.is_resolved:
                    self._apply_style(cell, highlight_type='dispute')
                else:
                    self._apply_style(cell)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15
        ws.column_dimensions['D'].width = 40
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 30
        ws.column_dimensions['J'].width = 60
        ws.freeze_panes = 'A2'
        wb.save(filepath)
        return filepath
