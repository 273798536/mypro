from datetime import datetime
from typing import List, Optional
from io import BytesIO
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.models.enums import RecordStatus, DirtyType
from app.models.models import Batch, AppointmentOrder, UserReview, DirtyRecord, User, StatusLog, ExternalReceipt
from app.schemas.schemas import ExportSummaryItem, SummaryStats


class ExportService:
    def __init__(self, db: Session):
        self.db = db

    def get_summary_stats(self, region: Optional[str] = None, use_cache: bool = True) -> SummaryStats:
        query = self.db.query(Batch)
        if region:
            query = query.filter(Batch.region == region)
        
        batches = query.all()
        
        total_orders = 0
        total_quantity = 0
        rescheduled_count = 0
        second_visit_count = 0
        bad_review_not_found_count = 0
        dirty_record_count = 0
        resolved_dirty_count = 0
        quantity_conflict_count = 0

        for batch in batches:
            if use_cache and batch.summary_cache:
                cache = batch.summary_cache
                total_orders += cache.get('total_orders', 0)
                total_quantity += cache.get('total_quantity', 0)
                rescheduled_count += cache.get('rescheduled_count', 0)
                second_visit_count += cache.get('second_visit_count', 0)
                bad_review_not_found_count += cache.get('bad_review_not_found_count', 0)
                dirty_record_count += cache.get('dirty_record_count', 0)
                resolved_dirty_count += cache.get('resolved_dirty_count', 0)
                quantity_conflict_count += cache.get('quantity_conflict_count', 0)
            else:
                orders = self.db.query(AppointmentOrder).filter(AppointmentOrder.batch_id == batch.id).all()
                total_orders += len(orders)
                total_quantity += sum(o.quantity or 0 for o in orders)
                rescheduled_count += sum(1 for o in orders if o.is_rescheduled)
                second_visit_count += sum(1 for o in orders if o.is_second_visit)

                reviews = self.db.query(UserReview).filter(UserReview.batch_id == batch.id).all()
                bad_review_not_found_count += sum(
                    1 for r in reviews 
                    if r.rating and r.rating <= 2 and not r.bad_review_found
                )

                dirty_records = self.db.query(DirtyRecord).filter(DirtyRecord.batch_id == batch.id).all()
                dirty_record_count += len(dirty_records)
                resolved_dirty_count += sum(1 for d in dirty_records if d.is_resolved)
                quantity_conflict_count += sum(1 for d in dirty_records if d.dirty_type == DirtyType.QUANTITY_CONFLICT)

        return SummaryStats(
            total_batches=len(batches),
            draft_count=sum(1 for b in batches if b.status == RecordStatus.DRAFT),
            pending_review_count=sum(1 for b in batches if b.status == RecordStatus.PENDING_REVIEW),
            reviewed_count=sum(1 for b in batches if b.status == RecordStatus.REVIEWED),
            frozen_count=sum(1 for b in batches if b.status == RecordStatus.FROZEN),
            settled_count=sum(1 for b in batches if b.status == RecordStatus.SETTLED),
            archived_count=sum(1 for b in batches if b.status == RecordStatus.ARCHIVED),
            total_orders=total_orders,
            total_quantity=total_quantity,
            rescheduled_count=rescheduled_count,
            second_visit_count=second_visit_count,
            bad_review_not_found_count=bad_review_not_found_count,
            dirty_record_count=dirty_record_count,
            resolved_dirty_count=resolved_dirty_count,
            quantity_conflict_count=quantity_conflict_count
        )

    def get_export_summary(self, region: Optional[str] = None, status: Optional[RecordStatus] = None, use_cache: bool = True) -> List[ExportSummaryItem]:
        query = self.db.query(Batch)
        if region:
            query = query.filter(Batch.region == region)
        if status:
            query = query.filter(Batch.status == status)
        
        batches = query.order_by(Batch.created_at.desc()).all()
        
        result = []
        for batch in batches:
            creator = self.db.query(User).filter(User.id == batch.created_by).first()
            
            if use_cache and batch.summary_cache:
                cache = batch.summary_cache
                total_orders = cache.get('total_orders', 0)
                total_quantity = cache.get('total_quantity', 0)
                rescheduled_count = cache.get('rescheduled_count', 0)
                second_visit_count = cache.get('second_visit_count', 0)
                bad_review_count = cache.get('bad_review_count', 0)
                bad_review_not_found_count = cache.get('bad_review_not_found_count', 0)
                dirty_record_count = cache.get('dirty_record_count', 0)
                quantity_conflict_count = cache.get('quantity_conflict_count', 0)
            else:
                orders = self.db.query(AppointmentOrder).filter(AppointmentOrder.batch_id == batch.id).all()
                reviews = self.db.query(UserReview).filter(UserReview.batch_id == batch.id).all()
                dirty_records = self.db.query(DirtyRecord).filter(DirtyRecord.batch_id == batch.id).all()
                
                total_orders = len(orders)
                total_quantity = sum(o.quantity or 0 for o in orders)
                rescheduled_count = sum(1 for o in orders if o.is_rescheduled)
                second_visit_count = sum(1 for o in orders if o.is_second_visit)
                bad_review_count = sum(1 for r in reviews if r.rating and r.rating <= 2)
                bad_review_not_found_count = sum(
                    1 for r in reviews 
                    if r.rating and r.rating <= 2 and not r.bad_review_found
                )
                dirty_record_count = len(dirty_records)
                quantity_conflict_count = sum(1 for d in dirty_records if d.dirty_type == DirtyType.QUANTITY_CONFLICT)
            
            latest_log = self.db.query(StatusLog).filter(
                StatusLog.batch_id == batch.id
            ).order_by(StatusLog.operation_time.desc()).first()

            item = ExportSummaryItem(
                batch_no=batch.batch_no,
                batch_name=batch.name or "",
                region=batch.region or "",
                status=batch.status.value,
                status_before_freeze=batch.status_before_freeze.value if batch.status_before_freeze else None,
                freeze_reason=batch.freeze_reason,
                manual_reason=latest_log.manual_reason if latest_log else None,
                total_orders=total_orders,
                total_quantity=total_quantity,
                rescheduled_count=rescheduled_count,
                second_visit_count=second_visit_count,
                bad_review_count=bad_review_count,
                bad_review_not_found_count=bad_review_not_found_count,
                quantity_conflict_count=quantity_conflict_count,
                dirty_record_count=dirty_record_count,
                created_at=batch.created_at,
                creator=creator.full_name if creator else ""
            )
            result.append(item)
        
        return result

    def export_to_excel(self, region: Optional[str] = None, status: Optional[RecordStatus] = None) -> BytesIO:
        summary_items = self.get_export_summary(region, status, use_cache=False)
        
        wb = Workbook()
        
        ws_summary = wb.active
        ws_summary.title = "区域售后汇总"
        
        headers = [
            "批次号", "批次名称", "区域", "状态", "冻结前状态",
            "冻结原因", "人工理由", "订单总数", "总数量",
            "改约数量", "二次上门数量", "差评数量", "差评原因未找到",
            "数量冲突数量", "脏记录数量", "创建时间", "创建人"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws_summary.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for row_idx, item in enumerate(summary_items, 2):
            ws_summary.cell(row=row_idx, column=1, value=item.batch_no)
            ws_summary.cell(row=row_idx, column=2, value=item.batch_name)
            ws_summary.cell(row=row_idx, column=3, value=item.region)
            ws_summary.cell(row=row_idx, column=4, value=item.status)
            ws_summary.cell(row=row_idx, column=5, value=item.status_before_freeze or "")
            ws_summary.cell(row=row_idx, column=6, value=item.freeze_reason or "")
            ws_summary.cell(row=row_idx, column=7, value=item.manual_reason or "")
            ws_summary.cell(row=row_idx, column=8, value=item.total_orders)
            ws_summary.cell(row=row_idx, column=9, value=item.total_quantity)
            ws_summary.cell(row=row_idx, column=10, value=item.rescheduled_count)
            ws_summary.cell(row=row_idx, column=11, value=item.second_visit_count)
            ws_summary.cell(row=row_idx, column=12, value=item.bad_review_count)
            ws_summary.cell(row=row_idx, column=13, value=item.bad_review_not_found_count)
            ws_summary.cell(row=row_idx, column=14, value=item.quantity_conflict_count)
            ws_summary.cell(row=row_idx, column=15, value=item.dirty_record_count)
            ws_summary.cell(row=row_idx, column=16, value=item.created_at.strftime("%Y-%m-%d %H:%M:%S"))
            ws_summary.cell(row=row_idx, column=17, value=item.creator)

        for col in range(1, len(headers) + 1):
            ws_summary.column_dimensions[chr(64 + col)].width = 15

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
