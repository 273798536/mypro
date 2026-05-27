import os
import json
import copy
import time
import csv
import io
from datetime import datetime


class DataHandler:

    def __init__(self, upload_dir, output_dir):
        self.upload_dir = upload_dir
        self.output_dir = output_dir
        self._trace = []

    def get_trace(self):
        return copy.deepcopy(self._trace)

    def _add_trace(self, action, detail):
        self._trace.append({
            'timestamp': datetime.now().isoformat(),
            'action': action,
            'detail': detail
        })
        if len(self._trace) > 200:
            self._trace = self._trace[-200:]

    def parse_upload(self, file_storage):
        filename = file_storage.filename
        ext = os.path.splitext(filename)[1].lower()

        self._add_trace('upload', {'filename': filename})

        content = file_storage.read()

        if ext in ('.csv',):
            return self._parse_csv(content, filename)
        elif ext in ('.xlsx', '.xls'):
            return self._parse_excel(content, filename)
        elif ext in ('.json',):
            return self._parse_json(content, filename)
        else:
            raise ValueError(f'不支持的文件格式: {ext}，请使用 CSV / Excel / JSON')

    def _parse_csv(self, content, filename):
        text = content.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        row_errors = []
        for i, row in enumerate(reader, start=2):
            parsed = self._normalize_row(row, i)
            if parsed.get('_error'):
                row_errors.append({'row': i, 'error': parsed['_error']})
            else:
                rows.append(parsed)

        if not rows:
            raise ValueError('CSV 解析后无有效行')

        return self._classify_and_package(rows, row_errors, filename)

    def _parse_excel(self, content, filename):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except ImportError:
            raise ValueError('缺少 openpyxl 库，无法解析 Excel 文件')

        rows = []
        row_errors = []
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        for i, excel_row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row_dict = {}
            for j, val in enumerate(excel_row):
                key = str(headers[j]).strip() if j < len(headers) and headers[j] else f'col_{j}'
                row_dict[key] = val
            parsed = self._normalize_row(row_dict, i)
            if parsed.get('_error'):
                row_errors.append({'row': i, 'error': parsed['_error']})
            else:
                rows.append(parsed)

        if not rows:
            raise ValueError('Excel 解析后无有效行')

        return self._classify_and_package(rows, row_errors, filename)

    def _parse_json(self, content, filename):
        try:
            data = json.loads(content.decode('utf-8'))
        except json.JSONDecodeError as e:
            raise ValueError(f'JSON 格式错误: {e}')

        if not isinstance(data, dict):
            raise ValueError('JSON 顶层必须是对象，包含 work_orders 和/或 technicians 字段')

        rows = []
        row_errors = []
        for list_key in ('work_orders', 'technicians', 'orders'):
            if list_key in data and isinstance(data[list_key], list):
                for i, item in enumerate(data[list_key], start=1):
                    parsed = self._normalize_row(item, i)
                    if parsed.get('_error'):
                        row_errors.append({'row': i, 'error': parsed['_error']})
                    else:
                        rows.append(parsed)

        if not rows:
            raise ValueError('JSON 中未找到工单行')

        return self._classify_and_package(rows, row_errors, filename)

    def _normalize_row(self, row, line_no):
        def g(*keys, default=None):
            for k in keys:
                for kk in (k, k.lower(), k.upper(), k.replace('_', ''), k.replace(' ', '_')):
                    if kk in row and row[kk] not in (None, ''):
                        return row[kk]
            return default

        result = {'_line': line_no}
        if line_no is not None:
            result['_source_row'] = line_no

        name = g('工单', '工单名称', '工单号', '订单号', '名称', 'name', 'order', 'order_id')
        tech_name = g('师傅', '师傅姓名', '技师', 'technician', 'tech')
        required_skill = g('所需技能', '技能', '技能需求', 'skill', 'required_skill', 'skill_required')
        skills = g('师傅技能', '技能', '技能列表', 'skills')

        if name:
            result['type'] = 'work_order'
            result['id'] = g('工单号', '工单编号', '订单号', 'id', 'order_id', default=f'WO_{line_no}')
            result['name'] = str(name).strip()
            result['location'] = str(g('工单位置', '位置', '地点', 'location', 'address', default='未知')).strip()
            result['lat'] = self._to_float(g('纬度', 'lat', 'latitude'))
            result['lng'] = self._to_float(g('经度', 'lng', 'lon', 'longitude'))
            urgency = g('紧急等级', '紧急程度', '紧急度', 'urgency', 'priority', default='普通')
            result['urgency'] = self._normalize_urgency(urgency)
            result['required_skill'] = str(required_skill).strip() if required_skill else ''
            win_start = g('预约开始', '开始时间', '预约窗口开始', 'window_start', 'appointment_start')
            win_end = g('预约结束', '结束时间', '预约窗口结束', 'window_end', 'appointment_end')
            result['window_start'] = self._parse_time(win_start)
            result['window_end'] = self._parse_time(win_end)
            result['service_time'] = self._to_float(g('服务时长', '工时', 'service_time', 'duration'), default=1.0)
            result['traffic_time'] = self._to_float(g('交通时间', 'travel_time', '交通'), default=0)
            result['assigned_tech'] = str(tech_name).strip() if tech_name else ''
            return result

        if skills or tech_name:
            result['type'] = 'technician'
            result['id'] = g('师傅编号', '技师编号', 'id', 'tech_id', default=f'T{line_no}')
            result['name'] = str(tech_name).strip() if tech_name else f'师傅_{line_no}'
            skills_raw = g('师傅技能', '技能', '技能列表', 'skills', default='')
            result['skills'] = [s.strip() for s in str(skills_raw).replace('，', ',').split(',') if s.strip()]
            result['location'] = str(g('师傅位置', '位置', 'location', default='未知')).strip()
            result['lat'] = self._to_float(g('师傅纬度', 'lat', 'latitude'))
            result['lng'] = self._to_float(g('师傅经度', 'lng', 'lon', 'longitude'))
            return result

        return {'_error': '未识别为工单或师傅，请检查列名'}

    def _normalize_urgency(self, val):
        v = str(val).strip()
        high = ('紧急', '高', 'high', 'h', 'urgent', 'critical', '3', '★★★', '★★★')
        mid = ('中', 'medium', 'm', 'normal', '一般', '2', '★★', '★★')
        low = ('低', 'low', 'l', '普通', '1', '★', '☆')
        if any(k in v.lower() for k in [h.lower() for h in high]):
            return 3
        if any(k in v.lower() for k in [m.lower() for m in mid]):
            return 2
        if any(k in v.lower() for k in [l.lower() for l in low]):
            return 1
        return 2

    def _parse_time(self, val):
        if val is None:
            return None
        if isinstance(val, (int, float)):
            return float(val)
        v = str(val).strip()
        if not v:
            return None
        for fmt in ('%Y-%m-%d %H:%M', '%Y-%m-%d %H:%M:%S', '%H:%M', '%H:%M:%S',
                     '%Y/%m/%d %H:%M', '%m/%d %H:%M'):
            try:
                dt = datetime.strptime(v, fmt)
                if '%Y' not in fmt:
                    return dt.hour * 60 + dt.minute
                return dt.timestamp()
            except ValueError:
                continue
        try:
            return float(v)
        except ValueError:
            return None

    def _to_float(self, val, default=None):
        if val is None:
            return default
        try:
            return float(val)
        except (ValueError, TypeError):
            return default

    def _classify_and_package(self, rows, row_errors, filename):
        work_orders = [r for r in rows if r.get('type') == 'work_order']
        technicians = [r for r in rows if r.get('type') == 'technician']

        if not work_orders and not technicians:
            raise ValueError(f'文件 {filename} 解析后无工单或师傅数据')

        return {
            'work_orders': work_orders,
            'technicians': technicians,
            'row_count': len(rows),
            'errors': row_errors,
            'source_file': filename
        }

    def get_sample_data(self):
        return {
            'work_orders': [
                {'id': 'WO-001', 'name': '空调维修-东区A栋', 'location': '东区A栋301',
                 'lat': 31.2304, 'lng': 121.4737, 'urgency': 3, 'required_skill': '空调维修',
                 'window_start': 8 * 60, 'window_end': 11 * 60, 'service_time': 1.5,
                 'traffic_time': 0.3, 'assigned_tech': ''},
                {'id': 'WO-002', 'name': '水管维修-西区B楼', 'location': '西区B楼205',
                 'lat': 31.2404, 'lng': 121.4637, 'urgency': 2, 'required_skill': '水管维修',
                 'window_start': 9 * 60, 'window_end': 13 * 60, 'service_time': 2.0,
                 'traffic_time': 0.5, 'assigned_tech': ''},
                {'id': 'WO-003', 'name': '电路检修-中区C楼', 'location': '中区C楼108',
                 'lat': 31.2254, 'lng': 121.4807, 'urgency': 1, 'required_skill': '电路检修',
                 'window_start': 14 * 60, 'window_end': 17 * 60, 'service_time': 1.0,
                 'traffic_time': 0.4, 'assigned_tech': ''},
                {'id': 'WO-004', 'name': '空调保养-东区D栋', 'location': '东区D栋502',
                 'lat': 31.2284, 'lng': 121.4757, 'urgency': 2, 'required_skill': '空调维修',
                 'window_start': 10 * 60, 'window_end': 14 * 60, 'service_time': 1.0,
                 'traffic_time': 0.2, 'assigned_tech': ''},
                {'id': 'WO-005', 'name': '紧急抢修-西区E楼', 'location': '西区E楼地下室',
                 'lat': 31.2454, 'lng': 121.4587, 'urgency': 3, 'required_skill': '水管维修',
                 'window_start': 8 * 60, 'window_end': 10 * 60, 'service_time': 0.5,
                 'traffic_time': 0.6, 'assigned_tech': ''},
                {'id': 'WO-006', 'name': '弱电安装-南区F楼', 'location': '南区F楼1503',
                 'lat': 31.2154, 'lng': 121.4887, 'urgency': 1, 'required_skill': '电路检修',
                 'window_start': 13 * 60, 'window_end': 16 * 60, 'service_time': 2.0,
                 'traffic_time': 0.8, 'assigned_tech': ''}
            ],
            'technicians': [
                {'id': 'T001', 'name': '张师傅', 'skills': ['空调维修', '空调保养'],
                 'location': '北区基地', 'lat': 31.2384, 'lng': 121.4817},
                {'id': 'T002', 'name': '李师傅', 'skills': ['水管维修', '紧急抢修'],
                 'location': '西区基地', 'lat': 31.2404, 'lng': 121.4607},
                {'id': 'T003', 'name': '王师傅', 'skills': ['电路检修', '弱电安装'],
                 'location': '中区基地', 'lat': 31.2284, 'lng': 121.4757}
            ]
        }

    def generate_report(self, payload):
        import pandas as pd

        routes = payload.get('routes', [])
        anomalies = payload.get('anomalies', [])
        schedule = payload.get('schedule', [])
        scored = payload.get('scored_orders', [])

        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'dispatch_report_{ts}.xlsx'
        filepath = os.path.join(self.output_dir, filename)

        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            schedule_rows = []
            for entry in schedule:
                r = dict(entry)
                r.pop('route_details', None)
                schedule_rows.append(r)
            if schedule_rows:
                pd.DataFrame(schedule_rows).to_excel(writer, sheet_name='派单日程', index=False)

            route_rows = []
            for route in routes:
                for order in route.get('orders', []):
                    route_rows.append({
                        '师傅': route.get('technician_name', ''),
                        '工单号': order.get('id', ''),
                        '工单名称': order.get('name', ''),
                        '综合评分': round(order.get('composite_score', 0), 2),
                        '距离分': round(order.get('score_distance', 0), 2),
                        '紧急分': round(order.get('score_urgency', 0), 2),
                        '技能分': round(order.get('score_skill', 0), 2),
                        '预计到达': order.get('arrival_min', ''),
                        '预约开始': order.get('window_start', ''),
                        '预约结束': order.get('window_end', ''),
                        '是否超窗': order.get('window_violation', '')
                    })
            if route_rows:
                pd.DataFrame(route_rows).to_excel(writer, sheet_name='路线详情', index=False)

            anomaly_rows = []
            for a in anomalies:
                anomaly_rows.append({
                    '类型': a.get('type', ''),
                    '工单号': a.get('order_id', ''),
                    '师傅': a.get('technician_name', ''),
                    '描述': a.get('description', ''),
                    '严重程度': a.get('severity', ''),
                    '来源行号': a.get('source_row', '')
                })
            if anomaly_rows:
                pd.DataFrame(anomaly_rows).to_excel(writer, sheet_name='异常清单', index=False)

            scored_rows = []
            for s in scored:
                scored_rows.append({
                    '工单号': s.get('id', ''),
                    '工单名称': s.get('name', ''),
                    '推荐师傅': s.get('best_technician', ''),
                    '综合评分': round(s.get('best_score', 0), 2),
                    '距离分': round(s.get('best_distance', 0), 2),
                    '紧急分': round(s.get('best_urgency', 0), 2),
                    '技能分': round(s.get('best_skill', 0), 2),
                    '异常标记': ', '.join(s.get('flags', []))
                })
            if scored_rows:
                pd.DataFrame(scored_rows).to_excel(writer, sheet_name='评分汇总', index=False)

        self._add_trace('export', {'filename': filename})
        return filepath
