from typing import Dict, Any, List, Optional
from dataclasses import asdict
import json
from io import BytesIO
import zipfile

from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .database import ExperimentRecord, BatchReport, CuringTimeResult, SpectrumData, ImportAuditLog
from .import_engine import query_canonical, query_by_batch, ActionableError
from .concentration_module import build_concentration_regression, ConcentrationRegression


class ConsistencyMismatchError(Exception):
    """界面摘要与导出数据不一致时抛出"""
    pass


def verify_consistency(db: Session, record_no: str, provided_summary: Dict[str, Any]) -> bool:
    """
    核心一致性校验:
    比较从数据库加载的 canonical 数据与前端页面上展示的 summary 是否一致。
    不一致时不允许导出,防止出现"页面说通过、文件里写待确认"。
    """
    canonical = query_canonical(db, record_no)

    db_judgment = canonical["judgment_badge"]
    provided_judgment = provided_summary.get("judgment_badge") or provided_summary.get("judgment")
    if db_judgment != provided_judgment:
        raise ConsistencyMismatchError(
            f"判定结论不一致:界面显示 {provided_judgment} / 数据库 {db_judgment}。"
            f"请刷新页面重新加载最新数据后再导出。"
        )

    db_hash = canonical.get("canonical_hash")
    provided_hash = provided_summary.get("canonical_hash")
    if db_hash is not None and provided_hash is not None and db_hash != provided_hash:
        raise ConsistencyMismatchError(
            "摘要数据哈希不一致,可能数据已被更新。请刷新页面后重新导出。"
        )
    return True


def build_export_excel_single(db: Session, record_no: str,
                              verify_summary: Optional[Dict[str, Any]] = None) -> bytes:
    """
    单条记录导出:
    Sheet1 = 摘要页(含结论徽标) - 与页面顶部摘要一致
    Sheet2 = 指标明细表 - 与页面表格一致
    Sheet3 = 文字说明 - 与页面文本一致
    保证三者结论字段完全相同。
    """
    if verify_summary:
        verify_consistency(db, record_no, verify_summary)

    canonical = query_canonical(db, record_no)
    exp = db.query(ExperimentRecord).filter(ExperimentRecord.record_no == record_no).first()
    br = db.query(BatchReport).filter(BatchReport.batch_no == exp.batch_no).first()

    wb = Workbook()
    fill_pass = PatternFill("solid", fgColor="22c55e")
    fill_hold = PatternFill("solid", fgColor="f59e0b")
    fill_fail = PatternFill("solid", fgColor="ef4444")
    fill_head = PatternFill("solid", fgColor="1e3a8a")
    white_font = Font(bold=True, color="FFFFFF", size=11)
    head_font = Font(bold=True, size=11)
    bold_font = Font(bold=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws1 = wb.active
    ws1.title = "摘要"
    ws1["A1"] = "树脂固化时间判定报告 - 摘要页"
    ws1["A1"].font = Font(bold=True, size=14)
    ws1.merge_cells("A1:D1")

    rows_summary = [
        ("记录编号", canonical["record_no"]),
        ("关联批号", canonical["batch_no"]),
        ("树脂型号", br.resin_type if br else ""),
        ("生产厂家", br.manufacturer if br else ""),
        ("标称浓度", f"{br.nominal_concentration}%" if br and br.nominal_concentration else ""),
        ("操作人员", exp.operator or ""),
        ("实验日期", exp.experiment_date or ""),
        ("", ""),
        ("最终结论", canonical["judgment_badge"]),
    ]
    r = 3
    for k, v in rows_summary:
        ws1.cell(row=r, column=1, value=k).font = head_font
        ws1.cell(row=r, column=2, value=v)
        ws1.cell(row=r, column=1).border = thin_border
        ws1.cell(row=r, column=2).border = thin_border
        if k == "最终结论":
            judgment_cell = ws1.cell(row=r, column=2)
            judgment_cell.font = Font(bold=True, color="FFFFFF", size=12)
            judgment_cell.alignment = center
            if canonical["judgment_badge"] == "PASS":
                judgment_cell.fill = fill_pass
            elif canonical["judgment_badge"] == "待确认":
                judgment_cell.fill = fill_hold
            else:
                judgment_cell.fill = fill_fail
        r += 1

    metrics = canonical["summary_numbers"]
    ws1.cell(row=r + 1, column=1, value="核心指标(与界面表格一致)").font = head_font
    ws1.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
    r += 2
    for k, v in metrics.items():
        ws1.cell(row=r, column=1, value=k).font = bold_font
        ws1.cell(row=r, column=2, value=v)
        ws1.cell(row=r, column=1).border = thin_border
        ws1.cell(row=r, column=2).border = thin_border
        r += 1

    ws2 = wb.create_sheet("指标明细表")
    ws2.cell(row=1, column=1, value="指标").fill = fill_head
    ws2.cell(row=1, column=2, value="数值").fill = fill_head
    ws2.cell(row=1, column=3, value="单位").fill = fill_head
    ws2.cell(row=1, column=4, value="说明").fill = fill_head
    for c in "ABCD":
        ws2[f"{c}1"].font = white_font
        ws2[f"{c}1"].alignment = center
        ws2[f"{c}1"].border = thin_border
    for i, row in enumerate(canonical["metrics_table"], start=2):
        ws2.cell(row=i, column=1, value=row["metric"]).border = thin_border
        ws2.cell(row=i, column=2, value=row["value"]).border = thin_border
        ws2.cell(row=i, column=3, value=row["unit"]).border = thin_border
        ws2.cell(row=i, column=4, value=row["description"]).border = thin_border
        if row["metric"] == "结论":
            jc = ws2.cell(row=i, column=2)
            jc.font = Font(bold=True, color="FFFFFF")
            jc.alignment = center
            if row["value"] == "PASS":
                jc.fill = fill_pass
            elif row["value"] == "待确认":
                jc.fill = fill_hold
            else:
                jc.fill = fill_fail

    ws3 = wb.create_sheet("文字说明")
    ws3["A1"] = "文字说明(与页面文本完全一致)"
    ws3["A1"].font = head_font
    for i, line in enumerate(canonical["text_report"].split("\n"), start=2):
        ws3.cell(row=i, column=1, value=line)

    ws4 = wb.create_sheet("谱图原始数据")
    chart = canonical["chart_data"]
    ws4["A1"] = "时间(秒)"
    ws4["B1"] = "原始信号"
    ws4["C1"] = "平滑信号"
    ws4["A1"].fill = fill_head
    ws4["B1"].fill = fill_head
    ws4["C1"].fill = fill_head
    for c in "ABC":
        ws4[f"{c}1"].font = white_font
    t_series = chart["time_series"]
    s_raw = chart["signal_series"]
    s_smooth = chart.get("smoothed_series", [])
    for i in range(len(t_series)):
        ws4.cell(row=i + 2, column=1, value=t_series[i])
        ws4.cell(row=i + 2, column=2, value=s_raw[i])
        ws4.cell(row=i + 2, column=3, value=s_smooth[i] if i < len(s_smooth) else "")

    for col_letter, width in [("A", 20), ("B", 20), ("C", 20), ("D", 50)]:
        ws2.column_dimensions[col_letter].width = width
    ws1.column_dimensions["A"].width = 18
    ws1.column_dimensions["B"].width = 40
    ws3.column_dimensions["A"].width = 60

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export_excel_batch(db: Session, batch_no: str,
                             verify_summaries: Optional[Dict[str, Dict]] = None) -> bytes:
    """批次级导出:每个子记录表单独sheet,并增加批次总览sheet保证结论一致"""
    verify_summaries = verify_summaries or {}
    canonicals = query_by_batch(db, batch_no)

    for c in canonicals:
        if c["record_no"] in verify_summaries:
            verify_consistency(db, c["record_no"], verify_summaries[c["record_no"]])

    wb = Workbook()
    ws = wb.active
    ws.title = f"批次{batch_no}-总览"
    fill_head = PatternFill("solid", fgColor="1e3a8a")
    fill_pass = PatternFill("solid", fgColor="22c55e")
    fill_hold = PatternFill("solid", fgColor="f59e0b")
    fill_fail = PatternFill("solid", fgColor="ef4444")
    white_font = Font(bold=True, color="FFFFFF")

    headers = ["记录编号", "凝胶时间(s)", "玻璃化时间(s)", "完全固化时间(s)",
               "固化度(%)", "置信度(%)", "结论"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fill_head
        c.font = white_font

    for r, c in enumerate(canonicals, start=2):
        sn = c["summary_numbers"]
        ws.cell(row=r, column=1, value=c["record_no"])
        ws.cell(row=r, column=2, value=sn["gel_time"])
        ws.cell(row=r, column=3, value=sn["vitrification_time"])
        ws.cell(row=r, column=4, value=sn["full_cure_time"])
        ws.cell(row=r, column=5, value=sn["curing_degree"])
        ws.cell(row=r, column=6, value=round(sn["confidence"] * 100, 2))
        jc = ws.cell(row=r, column=7, value=c["judgment_badge"])
        jc.font = Font(bold=True, color="FFFFFF")
        if c["judgment_badge"] == "PASS":
            jc.fill = fill_pass
        elif c["judgment_badge"] == "待确认":
            jc.fill = fill_hold
        else:
            jc.fill = fill_fail

    for idx, c in enumerate(canonicals):
        sn = c["summary_numbers"]
        sheet = wb.create_sheet(c["record_no"][:31])
        sheet["A1"] = "指标"
        sheet["B1"] = "数值"
        for col in ["A", "B"]:
            sheet[f"{col}1"].fill = fill_head
            sheet[f"{col}1"].font = white_font
        data_rows = [
            ("记录编号", c["record_no"]),
            ("批号", c["batch_no"]),
            *list(sn.items()),
            ("判定方法", sn.get("method_used", "")),
        ]
        for i, (k, v) in enumerate(data_rows, start=2):
            sheet.cell(row=i, column=1, value=str(k))
            sheet.cell(row=i, column=2, value=str(v))
        sheet.cell(row=len(data_rows) + 3, column=1, value="文字说明:")
        for j, line in enumerate(c["text_report"].split("\n"), start=len(data_rows) + 4):
            sheet.cell(row=j, column=1, value=line)

    for col in range(1, 8):
        ws.column_dimensions[chr(64 + col)].width = 18

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export_regression_excel(regression: ConcentrationRegression) -> bytes:
    """导出浓度校准回归的Excel,与页面图表一致"""
    wb = Workbook()
    fill_head = PatternFill("solid", fgColor="1e3a8a")
    white_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "浓度校准曲线"
    ws["A1"] = "浓度校准回归报告"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("方法", regression.method),
        ("回归方程", regression.equation),
        ("R²", regression.r_squared),
        ("斜率", regression.slope),
        ("截距", regression.intercept),
        ("X轴", regression.x_label),
        ("Y轴", regression.y_label),
        ("有效数据点", len(regression.points)),
        ("被排除异常点", ", ".join(regression.outliers) or "无"),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)

    start = len(info) + 5
    ws.cell(row=start, column=1, value="各数据点明细(与界面散点图一致)").font = Font(bold=True)
    point_headers = ["记录编号", "批号", "X值", "Y值(参考)", "预测Y值", "残差"]
    for i, h in enumerate(point_headers, start=1):
        c = ws.cell(row=start + 1, column=i, value=h)
        c.fill = fill_head
        c.font = white_font
    for r, p in enumerate(regression.points, start=start + 2):
        ws.cell(row=r, column=1, value=p.get("record_no"))
        ws.cell(row=r, column=2, value=p.get("batch_no"))
        ws.cell(row=r, column=3, value=p.get("_x_value"))
        ws.cell(row=r, column=4, value=p.get("_y_value"))
        ws.cell(row=r, column=5, value=round(p.get("_predicted_y", 0), 6))
        ws.cell(row=r, column=6, value=round(p.get("_residual", 0), 6))

    ws2 = wb.create_sheet("拟合曲线数据")
    ws2["A1"] = regression.x_label
    ws2["B1"] = f"拟合{regression.y_label}"
    ws2["A1"].fill = fill_head
    ws2["B1"].fill = fill_head
    ws2["A1"].font = white_font
    ws2["B1"].font = white_font
    for i in range(len(regression.plot_x)):
        ws2.cell(row=i + 2, column=1, value=regression.plot_x[i])
        ws2.cell(row=i + 2, column=2, value=regression.plot_y[i])

    for col in ["A", "B", "C", "D", "E", "F"]:
        ws.column_dimensions[col].width = 20

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
