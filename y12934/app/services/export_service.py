import json
from datetime import datetime
from pathlib import Path
from typing import Tuple, Optional
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.config import settings
from app.crud import crud
from app.models.enums import QuestionStatus, IssueType


ISSUE_TYPE_LABEL = {
    IssueType.MATERIAL_MISSING: "需补材料",
    IssueType.CALIBRATION_WRONG: "需改口径",
}

STATUS_LABEL = {
    QuestionStatus.IMPORTED: "已导入",
    QuestionStatus.UNDER_REVIEW: "复核中",
    QuestionStatus.REVIEW_PASSED: "复核通过",
    QuestionStatus.REVIEW_BLOCKED: "复核受阻",
    QuestionStatus.APPROVED: "评审会通过",
    QuestionStatus.REJECTED: "评审会拦截",
}


def _safe_filename(name: str) -> str:
    return "".join([c if c.isalnum() or c in "-_." else "_" for c in name])


def generate_export_files(db: Session, batch_id: int) -> Tuple[Optional[Path], Optional[Path], str]:
    db_batch = crud.get_batch(db, batch_id)
    if not db_batch:
        return None, None, "批次不存在"

    summary = crud.get_review_summary(db, batch_id)
    issue_breakdown = crud.get_batch_issue_breakdown(db, batch_id)
    bias = crud.analyze_bias(db, batch_id)
    plain_explanation = crud.generate_plain_explanation(db, batch_id)
    rejection_explanation = crud.get_rejection_explanation(db, batch_id)
    prompt_tracks = crud.get_batch_prompt_tracks(db, batch_id)

    questions_data = []
    for q in db_batch.questions:
        last_review = None
        if q.review_records:
            last_review = sorted(q.review_records, key=lambda r: r.review_time, reverse=True)[0]

        questions_data.append({
            "id": q.id,
            "question_id_external": q.question_id_external,
            "question_content": q.question_content,
            "standard_answer": q.standard_answer,
            "difficulty": q.difficulty,
            "knowledge_point": q.knowledge_point,
            "human_note": q.human_note,
            "current_status": q.current_status.value,
            "current_status_label": STATUS_LABEL.get(q.current_status, q.current_status.value),
            "last_review_passed": last_review.passed if last_review else None,
            "last_review_issue_type": last_review.issue_type.value if last_review and last_review.issue_type else None,
            "last_review_issue_type_label": ISSUE_TYPE_LABEL.get(last_review.issue_type) if last_review and last_review.issue_type else None,
            "last_review_issue_detail": last_review.issue_detail if last_review else None,
            "last_review_next_action": last_review.next_action if last_review else None,
            "last_review_human_note_preserved": last_review.human_note_preserved if last_review else None,
            "copyright_sources": [
                {
                    "copyright_type": cs.copyright_type.value,
                    "source_title": cs.source_title,
                    "source_author": cs.source_author,
                    "source_publisher": cs.source_publisher,
                    "source_url": cs.source_url,
                    "publication_date": cs.publication_date,
                    "authorization_number": cs.authorization_number,
                    "authorization_expiry": cs.authorization_expiry,
                    "fair_use_justification": cs.fair_use_justification,
                    "remark": cs.remark,
                }
                for cs in q.copyright_sources
            ],
            "review_records": [
                {
                    "reviewer": r.reviewer,
                    "review_time": r.review_time.isoformat(),
                    "issue_type": r.issue_type.value if r.issue_type else None,
                    "issue_type_label": ISSUE_TYPE_LABEL.get(r.issue_type) if r.issue_type else None,
                    "issue_detail": r.issue_detail,
                    "next_action": r.next_action,
                    "passed": r.passed,
                    "human_note_preserved": r.human_note_preserved,
                }
                for r in q.review_records
            ],
            "status_history": [
                {
                    "from_status": s.from_status.value if s.from_status else None,
                    "to_status": s.to_status.value,
                    "operator": s.operator,
                    "operate_time": s.operate_time.isoformat(),
                    "reason": s.reason,
                }
                for s in q.status_history
            ],
        })

    status_summary = {
        "batch_status": db_batch.current_status.value,
        "batch_status_label": STATUS_LABEL.get(db_batch.current_status, db_batch.current_status.value),
        "total_questions": db_batch.total_questions,
        "review_passed": summary.total_passed,
        "review_blocked": summary.total_blocked,
        "material_missing": summary.material_missing_count,
        "calibration_wrong": summary.calibration_wrong_count,
        "pending": summary.total_pending,
    }

    issue_breakdown_data = []
    for ib in issue_breakdown:
        issue_breakdown_data.append({
            "issue_type": ib.issue_type.value,
            "issue_type_label": ib.issue_type_label,
            "count": ib.count,
            "question_ids": ib.question_ids,
            "details": ib.details,
        })

    cleaned_status_history = []
    for s in db_batch.status_history:
        reason = s.reason or ""
        if reason and "待处理" in reason:
            import re as _re
            m = _re.search(r"待处理-?\d+题", reason)
            if m:
                total = db_batch.total_questions
                mm = status_summary["material_missing"]
                cw = status_summary["calibration_wrong"]
                ps = status_summary["review_passed"]
                pd = max(0, total - mm - cw - ps)
                reason = _re.sub(r"待处理-?\d+题", f"待处理{pd}题", reason)
        cleaned_status_history.append({
            "from_status": s.from_status.value if s.from_status else None,
            "from_status_label": STATUS_LABEL.get(s.from_status) if s.from_status else None,
            "to_status": s.to_status.value,
            "to_status_label": STATUS_LABEL.get(s.to_status, s.to_status.value),
            "operator": s.operator,
            "operate_time": s.operate_time.isoformat(),
            "reason": reason,
            "_original_reason": s.reason,
        })

    total = status_summary["total_questions"]
    check_sum = status_summary["review_passed"] + status_summary["review_blocked"] + status_summary["pending"]
    if check_sum != total:
        status_summary["pending"] = max(0, total - status_summary["review_passed"] - status_summary["review_blocked"])

    full_data = {
        "batch_info": {
            "id": db_batch.id,
            "batch_name": db_batch.batch_name,
            "import_time": db_batch.import_time.isoformat(),
            "importer": db_batch.importer,
            "description": db_batch.description,
            "subject_category": db_batch.subject_category,
            "rejection_reason": db_batch.rejection_reason,
            "current_status": db_batch.current_status.value,
            "current_status_label": STATUS_LABEL.get(db_batch.current_status, db_batch.current_status.value),
        },
        "export_time": datetime.utcnow().isoformat(),
        "plain_explanation": plain_explanation,
        "status_summary": status_summary,
        "issue_breakdown": issue_breakdown_data,
        "bias_check_result": bias,
        "rejection_explanation": rejection_explanation,
        "questions": questions_data,
        "status_history": cleaned_status_history,
        "prompt_version_tracks": [
            {
                "version_code": t.prompt_version.version_code if t.prompt_version else None,
                "version_name": t.prompt_version.version_name if t.prompt_version else None,
                "bind_time": t.bind_time.isoformat(),
                "operator": t.operator,
                "question_id": t.question_id,
                "remark": t.remark,
            }
            for t in prompt_tracks
        ],
    }

    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    base_name = _safe_filename(f"{db_batch.batch_name}_{ts}")
    json_path = settings.EXPORT_DIR / f"{base_name}.json"
    excel_path = settings.EXPORT_DIR / f"{base_name}.xlsx"

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(full_data, f, ensure_ascii=False, indent=2)

    _write_excel(excel_path, full_data)

    return json_path, excel_path, ""


def _write_excel(excel_path: Path, data: dict) -> None:
    wb = Workbook()
    bold = Font(bold=True)
    red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
    green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws_summary = wb.active
    ws_summary.title = "概览"
    bi = data["batch_info"]
    ss = data["status_summary"]
    rows = [
        ["批次名称", bi["batch_name"]],
        ["批次ID", bi["id"]],
        ["学科分类", bi["subject_category"] or ""],
        ["导入人", bi["importer"] or ""],
        ["导入时间", bi["import_time"]],
        ["当前状态", bi["current_status_label"]],
        ["总题数", ss["total_questions"]],
        ["复核通过", ss["review_passed"]],
        ["复核受阻", ss["review_blocked"]],
        ["  其中：需补材料", ss["material_missing"]],
        ["  其中：需改口径", ss["calibration_wrong"]],
        ["待复核", ss["pending"]],
        [],
        ["偏科检测", data["bias_check_result"].get("bias_explanation", "")],
        [],
        ["拦截原因（如有）", data.get("rejection_explanation") or ""],
        [],
        ["普通话解释（可直接转发）", ""],
    ]
    for row in rows:
        if row:
            ws_summary.append(row)
        else:
            ws_summary.append([""])
    ws_summary.append(["", data["plain_explanation"]])

    for row_idx in range(1, ws_summary.max_row + 1):
        cell = ws_summary.cell(row=row_idx, column=1)
        cell.font = bold
        cell.alignment = wrap
        c2 = ws_summary.cell(row=row_idx, column=2)
        c2.alignment = wrap

    ws_summary.column_dimensions["A"].width = 25
    ws_summary.column_dimensions["B"].width = 80

    ws_issues = wb.create_sheet("异常明细")
    headers = ["异常类型", "题ID", "题目内容", "问题描述", "下一步动作", "人工备注（原话保留）"]
    ws_issues.append(headers)
    for c_idx in range(1, len(headers) + 1):
        cell = ws_issues.cell(row=1, column=c_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    row = 2
    for ib in data["issue_breakdown"]:
        fill = red_fill if ib["issue_type"] == "material_missing" else yellow_fill
        for d in ib["details"]:
            ws_issues.cell(row=row, column=1, value=ib["issue_type_label"]).fill = fill
            ws_issues.cell(row=row, column=2, value=d["question_id"])
            ws_issues.cell(row=row, column=3, value=d["question_content"]).alignment = wrap
            ws_issues.cell(row=row, column=4, value=d.get("issue_detail") or "").alignment = wrap
            ws_issues.cell(row=row, column=5, value=d.get("next_action") or "").alignment = wrap
            ws_issues.cell(row=row, column=6, value=d.get("human_note") or "").alignment = wrap
            for c in range(1, 7):
                ws_issues.cell(row=row, column=c).fill = fill
            row += 1

    widths = [12, 8, 60, 50, 50, 40]
    for i, w in enumerate(widths):
        ws_issues.column_dimensions[get_column_letter(i + 1)].width = w

    ws_questions = wb.create_sheet("题目清单")
    q_headers = ["题ID", "外部题号", "知识点", "难度", "题目内容", "标准答案", "当前状态", "复核结果", "异常类型", "问题描述", "下一步动作", "人工备注", "版权类型", "版权来源"]
    ws_questions.append(q_headers)
    for c_idx in range(1, len(q_headers) + 1):
        cell = ws_questions.cell(row=1, column=c_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    for i, q in enumerate(data["questions"]):
        r = i + 2
        cs_text = ""
        if q["copyright_sources"]:
            cs = q["copyright_sources"][0]
            cs_text = f"{cs.get('source_title') or ''} | {cs.get('source_author') or ''} | {cs.get('source_publisher') or ''}"

        fill = green_fill if q.get("last_review_passed") else (
            red_fill if q.get("last_review_issue_type") == "material_missing" else (
                yellow_fill if q.get("last_review_issue_type") == "calibration_wrong" else None
            )
        )

        values = [
            q["id"],
            q["question_id_external"] or "",
            q["knowledge_point"] or "",
            q["difficulty"] or "",
            q["question_content"],
            q["standard_answer"] or "",
            q["current_status_label"],
            "通过" if q.get("last_review_passed") else ("未通过" if q.get("last_review_passed") is False else "未复核"),
            q.get("last_review_issue_type_label") or "",
            q.get("last_review_issue_detail") or "",
            q.get("last_review_next_action") or "",
            q.get("last_review_human_note_preserved") or q.get("human_note") or "",
            q["copyright_sources"][0]["copyright_type"] if q["copyright_sources"] else "",
            cs_text,
        ]
        for c_idx, v in enumerate(values):
            cell = ws_questions.cell(row=r, column=c_idx + 1, value=v)
            cell.alignment = wrap
            if fill:
                cell.fill = fill

    q_widths = [6, 12, 18, 10, 60, 30, 14, 10, 12, 40, 40, 40, 12, 60]
    for i, w in enumerate(q_widths):
        ws_questions.column_dimensions[get_column_letter(i + 1)].width = w

    ws_history = wb.create_sheet("状态历史")
    h_headers = ["时间", "操作人", "从状态", "到状态", "原因"]
    ws_history.append(h_headers)
    for c_idx in range(1, len(h_headers) + 1):
        cell = ws_history.cell(row=1, column=c_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
    for i, s in enumerate(data["status_history"]):
        r = i + 2
        ws_history.cell(row=r, column=1, value=s["operate_time"]).alignment = wrap
        ws_history.cell(row=r, column=2, value=s.get("operator") or "").alignment = wrap
        ws_history.cell(row=r, column=3, value=s.get("from_status_label") or "").alignment = wrap
        ws_history.cell(row=r, column=4, value=s["to_status_label"]).alignment = wrap
        ws_history.cell(row=r, column=5, value=s.get("reason") or "").alignment = wrap
    h_widths = [22, 16, 14, 14, 60]
    for i, w in enumerate(h_widths):
        ws_history.column_dimensions[get_column_letter(i + 1)].width = w

    ws_prompt = wb.create_sheet("提示词版本")
    p_headers = ["绑定时间", "操作人", "版本号", "版本名", "绑定题目ID", "备注"]
    ws_prompt.append(p_headers)
    for c_idx in range(1, len(p_headers) + 1):
        cell = ws_prompt.cell(row=1, column=c_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
    for i, t in enumerate(data["prompt_version_tracks"]):
        r = i + 2
        ws_prompt.cell(row=r, column=1, value=t["bind_time"]).alignment = wrap
        ws_prompt.cell(row=r, column=2, value=t.get("operator") or "").alignment = wrap
        ws_prompt.cell(row=r, column=3, value=t.get("version_code") or "").alignment = wrap
        ws_prompt.cell(row=r, column=4, value=t.get("version_name") or "").alignment = wrap
        ws_prompt.cell(row=r, column=5, value=str(t.get("question_id") or "（整批）")).alignment = wrap
        ws_prompt.cell(row=r, column=6, value=t.get("remark") or "").alignment = wrap
    p_widths = [22, 16, 14, 22, 14, 40]
    for i, w in enumerate(p_widths):
        ws_prompt.column_dimensions[get_column_letter(i + 1)].width = w

    wb.save(excel_path)
