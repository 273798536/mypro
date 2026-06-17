import json
import sys
import os
import zipfile
import io
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient
from app.main import app
from app.config import settings
from openpyxl import load_workbook

client = TestClient(app)


def _cleanup():
    for p in settings.EXPORT_DIR.glob("*"):
        try:
            p.unlink()
        except Exception:
            pass
    try:
        Path("data/copyright_ledger.db").unlink()
    except Exception:
        pass


def test_health():
    resp = client.get("/api/health")
    assert resp.status_code == 200
    assert resp.json()["status"] == "ok"
    print("[PASS] 健康检查接口正常")


def test_full_export_consistency():
    print("\n=== 完整导出链路验证：无负数、全链路一致、文件可正常打开 ===")

    batch_payload = {
        "batch_name": "2026-Q2-数学评测集-A",
        "importer": "张评测",
        "description": "二季度数学评测题库",
        "subject_category": "数学",
        "questions": [
            {
                "question_id_external": "MATH-001",
                "question_content": "求解方程：2x + 5 = 13",
                "standard_answer": "x = 4",
                "difficulty": "简单",
                "knowledge_point": "一元一次方程",
                "human_note": "这题去年也考过，学生容易把移项搞反，别漏了这个细节哈",
                "copyright_sources": [{"copyright_type": "authorized", "source_title": "初中数学精选题集（第3版）", "authorization_number": "AUTH-2024-00123"}],
            },
            {
                "question_id_external": "MATH-002",
                "question_content": "已知三角形ABC中，AB=3, BC=4, 角B=90度，求AC的长度",
                "standard_answer": "AC = 5",
                "difficulty": "中等",
                "knowledge_point": "勾股定理",
                "human_note": "勾股定理是高频考点，注意配图",
                "copyright_sources": [{"copyright_type": "public_domain", "source_title": "经典几何题库", "remark": "已进入公有领域"}],
            },
            {
                "question_id_external": "MATH-003",
                "question_content": "因式分解：x² - 6x + 9",
                "standard_answer": "(x-3)²",
                "difficulty": "简单",
                "knowledge_point": "因式分解",
                "copyright_sources": [{"copyright_type": "unknown", "remark": "原始来源待确认"}],
            },
            {
                "question_id_external": "MATH-004",
                "question_content": "求函数 y=x²-4x+3 的最小值",
                "standard_answer": "-1",
                "difficulty": "中等",
                "knowledge_point": "一元一次方程",
                "copyright_sources": [{"copyright_type": "fair_use", "source_title": "高中数学教案", "fair_use_justification": "教学评测合理使用"}],
            },
            {
                "question_id_external": "MATH-005",
                "question_content": "解不等式：3x - 7 > 2",
                "standard_answer": "x > 3",
                "difficulty": "简单",
                "knowledge_point": "一元一次方程",
                "copyright_sources": [{"copyright_type": "original"}],
            },
        ],
    }

    resp = client.post("/api/batches", json=batch_payload)
    assert resp.status_code == 200
    batch = resp.json()
    batch_id = batch["id"]
    question_ids = [q["id"] for q in batch["questions"]]
    print(f"[INFO] 批次已导入 batch_id={batch_id}")

    review_payload = {
        "reviewer": "李复核",
        "items": [
            {"question_id": question_ids[0], "passed": True},
            {
                "question_id": question_ids[1],
                "passed": False,
                "issue_type": "material_missing",
                "issue_detail": "公有领域证明材料缺失，需补充原始出版物扫描件",
                "next_action": "请题库同学补传原始出版物的版权页扫描件",
            },
            {
                "question_id": question_ids[2],
                "passed": False,
                "issue_type": "calibration_wrong",
                "issue_detail": "版权类型标记为unknown不符合入账要求，需重新确认",
                "next_action": "请评测负责人和出题人重新确认来源后改口径",
            },
            {"question_id": question_ids[3], "passed": True},
        ],
    }

    resp = client.post(f"/api/batches/{batch_id}/review", json=review_payload)
    assert resp.status_code == 200

    resp = client.get(f"/api/batches/{batch_id}/review-summary")
    after_review = resp.json()
    assert after_review["material_missing_count"] == 1
    assert after_review["calibration_wrong_count"] == 1
    assert after_review["total_passed"] == 2
    assert after_review["total_blocked"] == 2
    assert after_review["total_pending"] == 1
    print(f"[PASS] 复核后汇总正确：通过={after_review['total_passed']} 缺材料={after_review['material_missing_count']} 改口径={after_review['calibration_wrong_count']} 待处理={after_review['total_pending']}")

    status_payload = {
        "target_status": "rejected",
        "operator": "模型评审会-赵老师",
        "reason": "评测集知识点偏科严重，且仍有2题版权材料缺失，不符合上线要求",
    }
    resp = client.post(f"/api/batches/{batch_id}/status", json=status_payload)
    assert resp.status_code == 200
    print("[INFO] 评审会已将批次状态推进为 rejected")

    resp = client.get(f"/api/batches/{batch_id}/review-summary")
    after_reject = resp.json()
    assert after_reject["material_missing_count"] == 1, f"Bug: 缺材料被清零，实际={after_reject['material_missing_count']}"
    assert after_reject["calibration_wrong_count"] == 1, f"Bug: 改口径被清零，实际={after_reject['calibration_wrong_count']}"
    assert after_reject["total_pending"] == 1, f"Bug: 待处理被清零，实际={after_reject['total_pending']}"
    print("[PASS] 拦截后 review-summary 与拦截前一致")

    resp = client.get(f"/api/batches/{batch_id}/issue-breakdown")
    assert resp.status_code == 200
    breakdown = resp.json()
    mm_count = 0
    cw_count = 0
    for ib in breakdown:
        if ib["issue_type"] == "material_missing":
            mm_count = ib["count"]
        if ib["issue_type"] == "calibration_wrong":
            cw_count = ib["count"]
    assert mm_count == 1
    assert cw_count == 1
    print("[PASS] 拦截后 issue-breakdown 正确")

    resp = client.get(f"/api/batches/{batch_id}/export-report")
    assert resp.status_code == 200
    page_report = resp.json()
    ss = page_report["status_summary"]
    assert ss["material_missing"] == 1, f"报告 status_summary.material_missing={ss['material_missing']} 错误，应为1"
    assert ss["calibration_wrong"] == 1, f"报告 status_summary.calibration_wrong={ss['calibration_wrong']} 错误，应为1"
    assert ss["pending"] == 1, f"报告 status_summary.pending={ss['pending']} 错误，应为1"
    assert ss["review_passed"] == 2
    assert ss["review_blocked"] == 2
    check_sum = ss["review_passed"] + ss["review_blocked"] + ss["pending"]
    assert check_sum == ss["total_questions"], f"status_summary 数字加总 {check_sum} != 总题数 {ss['total_questions']}"
    print(f"[PASS] 页面报告 status_summary 正确且加总一致: {ss}")

    for ib in page_report["issue_breakdown"]:
        if ib["issue_type"] == "material_missing":
            assert ib["count"] == ss["material_missing"]
        if ib["issue_type"] == "calibration_wrong":
            assert ib["count"] == ss["calibration_wrong"]
    print("[PASS] 页面报告 issue_breakdown 与 status_summary 数字一致")

    for s in page_report["full_data"]["status_history"]:
        reason = s.get("reason") or ""
        assert "待处理-" not in reason, f"状态历史中出现负数待处理: {reason}"
    print("[PASS] 页面报告 status_history 中无负数待处理")

    rej_exp = page_report.get("rejection_explanation") or ""
    assert "不符合" in rej_exp or "偏科" in rej_exp
    plain = page_report["plain_explanation"]
    assert "需补材料 1 题" in plain
    assert "需改口径 1 题" in plain
    assert "未通过模型评审会审核" in plain
    print("[PASS] 页面报告 plain_explanation 和 rejection_explanation 正确可解释")

    _cleanup()
    resp = client.get(f"/api/batches/{batch_id}/download/json")
    assert resp.status_code == 200
    cd = resp.headers.get("content-disposition", "")
    assert "filename" in cd
    json_bytes = resp.content
    json_data = json.loads(json_bytes)
    print(f"[PASS] JSON 下载成功，大小={len(json_bytes)} 字节")

    assert json_data["status_summary"]["material_missing"] == 1
    assert json_data["status_summary"]["calibration_wrong"] == 1
    assert json_data["status_summary"]["pending"] == 1
    assert json_data["status_summary"]["review_passed"] == ss["review_passed"]
    assert json_data["status_summary"]["review_blocked"] == ss["review_blocked"]
    assert json_data["plain_explanation"] == page_report["plain_explanation"]
    assert json_data["rejection_explanation"] == page_report["rejection_explanation"]
    for s in json_data["status_history"]:
        reason = s.get("reason") or ""
        assert "待处理-" not in reason, f"JSON 导出状态历史中出现负数待处理: {reason}"
    print("[PASS] JSON 文件内容与页面报告完全一致，状态历史无负数")

    resp = client.get(f"/api/batches/{batch_id}/download/excel")
    assert resp.status_code == 200
    cd2 = resp.headers.get("content-disposition", "")
    assert "filename" in cd2 and ".xlsx" in cd2
    excel_bytes = resp.content
    assert len(excel_bytes) > 2000
    assert excel_bytes[:2] == b"PK"
    print(f"[PASS] Excel 下载成功，大小={len(excel_bytes)} 字节")

    wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
    sheet_names = wb.sheetnames
    expected_sheets = ["概览", "异常明细", "题目清单", "状态历史", "提示词版本"]
    for s in expected_sheets:
        assert s in sheet_names, f"Excel 缺失 sheet: {s}"
    print(f"[PASS] Excel 包含所有预期 sheet: {sheet_names}")

    ws_summary = wb["概览"]
    found_status_block = False
    status_dict = {}
    for row in ws_summary.iter_rows(min_row=1, max_row=ws_summary.max_row, values_only=True):
        if row and row[0]:
            key = str(row[0]).strip()
            value = str(row[1]).strip() if row[1] else ""
            if key == "总题数":
                status_dict["total"] = int(value)
            if key == "复核通过":
                status_dict["passed"] = int(value)
            if key == "复核受阻":
                status_dict["blocked"] = int(value)
            if "需补材料" in key:
                status_dict["mm"] = int(value)
            if "需改口径" in key:
                status_dict["cw"] = int(value)
            if key == "待复核":
                status_dict["pending"] = int(value)
            if "普通话解释" in key:
                found_status_block = True
    assert status_dict["mm"] == 1, f"Excel 概览页缺材料={status_dict['mm']} 错误"
    assert status_dict["cw"] == 1, f"Excel 概览页改口径={status_dict['cw']} 错误"
    assert status_dict["pending"] == 1, f"Excel 概览页待处理={status_dict['pending']} 错误"
    assert status_dict["mm"] + status_dict["cw"] == status_dict["blocked"]
    assert status_dict["total"] == status_dict["passed"] + status_dict["blocked"] + status_dict["pending"]
    print(f"[PASS] Excel 概览页数字正确且一致: {status_dict}")

    ws_issues = wb["异常明细"]
    mm_rows = 0
    cw_rows = 0
    for row_idx, row in enumerate(ws_issues.iter_rows(min_row=2, max_row=ws_issues.max_row, values_only=True)):
        if not row or not row[0]:
            continue
        issue_type = str(row[0]).strip()
        if issue_type == "需补材料":
            mm_rows += 1
            assert row[5] is not None, f"异常明细第 {row_idx+2} 行缺失人工备注"
        elif issue_type == "需改口径":
            cw_rows += 1
            assert row[4] is not None, f"异常明细第 {row_idx+2} 行缺失下一步动作"
    assert mm_rows == 1, f"Excel 异常明细页 缺材料行数={mm_rows} 错误"
    assert cw_rows == 1, f"Excel 异常明细页 改口径行数={cw_rows} 错误"
    print(f"[PASS] Excel 异常明细页正确，含 next_action 和人工备注原话保留")

    ws_history = wb["状态历史"]
    for row_idx, row in enumerate(ws_history.iter_rows(min_row=2, max_row=ws_history.max_row, values_only=True)):
        if row and row[4]:
            reason = str(row[4])
            assert "待处理-" not in reason, f"Excel 状态历史页第 {row_idx+2} 行出现负数待处理: {reason}"
    print("[PASS] Excel 状态历史页无负数待处理")

    ws_questions = wb["题目清单"]
    rows = list(ws_questions.iter_rows(min_row=2, max_row=ws_questions.max_row, values_only=True))
    assert len([r for r in rows if r and r[0]]) == 5, "Excel 题目清单行数不对"
    print(f"[PASS] Excel 题目清单页包含全部 {len([r for r in rows if r and r[0]])} 道题")

    json_files = list(settings.EXPORT_DIR.glob("*.json"))
    xlsx_files = list(settings.EXPORT_DIR.glob("*.xlsx"))
    assert len(json_files) >= 1
    assert len(xlsx_files) >= 1
    for f in json_files + xlsx_files:
        assert f.stat().st_size > 0
    print(f"[PASS] 导出文件已真实落盘: {[p.name for p in json_files]}, {[p.name for p in xlsx_files]}")

    print("\n=== 全部导出链路验证通过 ===")


if __name__ == "__main__":
    _cleanup()
    test_health()
    test_full_export_consistency()
