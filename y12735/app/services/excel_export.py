from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import ReviewRecord


def export_review_to_excel(record: ReviewRecord) -> bytes:
    buf = io.BytesIO()
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "复核结论"
    _style_header(ws_summary)

    rows = [
        ("复核编号", record.record_id),
        ("标题", record.title),
        ("来源文件", record.source_file),
        ("来源工作表", ", ".join(record.source_sheets)),
        ("当前状态", record.status.value),
        ("结果分类", record.classification.value),
        ("创建时间", record.created_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("更新时间", record.updated_at.strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if record.rejection_reason:
        rows.append(("驳回原因", record.rejection_reason))
    if record.editor_corrections:
        rows.append(("教研编辑修正", str(record.editor_corrections)))

    for i, (k, v) in enumerate(rows, start=1):
        ws_summary.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws_summary.cell(row=i, column=2, value=str(v))

    start = len(rows) + 2
    if record.explanation:
        ws_summary.cell(row=start, column=1, value="总体说明").font = Font(bold=True, size=12)
        ws_summary.cell(row=start + 1, column=1, value=record.explanation.summary).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        ws_summary.merge_cells(start_row=start + 1, start_column=1, end_row=start + 1, end_column=4)

        start += 3
        ws_summary.cell(row=start, column=1, value="投委会说明（直接可用 / 需复核）").font = Font(
            bold=True, size=12
        )
        cell = ws_summary.cell(row=start + 1, column=1, value=record.explanation.committee_summary)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws_summary.merge_cells(start_row=start + 1, start_column=1, end_row=start + 1, end_column=4)
        if record.classification.value == "可用":
            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        elif record.classification.value == "暂缓":
            cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        else:
            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        if record.explanation.editor_note:
            start += 3
            ws_summary.cell(row=start, column=1, value="教研编辑操作提示").font = Font(bold=True, size=12)
            cell = ws_summary.cell(row=start + 1, column=1, value=record.explanation.editor_note)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws_summary.merge_cells(start_row=start + 1, start_column=1, end_row=start + 4, end_column=4)

    for col in range(1, 5):
        ws_summary.column_dimensions[get_column_letter(col)].width = 30

    ws_data = wb.create_sheet("数据与拟合")
    headers = ["点号", "x", "y", "y_pred", "残差", "来源Sheet", "来源单元格", "标记", "简短说明"]
    for c, h in enumerate(headers, start=1):
        cell = ws_data.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    explanation_index = {o.point_index: o for o in record.outliers if o.point_index is not None}
    usable = set(record.explanation.usable_points) if record.explanation else set()
    deferred = set(record.explanation.deferred_points) if record.explanation else set()
    recollect = set(record.explanation.recollect_points) if record.explanation else set()

    for r, dp in enumerate(record.data_points, start=2):
        idx = dp.index
        ws_data.cell(row=r, column=1, value=idx + 1)
        ws_data.cell(row=r, column=2, value=dp.x)
        ws_data.cell(row=r, column=3, value=dp.y)
        if record.fit_result and idx < len(record.fit_result.y_predicted):
            ws_data.cell(row=r, column=4, value=record.fit_result.y_predicted[idx])
            ws_data.cell(row=r, column=5, value=record.fit_result.residuals[idx])
        ws_data.cell(row=r, column=6, value=dp.source_sheet or "")
        ws_data.cell(row=r, column=7, value=dp.source_cell or "")

        if idx in recollect:
            label = "建议重新采集"
            color = "FFC7CE"
        elif idx in deferred:
            label = "暂缓（需复核）"
            color = "FFEB9C"
        elif idx in usable:
            label = "可用"
            color = "C6EFCE"
        else:
            label = "-"
            color = None
        ws_data.cell(row=r, column=8, value=label)
        if color:
            for c in range(1, len(headers) + 1):
                ws_data.cell(row=r, column=c).fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid"
                )

        if idx in explanation_index:
            ws_data.cell(row=r, column=9, value=explanation_index[idx].short_explanation)

    for c in range(1, len(headers) + 1):
        ws_data.column_dimensions[get_column_letter(c)].width = 18

    if record.fit_result:
        row = len(record.data_points) + 3
        ws_data.cell(row=row, column=1, value="拟合方程").font = Font(bold=True)
        ws_data.cell(row=row, column=2, value=record.fit_result.model_formula)
        ws_data.cell(row=row + 1, column=1, value="R²").font = Font(bold=True)
        ws_data.cell(row=row + 1, column=2, value=record.fit_result.r_squared)
        for i, (name, val, err) in enumerate(
            zip(
                record.fit_result.param_names,
                record.fit_result.params,
                record.fit_result.param_errors,
            )
        ):
            ws_data.cell(row=row + 2 + i, column=1, value=name).font = Font(bold=True)
            ws_data.cell(row=row + 2 + i, column=2, value=val)
            ws_data.cell(row=row + 2 + i, column=3, value=f"±{err}")

    if record.constraint_violations:
        ws_cv = wb.create_sheet("约束违反")
        headers_cv = ["参数", "实际值", "下限", "上限", "简短说明", "详细说明"]
        for c, h in enumerate(headers_cv, start=1):
            cell = ws_cv.cell(row=1, column=c, value=h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        for r, cv in enumerate(record.constraint_violations, start=2):
            ws_cv.cell(row=r, column=1, value=cv.constraint.param_name)
            ws_cv.cell(row=r, column=2, value=cv.actual_value)
            ws_cv.cell(row=r, column=3, value=cv.constraint.lower if cv.constraint.lower is not None else "-")
            ws_cv.cell(row=r, column=4, value=cv.constraint.upper if cv.constraint.upper is not None else "-")
            ws_cv.cell(row=r, column=5, value=cv.short_explanation)
            ws_cv.cell(row=r, column=6, value=cv.detailed_explanation).alignment = Alignment(wrap_text=True)
        for c in range(1, len(headers_cv) + 1):
            ws_cv.column_dimensions[get_column_letter(c)].width = 22

    ws_audit = wb.create_sheet("操作审计")
    headers_audit = ["时间", "操作人", "角色", "原状态", "新状态", "备注", "字段变更"]
    for c, h in enumerate(headers_audit, start=1):
        cell = ws_audit.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for r, entry in enumerate(record.audit_log, start=2):
        ws_audit.cell(row=r, column=1, value=entry.timestamp.strftime("%Y-%m-%d %H:%M:%S"))
        ws_audit.cell(row=r, column=2, value=entry.reviewer)
        ws_audit.cell(row=r, column=3, value=entry.role.value)
        ws_audit.cell(row=r, column=4, value=entry.from_status.value if entry.from_status else "-")
        ws_audit.cell(row=r, column=5, value=entry.to_status.value)
        ws_audit.cell(row=r, column=6, value=entry.note)
        ws_audit.cell(row=r, column=7, value=str(entry.field_changes) if entry.field_changes else "")
    for c in range(1, len(headers_audit) + 1):
        ws_audit.column_dimensions[get_column_letter(c)].width = 20

    wb.save(buf)
    return buf.getvalue()


def _style_header(ws) -> None:
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
