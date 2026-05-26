from flask import Blueprint, request, jsonify, make_response
from datetime import datetime, timedelta
from dateutil.relativedelta import relativedelta
import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from app import db
from app.models import (
    PointsLedger, PointTransaction, OrderRefund, Coupon, 
    CouponRedemption, ActivityPlan, DataSource, LiabilityForecast
)
from app.services.points_service import PointsExpiryService, RefundService, CouponService
from app.services.forecast_engine import ForecastEngine

bp = Blueprint('api', __name__)

def parse_date(date_str):
    if not date_str:
        return None
    for fmt in ['%Y-%m-%d', '%Y-%m-%d %H:%M:%S']:
        try:
            return datetime.strptime(date_str, fmt)
        except ValueError:
            continue
    return None

def error_response(message, status_code=400, source_reference=None, details=None):
    response = {
        'error': message,
        'source_reference': source_reference
    }
    if details:
        response['details'] = details
    return jsonify(response), status_code

@bp.route('/help')
def api_help():
    return jsonify({
        'api_name': '会员积分负债预测 API',
        'version': '1.0.0',
        'endpoints': {
            '预测管理': {
                'POST /api/forecast': '创建新的负债预测',
                'GET /api/forecast': '获取预测列表',
                'GET /api/forecast/<id>': '获取单个预测详情',
                'PUT /api/forecast/<id>/status': '推进预测状态',
                'POST /api/forecast/<id>/correct': '修正预测数据',
                'GET /api/forecast/<id>/export': '导出预测报告'
            },
            '积分账本': {
                'GET /api/points/ledger': '查询积分账本',
                'GET /api/points/ledger/<member_id>': '查询会员积分明细',
                'GET /api/points/expiry/simulate': '积分过期试算'
            },
            '退货管理': {
                'GET /api/refunds/pending': '获取待处理退款列表',
                'POST /api/refunds/<id>/process': '处理退款返积分'
            },
            '兑换券管理': {
                'GET /api/coupons': '获取兑换券列表',
                'POST /api/coupons/redeem': '积分兑换券',
                'GET /api/coupons/pending-verify': '获取待核销列表',
                'POST /api/coupons/verify/<id>': '核销兑换券'
            },
            '活动计划': {
                'GET /api/activities': '获取活动计划列表',
                'POST /api/activities': '创建活动计划'
            },
            '数据导入': {
                'POST /api/data/import': '导入业务数据'
            }
        }
    })

@bp.route('/forecast', methods=['POST'])
def create_forecast():
    data = request.get_json()
    if not data:
        return error_response('请求体不能为空')
    
    required_fields = ['forecast_name', 'start_date', 'end_date']
    for field in required_fields:
        if field not in data:
            return error_response(f'缺少必填字段: {field}')
    
    start_date = parse_date(data['start_date'])
    end_date = parse_date(data['end_date'])
    
    if not start_date or not end_date:
        return error_response('日期格式错误，请使用 YYYY-MM-DD 格式')
    
    if start_date >= end_date:
        return error_response('开始日期必须早于结束日期')
    
    result = ForecastEngine.create_forecast(
        forecast_name=data['forecast_name'],
        start_date=start_date,
        end_date=end_date,
        period=data.get('forecast_period', 'monthly'),
        points_per_yuan=data.get('points_per_yuan', 0.01),
        created_by=data.get('created_by'),
        remark=data.get('remark')
    )
    
    if 'error' in result:
        return error_response(result['error'], source_reference=result.get('source_reference'))
    
    return jsonify(result), 201

@bp.route('/forecast', methods=['GET'])
def list_forecasts():
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 20))
    
    result = ForecastEngine.list_forecasts(page=page, per_page=per_page)
    return jsonify(result)

@bp.route('/forecast/<int:forecast_id>', methods=['GET'])
def get_forecast(forecast_id):
    result = ForecastEngine.get_forecast(forecast_id)
    if not result:
        return error_response(f'预测记录不存在: ID={forecast_id}', status_code=404)
    return jsonify(result)

@bp.route('/forecast/<int:forecast_id>/status', methods=['PUT'])
def update_forecast_status(forecast_id):
    data = request.get_json() or {}
    new_status = data.get('status')
    
    valid_statuses = ['draft', 'reviewing', 'approved', 'rejected', 'published', 'archived']
    if not new_status or new_status not in valid_statuses:
        return error_response(f'无效的状态值，有效值为: {valid_statuses}')
    
    result = ForecastEngine.update_forecast_status(
        forecast_id=forecast_id,
        new_status=new_status,
        updated_by=data.get('updated_by')
    )
    
    if 'error' in result:
        return error_response(result['error'], status_code=404)
    
    return jsonify(result)

@bp.route('/forecast/<int:forecast_id>/correct', methods=['POST'])
def correct_forecast(forecast_id):
    data = request.get_json()
    if not data:
        return error_response('请求体不能为空')
    
    required_fields = ['field_name', 'old_value', 'new_value', 'correction_reason']
    for field in required_fields:
        if field not in data:
            return error_response(f'缺少必填字段: {field}')
    
    result = ForecastEngine.apply_correction(
        forecast_id=forecast_id,
        field_name=data['field_name'],
        old_value=data['old_value'],
        new_value=data['new_value'],
        correction_reason=data['correction_reason'],
        corrected_by=data.get('corrected_by'),
        source_reference=data.get('source_reference')
    )
    
    if 'error' in result:
        return error_response(result['error'], source_reference=result.get('source_reference'), 
                             details=result.get('valid_fields'))
    
    return jsonify(result)

@bp.route('/forecast/<int:forecast_id>/export', methods=['GET'])
def export_forecast(forecast_id):
    forecast = LiabilityForecast.query.get(forecast_id)
    if not forecast:
        return error_response(f'预测记录不存在: ID={forecast_id}', status_code=404)
    
    forecast_data = ForecastEngine.get_forecast(forecast_id)
    format_type = request.args.get('format', 'xlsx')
    
    if format_type == 'csv':
        return _export_csv(forecast_data)
    else:
        return _export_excel(forecast_data)

def _export_csv(forecast_data):
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow(['会员积分负债预测报告'])
    writer.writerow(['预测名称', forecast_data['forecast_name']])
    writer.writerow(['预测日期', forecast_data['forecast_date']])
    writer.writerow(['预测期间', f"{forecast_data['start_date']} 至 {forecast_data['end_date']}"])
    writer.writerow([])
    
    writer.writerow(['=== 核心指标 ==='])
    writer.writerow(['指标', '数值'])
    writer.writerow(['当前积分总余额', forecast_data['total_points_balance']])
    writer.writerow(['预计过期积分', forecast_data['expected_expired_points']])
    writer.writerow(['预计退款返积分', forecast_data['expected_refund_points']])
    writer.writerow(['预计兑换消耗积分', forecast_data['expected_redemption_points']])
    writer.writerow(['预计兑换券成本(元)', forecast_data['expected_coupon_cost']])
    writer.writerow(['总预计负债(元)', forecast_data['total_estimated_liability']])
    writer.writerow([])
    
    writer.writerow(['=== 负债构成 ==='])
    writer.writerow(['项目', '金额(元)', '占比'])
    total = forecast_data['total_estimated_liability'] or 1
    summary = forecast_data['summary']
    writer.writerow(['积分过期负债', summary['expired_liability'], f"{summary['expired_liability']/total*100:.1f}%"])
    writer.writerow(['退款返积分负债', summary['refund_liability'], f"{summary['refund_liability']/total*100:.1f}%"])
    writer.writerow(['兑换券成本', summary['coupon_liability'], f"{summary['coupon_liability']/total*100:.1f}%"])
    writer.writerow([])
    
    writer.writerow(['=== 风险提示 ==='])
    writer.writerow(['级别', '类型', '提示信息', '来源参考'])
    for w in forecast_data['warnings']:
        writer.writerow([w['level'], w['type'], w['message'], w['source_reference']])
    writer.writerow([])
    
    writer.writerow(['=== 预测曲线 ==='])
    writer.writerow(['日期', '累计负债(元)', '累计过期积分', '累计兑换积分', '累计退款积分'])
    for c in forecast_data['curve']:
        writer.writerow([c['date'], c['cumulative_liability'], c['expired_points'], 
                        c['redemption_points'], c['refund_points']])
    
    output.seek(0)
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'text/csv; charset=utf-8'
    response.headers['Content-Disposition'] = f"attachment; filename=forecast_{forecast_data['id']}.csv"
    return response

def _export_excel(forecast_data):
    wb = Workbook()
    
    ws = wb.active
    ws.title = '预测报告'
    
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    center_align = Alignment(horizontal='center', vertical='center')
    
    row = 1
    ws.cell(row=row, column=1, value='会员积分负债预测报告').font = Font(bold=True, size=16)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 2
    
    info_data = [
        ['预测名称', forecast_data['forecast_name']],
        ['预测日期', forecast_data['forecast_date']],
        ['预测期间', f"{forecast_data['start_date']} 至 {forecast_data['end_date']}"],
        ['状态', forecast_data['status']],
        ['创建人', forecast_data['created_by'] or '-']
    ]
    for item in info_data:
        ws.cell(row=row, column=1, value=item[0]).font = Font(bold=True)
        ws.cell(row=row, column=2, value=item[1])
        row += 1
    row += 1
    
    ws.cell(row=row, column=1, value='核心指标').font = Font(bold=True, size=14)
    row += 1
    metrics_headers = ['指标', '数值']
    for col, header in enumerate(metrics_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    row += 1
    
    metrics = [
        ['当前积分总余额', forecast_data['total_points_balance']],
        ['预计过期积分', forecast_data['expected_expired_points']],
        ['预计退款返积分', forecast_data['expected_refund_points']],
        ['预计兑换消耗积分', forecast_data['expected_redemption_points']],
        ['预计兑换券成本(元)', forecast_data['expected_coupon_cost']],
        ['总预计负债(元)', forecast_data['total_estimated_liability']],
        ['积分折算比例(元/积分)', forecast_data['points_per_yuan']]
    ]
    for m in metrics:
        ws.cell(row=row, column=1, value=m[0])
        ws.cell(row=row, column=2, value=m[1])
        row += 1
    row += 1
    
    ws.cell(row=row, column=1, value='负债构成分析').font = Font(bold=True, size=14)
    row += 1
    comp_headers = ['项目', '金额(元)', '占比']
    for col, header in enumerate(comp_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    row += 1
    
    total = forecast_data['total_estimated_liability'] or 1
    summary = forecast_data['summary']
    components = [
        ['积分过期负债', summary['expired_liability'], summary['expired_liability']/total*100],
        ['退款返积分负债', summary['refund_liability'], summary['refund_liability']/total*100],
        ['兑换券成本', summary['coupon_liability'], summary['coupon_liability']/total*100]
    ]
    for comp in components:
        ws.cell(row=row, column=1, value=comp[0])
        ws.cell(row=row, column=2, value=comp[1])
        ws.cell(row=row, column=3, value=f"{comp[2]:.1f}%")
        row += 1
    row += 1
    
    ws.cell(row=row, column=1, value='风险提示').font = Font(bold=True, size=14)
    row += 1
    warn_headers = ['级别', '类型', '提示信息', '来源参考']
    for col, header in enumerate(warn_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    row += 1
    
    for w in forecast_data['warnings']:
        ws.cell(row=row, column=1, value=w['level'])
        ws.cell(row=row, column=2, value=w['type'])
        ws.cell(row=row, column=3, value=w['message'])
        ws.cell(row=row, column=4, value=w['source_reference'])
        row += 1
    row += 1
    
    ws.cell(row=row, column=1, value='修正记录').font = Font(bold=True, size=14)
    row += 1
    corr_headers = ['修正时间', '字段', '原值', '新值', '修正原因', '操作人']
    for col, header in enumerate(corr_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    row += 1
    
    for c in forecast_data['corrections']:
        ws.cell(row=row, column=1, value=c['created_at'])
        ws.cell(row=row, column=2, value=c['field_name'])
        ws.cell(row=row, column=3, value=c['old_value'])
        ws.cell(row=row, column=4, value=c['new_value'])
        ws.cell(row=row, column=5, value=c['correction_reason'])
        ws.cell(row=row, column=6, value=c['corrected_by'] or '-')
        row += 1
    row += 1
    
    ws.cell(row=row, column=1, value='预测曲线数据').font = Font(bold=True, size=14)
    row += 1
    curve_headers = ['日期', '累计负债(元)', '累计过期积分', '累计兑换积分', '累计退款积分']
    for col, header in enumerate(curve_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    row += 1
    
    for c in forecast_data['curve']:
        ws.cell(row=row, column=1, value=c['date'])
        ws.cell(row=row, column=2, value=c['cumulative_liability'])
        ws.cell(row=row, column=3, value=c['expired_points'])
        ws.cell(row=row, column=4, value=c['redemption_points'])
        ws.cell(row=row, column=5, value=c['refund_points'])
        row += 1
    
    for col in range(1, 7):
        ws.column_dimensions[chr(64 + col)].width = 20
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f"attachment; filename=forecast_{forecast_data['id']}.xlsx"
    return response

@bp.route('/points/ledger', methods=['GET'])
def list_ledgers():
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 20))
    member_id = request.args.get('member_id')
    
    query = PointsLedger.query
    if member_id:
        query = query.filter(PointsLedger.member_id.like(f'%{member_id}%'))
    
    total = query.count()
    ledgers = query.order_by(PointsLedger.last_updated.desc()).offset((page-1)*per_page).limit(per_page).all()
    
    return jsonify({
        'total': total,
        'page': page,
        'per_page': per_page,
        'items': [{
            'id': l.id,
            'member_id': l.member_id,
            'member_name': l.member_name,
            'total_points': l.total_points,
            'available_points': l.available_points,
            'frozen_points': l.frozen_points,
            'expired_points': l.expired_points,
            'last_updated': l.last_updated.strftime('%Y-%m-%d %H:%M:%S') if l.last_updated else None
        } for l in ledgers]
    })

@bp.route('/points/ledger/<member_id>', methods=['GET'])
def get_ledger_detail(member_id):
    ledger = PointsLedger.query.filter_by(member_id=member_id).first()
    if not ledger:
        return error_response(f'会员账本不存在: {member_id}', status_code=404)
    
    transactions = PointTransaction.query.filter_by(member_id=member_id).order_by(PointTransaction.transaction_time.desc()).limit(50).all()
    
    return jsonify({
        'ledger': {
            'id': ledger.id,
            'member_id': ledger.member_id,
            'member_name': ledger.member_name,
            'total_points': ledger.total_points,
            'available_points': ledger.available_points,
            'frozen_points': ledger.frozen_points,
            'expired_points': ledger.expired_points
        },
        'recent_transactions': [{
            'id': t.id,
            'type': t.transaction_type,
            'points': t.points,
            'balance_after': t.balance_after,
            'expire_date': t.expire_date.strftime('%Y-%m-%d') if t.expire_date else None,
            'order_no': t.order_no,
            'coupon_code': t.coupon_code,
            'remark': t.remark,
            'is_refund': t.is_refund,
            'transaction_time': t.transaction_time.strftime('%Y-%m-%d %H:%M:%S'),
            'source_line': t.source_line
        } for t in transactions]
    })

@bp.route('/points/expiry/simulate', methods=['GET'])
def simulate_expiry():
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    
    start_date = parse_date(start_date_str) or datetime.now()
    end_date = parse_date(end_date_str) or (start_date + relativedelta(months=3))
    
    result = PointsExpiryService.batch_calculate_expiry(start_date, end_date)
    return jsonify({
        'period': {
            'start_date': start_date.strftime('%Y-%m-%d'),
            'end_date': end_date.strftime('%Y-%m-%d')
        },
        'total_expired_points': result['total_expired_points'],
        'affected_ledger_count': result['ledger_count'],
        'warnings': result['warnings'],
        'details': result['results'][:100]
    })

@bp.route('/refunds/pending', methods=['GET'])
def list_pending_refunds():
    result = RefundService.get_pending_refunds()
    return jsonify({
        'count': len(result),
        'items': result
    })

@bp.route('/refunds/<int:refund_id>/process', methods=['POST'])
def process_refund(refund_id):
    result = RefundService.process_refund(refund_id)
    if 'error' in result:
        return error_response(result['error'], source_reference=result.get('source_reference'))
    return jsonify(result)

@bp.route('/coupons', methods=['GET'])
def list_coupons():
    coupons = Coupon.query.all()
    return jsonify({
        'count': len(coupons),
        'items': [{
            'id': c.id,
            'coupon_code': c.coupon_code,
            'coupon_name': c.coupon_name,
            'coupon_type': c.coupon_type,
            'points_cost': c.points_cost,
            'face_value': c.face_value,
            'total_quantity': c.total_quantity,
            'used_quantity': c.used_quantity,
            'remaining_quantity': c.total_quantity - c.used_quantity,
            'expire_date': c.expire_date.strftime('%Y-%m-%d'),
            'verify_fail_count': c.verify_fail_count,
            'source_line': c.source_line
        } for c in coupons]
    })

@bp.route('/coupons/redeem', methods=['POST'])
def redeem_coupon():
    data = request.get_json()
    if not data:
        return error_response('请求体不能为空')
    
    required = ['coupon_code', 'member_id']
    for f in required:
        if f not in data:
            return error_response(f'缺少必填字段: {f}')
    
    result = CouponService.redeem_coupon(data['coupon_code'], data['member_id'])
    if 'error' in result:
        return error_response(result['error'], source_reference=result.get('source_reference'))
    return jsonify(result)

@bp.route('/coupons/pending-verify', methods=['GET'])
def list_pending_verify():
    result = CouponService.get_pending_verifications()
    return jsonify({
        'count': len(result),
        'items': result
    })

@bp.route('/coupons/verify/<int:redemption_id>', methods=['POST'])
def verify_coupon(redemption_id):
    data = request.get_json() or {}
    success = data.get('success', True)
    fail_reason = data.get('fail_reason')
    
    if not success and not fail_reason:
        return error_response('核销失败时必须提供失败原因')
    
    result = CouponService.verify_redemption(redemption_id, success=success, fail_reason=fail_reason)
    if 'error' in result:
        return error_response(result['error'], source_reference=result.get('source_reference'))
    return jsonify(result)

@bp.route('/activities', methods=['GET'])
def list_activities():
    activities = ActivityPlan.query.all()
    return jsonify({
        'count': len(activities),
        'items': [{
            'id': a.id,
            'activity_name': a.activity_name,
            'activity_type': a.activity_type,
            'start_date': a.start_date.strftime('%Y-%m-%d'),
            'end_date': a.end_date.strftime('%Y-%m-%d'),
            'expected_points_issued': a.expected_points_issued,
            'expected_redemption_rate': a.expected_redemption_rate,
            'expected_coupon_cost': a.expected_coupon_cost,
            'status': a.status,
            'remark': a.remark
        } for a in activities]
    })

@bp.route('/activities', methods=['POST'])
def create_activity():
    data = request.get_json()
    if not data:
        return error_response('请求体不能为空')
    
    required = ['activity_name', 'activity_type', 'start_date', 'end_date']
    for f in required:
        if f not in data:
            return error_response(f'缺少必填字段: {f}')
    
    start_date = parse_date(data['start_date'])
    end_date = parse_date(data['end_date'])
    
    if not start_date or not end_date:
        return error_response('日期格式错误，请使用 YYYY-MM-DD 格式')
    
    activity = ActivityPlan(
        activity_name=data['activity_name'],
        activity_type=data['activity_type'],
        start_date=start_date,
        end_date=end_date,
        expected_points_issued=data.get('expected_points_issued', 0),
        expected_redemption_rate=data.get('expected_redemption_rate', 0.3),
        expected_coupon_cost=data.get('expected_coupon_cost', 0),
        status=data.get('status', 'planned'),
        remark=data.get('remark')
    )
    
    db.session.add(activity)
    db.session.commit()
    
    return jsonify({
        'success': True,
        'activity_id': activity.id,
        'activity_name': activity.activity_name
    }), 201

@bp.route('/data/import', methods=['POST'])
def import_data():
    data = request.get_json()
    if not data:
        return error_response('请求体不能为空')
    
    source_name = data.get('source_name', '手动导入')
    source_type = data.get('source_type', 'manual')
    file_name = data.get('file_name')
    records = data.get('records', [])
    
    if not records:
        return error_response('没有提供导入数据')
    
    source = DataSource(
        source_name=source_name,
        source_type=source_type,
        file_name=file_name,
        record_count=len(records)
    )
    db.session.add(source)
    db.session.flush()
    
    errors = []
    success_count = 0
    
    for idx, record in enumerate(records, start=1):
        record_type = record.get('type')
        try:
            if record_type == 'points_ledger':
                ledger = PointsLedger(
                    member_id=record['member_id'],
                    member_name=record.get('member_name'),
                    total_points=record.get('total_points', 0),
                    available_points=record.get('available_points', 0),
                    frozen_points=record.get('frozen_points', 0),
                    expired_points=record.get('expired_points', 0),
                    source_id=source.id,
                    source_line=idx
                )
                db.session.add(ledger)
                success_count += 1
            elif record_type == 'transaction':
                txn = PointTransaction(
                    ledger_id=record['ledger_id'],
                    member_id=record['member_id'],
                    transaction_type=record['transaction_type'],
                    points=record['points'],
                    balance_after=record.get('balance_after'),
                    expire_date=parse_date(record.get('expire_date')),
                    order_no=record.get('order_no'),
                    coupon_code=record.get('coupon_code'),
                    remark=record.get('remark'),
                    transaction_time=parse_date(record.get('transaction_time')) or datetime.now(),
                    source_id=source.id,
                    source_line=idx
                )
                db.session.add(txn)
                success_count += 1
            elif record_type == 'refund':
                refund = OrderRefund(
                    order_no=record['order_no'],
                    member_id=record['member_id'],
                    refund_time=parse_date(record['refund_time']) or datetime.now(),
                    original_points=record.get('original_points', 0),
                    return_points=record.get('return_points', 0),
                    remark=record.get('remark'),
                    source_id=source.id,
                    source_line=idx
                )
                db.session.add(refund)
                success_count += 1
            elif record_type == 'coupon':
                coupon = Coupon(
                    coupon_code=record['coupon_code'],
                    coupon_name=record['coupon_name'],
                    coupon_type=record.get('coupon_type', 'discount'),
                    points_cost=record.get('points_cost', 0),
                    face_value=record.get('face_value', 0),
                    total_quantity=record.get('total_quantity', 0),
                    expire_date=parse_date(record['expire_date']) or datetime.now(),
                    source_id=source.id,
                    source_line=idx
                )
                db.session.add(coupon)
                success_count += 1
            else:
                errors.append({
                    'line': idx,
                    'error': f'不支持的记录类型: {record_type}',
                    'record': record
                })
        except Exception as e:
            errors.append({
                'line': idx,
                'error': str(e),
                'record': record
            })
    
    db.session.commit()
    
    return jsonify({
        'success': True,
        'source_id': source.id,
        'total_records': len(records),
        'success_count': success_count,
        'error_count': len(errors),
        'errors': errors
    })
