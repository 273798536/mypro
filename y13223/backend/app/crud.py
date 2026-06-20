from __future__ import annotations

from sqlalchemy.orm import Session
from datetime import datetime
from . import models, schemas
from .humanize import anomaly_reason_to_human, bad_data_hint


def _log_change(db: Session, track_id: int, field_name: str,
                old_value: str | None, new_value: str | None,
                source: str, source_ref: str | None = None,
                operator: str | None = None):
    if old_value == new_value:
        return
    log = models.ChangeLog(
        track_id=track_id,
        field_name=field_name,
        old_value=old_value,
        new_value=new_value,
        source=source,
        source_ref=source_ref,
        operator=operator,
    )
    db.add(log)


def _enrich(track: models.Track) -> schemas.TrackOut:
    data = schemas.TrackOut.model_validate(track)
    data.anomaly_reason_human = anomaly_reason_to_human(
        track.anomaly_type, track.current_note,
        track.filename, track.track_name, track.contract_scan_ref
    )
    data.bad_data_hint = bad_data_hint(
        track.contract_scan_ref, track.current_source, track.is_anomaly
    )
    return data


def list_tracks(db: Session, only_anomaly: bool = False) -> list[schemas.TrackOut]:
    query = db.query(models.Track)
    if only_anomaly:
        query = query.filter(models.Track.is_anomaly == True)
    tracks = query.order_by(models.Track.updated_at.desc()).all()
    return [_enrich(t) for t in tracks]


def get_track(db: Session, track_id: int) -> schemas.TrackOut | None:
    track = db.query(models.Track).filter(models.Track.id == track_id).first()
    return _enrich(track) if track else None


def create_track(db: Session, data: schemas.TrackCreate,
                 operator: str | None = None, source_ref: str | None = None) -> schemas.TrackOut:
    track = models.Track(**data.model_dump())
    db.add(track)
    db.flush()
    _log_change(db, track.id, "track:create", None,
                f"{track.filename}|{track.track_name}",
                data.current_source or "system", source_ref, operator)
    db.commit()
    db.refresh(track)
    return _enrich(track)


def update_track(db: Session, track_id: int, data: schemas.TrackUpdate) -> schemas.TrackOut | None:
    track = db.query(models.Track).filter(models.Track.id == track_id).first()
    if not track:
        return None

    update_data = data.model_dump(exclude_unset=True)
    operator = update_data.pop("operator", None)
    source_ref = update_data.pop("source_ref", None)
    source = update_data.get("current_source") or getattr(track, "current_source", "system")

    for field, new_val in update_data.items():
        old_val = getattr(track, field, None)
        if old_val != new_val:
            _log_change(db, track.id, field,
                        str(old_val) if old_val is not None else None,
                        str(new_val) if new_val is not None else None,
                        source, source_ref, operator)
            setattr(track, field, new_val)

    track.updated_at = datetime.now()
    db.commit()
    db.refresh(track)
    return _enrich(track)


def delete_track(db: Session, track_id: int) -> bool:
    track = db.query(models.Track).filter(models.Track.id == track_id).first()
    if not track:
        return False
    db.delete(track)
    db.commit()
    return True


def human_verify(db: Session, track_id: int, data: schemas.HumanVerify) -> schemas.TrackOut | None:
    track = db.query(models.Track).filter(models.Track.id == track_id).first()
    if not track:
        return None

    op = data.operator or data.verifier

    if not data.verified and track.is_anomaly:
        _log_change(db, track.id, "is_anomaly",
                    str(track.is_anomaly), "False",
                    "human_override", None, op)
        _log_change(db, track.id, "anomaly_type",
                    track.anomaly_type or "", "",
                    "human_override", None, op)
        track.is_anomaly = False
        track.anomaly_type = None
    elif data.verified and not track.is_anomaly:
        _log_change(db, track.id, "is_anomaly",
                    str(track.is_anomaly), "True",
                    "human_override", None, op)
        track.is_anomaly = True

    _log_change(db, track.id, "human_verified",
                str(track.human_verified), str(data.verified),
                "human_override", None, op)
    _log_change(db, track.id, "human_verify_reason",
                track.human_verify_reason, data.reason,
                "human_override", None, op)

    track.human_verified = data.verified
    track.human_verifier = data.verifier
    track.human_verify_reason = data.reason
    track.next_step = data.next_step
    track.updated_at = datetime.now()
    db.commit()
    db.refresh(track)
    return _enrich(track)


def list_change_logs(db: Session, track_id: int) -> list[schemas.ChangeLogOut]:
    logs = (db.query(models.ChangeLog)
              .filter(models.ChangeLog.track_id == track_id)
              .order_by(models.ChangeLog.changed_at.desc())
              .all())
    return [schemas.ChangeLogOut.model_validate(l) for l in logs]


def export_excel(db: Session, only_anomaly: bool = False) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    import io

    tracks = list_tracks(db, only_anomaly=only_anomaly)

    wb = Workbook()
    ws = wb.active
    ws.title = "鼓组节拍异常提醒"

    headers = [
        "ID", "文件名", "曲目名称", "曲目编号", "艺人/演奏者",
        "合同扫描件定位", "是否异常", "异常类型", "异常说明(人话)",
        "当前备注", "备注来源", "材料定位提示",
        "人工复核", "复核人", "复核原因", "下一步",
        "创建时间", "最后更新"
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    anomaly_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for t in tracks:
        row = [
            t.id, t.filename, t.track_name, t.track_no or "", t.artist or "",
            t.contract_scan_ref or "",
            "是" if t.is_anomaly else "否",
            t.anomaly_type or "",
            t.anomaly_reason_human or "",
            t.current_note or "",
            t.current_source or "",
            t.bad_data_hint or "",
            "已复核" if t.human_verified else "待复核",
            t.human_verifier or "",
            t.human_verify_reason or "",
            t.next_step or "",
            t.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            t.updated_at.strftime("%Y-%m-%d %H:%M:%S"),
        ]
        ws.append(row)
        if t.is_anomaly:
            for cell in ws[ws.max_row]:
                cell.fill = anomaly_fill

    widths = [6, 28, 28, 12, 20, 24, 10, 14, 40, 40, 14, 40, 10, 12, 30, 30, 20, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
